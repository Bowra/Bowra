"""Bounded evidence preparation for the isolated document worker only.

Never evaluate spreadsheet formulas or fetch external links. Original uploads
remain frozen separately; this normalized representation is also hashed into the
planning context. CSV numeric syntax is deliberately narrow to preserve IDs.
"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from bowhead_core.documents.package_io import MAX_EXPANDED_BYTES, MAX_PART_BYTES, MAX_PARTS, OfficePackageError, _safe_part, _xml, _XmlBudget
from bowhead_core.documents.pptx_plan import JsonEvidenceSource, _json

MAX_DATA_BYTES = 4 * 1024 * 1024
MAX_CELLS = 512
_NUMBER = re.compile(r"-?(?:0|[1-9][0-9]{0,99})(?:\.[0-9]{1,100})?\Z")


def _encode(value: object) -> str:
    if isinstance(value, Decimal):
        if not value.is_finite():
            raise OfficePackageError("document_data_invalid")
        return str(value)
    if isinstance(value, list):
        return "[" + ",".join(_encode(item) for item in value) + "]"
    if isinstance(value, dict):
        return "{" + ",".join(json.dumps(key, ensure_ascii=False) + ":" + _encode(item) for key, item in value.items()) + "}"
    return json.dumps(value, ensure_ascii=False, allow_nan=False)


def _xlsx(content: bytes) -> object:
    from openpyxl import load_workbook

    try:
        with ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_PARTS or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                raise OfficePackageError("document_data_limit")
            seen = set()
            budget = _XmlBudget()
            for entry in entries:
                name = entry.filename
                if entry.is_dir():
                    if not _safe_part(name.rstrip("/")) or entry.file_size:
                        raise OfficePackageError("document_data_invalid")
                    continue
                if not _safe_part(name) or name.casefold() in seen or entry.flag_bits & 1 or entry.file_size > MAX_PART_BYTES:
                    raise OfficePackageError("document_data_invalid")
                seen.add(name.casefold())
                if any(marker in name.casefold() for marker in ("vbaproject", "externallinks/", "embeddings/", "connections.xml")):
                    raise OfficePackageError("document_data_unsupported")
                raw = archive.read(entry)
                if name.endswith((".xml", ".rels")):
                    root = _xml(raw, budget)
                    if name.endswith(".rels") and any(node.get("TargetMode") == "External" for node in root):
                        raise OfficePackageError("document_data_unsupported")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, ValueError, OSError) as exc:
        if isinstance(exc, OfficePackageError):
            raise
        raise OfficePackageError("document_data_invalid") from exc
    result, count = {}, 0
    try:
        if len(workbook.worksheets) > 8:
            raise OfficePackageError("document_data_limit")
        for sheet in workbook.worksheets:
            if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_CELLS:
                raise OfficePackageError("document_data_limit")
            rows = []
            for row in sheet.iter_rows():
                values = []
                for cell in row:
                    count += 1
                    if count > MAX_CELLS:
                        raise OfficePackageError("document_data_limit")
                    if cell.data_type in {"f", "e"}:
                        raise OfficePackageError("document_data_formula_unsupported")
                    value = cell.value
                    if isinstance(value, (datetime, date, time)):
                        value = value.isoformat()
                    elif isinstance(value, float):
                        value = Decimal(str(value))
                    values.append(value)
                rows.append(values)
            result[sheet.title] = rows
    finally:
        workbook.close()
    return result


def prepare_evidence(name: str, content: bytes) -> JsonEvidenceSource:
    """Normalize a supported upload without truncating cells or inventing data."""
    if not name or len(name) > 256 or len(content) > MAX_DATA_BYTES:
        raise OfficePackageError("document_data_limit")
    suffix = PurePosixPath(name).suffix.lower()
    if suffix == ".json":
        if len(content) > 128_000:
            raise OfficePackageError("document_data_limit")
        _json(content)
        return JsonEvidenceSource(name, content)
    if suffix == ".xlsx":
        value = _xlsx(content)
    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeError as exc:
            raise OfficePackageError("document_data_encoding_unsupported") from exc
        if suffix in {".txt", ".md"}:
            value = {"text": text}
        elif suffix == ".csv":
            rows, count = [], 0
            try:
                for row in csv.reader(io.StringIO(text), strict=True):
                    count += len(row)
                    if count > MAX_CELLS:
                        raise OfficePackageError("document_data_limit")
                    rows.append([Decimal(cell) if _NUMBER.fullmatch(cell) else cell for cell in row])
            except csv.Error as exc:
                raise OfficePackageError("document_data_invalid") from exc
            value = rows
        else:
            raise OfficePackageError("document_data_format_unsupported")
    encoded = _encode(value).encode()
    if len(encoded) > 128_000:
        raise OfficePackageError("document_data_limit")
    return JsonEvidenceSource(name, encoded)
