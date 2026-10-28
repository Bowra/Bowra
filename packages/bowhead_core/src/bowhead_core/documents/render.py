"""Exact-source LibreOffice preview and bounded PDF checks inside the worker."""

from __future__ import annotations

import hashlib
import io
import json
import platform
import posixpath
import subprocess
import tempfile
from pathlib import Path

from bowhead_core.documents.package_io import OfficePackageError
from bowhead_core.documents.pptx_text import REL, A, P, R, _children, _parse, _parts, inventory_pptx_text

RENDER_PROFILE = "ordinary-pptx-preview-v1"
PDF_FILTER = 'pdf:impress_pdf_Export:{"ExportHiddenSlides":{"type":"boolean","value":"true"}}'
MAX_PDF_BYTES = 32 * 1024 * 1024


def _slide_numbers(parts: dict[str, bytes]) -> dict[str, int]:
    """Map each slide part to its 1-based position in the presentation's slide order (empty when unreadable)."""
    try:
        relationships = _parse(parts.get("ppt/_rels/presentation.xml.rels", b"<Relationships/>"))
        targets = {
            rel.getAttribute("Id"): rel.getAttribute("Target").lstrip("/")
            if rel.getAttribute("Target").startswith("/")
            else posixpath.normpath(posixpath.join("ppt", rel.getAttribute("Target")))
            for rel in relationships.getElementsByTagNameNS(REL, "Relationship")
            if rel.getAttribute("Type") == R + "/slide" and rel.getAttribute("TargetMode") != "External"
        }
        presentation = _parse(parts["ppt/presentation.xml"]).documentElement
        lists = presentation.getElementsByTagNameNS(P, "sldIdLst")
        ordered = _children(lists[0], P, "sldId") if lists else []
        return {
            targets[slide.getAttributeNS(R, "id")]: index for index, slide in enumerate(ordered, start=1) if slide.getAttributeNS(R, "id") in targets
        }
    except Exception:  # noqa: BLE001 - page hints are optional and must never block rendering
        return {}


def font_fallbacks(source: bytes) -> list[dict[str, object]]:
    """Report unresolved font declarations, without claiming they are all used.

    `slides` lists the slides whose own XML declares the face (layout/master declarations add no slide).
    """
    requested: dict[str, set[int]] = {}
    parts = _parts(source)
    slide_numbers = _slide_numbers(parts)
    for name, content in parts.items():
        if name.startswith("ppt/") and name.endswith(".xml"):
            document = _parse(content)
            try:
                for tag in ("latin", "ea", "cs"):
                    for node in document.getElementsByTagNameNS(A, tag):
                        face = node.getAttribute("typeface")
                        if face and not face.startswith("+"):
                            slides = requested.setdefault(face, set())
                            if name in slide_numbers:
                                slides.add(slide_numbers[name])
            finally:
                document.unlink()
    if len(requested) > 128 or any(len(face) > 256 or "\n" in face for face in requested):
        raise OfficePackageError("document_font_inventory_unavailable")
    available = {
        family.strip().casefold()
        for line in run_binary(["fc-list", "-f", "%{family}\n"], timeout=15, output_limit=1_048_576, merge_stderr=False).splitlines()
        for family in line.split(",")
    }
    result = []
    for face in sorted(requested):
        if face.casefold() not in available:
            escaped = "".join("\\" + char if char in "\\-:,=" else char for char in face)
            replacement = run_binary(["fc-match", "-f", "%{family}", "--", escaped], timeout=5, merge_stderr=False)
            result.append({"declared_font": face, "environment_fallback": replacement, "slides": sorted(requested[face])})
    return result


def run_binary(argv: list[str], *, timeout: int = 90, output_limit: int = 8192, merge_stderr: bool = True) -> str:
    """No shell; bounded diagnostics never become user-visible exception text."""
    with tempfile.TemporaryFile() as output:
        try:
            result = subprocess.run(argv, stdout=output, stderr=output if merge_stderr else subprocess.DEVNULL, timeout=timeout, check=False)  # noqa: S603
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise OfficePackageError("document_renderer_unavailable") from exc
        if result.returncode:
            raise OfficePackageError("document_render_failed")
        output.seek(0)
        content = output.read(output_limit + 1)
        if len(content) > output_limit:
            raise OfficePackageError("document_renderer_output_limit")
        return content.decode("utf-8", errors="replace").strip()


def runtime_identity(*, soffice: str = "soffice", pdftoppm: str = "pdftoppm") -> dict[str, object]:
    """Fingerprint actual renderer, adapter code and installed font bytes."""
    font_list = run_binary(["fc-list", "-f", "%{file}\n"], timeout=15, output_limit=256_000, merge_stderr=False)
    fonts = sorted(set(font_list.splitlines()))
    if not fonts or len(fonts) > 4096:
        raise OfficePackageError("document_font_inventory_unavailable")
    font_hashes = []
    total_font_bytes = 0
    for name in fonts:
        path = Path(name)
        if not path.is_file() or path.stat().st_size > 256 * 1024 * 1024:
            raise OfficePackageError("document_font_inventory_unavailable")
        total_font_bytes += path.stat().st_size
        if total_font_bytes > 2 * 1024 * 1024 * 1024:
            raise OfficePackageError("document_font_inventory_unavailable")
        with path.open("rb") as stream:
            font_hashes.append(hashlib.file_digest(stream, "sha256").hexdigest())
    module_root = Path(__file__).parent
    code = {
        name: hashlib.sha256((module_root / name).read_bytes()).hexdigest()
        for name in ("package_io.py", "pptx_text.py", "pptx_plan.py", "pptx_draft.py", "evidence.py", "render.py", "worker.py")
    }
    record = {
        "profile": RENDER_PROFILE,
        "engine": run_binary([soffice, "--version"], timeout=15),
        "rasterizer": run_binary([pdftoppm, "-v"], timeout=15),
        "python": platform.python_version(),
        "font_files_sha256": font_hashes,
        "adapter_files": code,
        "filter": PDF_FILTER,
        "locale": "C.UTF-8",
        "timezone": "UTC",
    }
    return {**record, "digest": hashlib.sha256(json.dumps(record, sort_keys=True, separators=(",", ":")).encode()).hexdigest()}


def _text_boxes(content: bytes) -> tuple[tuple[float, float], dict[tuple, tuple[float, float, float, float, bool, bool]]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(io.BytesIO(content))
    if not 72 <= deck.slide_width / 12700 <= 4000 or not 72 <= deck.slide_height / 12700 <= 4000:
        raise OfficePackageError("document_layout_unsupported")
    boxes = {}

    def visit(shapes, slide, transform=(1.0, 1.0, 0.0, 0.0)):
        sx, sy, tx, ty = transform
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                xfrm = shape._element.grpSpPr.xfrm
                if xfrm is None or not xfrm.chExt.cx or not xfrm.chExt.cy or shape.rotation or xfrm.flipH or xfrm.flipV:
                    raise OfficePackageError("document_layout_unsupported")
                gx, gy = shape.width / xfrm.chExt.cx, shape.height / xfrm.chExt.cy
                visit(shape.shapes, slide, (sx * gx, sy * gy, tx + sx * (shape.left - gx * xfrm.chOff.x), ty + sy * (shape.top - gy * xfrm.chOff.y)))
                continue
            if not shape.has_text_frame and not shape.has_table:
                continue
            if shape.rotation:
                raise OfficePackageError("document_layout_unsupported")
            # The first release verifies horizontal, unflipped text geometry.
            # Reject other transforms instead of claiming a box-only check.
            if any(node.get("flipH") in {"1", "true"} or node.get("flipV") in {"1", "true"} for node in shape._element.iter()):
                raise OfficePackageError("document_layout_unsupported")
            if any(node.get("vert") not in {None, "horz"} for node in shape._element.iter()):
                raise OfficePackageError("document_layout_unsupported")
            left, top = tx + sx * shape.left, ty + sy * shape.top
            if shape.has_table:
                y = top
                for r, row in enumerate(shape.table.rows):
                    x = left
                    for c, column in enumerate(shape.table.columns):
                        boxes[(slide, str(shape.shape_id), (r, c))] = (
                            x / 12700,
                            y / 12700,
                            (x + sx * column.width) / 12700,
                            (y + sy * row.height) / 12700,
                            False,
                            False,
                        )
                        x += sx * column.width
                    y += sy * row.height
            else:
                body = shape.text_frame._txBody.bodyPr
                boxes[(slide, str(shape.shape_id), None)] = (
                    left / 12700,
                    top / 12700,
                    (left + sx * shape.width) / 12700,
                    (top + sy * shape.height) / 12700,
                    body.get("horzOverflow", "overflow") == "overflow",
                    body.get("vertOverflow", "overflow") == "overflow",
                )

    for number, slide in enumerate(deck.slides, 1):
        visit(slide.shapes, number)
    return (deck.slide_width / 12700, deck.slide_height / 12700), boxes


def validate_pdf(source: bytes, pdf: bytes) -> dict[str, object]:
    """Check every visible source text at its original shape/table position.

    Respect DrawingML overflow: a text frame is not necessarily a clip boundary.
    Keep page bounds, source-position association and explicit clipping checks.
    This is not native PowerPoint certification or a semantic
    review, and it does not replace the user's actual preview confirmation.
    """
    import pdfplumber

    if len(pdf) > MAX_PDF_BYTES or not pdf.startswith(b"%PDF-"):
        raise OfficePackageError("document_pdf_invalid")
    inventory = inventory_pptx_text(source)
    dimensions, boxes = _text_boxes(source)
    fonts = set()
    with pdfplumber.open(io.BytesIO(pdf)) as document:
        if len(document.pages) != inventory.slide_count:
            raise OfficePackageError("document_pdf_page_mismatch")
        for number, page in enumerate(document.pages, 1):
            if any(abs(actual - expected) > 0.1 for actual, expected in zip((page.width, page.height), dimensions, strict=True)):
                raise OfficePackageError("document_pdf_page_mismatch")
            chars = [char for char in page.chars if char["text"].strip()]
            if len(chars) > 100_000 or any("\ufffd" in char["text"] or "\x00" in char["text"] or "(cid:" in char["text"] for char in chars):
                raise OfficePackageError("document_pdf_glyph_unverifiable")
            fonts.update(char["fontname"] for char in chars)
            values, positions = [], []
            for char in chars:
                value = "".join(char["text"].split())
                values.append(value)
                positions.extend([char] * len(value))
            text = "".join(values)
            used = set()
            for slot in inventory.slots:
                if slot.slide != number or not slot.text.strip():
                    continue
                expected = "".join(slot.text.split())
                left, top, right, bottom, overflow_x, overflow_y = boxes[(number, slot.shape_id, slot.cell)]
                start, found = 0, False
                while (index := text.find(expected, start)) >= 0:
                    start = index + 1
                    selected = positions[index : index + len(expected)]
                    if (
                        index not in used
                        and selected[0]["x0"] <= right + 2
                        and selected[0]["top"] <= bottom + 2
                        and all(
                            char["x0"] >= max(0, left) - 2
                            and char["x1"] <= (page.width if overflow_x else min(page.width, right)) + 2
                            and char["top"] >= max(0, top) - 2
                            and char["bottom"] <= (page.height if overflow_y else min(page.height, bottom)) + 2
                            for char in selected
                        )
                    ):
                        used.add(index)
                        found = True
                        break
                if not found:
                    raise OfficePackageError("document_pdf_text_missing_or_clipped")
    return {
        "profile": RENDER_PROFILE,
        "source_sha256": hashlib.sha256(source).hexdigest(),
        "pdf_sha256": hashlib.sha256(pdf).hexdigest(),
        "page_count": inventory.slide_count,
        "fonts_used": sorted(fonts),
        "text_and_geometry": "passed",
        "native_powerpoint_certified": False,
        "visual_review": "user_confirmation_required",
    }


def render_candidate(candidate: Path, output: Path, *, soffice: str = "soffice", pdftoppm: str = "pdftoppm") -> dict[str, object]:
    source = candidate.read_bytes()
    identity = runtime_identity(soffice=soffice, pdftoppm=pdftoppm)
    fallbacks = font_fallbacks(source)
    with tempfile.TemporaryDirectory(prefix="document-lo-") as temporary:
        run_binary(
            [
                soffice,
                "-env:UserInstallation=" + Path(temporary).as_uri(),
                "--headless",
                "--convert-to",
                PDF_FILTER,
                "--outdir",
                str(output),
                str(candidate),
            ]
        )
    if candidate.read_bytes() != source:
        raise OfficePackageError("source_hash_mismatch")
    pdf = output / (candidate.stem + ".pdf")
    if not pdf.is_file() or pdf.stat().st_size > MAX_PDF_BYTES:
        raise OfficePackageError("document_pdf_invalid")
    receipt = validate_pdf(source, pdf.read_bytes())
    run_binary([pdftoppm, "-png", "-r", "96", str(pdf), str(output / "page")])
    pages = sorted(output.glob("page-*.png"), key=lambda page: int(page.stem.removeprefix("page-")))
    if candidate.read_bytes() != source:
        raise OfficePackageError("source_hash_mismatch")
    if len(pages) != receipt["page_count"] or sum(page.stat().st_size for page in pages) > 128 * 1024 * 1024:
        raise OfficePackageError("document_preview_limit")
    return {
        **receipt,
        "runtime": identity,
        "declared_font_fallbacks": fallbacks,
        "pages": [{"name": page.name, "sha256": hashlib.sha256(page.read_bytes()).hexdigest()} for page in pages],
    }
