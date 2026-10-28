"""Clean-shell PPTX renderer and independent closed-file evidence extractor.

The renderer consumes frozen ExportIR plus materialized PNGs. It creates a
presentation from an empty shell, never copies sample slides, and never emits
native charts or embedded workbooks. The extractor reopens the completed
package, reconstructs evidence from visible shapes and image pixels, then
renders every slide through LibreOffice before returning validation evidence.
"""

from __future__ import annotations

import hashlib
import io
import json
import math
import os
import re
import subprocess
import sys
import unicodedata
import zipfile
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any, Literal, cast
from xml.etree import ElementTree as ET

from PIL import Image, ImageChops, ImageStat, UnidentifiedImageError
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.presentation import Presentation as PresentationType
from pptx.shapes.base import BaseShape
from pptx.slide import Slide
from pptx.util import Inches, Pt
from pypdf import PdfReader

from bowhead_core.error_catalog import CodedError
from bowhead_core.export_capabilities import CapabilityAction, ExportTarget
from bowhead_core.export_ir import (
    CapabilityPlanEntry,
    ExportIR,
    ExportIRBlock,
    build_capability_plan,
    iter_export_blocks,
)
from bowhead_core.export_manifest import (
    ExtractedOutputEvidence,
    ExtractedSectionEvidence,
    ManifestDisposition,
    ManifestReasonCode,
    OutputTargetKind,
    OutputTargetRef,
)
from bowhead_core.export_materializer import MaterializedVisual
from bowhead_core.export_native_palette import NativeExportPalette, native_export_palette
from bowhead_core.print_theme import PrintThemeError

PPTX_RENDERER_NAME = "canvas-clean-pptx-v2"
# "3": colors come from the job's print-theme snapshot (product accent, cool compiled
# neutrals) instead of the fixed forest-green/cream/terracotta palette of "2".
PPTX_RENDERER_VERSION = "3"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MAX_PPTX_BYTES = 100 * 1024 * 1024
MAX_PPTX_SLIDES = 500
MAX_ZIP_MEMBERS = 5_000
MAX_ZIP_MEMBER_BYTES = 100 * 1024 * 1024
MAX_ZIP_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
MAX_ZIP_COMPRESSION_RATIO = 100
MAX_XML_MEMBER_BYTES = 10 * 1024 * 1024
MAX_TABLE_ROWS = 100_000
MAX_GROUPED_TEXT_CHARS = 180
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Text-box geometry shared by the renderer and the closed-file capacity check.
# `_set_shape_text` applies these fixed inner margins; `_shape_text_within_capacity`
# and the pagination chunker derive their line budgets from the same numbers, so a
# packed chunk provably satisfies the overflow gate by construction rather than by
# a looser char-count heuristic.
_TEXT_MARGIN_H_IN = 0.12
_TEXT_MARGIN_V_IN = 0.08
_LINE_HEIGHT_FACTOR = 1.35
# Business narrative boxes: (left, top, width, height) in inches.
_BODY_BOX_IN = (0.85, 1.3, 11.55, 5.25)
_FALLBACK_CARD_BOX_IN = (0.82, 1.42, 11.7, 4.65)
# Cover masthead: business blocks stack from top down to a bounded bottom so a
# report with many h1/meta sections cannot push shapes off the slide.
_COVER_KINDS = frozenset({"eyebrow", "h1", "deck", "meta"})
_COVER_META_KINDS = frozenset({"eyebrow", "meta"})
_COVER_TEXT_WIDTH_IN = 10.9
_COVER_TOP_IN = 3.05
_COVER_BOTTOM_IN = 7.0
_COVER_GAP_IN = 0.18

_BUSINESS_RE = re.compile(r"^BH1B_([0-9a-f]{24})_([0-9a-f]{64})_([0-9]{4})$")
_VISUAL_RE = re.compile(
    r"^BH1V_([0-9a-f]{24})_([0-9a-f]{64})_"
    r"(sha256:[0-9a-f]{64})_([0-9]{4})$"
)
_EXCLUSION_RE = re.compile(
    r"^BH1X_([0-9a-f]{24})_([0-9a-f]{64})_"
    r"(summary_detail_omitted|summary_visual_limit)$"
)
_DECOR_RE = re.compile(r"^BH1D_[0-9]{4}_[0-9]{4}$")
_DEFAULT_DEMO_FINGERPRINTS = (
    "CMP-2604-001",
    "Microsoft_Excel_Worksheet2.xlsx",
    "整体合规率",
)
_FORBIDDEN_MEMBER_PARTS = (
    "/embeddings/",
    "/activex/",
    "/oleobjects/",
)
_FORBIDDEN_MEMBER_SUFFIXES = (
    "vbaproject.bin",
    ".xls",
    ".xlsx",
    ".xlsm",
    ".exe",
    ".dll",
)
_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"

_FONT = "Arial Unicode MS" if sys.platform == "darwin" else "Noto Sans CJK SC"
_MONO_FONT = _FONT

PptxProfile = Literal["full_report", "executive_summary"]
_TaggedShapeKind = Literal["business", "visual", "exclusion"]


class PptxExportError(CodedError):
    """The PPTX cannot be rendered or validated without losing parity."""


@dataclass(frozen=True, slots=True)
class PptxRenderResult:
    relative_path: str
    content_sha256: str
    size_bytes: int
    slide_count: int
    core_slide_count: int
    excluded_section_count: int
    renderer_name: str = PPTX_RENDERER_NAME
    renderer_version: str = PPTX_RENDERER_VERSION


@dataclass(frozen=True, slots=True)
class _SummarySelection:
    rendered: frozenset[str]
    exclusions: dict[str, ManifestReasonCode]


@dataclass(slots=True)
class _TaggedShape:
    marker: str
    source_digest: str
    part: int
    slide_number: int
    kind: _TaggedShapeKind
    text: str = ""
    pixel_sha256: str | None = None
    pixel_variance: float | None = None


def render_export_pptx(
    *,
    export_ir: ExportIR,
    materialized_visuals: tuple[MaterializedVisual, ...],
    attempt_dir: Path,
    profile: PptxProfile,
    include_appendix: bool,
    print_theme: dict[str, Any],
    output_name: str = "output.pptx",
) -> PptxRenderResult:
    """Render a frozen ExportIR into a new clean-shell presentation.

    ``print_theme`` is the job's frozen print-theme snapshot; every slide color is resolved
    from it (``export_native_palette``), so a different theme needs no renderer change.
    """
    try:
        palette = native_export_palette(print_theme)
    except PrintThemeError as exc:
        raise PptxExportError("pptx_print_theme_invalid") from exc
    target = _target_for_profile(profile)
    output_path = _prepare_output_path(
        attempt_dir=attempt_dir,
        output_name=output_name,
    )
    visuals = _index_visuals(
        export_ir=export_ir,
        materialized_visuals=materialized_visuals,
        attempt_dir=attempt_dir,
    )
    plan = build_capability_plan(
        export_ir,
        target=target,
        include_appendix=include_appendix,
    )
    plan_by_key = {entry.block_key: entry for entry in plan}
    if len(plan_by_key) != len(plan):
        raise PptxExportError("duplicate_ir_block_key")
    selection = (
        _select_summary(export_ir=export_ir, plan=plan)
        if profile == "executive_summary"
        else _SummarySelection(
            rendered=frozenset(plan_by_key),
            exclusions={},
        )
    )

    blocks = iter_export_blocks(export_ir.blocks)
    cover_blocks = _select_cover_blocks(
        blocks,
        rendered=frozenset(selection.rendered),
        plan_by_key=plan_by_key,
    )

    render_items: list[tuple[ExportIRBlock, CapabilityPlanEntry]] = []
    excluded_items: list[tuple[ExportIRBlock, ManifestReasonCode]] = []
    for block in blocks:
        entry = plan_by_key[block.block_key]
        if block in cover_blocks:
            continue
        reason = selection.exclusions.get(block.block_key)
        if reason is not None or entry.action == "exclude":
            excluded_items.append(
                (
                    block,
                    reason or "summary_detail_omitted",
                )
            )
            continue
        if block.block_key not in selection.rendered:
            raise PptxExportError("pptx_summary_selection_unaccounted")
        render_items.append((block, entry))
    if profile == "executive_summary":
        render_items.sort(
            key=lambda item: (
                _summary_render_rank(
                    block=item[0],
                    entry=item[1],
                ),
                item[0].source.section_path,
                item[0].block_key,
            )
        )

    slide_budget = _planned_slide_count(
        cover_blocks=cover_blocks,
        render_items=render_items,
    )
    if slide_budget > MAX_PPTX_SLIDES:
        raise PptxExportError("pptx_slide_count_invalid")

    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    presentation.core_properties.title = f"管理层摘要 — {export_ir.document.title}" if profile == "executive_summary" else export_ir.document.title
    presentation.core_properties.subject = "Canvas 管理层摘要" if profile == "executive_summary" else "Canvas 完整报告"
    # feedback-triage-2026-08-24 A-1 continued: the deck is the customer's deliverable; omit vendor attribution.
    # Same decision as canvas-pptx skill's `_scrub_document_properties`: it clears template-inherited
    # "Bowhead Analytics", while this clears our hardcoded "Bowhead".
    # Keep `comments`: renderer id + snapshot id support troubleshooting and are not branding.
    presentation.core_properties.author = ""
    presentation.core_properties.last_modified_by = ""
    presentation.core_properties.comments = f"{PPTX_RENDERER_NAME}/{PPTX_RENDERER_VERSION};snapshot={export_ir.snapshot_id}"

    core_slide_count = 0
    if cover_blocks:
        _render_cover(
            palette=palette,
            presentation=presentation,
            export_ir=export_ir,
            blocks=cover_blocks,
            profile=profile,
        )
        core_slide_count += 1

    index = 0
    while index < len(render_items):
        block, entry = render_items[index]
        appendix = _is_appendix(
            block=block,
            action=entry.action,
        )
        if _is_groupable_text(block=block, entry=entry):
            group: list[ExportIRBlock] = [block]
            index += 1
            while index < len(render_items) and len(group) < 3:
                candidate, candidate_entry = render_items[index]
                if (
                    not _is_groupable_text(
                        block=candidate,
                        entry=candidate_entry,
                    )
                    or _is_appendix(
                        block=candidate,
                        action=candidate_entry.action,
                    )
                    != appendix
                ):
                    break
                group.append(candidate)
                index += 1
            _render_text_group(
                palette=palette,
                presentation=presentation,
                blocks=tuple(group),
                appendix=appendix,
            )
            if not appendix:
                core_slide_count += 1
            continue
        if entry.action == "fallback" and _is_groupable_fallback(block):
            group = [block]
            index += 1
            while index < len(render_items) and len(group) < 3:
                candidate, candidate_entry = render_items[index]
                if candidate_entry.action != "fallback" or not _is_groupable_fallback(candidate):
                    break
                group.append(candidate)
                index += 1
            _render_fallback_group(
                palette=palette,
                presentation=presentation,
                blocks=tuple(group),
            )
            core_slide_count += 1
            continue
        before = len(presentation.slides)
        _render_block(
            palette=palette,
            presentation=presentation,
            block=block,
            visual=visuals.get(block.block_key),
            attempt_dir=attempt_dir,
            appendix=appendix,
            fallback=entry.action == "fallback",
        )
        index += 1
        added = len(presentation.slides) - before
        if added <= 0:
            raise PptxExportError("pptx_block_not_rendered")
        if not appendix:
            core_slide_count += added

    if not presentation.slides:
        slide = _add_slide(
            presentation,
            palette=palette,
            title=export_ir.document.title,
            appendix=False,
        )
        _add_decor_text(
            slide,
            text="该报告当前没有可导出的内容",
            left=0.9,
            top=2.7,
            width=11.5,
            height=0.8,
            size=24,
            color=_rgb(palette.ink),
        )
        core_slide_count = 1
    for block, reason in excluded_items:
        _add_exclusion_marker(
            palette=palette,
            presentation=presentation,
            block=block,
            reason=reason,
        )

    if len(presentation.slides) != slide_budget or len(presentation.slides) > MAX_PPTX_SLIDES:
        raise PptxExportError("pptx_slide_count_invalid")
    _add_slide_numbers(
        presentation,
        palette=palette,
        cover_slide_count=1 if cover_blocks else 0,
    )
    part_path = output_path.with_suffix(".pptx.part")
    if part_path.exists() or part_path.is_symlink():
        raise PptxExportError("pptx_staging_conflict")
    try:
        presentation.save(str(part_path))
    except Exception as exc:
        part_path.unlink(missing_ok=True)
        raise PptxExportError("pptx_render_failed") from exc
    size_bytes = part_path.stat().st_size
    if size_bytes <= 0 or size_bytes > MAX_PPTX_BYTES:
        part_path.unlink(missing_ok=True)
        raise PptxExportError("pptx_size_invalid")
    try:
        closed = Presentation(str(part_path))
        slide_count = len(closed.slides)
    except Exception as exc:
        part_path.unlink(missing_ok=True)
        raise PptxExportError("pptx_parser_failed") from exc
    if slide_count != len(presentation.slides) or slide_count <= 0:
        part_path.unlink(missing_ok=True)
        raise PptxExportError("pptx_slide_count_invalid")
    part_path.replace(output_path)
    content = output_path.read_bytes()
    return PptxRenderResult(
        relative_path=output_path.relative_to(attempt_dir).as_posix(),
        content_sha256=_sha256_bytes(content),
        size_bytes=len(content),
        slide_count=slide_count,
        core_slide_count=core_slide_count,
        excluded_section_count=len(selection.exclusions),
    )


def extract_pptx_evidence(
    *,
    pptx_path: Path,
    export_ir: ExportIR,
    materialized_visuals: tuple[MaterializedVisual, ...],
    profile: PptxProfile,
    include_appendix: bool,
    validation_dir: Path,
    soffice_path: Path,
    pdftoppm_path: Path,
    demo_fingerprints: tuple[str, ...] = _DEFAULT_DEMO_FINGERPRINTS,
) -> ExtractedOutputEvidence:
    """Reopen a closed PPTX and prove visible content plus slide rendering."""
    size_bytes = pptx_path.stat().st_size if pptx_path.is_file() else 0
    mime_magic_valid = size_bytes > 0 and pptx_path.read_bytes()[:4] == b"PK\x03\x04"
    if not mime_magic_valid:
        return ExtractedOutputEvidence(
            mime_magic_valid=False,
            parser_opened=False,
            size_bytes=size_bytes,
            sections=(),
            hard_check_failures=("pptx_mime_magic_invalid",),
        )
    package_failures = _inspect_package(pptx_path)
    try:
        presentation = Presentation(str(pptx_path))
    except Exception:
        return ExtractedOutputEvidence(
            mime_magic_valid=True,
            parser_opened=False,
            size_bytes=size_bytes,
            sections=(),
            hard_check_failures=tuple(sorted(package_failures | {"pptx_parser_failed"})),
        )

    hard_failures = set(package_failures)
    if presentation.slide_width != SLIDE_WIDTH or presentation.slide_height != SLIDE_HEIGHT:
        hard_failures.add("pptx_slide_size_invalid")
    if not 0 < len(presentation.slides) <= MAX_PPTX_SLIDES:
        hard_failures.add("pptx_slide_count_invalid")

    blocks = iter_export_blocks(export_ir.blocks)
    block_by_marker = {_marker_id(block.block_key): block for block in blocks}
    if len(block_by_marker) != len(blocks):
        hard_failures.add("pptx_marker_collision")
    business, visuals, exclusions, unexpected, full_text = _read_tagged_shapes(
        presentation=presentation,
        known_markers=frozenset(block_by_marker),
        hard_failures=hard_failures,
    )
    expected_visuals = {visual.block_key: visual for visual in materialized_visuals}
    if len(expected_visuals) != len(materialized_visuals):
        hard_failures.add("duplicate_materialized_visual")
    expected_visual_keys = {block.block_key for block in blocks if block.visual_input_sha256 is not None}
    if set(expected_visuals) != expected_visual_keys:
        hard_failures.add("pptx_visual_materialization_mismatch")

    target = _target_for_profile(profile)
    plan = build_capability_plan(
        export_ir,
        target=target,
        include_appendix=include_appendix,
    )
    plan_by_key = {entry.block_key: entry for entry in plan}
    selection = (
        _select_summary(export_ir=export_ir, plan=plan)
        if profile == "executive_summary"
        else _SummarySelection(
            rendered=frozenset(plan_by_key),
            exclusions={},
        )
    )
    cover_keys = {
        block.block_key
        for block in blocks
        if block.source.section_kind in {"eyebrow", "h1", "deck", "meta"}
        and block.block_key in selection.rendered
        and plan_by_key[block.block_key].action != "exclude"
    }
    sections: list[ExtractedSectionEvidence] = []
    for marker, block in block_by_marker.items():
        evidence, failures = _extract_block_evidence(
            block=block,
            entry=plan_by_key[block.block_key],
            cover=block.block_key in cover_keys,
            business=business.get(marker, ()),
            visual_shapes=visuals.get(marker, ()),
            exclusion=exclusions.get(marker),
            expected_visual=expected_visuals.get(block.block_key),
        )
        hard_failures.update(failures)
        if evidence is not None:
            sections.append(evidence)
    if len(sections) != len(blocks):
        hard_failures.add("pptx_section_accounting_mismatch")

    render_failures = _render_all_slides(
        pptx_path=pptx_path,
        validation_dir=validation_dir,
        soffice_path=soffice_path,
        pdftoppm_path=pdftoppm_path,
        expected_slide_count=len(presentation.slides),
        expected_cjk_text=full_text,
    )
    hard_failures.update(render_failures)
    expected_text = _expected_visible_text(export_ir)
    demo_residue_count = _demo_residue_count(
        full_text=full_text,
        expected_text=expected_text,
        fingerprints=demo_fingerprints,
    )
    warnings = tuple(sorted({warning for visual in materialized_visuals for warning in visual.warnings}))
    return ExtractedOutputEvidence(
        mime_magic_valid=True,
        parser_opened=True,
        size_bytes=size_bytes,
        sections=tuple(sections),
        unexpected_business_section_count=unexpected,
        demo_residue_count=demo_residue_count,
        hard_check_failures=tuple(sorted(hard_failures)),
        warnings=warnings,
    )


def _render_cover(
    *,
    palette: NativeExportPalette,
    presentation: PresentationType,
    export_ir: ExportIR,
    blocks: tuple[ExportIRBlock, ...],
    profile: PptxProfile,
) -> None:
    slide = _add_slide(
        presentation,
        palette=palette,
        title=("管理层摘要" if profile == "executive_summary" else "完整报告"),
        appendix=False,
        dark=True,
    )
    _add_decor_text(
        slide,
        text=("MANAGEMENT BRIEF" if profile == "executive_summary" else "FULL CANVAS REPORT"),
        left=0.9,
        top=0.8,
        width=5.2,
        height=0.35,
        size=12,
        color=_rgb(palette.cover_eyebrow),
        bold=True,
    )
    _add_decor_text(
        slide,
        text=export_ir.document.title,
        left=0.9,
        top=1.5,
        width=11.2,
        height=1.25,
        size=34,
        color=_rgb(palette.cover_title),
        bold=True,
    )
    top = _COVER_TOP_IN
    for block in blocks:
        text, size, height = _cover_block_layout(block)
        _add_business_text(
            slide,
            block=block,
            part=0,
            text=text,
            left=0.95,
            top=top,
            width=_COVER_TEXT_WIDTH_IN,
            height=height,
            size=size,
            color=_rgb(palette.cover_body),
            bold=block.source.section_kind == "h1",
        )
        top += height + _COVER_GAP_IN


def _render_block(
    *,
    palette: NativeExportPalette,
    presentation: PresentationType,
    block: ExportIRBlock,
    visual: MaterializedVisual | None,
    attempt_dir: Path,
    appendix: bool,
    fallback: bool,
) -> None:
    if block.block_kind == "visual":
        if visual is None:
            raise PptxExportError("visual_materialization_missing")
        _render_visual_block(
            palette=palette,
            presentation=presentation,
            block=block,
            visual=visual,
            attempt_dir=attempt_dir,
            appendix=appendix,
        )
        return
    if block.block_kind == "table":
        _render_table_block(
            palette=palette,
            presentation=presentation,
            block=block,
            appendix=appendix,
        )
        return
    if block.block_kind == "kpi_group":
        _render_kpi_block(
            palette=palette,
            presentation=presentation,
            block=block,
            appendix=appendix,
        )
        return
    _render_text_block(
        palette=palette,
        presentation=presentation,
        block=block,
        appendix=appendix,
        fallback=fallback,
    )


def _render_text_group(
    *,
    palette: NativeExportPalette,
    presentation: PresentationType,
    blocks: tuple[ExportIRBlock, ...],
    appendix: bool,
) -> None:
    slide = _add_slide(
        presentation,
        palette=palette,
        title=("技术附录" if appendix else "报告要点"),
        appendix=appendix,
    )
    count = len(blocks)
    gap = 0.18
    available_height = 5.15
    card_height = (available_height - gap * (count - 1)) / count
    top = 1.3
    for index, block in enumerate(blocks):
        tokens = _block_display_tokens(block)
        text = "\n".join(tokens)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.88),
            Inches(top),
            Inches(11.5),
            Inches(card_height),
        )
        shape.name = _business_shape_name(
            block=block,
            part=0,
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(palette.card if index % 2 == 0 else palette.card_alt)
        shape.line.color.rgb = _rgb(palette.rule)
        _set_shape_text(
            shape,
            text=text,
            size=_fit_font_pt(
                text,
                width_in=11.5,
                height_in=card_height,
                max_pt=(12 if block.block_kind == "code" else 16 if count > 1 else 20),
                min_pt=11,
            ),
            color=_rgb(palette.ink),
            bold=block.block_kind in {"heading", "callout"},
            vertical=MSO_ANCHOR.MIDDLE,
            font=(_MONO_FONT if block.block_kind == "code" else _FONT),
        )
        top += card_height + gap


def _render_fallback_group(
    *,
    palette: NativeExportPalette,
    presentation: PresentationType,
    blocks: tuple[ExportIRBlock, ...],
) -> None:
    slide = _add_slide(
        presentation,
        palette=palette,
        title="在线交互内容",
        appendix=False,
    )
    count = len(blocks)
    gap = 0.18
    available_height = 5.15
    card_height = (available_height - gap * (count - 1)) / count
    top = 1.3
    for block in blocks:
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.88),
            Inches(top),
            Inches(11.5),
            Inches(card_height),
        )
        card.name = _business_shape_name(
            block=block,
            part=0,
        )
        card.fill.solid()
        card.fill.fore_color.rgb = _rgb(palette.callout_fill)
        card.line.color.rgb = _rgb(palette.callout_border)
        text = "\n".join(_block_display_tokens(block))
        _set_shape_text(
            card,
            text=text,
            size=_fit_font_pt(
                text,
                width_in=11.5,
                height_in=card_height,
                max_pt=15 if count > 1 else 19,
                min_pt=11,
            ),
            color=_rgb(palette.callout_ink),
            vertical=MSO_ANCHOR.MIDDLE,
        )
        top += card_height + gap


def _render_visual_block(
    *,
    palette: NativeExportPalette,
    presentation: PresentationType,
    block: ExportIRBlock,
    visual: MaterializedVisual,
    attempt_dir: Path,
    appendix: bool,
) -> None:
    slide = _add_slide(
        presentation,
        palette=palette,
        title=_slide_title(block),
        appendix=appendix,
    )
    part = 0
    if block.semantic_tokens:
        title_text = "\n".join(block.semantic_tokens)
        _add_business_text(
            slide,
            block=block,
            part=part,
            text=title_text,
            left=0.85,
            top=0.82,
            width=11.65,
            height=0.8,
            size=_fit_font_pt(title_text, width_in=11.65, height_in=0.8, max_pt=20, min_pt=11),
            color=_rgb(palette.ink),
            bold=True,
        )
        part += 1
    visual_path = attempt_dir / visual.relative_path
    left, top, width, height = _aspect_fit(
        source_width=visual.width,
        source_height=visual.height,
        box=(0.85, 1.68, 11.65, 4.95),
    )
    picture = slide.shapes.add_picture(
        str(visual_path),
        Inches(left),
        Inches(top),
        width=Inches(width),
        height=Inches(height),
    )
    picture.name = _visual_shape_name(
        block=block,
        pixel_sha256=visual.pixel_sha256,
        part=part,
    )


def _render_table_block(
    *,
    palette: NativeExportPalette,
    presentation: PresentationType,
    block: ExportIRBlock,
    appendix: bool,
) -> None:
    columns = block.content.get("columns")
    rows = block.content.get("rows")
    total_row = block.content.get("total_row")
    if not isinstance(columns, list) or not isinstance(rows, list):
        raise PptxExportError("pptx_table_invalid")
    if len(rows) > MAX_TABLE_ROWS:
        raise PptxExportError("pptx_table_row_limit")
    normalized_columns = [column for column in columns if isinstance(column, dict)]
    if not normalized_columns:
        _render_text_block(
            palette=palette,
            presentation=presentation,
            block=block,
            appendix=appendix,
            fallback=False,
        )
        return
    keys = [str(column.get("key", f"column_{index + 1}")) for index, column in enumerate(normalized_columns)]
    labels = [str(column.get("label", keys[index])) for index, column in enumerate(normalized_columns)]
    normalized_rows = [row for row in rows if isinstance(row, dict)]
    if isinstance(total_row, dict):
        normalized_rows.append(total_row)
    page_size = max(5, min(12, 48 // max(1, len(normalized_columns))))
    pages = [normalized_rows[index : index + page_size] for index in range(0, len(normalized_rows), page_size)] or [[]]
    part = 0
    for page_rows in pages:
        slide = _add_slide(
            presentation,
            palette=palette,
            title=_slide_title(block),
            appendix=appendix,
        )
        title_tokens = block.semantic_tokens[: max(0, len(block.semantic_tokens) - len(labels))]
        if title_tokens:
            title_text = "\n".join(title_tokens)
            _add_business_text(
                slide,
                block=block,
                part=part,
                text=title_text,
                left=0.8,
                top=0.78,
                width=11.75,
                height=0.72,
                size=_fit_font_pt(title_text, width_in=11.75, height_in=0.72, max_pt=18, min_pt=11),
                color=_rgb(palette.ink),
                bold=True,
            )
            part += 1
        table_height = min(
            4.75,
            max(1.1, 0.55 * (len(page_rows) + 1)),
        )
        table_top = 1.58 + (4.95 - table_height) / 2
        table = slide.shapes.add_table(
            max(2, len(page_rows) + 1),
            len(labels),
            Inches(0.8),
            Inches(table_top),
            Inches(11.75),
            Inches(table_height),
        )
        table.name = _business_shape_name(block=block, part=part)
        part += 1
        for column_index, label in enumerate(labels):
            cell = table.table.cell(0, column_index)
            cell.text = _safe_text(label)
            cell.fill.solid()  # type: ignore[no-untyped-call]
            cell.fill.fore_color.rgb = _rgb(palette.table_header_fill)
            _style_cell(
                cell,
                size=14 if len(page_rows) <= 5 else 11,
                color=_rgb(palette.table_header_ink),
                bold=True,
            )
        values = page_rows or [{}]
        for row_index, row in enumerate(values, start=1):
            for column_index, key in enumerate(keys):
                cell = table.table.cell(row_index, column_index)
                cell.text = _scalar_text(row.get(key)) if page_rows else "—"
                cell.fill.solid()  # type: ignore[no-untyped-call]
                cell.fill.fore_color.rgb = _rgb(palette.card if row_index % 2 else palette.card_alt)
                _style_cell(
                    cell,
                    size=13 if len(page_rows) <= 5 else 10,
                    color=_rgb(palette.ink),
                )


def _render_kpi_block(
    *,
    palette: NativeExportPalette,
    presentation: PresentationType,
    block: ExportIRBlock,
    appendix: bool,
) -> None:
    cells = block.content.get("cells")
    values = cells if isinstance(cells, list) and cells else [{}]
    for page_start in range(0, len(values), 6):
        page_values = values[page_start : page_start + 6]
        slide = _add_slide(
            presentation,
            palette=palette,
            title=_slide_title(block),
            appendix=appendix,
        )
        count = len(page_values)
        columns = 3 if count > 2 else count
        rows = (count + columns - 1) // columns
        card_width = 11.5 / columns
        card_height = 3.4 if count == 1 else min(2.15, 4.8 / rows)
        for page_index, cell in enumerate(page_values):
            part = page_start + page_index
            col = page_index % columns
            row = page_index // columns
            text = "\n".join(_kpi_visible_tokens(cell)) or "—"
            card_w = 9.0 if count == 1 else card_width - 0.25
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(2.15 if count == 1 else 0.9 + col * card_width),
                Inches(2.05 if count == 1 else 1.6 + row * (card_height + 0.25)),
                Inches(card_w),
                Inches(card_height),
            )
            shape.name = _business_shape_name(block=block, part=part)
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(palette.card)
            shape.line.color.rgb = _rgb(palette.rule)
            _set_shape_text(
                shape,
                text=text,
                size=_fit_font_pt(text, width_in=card_w, height_in=card_height, max_pt=30 if count == 1 else 18, min_pt=11),
                color=_rgb(palette.ink),
                bold=part < 3,
                vertical=MSO_ANCHOR.MIDDLE,
            )


def _render_text_block(
    *,
    palette: NativeExportPalette,
    presentation: PresentationType,
    block: ExportIRBlock,
    appendix: bool,
    fallback: bool,
) -> None:
    tokens = list(_block_display_tokens(block))
    units_per_line, max_lines = _text_block_geometry(block, fallback=fallback)
    chunks = _chunk_tokens(tokens, units_per_line=units_per_line, max_lines=max_lines)
    font_pt = _text_block_font_pt(block, fallback=fallback)
    for chunk_index, chunk in enumerate(chunks):
        slide = _add_slide(
            presentation,
            palette=palette,
            title=_slide_title(block),
            appendix=appendix,
        )
        if fallback:
            left, top, width, height = _FALLBACK_CARD_BOX_IN
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(left),
                Inches(top),
                Inches(width),
                Inches(height),
            )
            card.fill.solid()
            card.fill.fore_color.rgb = _rgb(palette.callout_fill)
            card.line.color.rgb = _rgb(palette.callout_border)
            card.name = _business_shape_name(
                block=block,
                part=chunk_index,
            )
            _set_shape_text(
                card,
                text=chunk,
                size=font_pt,
                color=_rgb(palette.callout_ink),
                vertical=MSO_ANCHOR.TOP,
            )
        else:
            left, top, width, height = _BODY_BOX_IN
            _add_business_text(
                slide,
                block=block,
                part=chunk_index,
                text=chunk,
                left=left,
                top=top,
                width=width,
                height=height,
                size=font_pt,
                color=_rgb(palette.ink),
                bold=block.block_kind in {"heading", "callout"},
                font=_MONO_FONT if block.block_kind == "code" else _FONT,
            )


def _add_slide(
    presentation: PresentationType,
    *,
    palette: NativeExportPalette,
    title: str,
    appendix: bool,
    dark: bool = False,
) -> Slide:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide_number = len(presentation.slides)
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0,
        0,
        presentation.slide_width,
        presentation.slide_height,
    )
    background.name = _decor_shape_name(slide_number, 0)
    background.fill.solid()
    background.fill.fore_color.rgb = _rgb(palette.cover_fill if dark else palette.slide)
    background.line.fill.background()
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.58),
        Inches(0.55),
        Inches(0.08),
        Inches(0.48),
    )
    accent.name = _decor_shape_name(slide_number, 1)
    accent.fill.solid()
    accent.fill.fore_color.rgb = _rgb(palette.cover_bar if dark else palette.accent_bar)
    accent.line.fill.background()
    # feedback-triage-2026-08-24 A-1: headers previously showed the vendor wordmark. In private deployments,
    # the deck belongs to the customer and must not carry our attribution (the same batch removed canvas-pptx
    # template footers and FIN-page marketing copy). Non-appendix pages keep only the left accent stripe, with no text.
    _add_decor_text(
        slide,
        text="附录" if appendix else "",
        left=0.78,
        top=0.56,
        width=4.8,
        height=0.32,
        size=10,
        color=_rgb(palette.cover_eyebrow if dark else palette.accent_text),
        bold=True,
    )
    if title:
        _add_decor_text(
            slide,
            text=title,
            left=6.2,
            top=0.56,
            width=6.1,
            height=0.32,
            size=9,
            color=_rgb(palette.cover_muted if dark else palette.muted),
            align=PP_ALIGN.RIGHT,
        )
    return cast(Slide, slide)


def _add_slide_numbers(
    presentation: PresentationType,
    *,
    palette: NativeExportPalette,
    cover_slide_count: int,
) -> None:
    total = len(presentation.slides)
    for slide_number, slide in enumerate(presentation.slides, start=1):
        on_cover = slide_number <= cover_slide_count
        _add_decor_text(
            slide,
            text=f"{slide_number:02d} / {total:02d}",
            left=11.45,
            top=7.05,
            width=0.95,
            height=0.22,
            size=8,
            color=_rgb(palette.cover_muted if on_cover else palette.muted),
            align=PP_ALIGN.RIGHT,
        )


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.removeprefix("#"))


def _add_business_text(
    slide: Slide,
    *,
    block: ExportIRBlock,
    part: int,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    size: int,
    color: RGBColor,
    bold: bool = False,
    font: str = _FONT,
) -> None:
    shape = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.name = _business_shape_name(block=block, part=part)
    _set_shape_text(
        shape,
        text=text,
        size=size,
        color=color,
        bold=bold,
        font=font,
    )


def _add_decor_text(
    slide: Slide,
    *,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    size: int,
    color: RGBColor,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    shape = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.name = _decor_shape_name(
        _slide_number(slide),
        len(slide.shapes),
    )
    _set_shape_text(
        shape,
        text=text,
        size=size,
        color=color,
        bold=bold,
        align=align,
    )


def _set_shape_text(
    shape: Any,
    *,
    text: str,
    size: int,
    color: RGBColor,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    vertical: MSO_ANCHOR = MSO_ANCHOR.TOP,
    font: str = _FONT,
) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(_TEXT_MARGIN_H_IN)
    frame.margin_right = Inches(_TEXT_MARGIN_H_IN)
    frame.margin_top = Inches(_TEXT_MARGIN_V_IN)
    frame.margin_bottom = Inches(_TEXT_MARGIN_V_IN)
    frame.vertical_anchor = vertical
    lines = _safe_text(text).splitlines() or ["—"]
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line or " "
        paragraph.alignment = align
        paragraph.space_after = Pt(5)
        paragraph.line_spacing = 1.08
        for run in paragraph.runs:
            _style_run(
                run,
                font=font,
                size=size,
                color=color,
                bold=bold,
            )


def _style_cell(
    cell: Any,
    *,
    size: int,
    color: RGBColor,
    bold: bool = False,
) -> None:
    cell.margin_left = Inches(0.06)
    cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in cell.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        for run in paragraph.runs:
            _style_run(
                run,
                font=_FONT,
                size=size,
                color=color,
                bold=bold,
            )


def _style_run(
    run: Any,
    *,
    font: str,
    size: int,
    color: RGBColor,
    bold: bool,
) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    properties = run._r.get_or_add_rPr()
    east_asian = properties.find(qn("a:ea"))
    if east_asian is None:
        east_asian = OxmlElement("a:ea")
        properties.append(east_asian)
    east_asian.set("typeface", font)
    complex_script = properties.find(qn("a:cs"))
    if complex_script is None:
        complex_script = OxmlElement("a:cs")
        properties.append(complex_script)
    complex_script.set("typeface", font)
    properties.set("lang", "zh-CN")


def _add_exclusion_marker(
    *,
    palette: NativeExportPalette,
    presentation: PresentationType,
    block: ExportIRBlock,
    reason: ManifestReasonCode,
) -> None:
    if not presentation.slides:
        _add_slide(
            presentation,
            palette=palette,
            title="管理层摘要",
            appendix=False,
        )
    slide = presentation.slides[0]
    shape = slide.shapes.add_textbox(
        -Inches(2),
        -Inches(2),
        Inches(0.1),
        Inches(0.1),
    )
    shape.name = f"BH1X_{_marker_id(block.block_key)}_{_source_digest(block.source_refs)}_{reason}"
    shape.text = ""


def _select_summary(
    *,
    export_ir: ExportIR,
    plan: tuple[CapabilityPlanEntry, ...],
) -> _SummarySelection:
    blocks = iter_export_blocks(export_ir.blocks)
    block_by_key = {block.block_key: block for block in blocks}
    rendered = {entry.block_key for entry in plan if entry.action in {"render", "fallback", "appendix"}}
    exclusions: dict[str, ManifestReasonCode] = {entry.block_key: "summary_detail_omitted" for entry in plan if entry.action == "exclude"}
    candidates = [block_by_key[entry.block_key] for entry in plan if entry.action == "summary_candidate"]
    text_priority = {
        "narrative": 0,
        "section": 1,
    }
    text_candidates = sorted(
        (block for block in candidates if block.visual_input_sha256 is None),
        key=lambda block: (
            text_priority.get(block.source.section_kind, 99),
            block.source.section_path,
            block.block_key,
        ),
    )
    if text_candidates:
        rendered.add(text_candidates[0].block_key)
        for block in text_candidates[1:]:
            exclusions[block.block_key] = "summary_detail_omitted"
    visual_priority = {
        "chart": 1,
        "spec_block": 3,
        "image": 4,
        "design": 5,
    }
    visual_candidates = sorted(
        (block for block in candidates if block.visual_input_sha256 is not None),
        key=lambda block: (
            visual_priority.get(block.source.section_kind, 99),
            block.source.section_path,
            block.block_key,
        ),
    )
    for block in visual_candidates[:2]:
        rendered.add(block.block_key)
    for block in visual_candidates[2:]:
        exclusions[block.block_key] = "summary_visual_limit"
    for block in candidates:
        if block.block_key not in rendered and block.block_key not in exclusions:
            exclusions[block.block_key] = "summary_detail_omitted"
    if rendered & exclusions.keys():
        raise PptxExportError("pptx_summary_selection_invalid")
    return _SummarySelection(
        rendered=frozenset(rendered),
        exclusions=exclusions,
    )


def _read_tagged_shapes(
    *,
    presentation: PresentationType,
    known_markers: frozenset[str],
    hard_failures: set[str],
) -> tuple[
    dict[str, tuple[_TaggedShape, ...]],
    dict[str, tuple[_TaggedShape, ...]],
    dict[str, ManifestReasonCode],
    int,
    str,
]:
    business: dict[str, list[_TaggedShape]] = defaultdict(list)
    visuals: dict[str, list[_TaggedShape]] = defaultdict(list)
    exclusions: dict[str, ManifestReasonCode] = {}
    unexpected = 0
    all_text: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            name = shape.name
            text = _shape_text(shape)
            if text:
                all_text.append(text)
            business_match = _BUSINESS_RE.fullmatch(name)
            visual_match = _VISUAL_RE.fullmatch(name)
            exclusion_match = _EXCLUSION_RE.fullmatch(name)
            if business_match:
                marker, source_digest, part_text = business_match.groups()
                if marker not in known_markers:
                    unexpected += 1
                    continue
                if not _shape_on_slide(
                    shape=shape,
                    presentation=presentation,
                ):
                    hard_failures.add("pptx_shape_overflow")
                if not _shape_text_within_capacity(shape):
                    hard_failures.add("pptx_text_overflow")
                if not text:
                    hard_failures.add("pptx_business_shape_empty")
                business[marker].append(
                    _TaggedShape(
                        marker=marker,
                        source_digest=source_digest,
                        part=int(part_text),
                        slide_number=slide_number,
                        kind="business",
                        text=text,
                    )
                )
                continue
            if visual_match:
                marker, source_digest, pixel_sha256, part_text = visual_match.groups()
                if marker not in known_markers:
                    unexpected += 1
                    continue
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    hard_failures.add("pptx_visual_shape_invalid")
                    continue
                if not _shape_on_slide(
                    shape=shape,
                    presentation=presentation,
                ):
                    hard_failures.add("pptx_shape_overflow")
                try:
                    observed_image = Image.open(io.BytesIO(shape.image.blob))
                    observed_pixel_sha256 = _visible_pixel_sha256(observed_image)
                    pixel_variance = max(ImageStat.Stat(observed_image.convert("RGB")).var)
                except (AttributeError, UnidentifiedImageError, OSError):
                    hard_failures.add("pptx_picture_decode_failed")
                    continue
                if observed_pixel_sha256 != pixel_sha256:
                    hard_failures.add("pptx_visual_pixel_marker_mismatch")
                visuals[marker].append(
                    _TaggedShape(
                        marker=marker,
                        source_digest=source_digest,
                        part=int(part_text),
                        slide_number=slide_number,
                        kind="visual",
                        pixel_sha256=observed_pixel_sha256,
                        pixel_variance=pixel_variance,
                    )
                )
                continue
            if exclusion_match:
                marker, source_digest, reason_text = exclusion_match.groups()
                if marker not in known_markers:
                    unexpected += 1
                    continue
                if _shape_on_slide(
                    shape=shape,
                    presentation=presentation,
                ):
                    hard_failures.add("pptx_exclusion_marker_visible")
                if marker in exclusions:
                    hard_failures.add("pptx_duplicate_exclusion_marker")
                else:
                    exclusions[marker] = _reason_code(reason_text)
                business[marker].append(
                    _TaggedShape(
                        marker=marker,
                        source_digest=source_digest,
                        part=0,
                        slide_number=slide_number,
                        kind="exclusion",
                    )
                )
                continue
            if _DECOR_RE.fullmatch(name):
                continue
            if _shape_has_business_content(shape):
                unexpected += 1

    return (
        {marker: tuple(sorted(items, key=lambda item: (item.part, item.slide_number))) for marker, items in business.items()},
        {marker: tuple(sorted(items, key=lambda item: (item.part, item.slide_number))) for marker, items in visuals.items()},
        exclusions,
        unexpected,
        "\n".join(all_text),
    )


def _extract_block_evidence(
    *,
    block: ExportIRBlock,
    entry: CapabilityPlanEntry,
    cover: bool,
    business: tuple[_TaggedShape, ...],
    visual_shapes: tuple[_TaggedShape, ...],
    exclusion: ManifestReasonCode | None,
    expected_visual: MaterializedVisual | None,
) -> tuple[ExtractedSectionEvidence | None, set[str]]:
    failures: set[str] = set()
    expected_source = _source_digest(block.source_refs)
    all_shapes = (*business, *visual_shapes)
    parts = [item.part for item in all_shapes]
    if len(parts) != len(set(parts)):
        failures.add("pptx_tagged_shape_duplicate")
    if sorted(set(parts)) != list(range(len(set(parts)))):
        failures.add("pptx_tagged_shape_sequence_invalid")
    observed_kinds = tuple(
        item.kind
        for item in sorted(
            all_shapes,
            key=lambda item: (item.part, item.slide_number, item.kind),
        )
    )
    expected_kinds = _expected_tagged_shape_kinds(
        block=block,
        entry=entry,
        cover=cover,
        excluded=exclusion is not None,
    )
    if observed_kinds != expected_kinds:
        failures.add("pptx_tagged_shape_type_invalid")
    if any(item.source_digest != expected_source for item in all_shapes):
        failures.add("pptx_source_receipt_mismatch")
    if exclusion is not None:
        if any(item.text for item in business) or visual_shapes:
            failures.add("pptx_excluded_section_visible")
        return (
            ExtractedSectionEvidence(
                block_key=block.block_key,
                disposition="excluded_by_policy",
                visible=False,
                targets=(),
                source_refs=(),
                reason_code=exclusion,
            ),
            failures,
        )
    if not all_shapes:
        return None, {"pptx_section_marker_missing"}

    visible_text = "".join(item.text for item in business)
    semantic_match = _contains_tokens_in_order(
        visible_text,
        block.semantic_tokens,
    )
    typed_tokens = _typed_visible_tokens(block)
    typed_match = _contains_tokens_in_order(
        visible_text,
        typed_tokens,
    )
    visual_match = False
    artifact_hashes: tuple[tuple[str, str], ...] = ()
    if block.visual_input_sha256 is not None:
        if expected_visual is None:
            failures.add("pptx_visual_missing")
        else:
            matching = [item for item in visual_shapes if item.pixel_sha256 == expected_visual.pixel_sha256]
            visual_match = len(matching) == 1
            if not visual_match:
                failures.add("pptx_visual_pixel_mismatch")
            else:
                artifact_hashes = expected_visual.source_artifact_hashes
                if (
                    block.source.section_kind
                    in {
                        "chart",
                        "spec_block",
                    }
                    and matching[0].pixel_variance is not None
                    and matching[0].pixel_variance < 10.0
                ):
                    failures.add("pptx_visual_blank")
    elif visual_shapes:
        failures.add("unexpected_pptx_visual")
    visible = bool(_normalize_visible(visible_text)) or visual_match
    if not visible:
        failures.add("pptx_section_not_visible")

    semantic_digest = block.semantic_content_sha256 if semantic_match and visible else None
    if block.typed_data_sha256 is None:
        typed_digest = None
    elif block.visual_input_sha256 is not None:
        typed_digest = block.typed_data_sha256 if visual_match else None
    else:
        typed_digest = block.typed_data_sha256 if typed_match and visible else None
    visual_digest = block.visual_input_sha256 if visual_match else None
    slide_numbers = sorted({item.slide_number for item in all_shapes if item.text or item.pixel_sha256 is not None})
    if entry.action == "fallback":
        disposition: ManifestDisposition = "fallback"
        reason: ManifestReasonCode | None = "interactive_not_executable"
        target_kind: OutputTargetKind = "interaction_fallback"
        region = "body"
    elif _is_appendix(block=block, action=entry.action):
        disposition = "rendered"
        reason = None
        target_kind = "appendix"
        region = "appendix"
    else:
        disposition = "rendered"
        reason = None
        target_kind = "slide"
        region = "body"
    targets = tuple(
        OutputTargetRef(
            kind=target_kind,
            id=f"slide-{slide_number}",
            region=region,
        )
        for slide_number in slide_numbers
    )
    return (
        ExtractedSectionEvidence(
            block_key=block.block_key,
            disposition=disposition,
            visible=visible,
            targets=targets,
            semantic_content_sha256=semantic_digest,
            typed_data_sha256=typed_digest,
            visual_input_sha256=visual_digest,
            visual_artifact_hashes=artifact_hashes,
            source_refs=block.source_refs if all(item.source_digest == expected_source for item in all_shapes) else (),
            reason_code=reason,
        ),
        failures,
    )


def _expected_tagged_shape_kinds(
    *,
    block: ExportIRBlock,
    entry: CapabilityPlanEntry,
    cover: bool,
    excluded: bool,
) -> tuple[_TaggedShapeKind, ...]:
    business: _TaggedShapeKind = "business"
    visual: _TaggedShapeKind = "visual"
    if excluded:
        return ("exclusion",)
    if cover or entry.action == "fallback" or _is_groupable_text(block=block, entry=entry):
        return (business,)
    if block.block_kind == "visual":
        return ((business,) if block.semantic_tokens else ()) + (visual,)
    if block.block_kind == "table":
        columns = block.content.get("columns")
        if not isinstance(columns, list):
            raise PptxExportError("pptx_table_invalid")
        normalized_columns = [column for column in columns if isinstance(column, dict)]
        if normalized_columns:
            title_tokens = block.semantic_tokens[: max(0, len(block.semantic_tokens) - len(normalized_columns))]
            shapes_per_slide = 2 if title_tokens else 1
            return (business,) * (_block_slide_count(block, fallback=False) * shapes_per_slide)
    if block.block_kind == "kpi_group":
        cells = block.content.get("cells")
        count = len(cells) if isinstance(cells, list) and cells else 1
        return (business,) * count
    return (business,) * _block_slide_count(block, fallback=False)


def _inspect_package(path: Path) -> set[str]:
    failures: set[str] = set()
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS:
                failures.add("pptx_zip_member_limit")
            total_uncompressed = 0
            seen_names: set[str] = set()
            for member in members:
                name = member.filename
                if name in seen_names:
                    failures.add("pptx_zip_duplicate_member")
                seen_names.add(name)
                pure = PurePosixPath(name)
                if pure.is_absolute() or ".." in pure.parts or "\\" in name:
                    failures.add("pptx_zip_path_invalid")
                total_uncompressed += member.file_size
                if member.file_size > MAX_ZIP_MEMBER_BYTES:
                    failures.add("pptx_zip_member_size_limit")
                if member.compress_size > 0 and member.file_size / member.compress_size > MAX_ZIP_COMPRESSION_RATIO:
                    failures.add("pptx_zip_ratio_limit")
                lowered = f"/{name.lower()}"
                if any(part in lowered for part in _FORBIDDEN_MEMBER_PARTS) or lowered.endswith(_FORBIDDEN_MEMBER_SUFFIXES) or "/charts/" in lowered:
                    failures.add("pptx_forbidden_package_member")
                if name.endswith((".xml", ".rels")):
                    if member.file_size > MAX_XML_MEMBER_BYTES:
                        failures.add("pptx_xml_size_limit")
                        continue
                    content = archive.read(member)
                    if b"<!DOCTYPE" in content or b"<!ENTITY" in content:
                        failures.add("pptx_xml_entity_forbidden")
                    if name.endswith(".rels"):
                        failures.update(_relationship_failures(content))
            if total_uncompressed > MAX_ZIP_UNCOMPRESSED_BYTES:
                failures.add("pptx_zip_uncompressed_limit")
            names = {member.filename for member in members}
            if "[Content_Types].xml" not in names or "ppt/presentation.xml" not in names:
                failures.add("pptx_required_member_missing")
    except (OSError, zipfile.BadZipFile, RuntimeError):
        failures.add("pptx_zip_invalid")
    return failures


def _relationship_failures(content: bytes) -> set[str]:
    failures: set[str] = set()
    try:
        # The caller rejects DTD/entity declarations and bounds XML to 10 MiB
        # before this parser is reached.
        root = ET.fromstring(content)  # noqa: S314
    except ET.ParseError:
        return {"pptx_relationship_xml_invalid"}
    for relationship in root.findall(f"{_REL_NS}Relationship"):
        if relationship.attrib.get("TargetMode") == "External":
            failures.add("pptx_external_relationship")
    return failures


def _render_all_slides(
    *,
    pptx_path: Path,
    validation_dir: Path,
    soffice_path: Path,
    pdftoppm_path: Path,
    expected_slide_count: int,
    expected_cjk_text: str,
) -> set[str]:
    failures: set[str] = set()
    try:
        _prepare_validation_dir(validation_dir)
    except PptxExportError:
        return {"pptx_validation_staging_invalid"}
    if not soffice_path.is_file() or not pdftoppm_path.is_file():
        return {"pptx_render_dependency_missing"}
    profile_dir = validation_dir / "lo-profile"
    profile_dir.mkdir(mode=0o700)
    command = [
        str(soffice_path),
        "--headless",
        f"-env:UserInstallation={profile_dir.as_uri()}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(validation_dir),
        str(pptx_path),
    ]
    try:
        cache_dir = validation_dir / "cache"
        cache_dir.mkdir(mode=0o700)
        render_environment = {
            **dict(os.environ),
            "XDG_CACHE_HOME": str(cache_dir),
        }
        if sys.platform == "darwin":
            fontconfig_path = validation_dir / "fonts.conf"
            fontconfig = ET.Element("fontconfig")
            for font_dir in (
                "/System/Library/Fonts",
                "/System/Library/Fonts/Supplemental",
                "/Library/Fonts",
            ):
                ET.SubElement(
                    fontconfig,
                    "dir",
                ).text = font_dir
            ET.SubElement(
                fontconfig,
                "cachedir",
            ).text = str(cache_dir / "fontconfig")
            ET.ElementTree(fontconfig).write(
                fontconfig_path,
                encoding="utf-8",
                xml_declaration=True,
            )
            render_environment["SAL_FONTPATH"] = "/System/Library/Fonts:/System/Library/Fonts/Supplemental"
            render_environment["FONTCONFIG_FILE"] = str(fontconfig_path)
        completed = subprocess.run(
            command,
            check=False,
            capture_output=True,
            timeout=180,
            env=render_environment,
        )
    except (OSError, subprocess.TimeoutExpired):
        return {"pptx_libreoffice_render_failed"}
    pdf_path = validation_dir / f"{pptx_path.stem}.pdf"
    if completed.returncode != 0 or not pdf_path.is_file():
        return {"pptx_libreoffice_render_failed"}
    try:
        reader = PdfReader(str(pdf_path))
        if len(reader.pages) != expected_slide_count:
            failures.add("pptx_rendered_slide_count_mismatch")
        rendered_text = "\n".join(page.extract_text() or "" for page in reader.pages)
        expected_cjk = _cjk_sequence(expected_cjk_text)
        rendered_cjk = _cjk_sequence(rendered_text)
        if not _is_subsequence(expected_cjk, rendered_cjk):
            failures.add("pptx_cjk_glyph_mismatch")
    except Exception:
        return {"pptx_libreoffice_pdf_invalid"}
    prefix = validation_dir / "slide"
    try:
        raster = subprocess.run(
            [
                str(pdftoppm_path),
                "-png",
                "-r",
                "96",
                str(pdf_path),
                str(prefix),
            ],
            check=False,
            capture_output=True,
            timeout=180,
        )
    except (OSError, subprocess.TimeoutExpired):
        return failures | {"pptx_slide_raster_failed"}
    images = sorted(validation_dir.glob("slide-*.png"))
    if raster.returncode != 0 or len(images) != expected_slide_count:
        failures.add("pptx_slide_raster_failed")
        return failures
    for image_path in images:
        try:
            with Image.open(image_path) as image:
                rgb = image.convert("RGB")
                if _pptx_body_is_blank(rgb):
                    failures.add("pptx_blank_slide")
        except (OSError, UnidentifiedImageError):
            failures.add("pptx_slide_raster_invalid")
    return failures


def _pptx_body_is_blank(image: Image.Image) -> bool:
    content = image.crop(
        (
            0,
            round(image.height * 0.16),
            image.width,
            round(image.height * 0.92),
        )
    )
    background = Image.new(
        "RGB",
        content.size,
        content.getpixel((0, content.height - 1)),
    )
    return ImageChops.difference(content, background).getbbox() is None


def _prepare_output_path(
    *,
    attempt_dir: Path,
    output_name: str,
) -> Path:
    if not attempt_dir.is_absolute() or Path(output_name).name != output_name or not output_name.endswith(".pptx"):
        raise PptxExportError("pptx_staging_invalid")
    attempt_dir.mkdir(mode=0o700, parents=True, exist_ok=True)
    if attempt_dir.is_symlink() or not attempt_dir.is_dir():
        raise PptxExportError("pptx_staging_invalid")
    output_dir = attempt_dir / "output"
    output_dir.mkdir(mode=0o700, exist_ok=True)
    if output_dir.is_symlink() or not output_dir.is_dir():
        raise PptxExportError("pptx_staging_invalid")
    output_path = output_dir / output_name
    if output_path.exists() or output_path.is_symlink():
        raise PptxExportError("pptx_staging_conflict")
    return output_path


def _prepare_validation_dir(path: Path) -> None:
    if not path.is_absolute() or path.exists() or path.is_symlink():
        raise PptxExportError("pptx_validation_staging_invalid")
    path.mkdir(mode=0o700, parents=True)


def _index_visuals(
    *,
    export_ir: ExportIR,
    materialized_visuals: tuple[MaterializedVisual, ...],
    attempt_dir: Path,
) -> dict[str, MaterializedVisual]:
    visuals = {visual.block_key: visual for visual in materialized_visuals}
    if len(visuals) != len(materialized_visuals):
        raise PptxExportError("duplicate_materialized_visual")
    expected = {block.block_key for block in iter_export_blocks(export_ir.blocks) if block.visual_input_sha256 is not None}
    if set(visuals) != expected:
        raise PptxExportError("visual_materialization_accounting_mismatch")
    root = attempt_dir.resolve()
    for visual in materialized_visuals:
        relative = Path(visual.relative_path)
        path = attempt_dir / relative
        if relative.is_absolute() or ".." in relative.parts or path.is_symlink() or not path.is_file() or not path.resolve().is_relative_to(root):
            raise PptxExportError("visual_staging_invalid")
        if _sha256_bytes(path.read_bytes()) != visual.content_sha256:
            raise PptxExportError("visual_content_changed")
    return visuals


def _target_for_profile(profile: PptxProfile) -> ExportTarget:
    if profile == "full_report":
        return "pptx_full"
    if profile == "executive_summary":
        return "pptx_summary"
    raise PptxExportError("pptx_profile_invalid")


def _is_appendix(
    *,
    block: ExportIRBlock,
    action: CapabilityAction,
) -> bool:
    return action == "appendix" or block.source.surface_placement in {"data_tab", "code_tab"}


def _is_groupable_text(
    *,
    block: ExportIRBlock,
    entry: CapabilityPlanEntry,
) -> bool:
    if entry.action == "fallback":
        return False
    if block.block_kind not in {
        "heading",
        "paragraph",
        "callout",
        "recommendations",
        "code",
        "source_list",
        "container",
    }:
        return False
    tokens = _block_display_tokens(block)
    return bool(tokens) and sum(len(token) for token in tokens) <= MAX_GROUPED_TEXT_CHARS


def _is_groupable_fallback(block: ExportIRBlock) -> bool:
    tokens = _block_display_tokens(block)
    return bool(tokens) and sum(len(token) for token in tokens) <= MAX_GROUPED_TEXT_CHARS


def _planned_slide_count(
    *,
    cover_blocks: tuple[ExportIRBlock, ...],
    render_items: list[tuple[ExportIRBlock, CapabilityPlanEntry]],
) -> int:
    slide_count = 1 if cover_blocks else 0
    index = 0
    while index < len(render_items):
        block, entry = render_items[index]
        appendix = _is_appendix(
            block=block,
            action=entry.action,
        )
        if _is_groupable_text(block=block, entry=entry):
            index += 1
            grouped = 1
            while index < len(render_items) and grouped < 3:
                candidate, candidate_entry = render_items[index]
                if (
                    not _is_groupable_text(
                        block=candidate,
                        entry=candidate_entry,
                    )
                    or _is_appendix(
                        block=candidate,
                        action=candidate_entry.action,
                    )
                    != appendix
                ):
                    break
                grouped += 1
                index += 1
            slide_count += 1
            continue
        if entry.action == "fallback" and _is_groupable_fallback(block):
            index += 1
            grouped = 1
            while index < len(render_items) and grouped < 3:
                candidate, candidate_entry = render_items[index]
                if candidate_entry.action != "fallback" or not _is_groupable_fallback(candidate):
                    break
                grouped += 1
                index += 1
            slide_count += 1
            continue
        slide_count += _block_slide_count(block, fallback=entry.action == "fallback")
        index += 1
    return max(1, slide_count)


def _block_slide_count(block: ExportIRBlock, *, fallback: bool) -> int:
    if block.block_kind == "visual":
        return 1
    if block.block_kind == "table":
        columns = block.content.get("columns")
        rows = block.content.get("rows")
        total_row = block.content.get("total_row")
        if not isinstance(columns, list) or not isinstance(rows, list):
            raise PptxExportError("pptx_table_invalid")
        if len(rows) > MAX_TABLE_ROWS:
            raise PptxExportError("pptx_table_row_limit")
        normalized_columns = [column for column in columns if isinstance(column, dict)]
        if not normalized_columns:
            # A column-less table renders through `_render_text_block(fallback=False)`.
            return _text_block_slide_count(block, fallback=False)
        normalized_row_count = sum(isinstance(row, dict) for row in rows)
        if isinstance(total_row, dict):
            normalized_row_count += 1
        page_size = max(5, min(12, 48 // len(normalized_columns)))
        return max(1, (normalized_row_count + page_size - 1) // page_size)
    if block.block_kind == "kpi_group":
        cells = block.content.get("cells")
        count = len(cells) if isinstance(cells, list) and cells else 1
        return (count + 5) // 6
    return _text_block_slide_count(block, fallback=fallback)


def _text_block_slide_count(block: ExportIRBlock, *, fallback: bool) -> int:
    units_per_line, max_lines = _text_block_geometry(block, fallback=fallback)
    return len(
        _chunk_tokens(
            list(_block_display_tokens(block)),
            units_per_line=units_per_line,
            max_lines=max_lines,
        )
    )


def _summary_render_rank(
    *,
    block: ExportIRBlock,
    entry: CapabilityPlanEntry,
) -> int:
    if _is_appendix(block=block, action=entry.action):
        return 90
    if block.block_kind == "kpi_group":
        return 10
    if _is_groupable_text(block=block, entry=entry) and block.block_kind != "recommendations":
        return 20
    if block.block_kind == "visual":
        return 30
    if block.block_kind == "recommendations":
        return 40
    return 50


def _block_display_tokens(
    block: ExportIRBlock,
) -> tuple[str, ...]:
    tokens = list(block.semantic_tokens)
    for token in _typed_visible_tokens(block):
        if token not in tokens:
            tokens.append(token)
    return tuple(tokens or ["—"])


def _business_shape_name(
    *,
    block: ExportIRBlock,
    part: int,
) -> str:
    return f"BH1B_{_marker_id(block.block_key)}_{_source_digest(block.source_refs)}_{part:04d}"


def _visual_shape_name(
    *,
    block: ExportIRBlock,
    pixel_sha256: str,
    part: int,
) -> str:
    return f"BH1V_{_marker_id(block.block_key)}_{_source_digest(block.source_refs)}_{pixel_sha256}_{part:04d}"


def _decor_shape_name(slide_number: int, part: int) -> str:
    return f"BH1D_{slide_number:04d}_{part:04d}"


def _marker_id(block_key: str) -> str:
    return hashlib.sha256(block_key.encode("utf-8")).hexdigest()[:24]


def _source_digest(source_refs: tuple[str, ...]) -> str:
    encoded = json.dumps(
        list(source_refs),
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _shape_text(shape: BaseShape) -> str:
    if getattr(shape, "has_table", False):
        table = cast(Any, shape).table
        return "".join(cell.text for row in table.rows for cell in row.cells)
    if getattr(shape, "has_text_frame", False):
        return str(cast(Any, shape).text)
    return ""


def _shape_has_business_content(shape: BaseShape) -> bool:
    return bool(_normalize_visible(_shape_text(shape))) or getattr(shape, "has_table", False) or shape.shape_type == MSO_SHAPE_TYPE.PICTURE


def _shape_on_slide(
    *,
    shape: BaseShape,
    presentation: PresentationType,
) -> bool:
    slide_width = presentation.slide_width
    slide_height = presentation.slide_height
    if slide_width is None or slide_height is None:
        return False
    return (
        shape.left >= 0
        and shape.top >= 0
        and shape.width > 0
        and shape.height > 0
        and shape.left + shape.width <= slide_width
        and shape.top + shape.height <= slide_height
    )


def _line_unit_width(text: str) -> float:
    """Rendered width of one line, measured in CJK-em units.

    A full-width ideograph counts as 1.0, whitespace as 0.35 and any other
    (typically half-width) glyph as 0.58. Shared by the overflow check and the
    pagination chunker so both estimate wrapping identically.
    """
    return sum(1.0 if _cjk_sequence(character) else 0.35 if character.isspace() else 0.58 for character in text)


def _line_budget(
    width_points: float,
    height_points: float,
    font_pt: float,
) -> tuple[float, int]:
    """Units admitted per line and the number of whole lines a box holds."""
    if font_pt <= 0:
        return (0.0, 0)
    return (width_points / font_pt, int(height_points / (font_pt * _LINE_HEIGHT_FACTOR)))


def _box_line_budget(
    *,
    width_in: float,
    height_in: float,
    font_pt: float,
) -> tuple[float, int]:
    """`_line_budget` for a box declared in inches with the fixed text margins.

    A shape built from `Inches(width_in)`/`Inches(height_in)` and rendered by
    `_set_shape_text` yields exactly these content points, so this matches the
    EMU-based measurement `_shape_text_within_capacity` performs on the closed
    file.
    """
    width_points = max(0.0, (width_in - 2 * _TEXT_MARGIN_H_IN) * 72.0)
    height_points = max(0.0, (height_in - 2 * _TEXT_MARGIN_V_IN) * 72.0)
    return _line_budget(width_points, height_points, font_pt)


def _paragraph_line_count(line: str, units_per_line: float) -> int:
    if units_per_line <= 0:
        return 1
    return max(1, math.ceil(_line_unit_width(line or " ") / units_per_line))


def _required_lines_for_text(text: str, units_per_line: float) -> int:
    """Lines `_set_shape_text` will consume for `text` at the given wrap width.

    Mirrors the renderer exactly: `_safe_text` then `splitlines`, each resulting
    paragraph (blank ones included) taking at least one line.
    """
    lines = _safe_text(text).splitlines() or ["—"]
    return sum(_paragraph_line_count(line, units_per_line) for line in lines)


def _shape_text_within_capacity(shape: BaseShape) -> bool:
    """Conservatively prove text fits its declared business text box."""
    if getattr(shape, "has_table", False):
        return True
    if not getattr(shape, "has_text_frame", False):
        return True
    frame = cast(Any, shape).text_frame
    font_sizes = [float(run.font.size.pt) for paragraph in frame.paragraphs for run in paragraph.runs if run.font.size is not None]
    font_size = max(font_sizes, default=18.0)
    width_points = max(
        0.0,
        (float(shape.width) - float(frame.margin_left) - float(frame.margin_right)) / 12_700,
    )
    height_points = max(
        0.0,
        (float(shape.height) - float(frame.margin_top) - float(frame.margin_bottom)) / 12_700,
    )
    max_units_per_line, max_lines = _line_budget(width_points, height_points, font_size)
    if max_units_per_line <= 0 or max_lines <= 0:
        return False
    required_lines = sum(_paragraph_line_count(paragraph.text, max_units_per_line) for paragraph in frame.paragraphs)
    return required_lines <= max_lines


def _slide_number(slide: Slide) -> int:
    presentation = slide.part.package.presentation_part.presentation
    for index, candidate in enumerate(presentation.slides, start=1):
        if candidate == slide:
            return index
    raise PptxExportError("pptx_slide_not_attached")


def _slide_title(block: ExportIRBlock) -> str:
    return {
        "kpi_group": "关键指标",
        "table": "数据明细",
        "visual": "核心图表",
        "recommendations": "行动建议",
        "source_list": "来源",
        "code": "SQL 与代码",
        "interactive_fallback": "在线交互内容",
    }.get(block.block_kind, block.source.section_kind)


def _split_to_line_budget(
    value: str,
    *,
    units_per_line: float,
    max_lines: int,
) -> list[str]:
    """Split one paragraph so each piece fits `max_lines` lines on its own.

    Used only when a single token is too long for one slide (a narrative
    section that flows across pages in the PDF). Pieces are cut on rendered
    width via a monotonic prefix search, so every piece — and therefore every
    downstream chunk — satisfies the geometric capacity check.
    """
    if _required_lines_for_text(value, units_per_line) <= max_lines:
        return [value]
    pieces: list[str] = []
    remaining = value
    while _required_lines_for_text(remaining, units_per_line) > max_lines:
        low, high = 1, len(remaining) - 1
        best = 1
        while low <= high:
            mid = (low + high) // 2
            if _required_lines_for_text(remaining[:mid], units_per_line) <= max_lines:
                best = mid
                low = mid + 1
            else:
                high = mid - 1
        pieces.append(remaining[:best])
        remaining = remaining[best:]
    if remaining:
        pieces.append(remaining)
    return pieces


def _chunk_tokens(
    tokens: list[str],
    *,
    units_per_line: float,
    max_lines: int,
) -> list[str]:
    """Pack paragraphs into slide-sized chunks that fit the text box's geometry.

    Each token is one paragraph; consecutive tokens are joined the same way
    `_render_text_block` feeds `_set_shape_text` (a blank line between them). A
    chunk is closed as soon as adding the next paragraph would exceed the box's
    line budget — the identical accounting `_shape_text_within_capacity`
    performs on the closed file — so no emitted slide can overflow. Token order
    is preserved, so the concatenated part text still contains every semantic
    token in order.
    """
    if units_per_line <= 0 or max_lines <= 0:
        joined = "\n\n".join(text for text in (_safe_text(token) for token in tokens) if text)
        return [joined] if joined else ["—"]
    chunks: list[str] = []
    current = ""
    for token in tokens:
        value = _safe_text(token)
        if not value:
            continue
        for piece in _split_to_line_budget(value, units_per_line=units_per_line, max_lines=max_lines):
            candidate = f"{current}\n\n{piece}" if current else piece
            if current and _required_lines_for_text(candidate, units_per_line) > max_lines:
                chunks.append(current)
                current = piece
            else:
                current = candidate
    if current:
        chunks.append(current)
    return chunks or ["—"]


def _text_block_font_pt(block: ExportIRBlock, *, fallback: bool) -> int:
    if fallback:
        return 19
    if block.block_kind == "code":
        return 13
    if block.block_kind == "heading":
        return 20
    return 17


def _text_block_geometry(block: ExportIRBlock, *, fallback: bool) -> tuple[float, int]:
    """(units_per_line, max_lines) for the box a text block renders into.

    The planner and the renderer both call this so their slide counts agree,
    which the `pptx_slide_count_invalid` gate requires.
    """
    _, _, width_in, height_in = _FALLBACK_CARD_BOX_IN if fallback else _BODY_BOX_IN
    return _box_line_budget(
        width_in=width_in,
        height_in=height_in,
        font_pt=float(_text_block_font_pt(block, fallback=fallback)),
    )


def _fit_font_pt(
    text: str,
    *,
    width_in: float,
    height_in: float,
    max_pt: int,
    min_pt: int,
) -> int:
    """Largest integer point size in [min_pt, max_pt] at which `text` fits a card.

    Used by the non-paginating grouped-card renderers, whose fixed layout can
    hold at most a few short blocks: shrinking within a readable floor keeps the
    card on one slide while still satisfying the overflow gate. Returns `max_pt`
    unchanged whenever it already fits, so fitting content is never restyled.
    """
    for candidate in range(max_pt, min_pt - 1, -1):
        units_per_line, max_lines = _box_line_budget(
            width_in=width_in,
            height_in=height_in,
            font_pt=float(candidate),
        )
        if max_lines > 0 and _required_lines_for_text(text, units_per_line) <= max_lines:
            return candidate
    return min_pt


def _fit_box_height_in(
    text: str,
    *,
    width_in: float,
    font_pt: int,
    min_height_in: float,
) -> float:
    """Height (inches) a text box needs so `text` fits at `font_pt`, at least min.

    The cover stacks fixed-height metadata boxes; a multi-line block (e.g. a
    three-line meta block) otherwise clips and the closed-file overflow check
    rejects the whole deck. Growing the box to its content keeps a readable font
    while satisfying the same geometric invariant `_shape_text_within_capacity`
    enforces. One point of slack absorbs the `int()` truncation in the gate's
    max-lines computation.
    """
    width_points = max(0.0, (width_in - 2 * _TEXT_MARGIN_H_IN) * 72.0)
    if font_pt <= 0 or width_points <= 0:
        return min_height_in
    required = _required_lines_for_text(text, width_points / font_pt)
    needed_in = (required * font_pt * _LINE_HEIGHT_FACTOR + 1.0) / 72.0 + 2 * _TEXT_MARGIN_V_IN
    return max(min_height_in, needed_in)


def _cover_block_layout(block: ExportIRBlock) -> tuple[str, int, float]:
    """(text, font_pt, box_height_in) for one cover masthead block.

    Shared by the cover-block selection (which bounds the stack to the slide)
    and `_render_cover` (which draws it), so the height budget the selection
    reserves is exactly the height the renderer uses.
    """
    text = "\n".join(block.semantic_tokens or ("—",))
    is_meta = block.source.section_kind in _COVER_META_KINDS
    size = 12 if is_meta else 19
    height = _fit_box_height_in(
        text,
        width_in=_COVER_TEXT_WIDTH_IN,
        font_pt=size,
        min_height_in=0.45 if is_meta else 0.85,
    )
    return text, size, height


def _select_cover_blocks(
    blocks: tuple[ExportIRBlock, ...],
    *,
    rendered: frozenset[str],
    plan_by_key: dict[str, CapabilityPlanEntry],
) -> tuple[ExportIRBlock, ...]:
    """The leading masthead run that fits the cover; the rest flows to the body.

    Only the contiguous run of cover-kind blocks at the very start is masthead;
    later h1 blocks are chapter headings and belong in the body. The run is
    further capped by the cover's vertical budget so grown boxes never push a
    shape off the slide (regression seen on a 75-section report). Blocks left
    out stay in the normal render items, so nothing is dropped.
    """
    selected: list[ExportIRBlock] = []
    top = _COVER_TOP_IN
    for block in blocks:
        if block.source.section_kind not in _COVER_KINDS:
            break  # end of the leading masthead run
        if block.block_key not in rendered or plan_by_key[block.block_key].action == "exclude":
            continue
        _, _, height = _cover_block_layout(block)
        if selected and top + height > _COVER_BOTTOM_IN:
            break  # cover region full — remaining masthead blocks render in the body
        selected.append(block)
        top += height + _COVER_GAP_IN
    return tuple(selected)


def _aspect_fit(
    *,
    source_width: int,
    source_height: int,
    box: tuple[float, float, float, float],
) -> tuple[float, float, float, float]:
    left, top, max_width, max_height = box
    scale = min(
        max_width / source_width,
        max_height / source_height,
    )
    width = source_width * scale
    height = source_height * scale
    return (
        left + (max_width - width) / 2,
        top + (max_height - height) / 2,
        width,
        height,
    )


def _contains_tokens_in_order(
    haystack: str,
    tokens: Iterable[str],
) -> bool:
    normalized = _normalize_visible(haystack)
    offset = 0
    for token in tokens:
        needle = _normalize_visible(token)
        if not needle:
            continue
        found = normalized.find(needle, offset)
        if found < 0:
            return False
        offset = found + len(needle)
    return True


def _typed_visible_tokens(block: ExportIRBlock) -> tuple[str, ...]:
    if block.typed_data_sha256 is None or block.block_kind == "visual":
        return ()
    if block.block_kind == "table":
        columns = block.content.get("columns")
        rows = block.content.get("rows")
        total_row = block.content.get("total_row")
        if not isinstance(columns, list) or not isinstance(rows, list):
            return ()
        keys = [column.get("key") for column in columns if isinstance(column, dict) and isinstance(column.get("key"), str)]
        values: list[str] = []
        for row in rows:
            if isinstance(row, dict):
                values.extend(_scalar_text(row.get(key)) for key in keys)
        if isinstance(total_row, dict):
            values.extend(_scalar_text(total_row.get(key)) for key in keys)
        return tuple(values)
    if block.block_kind == "kpi_group":
        return block.semantic_tokens
    if block.block_kind == "interactive_fallback":
        return tuple(_flatten_visible_scalars(block.content.get("data")))
    return ()


def _kpi_visible_tokens(value: Any) -> list[str]:
    if not isinstance(value, dict):
        return []
    return _flatten_visible_scalars(
        {
            key: value[key]
            for key in (
                "lbl",
                "label",
                "val",
                "value",
                "unit",
                "sub",
                "delta",
            )
            if key in value
        }
    )


def _flatten_visible_scalars(value: Any) -> list[str]:
    values: list[str] = []
    if value is None:
        return values
    if isinstance(value, dict):
        for key in sorted(value):
            values.extend(_flatten_visible_scalars(value[key]))
        return values
    if isinstance(value, list | tuple):
        for item in value:
            values.extend(_flatten_visible_scalars(item))
        return values
    values.append(_scalar_text(value))
    return values


def _scalar_text(value: Any) -> str:
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return json.dumps(
            value,
            ensure_ascii=False,
            allow_nan=False,
            separators=(",", ":"),
        )
    return _normalize_visible(str(value))


def _expected_visible_text(export_ir: ExportIR) -> str:
    values: list[str] = [export_ir.document.title]
    for block in iter_export_blocks(export_ir.blocks):
        values.extend(block.semantic_tokens)
        values.extend(_typed_visible_tokens(block))
    return "\n".join(values)


def _demo_residue_count(
    *,
    full_text: str,
    expected_text: str,
    fingerprints: tuple[str, ...],
) -> int:
    normalized_output = _normalize_visible(full_text)
    normalized_expected = _normalize_visible(expected_text)
    return sum(
        normalized_output.count(normalized)
        for fingerprint in fingerprints
        if (normalized := _normalize_visible(fingerprint)) and normalized not in normalized_expected
    )


def _normalize_visible(value: str) -> str:
    # Mirror export_pdf._normalize_visible: NFKC-fold, then drop ALL
    # whitespace. LibreOffice/python-pptx text extraction inserts stray
    # spaces at word-wrap boundaries that, in a CJK-dominant report, can
    # land at any position -- between ideographs, inside a number/citation,
    # or inside an ASCII acronym. None are semantic. Applied symmetrically
    # to the expected token and the observed text, dropping whitespace can
    # only let a true match succeed, never fabricate a false one (structural
    # fidelity is enforced by the visual and typed-data checks).
    return "".join(unicodedata.normalize("NFKC", value).split())


def _cjk_sequence(value: str) -> str:
    return "".join(
        character
        for character in value
        if ("\u3400" <= character <= "\u4dbf" or "\u4e00" <= character <= "\u9fff" or "\uf900" <= character <= "\ufaff")
    )


def _is_subsequence(expected: str, observed: str) -> bool:
    if not expected:
        return True
    iterator = iter(observed)
    return all(any(candidate == character for candidate in iterator) for character in expected)


def _safe_text(value: str) -> str:
    return "".join(character for character in unicodedata.normalize("NFC", value) if character in "\n\t" or ord(character) >= 32)


def _visible_pixel_sha256(image: Image.Image) -> str:
    rgba = image.convert("RGBA")
    visible = Image.new("RGB", rgba.size, "white")
    visible.paste(rgba, mask=rgba.getchannel("A"))
    identity = visible.width.to_bytes(4, "big") + visible.height.to_bytes(4, "big") + visible.tobytes()
    return _sha256_bytes(identity)


def _sha256_bytes(content: bytes) -> str:
    return f"sha256:{hashlib.sha256(content).hexdigest()}"


def _reason_code(value: str) -> ManifestReasonCode:
    if value == "summary_visual_limit":
        return "summary_visual_limit"
    return "summary_detail_omitted"


__all__ = [
    "PPTX_MIME",
    "PPTX_RENDERER_NAME",
    "PPTX_RENDERER_VERSION",
    "PptxExportError",
    "PptxRenderResult",
    "extract_pptx_evidence",
    "render_export_pptx",
]
