"""DOCX loader (python-docx L0).

Tech design v3.6 §3.1 + image-slim sprint.

V3: 砍 unstructured[docx] (~500MB spacy+blis+weasel+nltk), 直接走 python-docx
paragraphs + tables. unstructured 的 element-category 元数据用得少, python-docx 够用.
"""

from __future__ import annotations

import asyncio
import logging
import time

from bowhead_core.kb.document import KbDocument
from bowhead_core.kb.loader.base import BaseLoader, LoadResult

logger = logging.getLogger(__name__)


class DocxLoader(BaseLoader):
    """python-docx based DOCX loader (paragraphs + tables, cell-joined)."""

    async def load(self, file_path: str, *, use_llm: bool = False) -> LoadResult:
        t0 = time.time()
        try:
            docs = await asyncio.to_thread(self._extract, file_path)
            return LoadResult(
                documents=docs,
                layer_used="L0",
                duration_ms=int((time.time() - t0) * 1000),
                metadata={"element_count": len(docs)},
            )
        except Exception as exc:
            logger.warning("DOCX load failed for %s: %s", file_path, exc)
            return LoadResult(
                documents=[
                    KbDocument(
                        page_content="(DOCX parse failed)",
                        metadata={"page": None, "parse_failed": True},
                    )
                ],
                layer_used="L0",
                duration_ms=int((time.time() - t0) * 1000),
                metadata={"parse_failed": True},
            )

    @staticmethod
    def _extract(file_path: str) -> list[KbDocument]:
        """Sync extract: python-docx paragraphs + tables (cell-joined per row)."""
        from docx import Document as DocxDoc

        doc = DocxDoc(file_path)
        docs: list[KbDocument] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                docs.append(
                    KbDocument(
                        page_content=text,
                        metadata={"page": None, "element_category": "Text"},
                    )
                )
        # tables (cell-joined)
        for tbl in doc.tables:
            for row in tbl.rows:
                cells_text = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(c for c in cells_text if c)
                if row_text:
                    docs.append(
                        KbDocument(
                            page_content=row_text,
                            metadata={"page": None, "element_category": "Table"},
                        )
                    )
        return docs
