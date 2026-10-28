"""PPTX loader (python-pptx L0).

pptx-loader sprint v1.0 — 补齐 KB 摄入的 PowerPoint 缺口. 改动前是**两层堵**:
- 服务端 `routes/files.py` 的 ALLOWED_MIME **已**放行 pptx, 但 dispatch 无 loader
  → `UnsupportedMimeError` → 文件带 parse_error 躺库里, 不分块/不索引/chat 查不到/
  进不了文档矩阵抽取;
- 前端 `kb-upload-modal.tsx` 的 `ACCEPTED_EXT` 又漏了 `.pptx` → 用户在 KB 上传框里
  **连选都选不中** (故上面那条服务端路径实际只有 API 直调才走得到).
marketing 素材大量是 deck, 这是硬缺口. 本 sprint 两层一起补 (+ `kb-files-controls.tsx`
的 "slide" 类型筛选归位 —— 当年注释写明 "Re-add if that changes").

设计 (spike 实证, 见下):
- **一 slide 一 KbDocument**, `metadata["page"] = slide 序号` → worker `_split` 走默认
  段落分割器并逐 chunk 保留 metadata (`worker.py:1044`), `_build_kb_chunks` 把
  `page` 落进 `ChunkMetadata.page` (`worker.py:1085`) → 引用能溯源到"第几页".
- 抽取面 = 文本框 + 表格 + **图表(标题/分类/系列/数值)** + 演讲者备注.
- L0 only (无 LLM). 幻灯片内嵌图片不做 OCR — 同 DocxLoader 不 OCR 图, vision 留后续.
  **SmartArt 同样抽不到** (python-pptx 不暴露 diagram part), 故失败文案不写死
  "image-only" —— 见下.

Spike 实证的两个坑 (改这个文件前先读):
1. **组合形状 (GROUP) 不自动递归** —— `slide.shapes` 只吐 GROUP 本身 (实测顶层=1/
   递归=3), naive `for shape in slide.shapes` 会**静默丢掉整个组合里的文本**.
   必须走 `_iter_leaf_shapes` 手动递归 `group.shapes`.
2. **`slide.notes_slide` 惰性创建** —— 实测无备注时直接访问, `has_notes_slide`
   False→True (凭空造出空备注页). 本 loader 只读不 save 故**不落盘**, 但仍先
   `has_notes_slide` 门控再碰: 省掉无谓对象 + 不走空串分支.

失败模式 (spike 实测) —— 一律返 `parse_failed` 不 raise (同 DocxLoader/XlsxLoader,
worker.py:194 据此标 failed 并把 reason 给用户看):
- 损坏字节 → `pptx.exc.PackageNotFoundError`
- 合法 zip 但非 pptx → `KeyError: no member '/[Content_Types].xml'`
- 空 deck / 纯图 deck → 抽不出文本 → `parse_failed("no extractable text")`
  (同 PdfLoader 对 scan-only 无 vision 的约定 —— 静默存 junk chunk 会让用户以为
  文件可查, 实际搜不到.)
"""

from __future__ import annotations

import asyncio
import logging
import os
import time
from typing import TYPE_CHECKING

from bowhead_core.kb.document import KbDocument
from bowhead_core.kb.loader.base import BaseLoader, LoadResult

if TYPE_CHECKING:
    from collections.abc import Iterable, Iterator

    from pptx.chart.chart import Chart
    from pptx.shapes.base import BaseShape
    from pptx.slide import Slide
    from pptx.table import Table

logger = logging.getLogger(__name__)

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

MAX_SLIDES = int(os.environ.get("BOWHEAD_KB_PPTX_MAX_SLIDES", "500"))
"""Slide cap —— **只是 embed 成本闸, 不是内存闸**(同 XlsxLoader 的 row cap 范式).

真实 deck 20-80 页, 年度汇报 ~200 → 500 留足余量. 超出**记日志 + metadata 标
`truncated`**, 不静默截断.

⚠ **别指望它挡内存**: `Presentation(file_path)` 在切片之前就把 zip 里所有部件
(含 media blob)eager 读进内存(python-pptx `opc/serialized.py`), 所以 1 页的
deck 照样能吃几个 G —— 解压炸弹面归 `untrusted_parse` 的「簇 B」follow-on
(见 docs/sprints/skill-multiruntime/s3-followup-untrusted-parse-design.md;
docx/xlsx 同样暴露, 非 pptx 独有).
"""


class PptxLoader(BaseLoader):
    """python-pptx based PPTX loader (per-slide: 文本框 + 表格 + 演讲者备注)."""

    SUPPORTED_MIMES = frozenset({PPTX_MIME})

    async def load(self, file_path: str, *, use_llm: bool = False) -> LoadResult:
        t0 = time.time()
        try:
            docs, slide_count, truncated = await asyncio.to_thread(self._extract, file_path)
        except Exception as exc:
            # Corrupt pptx / non-pptx zip / lxml parse error—return parse_failed so worker marks failed.
            logger.warning("pptx_loader: load failed for %s: %s", file_path, exc)
            return LoadResult(
                documents=[
                    KbDocument(
                        page_content=f"(pptx parse failed: {exc})",
                        metadata={"page": None, "parse_failed": True},
                    )
                ],
                layer_used="L0",
                duration_ms=int((time.time() - t0) * 1000),
                metadata={"parse_failed": True, "reason": str(exc)},
            )

        if not docs:
            # No text extracted: saving as done suggests searchable content that cannot be found → mark failed.
            # ⚠ Do not hardcode "image-only": all-SmartArt (python-pptx does not expose diagram parts)
            #   or MAX_SLIDES truncation also reaches this branch; asserting image-only misdirects users to nonexistent images.
            reason = f"no extractable text in {slide_count} slide(s) (image-only / SmartArt-only / empty?)"
            if truncated:
                reason += f" — 注意: 已按 MAX_SLIDES={MAX_SLIDES} 截断"
            logger.info("pptx_loader: %s → %s", file_path, reason)
            return LoadResult(
                documents=[
                    KbDocument(
                        page_content="(pptx: no extractable text)",
                        metadata={"page": None, "parse_failed": True},
                    )
                ],
                layer_used="L0",
                duration_ms=int((time.time() - t0) * 1000),
                metadata={
                    "parse_failed": True,
                    "reason": reason,
                    "slide_count": slide_count,
                    "truncated": truncated,
                },
            )

        meta: dict[str, object] = {
            "slide_count": slide_count,
            # page_count: same key as PdfLoader (worker._split reads it to detect long docs;
            # that branch gates on mime==pdf, so pptx never reaches it; same name only for consistency/observability).
            "page_count": slide_count,
            "chunk_count": len(docs),
        }
        if truncated:
            meta["truncated"] = True
            meta["max_slides"] = MAX_SLIDES

        return LoadResult(
            documents=docs,
            layer_used="L0",
            duration_ms=int((time.time() - t0) * 1000),
            metadata=meta,
        )

    @staticmethod
    def _extract(file_path: str) -> tuple[list[KbDocument], int, bool]:
        """Sync extract: slides → KbDocument (一页一 doc). 返 (docs, slide_count, truncated)."""
        from pptx import Presentation

        prs = Presentation(file_path)
        slides = list(prs.slides)
        slide_count = len(slides)
        truncated = slide_count > MAX_SLIDES
        if truncated:
            logger.warning(
                "pptx_loader: deck has %d slides, capping at MAX_SLIDES=%d (dropped %d)",
                slide_count,
                MAX_SLIDES,
                slide_count - MAX_SLIDES,
            )
            slides = slides[:MAX_SLIDES]

        docs: list[KbDocument] = []
        for idx, slide in enumerate(slides, start=1):
            text = _slide_to_text(slide, idx)
            if text is None:
                continue  # Image-only / empty slide—no indexable text, skip
            docs.append(
                KbDocument(
                    page_content=text,
                    metadata={"page": idx, "slide": idx},
                )
            )
        return docs, slide_count, truncated


def _slide_to_text(slide: Slide, idx: int) -> str | None:
    """一页 → 文本块 (文本框 + 表格 markdown + 备注). 全空返 None.

    isinstance 判型而非 `shp.has_text_frame` / `shp.has_table`: 两者**语义等价**
    (穷举 pptx.shapes.* 实证 —— `has_text_frame` 只在 `Shape` 上定义为 `return True`,
    其余全继承 BaseShape 的 False; `has_table` 只可能在 `GraphicFrame` 血统为 True),
    但 isinstance 能给 类型检查器 窄化出 `.text_frame` / `.table` (它们不在 BaseShape 上).
    """
    from pptx.shapes.autoshape import Shape
    from pptx.shapes.graphfrm import GraphicFrame

    parts: list[str] = []
    skipped = 0
    # ⚠ Iteration itself must be inside try: python-pptx shape factory is lazy; malformed shapes throw
    # **when advancing the generator**, not in the loop body; outside try, one bad shape ruins the whole deck.
    shapes_iter = _iter_leaf_shapes(slide.shapes)
    while True:
        try:
            shp = next(shapes_iter)
        except StopIteration:
            break
        except Exception as exc:  # noqa: BLE001
            logger.warning("pptx_loader: slide %d shape iteration aborted: %s", idx, exc)
            break
        try:
            if isinstance(shp, GraphicFrame) and shp.has_table:
                md = _table_to_md(shp.table)
                if md:
                    parts.append(md)
            elif isinstance(shp, GraphicFrame) and shp.has_chart:
                # Charts: titles/categories/series are all **text**, but charts have no text_frame—
                # checking only has_text_frame silently loses all chart information (marketing decks are chart-heavy).
                cm = _chart_to_text(shp.chart)
                if cm:
                    parts.append(cm)
            elif isinstance(shp, Shape) and shp.has_text_frame:
                text = shp.text_frame.text.strip()
                if text:
                    parts.append(text)
        except Exception as exc:  # noqa: BLE001
            # One malformed shape must not ruin the whole slide (PPT has many shape types).
            skipped += 1
            logger.warning("pptx_loader: slide %d shape skipped: %s", idx, exc)
            continue
    if skipped:
        logger.warning("pptx_loader: slide %d skipped %d malformed shape(s)", idx, skipped)

    # ⚠ Gate on has_notes_slide first—accessing notes_slide lazily creates an empty notes page (observed
    # False→True; read-only without save does not persist, but creation is unnecessary).
    try:
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            note = (notes_tf.text or "").strip() if notes_tf is not None else ""
            if note:
                parts.append(f"Notes: {note}")
    except Exception as exc:  # noqa: BLE001
        logger.debug("pptx_loader: slide %d notes skipped: %s", idx, exc)

    if not parts:
        return None
    return f"## Slide {idx}\n\n" + "\n\n".join(parts)


def _iter_leaf_shapes(shapes: Iterable[BaseShape]) -> Iterator[BaseShape]:
    """深度优先递归展开组合形状 → 叶子 shape.

    ⚠ spike 实证: `slide.shapes` **不**递归 GROUP (只吐 GROUP 本身, 实测顶层=1/
    递归=3), 不手动进 `group.shapes` 会静默丢掉组合内全部文本.

    用 `isinstance(GroupShape)` 而非 `shape_type == MSO_SHAPE_TYPE.GROUP`: 异种
    shape 上访问 `shape_type` 可能抛, isinstance 不会 (且给 类型检查器 窄化出 `.shapes`).
    """
    from pptx.shapes.group import GroupShape

    for shp in shapes:
        if isinstance(shp, GroupShape):
            yield from _iter_leaf_shapes(shp.shapes)
        else:
            yield shp


def _chart_to_text(chart: Chart) -> str:
    """图表 → 文本 (标题 + 分类×系列 的 pipe 表).

    图表是 `GraphicFrame` 且 `has_text_frame=False` —— 光判 text_frame 会把标题
    (常是该页的核心结论)、系列名、分类标签全部静默丢掉. 顺带把数值也铺出来:
    marketing deck 的图表数据正是矩阵抽取想要的.

    各部分独立兜底: 不同 chart 类型 API 行为差异大(XY/bubble 无 categories,
    无名系列的 `.name` 会抛), 缺一块不该废掉整张图.
    """
    parts: list[str] = []
    try:
        if chart.has_title:
            title = chart.chart_title.text_frame.text.strip()
            if title:
                parts.append(f"Chart: {title}")
    except Exception as exc:  # noqa: BLE001
        logger.debug("pptx_loader: chart title skipped: %s", exc)

    try:
        cats = [str(c).strip() for c in chart.plots[0].categories]
    except Exception:  # noqa: BLE001
        cats = []  # XY / bubble has no categories

    series: list[tuple[str, list[object]]] = []
    try:
        for s in chart.series:
            try:
                name = str(s.name or "").strip()
            except Exception:  # noqa: BLE001
                name = ""
            try:
                vals = list(s.values)
            except Exception:  # noqa: BLE001
                vals = []
            series.append((name, vals))
    except Exception as exc:  # noqa: BLE001
        logger.debug("pptx_loader: chart series skipped: %s", exc)

    if cats and series:
        parts.append("| " + " | ".join(["", *(n for n, _ in series)]) + " |")
        for i, cat in enumerate(cats):
            row = [cat] + [_num_to_str(v[i]) if i < len(v) else "" for _, v in series]
            parts.append("| " + " | ".join(row) + " |")
    elif series:  # No category axis (XY/bubble): at least keep series names searchable
        names = [n for n, _ in series if n]
        if names:
            parts.append("Series: " + " / ".join(names))
    return "\n".join(parts)


def _num_to_str(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _table_to_md(table: Table) -> str:
    """表格 → pipe 行 (同 DocxLoader 的 cell-joined 约定).

    不造 markdown header 分隔行 —— pptx 表格不声明表头语义, 不臆造
    (XlsxLoader 靠启发式判表头是因为它面对的就是纯表格; pptx 更接近 docx).
    """
    lines: list[str] = []
    for row in table.rows:
        cells = [(c.text or "").replace("|", "\\|").replace("\n", " ").strip() for c in row.cells]
        if any(cells):
            lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)
