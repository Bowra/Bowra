"""V2.9 A5 · Template-clone PPTX render.

用户传 .pptx 母版 + ComposedSection[] (来自 A4) → python-pptx unpack→fill→pack
→ 输出新 .pptx 完整保留母版.

跟 V2.6 canvas-pptx 共存:
  - canvas-pptx: 用 canvas-default.pptx, 7-slide 固定结构, "导出 canvas 为 PPT"
  - template-clone (本模块): 用用户 .pptx, 任意结构完整保留, "用这个模板做新一期"

Workflow (Anthropic pptx editing.md 模式):
  1. python-pptx Presentation(template_path)
  2. 逐 slide 找 anchors (text frames / tables / charts)
  3. 按 title/chapter_tag 把 ComposedSection 匹配到 slide
  4. 替换 text / table cells / chart data (chart 是真 chart 才能换)
  5. Save new .pptx

设计文档: docs/template-clone-design.md § 4 [render] PPTX 路径
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from pathlib import Path

    from bowhead_core.template_clone_workflow import ComposedSection

logger = logging.getLogger(__name__)


@dataclass
class FillResult:
    output_path: str
    slides_filled: int
    slides_skipped: int
    sections_unmatched: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "output_path": self.output_path,
            "slides_filled": self.slides_filled,
            "slides_skipped": self.slides_skipped,
            "sections_unmatched": self.sections_unmatched,
            "errors": self.errors,
        }


# ─────────────────────────────────────────────────────────────
# Section ↔ slide matching (heuristic)
# ─────────────────────────────────────────────────────────────


def _slide_text_blob(slide: Any) -> str:
    """Slide 所有 text frame 拼起来 (给 title 匹配用)."""
    parts: list[str] = []
    for shape in slide.shapes:
        if hasattr(shape, "has_text_frame") and shape.has_text_frame:
            t = shape.text_frame.text
            if t:
                parts.append(t)
    return "\n".join(parts)


def _match_sections_to_slides(
    prs: Any,
    sections: list[ComposedSection],
) -> dict[int, ComposedSection | None]:
    """Heuristic: title substring + chapter_tag prefix.

    返 {slide_idx -> ComposedSection or None}.
    每个 section 最多匹配 1 个 slide (优先标题完整命中, 否则 chapter_tag fallback).
    """
    out: dict[int, ComposedSection | None] = {}
    used: set[str] = set()  # Already-used section IDs
    slide_texts: list[str] = []
    for slide in prs.slides:
        slide_texts.append(_slide_text_blob(slide))

    # Pass 1: full title match
    for slide_idx, txt in enumerate(slide_texts):
        for sec in sections:
            if sec.id in used:
                continue
            if sec.title and sec.title in txt:
                out[slide_idx] = sec
                used.add(sec.id)
                break

    # Pass 2: chapter_tag fallback (unmatched slide + unused section)
    for slide_idx, txt in enumerate(slide_texts):
        if slide_idx in out:
            continue
        for sec in sections:
            if sec.id in used:
                continue
            if sec.chapter_tag and sec.chapter_tag in txt[:80]:
                out[slide_idx] = sec
                used.add(sec.id)
                break
    return out


# ─────────────────────────────────────────────────────────────
# Fill one slide
# ─────────────────────────────────────────────────────────────


def _fill_table_with_rows(tbl: Any, rows: list[dict[str, Any]]) -> None:
    """填 table cell. row 0 (表头) 保留, row 1+ 替换成 rows 数据.

    cols 取 rows[0].keys() 顺序. 超出 table 容量截断."""
    if not rows or not isinstance(rows[0], dict):
        return
    cols = list(rows[0].keys())
    n_table_rows = len(tbl.rows)
    if n_table_rows < 2:
        return
    # Skip header row(s): simply treat row 0 as header. Nested headers need Phase B2 improvements.
    for r_offset, row in enumerate(rows[: n_table_rows - 1]):
        target_row = tbl.rows[r_offset + 1]
        for ci, col in enumerate(cols):
            if ci >= len(target_row.cells):
                break
            cell = target_row.cells[ci]
            v = row.get(col, "")
            cell.text = str(v) if v is not None else ""


def _fill_chart_data(chart: Any, section: ComposedSection) -> bool:
    """Chart 数据替换. 只支持 Bowhead `data` 是 list[dict] 形态.

    返 True 已替换, False 跳过 (chart 不是 categorical / 数据形态不匹配)."""
    if section.data is None:
        return False
    if not isinstance(section.data, list) or not section.data:
        return False
    if not isinstance(section.data[0], dict):
        return False
    rows = section.data
    cols = list(rows[0].keys())
    if len(cols) < 2:
        return False  # Requires at least 1 categorical + 1 numeric column

    try:
        from pptx.chart.data import CategoryChartData
    except ImportError:
        return False

    # First column = category; subsequent columns = series
    cat_col = cols[0]
    series_cols = cols[1:]
    chart_data = CategoryChartData()
    chart_data.categories = [str(r.get(cat_col, "")) for r in rows]
    for sc in series_cols:
        try:
            values = [float(r.get(sc, 0)) if r.get(sc) is not None else 0.0 for r in rows]
        except (TypeError, ValueError):
            continue
        chart_data.add_series(sc, values)

    try:
        chart.replace_data(chart_data)
        return True
    except Exception as e:
        logger.debug("chart.replace_data failed: %s", e)
        return False


def _fill_narrative_blocks(slide: Any, section: ComposedSection) -> int:
    """填 narrative 槽位 (客观结论 Facts / 关键发现 Key Findings).

    简单实现: 找含"客观结论" 或 "关键发现" 的 text frame, 替换其内容为 section.narrative.
    返替换次数."""
    if not section.narrative:
        return 0
    n = 0
    for shape in slide.shapes:
        if not (hasattr(shape, "has_text_frame") and shape.has_text_frame):
            continue
        tf = shape.text_frame
        existing = tf.text or ""
        if any(kw in existing for kw in ("客观结论", "关键发现", "Facts", "Findings")):
            # Simple approach: replace all frame text with narrative.
            # NB: retaining bullet style needs finer paragraph operations (left for Phase B optimization).
            tf.text = section.narrative
            n += 1
    return n


def _fill_slide(slide: Any, section: ComposedSection) -> dict[str, int]:
    """单 slide 填表 + chart + narrative. 返各类填充次数."""
    counts = {"tables": 0, "charts": 0, "narratives": 0}

    # Tabular data → fill the slide's first table.
    if isinstance(section.data, list) and section.data and isinstance(section.data[0], dict):
        for shape in slide.shapes:
            if hasattr(shape, "has_table") and shape.has_table:
                try:
                    _fill_table_with_rows(shape.table, section.data)
                    counts["tables"] += 1
                except Exception as e:
                    logger.warning("fill_table_failed: %s", e)
                break  # Fill one table per slide; multiple tables left for later

    # Replace chart data (try every chart in slide; fail-soft)
    for shape in slide.shapes:
        if hasattr(shape, "has_chart") and shape.has_chart:
            try:
                if _fill_chart_data(shape.chart, section):
                    counts["charts"] += 1
            except Exception as e:
                logger.debug("chart_fill_skip: %s", e)

    # narrative
    counts["narratives"] = _fill_narrative_blocks(slide, section)
    return counts


# ─────────────────────────────────────────────────────────────
# Main entry point
# ─────────────────────────────────────────────────────────────


def render_filled_pptx(
    template_path: str | Path,
    composed_sections: list[ComposedSection],
    output_path: str | Path,
) -> FillResult:
    """用户 .pptx 当 baseline, 把 ComposedSection 数据填进去 → 输出新 .pptx."""
    try:
        from pptx import Presentation
    except ImportError as e:
        raise RuntimeError("python-pptx 未装") from e

    template_path = str(template_path)
    output_path = str(output_path)
    prs = Presentation(template_path)

    matches = _match_sections_to_slides(prs, composed_sections)
    filled = 0
    skipped = 0
    errors: list[str] = []

    slides = list(prs.slides)
    for slide_idx, slide in enumerate(slides):
        section = matches.get(slide_idx)
        if section is None:
            skipped += 1
            continue
        if section.needs_user_input:
            # Deferred/failed section: do not fill data (user fills later); preserve template unchanged.
            skipped += 1
            continue
        try:
            _fill_slide(slide, section)
            filled += 1
        except Exception as e:
            errors.append(f"slide {slide_idx}: {e}")
            logger.warning("fill_slide_failed slide=%d err=%s", slide_idx, e)

    # Find sections unmatched to slides.
    matched_ids = {s.id for s in matches.values() if s is not None}
    unmatched = [s.id for s in composed_sections if s.id not in matched_ids]

    prs.save(output_path)
    logger.info(
        "template_clone.render template=%s out=%s filled=%d skipped=%d unmatched=%d",
        template_path,
        output_path,
        filled,
        skipped,
        len(unmatched),
    )
    return FillResult(
        output_path=output_path,
        slides_filled=filled,
        slides_skipped=skipped,
        sections_unmatched=unmatched,
        errors=errors,
    )
