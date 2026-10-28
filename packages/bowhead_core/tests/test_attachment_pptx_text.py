"""Bounded Office reads retain table relationships and nested group text."""

import io
import json

import pytest
from bowhead_core.attachments_injector import _PPTX_DEEPREAD_CODE, _extract_pptx
from pptx import Presentation
from pptx.util import Inches

pytestmark = pytest.mark.unit


def deck() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "Dispatch review"
    table = slide.shapes.add_table(3, 2, 0, Inches(1), Inches(4), Inches(2)).table
    for row, values in zip(table.rows, (("Hub", "Parcels"), ("West", "100"), ("East", "50")), strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    group = slide.shapes.add_group_shape().shapes.add_group_shape()
    group.shapes.add_textbox(0, Inches(3), Inches(4), Inches(1)).text = "Nested footer"
    table_shape = slide.shapes.add_table(2, 2, 0, 0, Inches(2), Inches(2))
    group.shapes.add_group_shape([table_shape])
    grouped_table = table_shape.table
    grouped_table.cell(0, 0).text = "Merged heading"
    grouped_table.cell(0, 0).merge(grouped_table.cell(0, 1))
    grouped_table.cell(1, 1).text = "Two\nlines"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def assert_structure(text: str) -> None:
    assert "Dispatch review" in text and "Nested footer" in text
    assert '["Hub", "Parcels"]' in text
    assert '["West", "100"]' in text and '["East", "50"]' in text
    assert '["Merged heading", ""]' in text
    assert json.dumps(["", "Two\nlines"]) in text
    assert text.count("Merged heading") == 1
    assert text.index("Dispatch review") < text.index("West") < text.index("Nested footer")


def test_bounded_pptx_read_preserves_tables_and_nested_groups() -> None:
    text, note = _extract_pptx(deck())
    assert_structure(text)
    assert "charts" in note.lower() and "images" in note.lower()


def test_pptx_deep_read_example_executes_with_the_same_structure(tmp_path, capsys) -> None:
    path = tmp_path / "template.pptx"
    path.write_bytes(deck())
    exec(_PPTX_DEEPREAD_CODE, {"INPUTS": {"doc": {"path": str(path)}}})  # noqa: S102 - Execute the repository-owned sandbox example.
    assert_structure(capsys.readouterr().out)


def test_bounded_read_keeps_slide_limit() -> None:
    presentation = Presentation()
    for index in range(31):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = f"Slide marker {index + 1}"
    output = io.BytesIO()
    presentation.save(output)
    text, note = _extract_pptx(output.getvalue())
    assert "Slide marker 30" in text and "Slide marker 31" not in text
    assert "31" in note and "30" in note
