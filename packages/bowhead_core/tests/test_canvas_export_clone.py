"""V2.9 A6+ #60: canvas_export use_clone_render path 单测."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pytest
from bowhead_core.artifact_store import ArtifactStore, LocalBackend
from bowhead_core.canvas_export import (
    ExportJob,
    _load_owned_pptx_template,
    _run_clone_render,
    _run_export,
)

from bowhead_core import canvas_export


@dataclass
class _FakeArt:
    id: str
    filename: str
    mime: str
    size: int = 1024
    user_id: str = "u_test"
    bytes_path: Path | None = None


class _FakeArtStore:
    def __init__(self) -> None:
        self.gets: dict[str, _FakeArt] = {}
        self.get_calls: list[dict[str, object]] = []
        self.put_calls: list[dict] = []
        self.next_put_id = "art_clone_out_1"

    async def get(
        self,
        aid: str,
        *,
        check_created_by: str | None = None,
        reject_null_owner: bool = False,
    ) -> _FakeArt | None:
        self.get_calls.append(
            {
                "artifact_id": aid,
                "check_created_by": check_created_by,
                "reject_null_owner": reject_null_owner,
            }
        )
        return self.gets.get(aid)

    async def open_read(self, art: _FakeArt) -> bytes:
        if art.bytes_path:
            return art.bytes_path.read_bytes()
        return b""

    async def put_file(self, *, src_path: Path, filename: str, mime: str, source: str, canvas_id: str, created_by: str) -> _FakeArt:
        self.put_calls.append(
            {
                "src_path": str(src_path),
                "filename": filename,
                "mime": mime,
                "source": source,
                "canvas_id": canvas_id,
            }
        )
        size = src_path.stat().st_size if src_path.exists() else 0
        return _FakeArt(
            id=self.next_put_id,
            filename=filename,
            mime=mime,
            size=size,
            user_id=created_by,
        )


@dataclass
class _FakeCanvas:
    id: str
    title: str
    user_id: str = "u_test"


@dataclass
class _FakeReportSection:
    section_id: str
    title: str = ""
    payload: dict = field(default_factory=dict)

    def to_report_section(self) -> dict:
        return {
            "id": self.section_id,
            "title": self.title,
            "chapter_tag": self.payload.get("chapter_tag", ""),
            "kind": self.payload.get("kind", "narrative"),
            "data": self.payload.get("data"),
            "narrative": self.payload.get("narrative", ""),
            "needs_user_input": self.payload.get("needs_user_input", False),
        }

    @property
    def id(self) -> str:
        return self.section_id


class _FakeStore:
    def __init__(self, canvas: _FakeCanvas, sections: list[_FakeReportSection]) -> None:
        self._c = canvas
        self._s = sections

    async def get_canvas(self, _id: str) -> _FakeCanvas:
        return self._c

    async def list_sections(self, _id: str) -> list[_FakeReportSection]:
        return self._s


def _make_template_pptx(out: Path) -> None:
    """合成 1-slide 模板."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    s0 = prs.slides.add_slide(blank)
    tb = s0.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    tb.text_frame.text = "测试章节"
    tbl = s0.shapes.add_table(
        rows=2,
        cols=2,
        left=Inches(0.5),
        top=Inches(1.5),
        width=Inches(8),
        height=Inches(2),
    ).table
    tbl.rows[0].cells[0].text = "区域"
    tbl.rows[0].cells[1].text = "销售额"
    tbl.rows[1].cells[0].text = "(填充)"
    tbl.rows[1].cells[1].text = "(填充)"
    prs.save(str(out))


@pytest.mark.asyncio
async def test_run_clone_render_fills_user_template(tmp_path: Path) -> None:
    # 1. Synthesize a template and register a fake artifact_store.
    tpl_path = tmp_path / "tpl.pptx"
    _make_template_pptx(tpl_path)

    art_store = _FakeArtStore()
    art_store.gets["art_tpl_1"] = _FakeArt(
        id="art_tpl_1",
        filename="user_template.pptx",
        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        bytes_path=tpl_path,
    )

    # 2. Canvas has one section matching the template chapter.
    canvas = _FakeCanvas(id="cv1", title="测试 canvas")
    sections = [
        _FakeReportSection(
            section_id="s1",
            title="测试章节",
            payload={"kind": "breakdown", "data": [{"区域": "华东", "销售额": 100}]},
        )
    ]
    store = _FakeStore(canvas, sections)

    # 3. job
    job = ExportJob(id="exp_test", canvas_id="cv1", fmt="pptx", status="queued")

    # 4. run
    await _run_clone_render(
        job=job,
        store=store,
        artifact_store=art_store,
        template_artifact_id="art_tpl_1",
    )

    # 5. Verify.
    assert job.status == "done", f"status={job.status} err={job.error}"
    assert job.artifact_id == "art_clone_out_1"
    assert art_store.get_calls == [
        {
            "artifact_id": "art_tpl_1",
            "check_created_by": "u_test",
            "reject_null_owner": True,
        }
    ]
    assert len(art_store.put_calls) == 1
    put = art_store.put_calls[0]
    assert put["source"] == "canvas_export_clone"
    assert put["canvas_id"] == "cv1"

    # 6. Output .pptx should fill row 1.
    saved_path = Path(put["src_path"])
    assert saved_path.exists()
    from pptx import Presentation

    prs = Presentation(str(saved_path))
    slide0 = next(iter(prs.slides))
    tbl = next((sh.table for sh in slide0.shapes if sh.has_table), None)
    assert tbl is not None
    assert tbl.rows[1].cells[0].text == "华东"
    assert tbl.rows[1].cells[1].text == "100"


@pytest.mark.asyncio
async def test_run_clone_render_missing_template_fails(tmp_path: Path) -> None:
    art_store = _FakeArtStore()  # Empty; art_tpl_1 does not exist
    canvas = _FakeCanvas(id="cv1", title="test")
    sections = [_FakeReportSection(section_id="s1", title="x")]
    store = _FakeStore(canvas, sections)

    job = ExportJob(id="exp_test", canvas_id="cv1", fmt="pptx", status="queued")
    await _run_clone_render(
        job=job,
        store=store,
        artifact_store=art_store,
        template_artifact_id="art_missing",
    )
    assert job.status == "failed"
    assert job.error == "template artifact missing_or_forbidden"


@pytest.mark.asyncio
async def test_run_clone_render_no_template_id_fails(tmp_path: Path) -> None:
    art_store = _FakeArtStore()
    canvas = _FakeCanvas(id="cv1", title="test")
    store = _FakeStore(canvas, [])

    job = ExportJob(id="exp_test", canvas_id="cv1", fmt="pptx", status="queued")
    await _run_clone_render(
        job=job,
        store=store,
        artifact_store=art_store,
        template_artifact_id=None,
    )
    assert job.status == "failed"
    assert "template_artifact_id" in (job.error or "")


@pytest.mark.asyncio
async def test_owned_pptx_template_security_matrix(tmp_path: Path) -> None:
    """W0-B: cross-user/null-owner/deleted/non-PPTX all look missing."""
    artifact_store = ArtifactStore(
        backend=LocalBackend(root_dir=tmp_path / "artifacts"),
        db_path=tmp_path / "artifacts.db",
    )
    await artifact_store.initialize()
    pptx_mime = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    allowed = await artifact_store.put_bytes(
        content=b"pptx",
        filename="allowed.pptx",
        mime=pptx_mime,
        created_by="u_owner",
    )
    cross_user = await artifact_store.put_bytes(
        content=b"pptx",
        filename="cross-user.pptx",
        mime=pptx_mime,
        created_by="u_other",
    )
    owner_null = await artifact_store.put_bytes(
        content=b"pptx",
        filename="owner-null.pptx",
        mime=pptx_mime,
        created_by=None,
    )
    non_pptx = await artifact_store.put_bytes(
        content=b"not a deck",
        filename="not-a-deck.txt",
        mime="text/plain",
        created_by="u_owner",
    )
    deleted = await artifact_store.put_bytes(
        content=b"pptx",
        filename="deleted.pptx",
        mime=pptx_mime,
        created_by="u_owner",
    )
    assert await artifact_store.soft_delete(deleted.id, user_id="u_owner")

    loaded = await _load_owned_pptx_template(
        artifact_store=artifact_store,
        template_artifact_id=allowed.id,
        owner_user_id="u_owner",
    )
    assert loaded is not None and loaded.id == allowed.id

    for denied_id in (
        cross_user.id,
        owner_null.id,
        non_pptx.id,
        deleted.id,
        "art_missing",
    ):
        assert (
            await _load_owned_pptx_template(
                artifact_store=artifact_store,
                template_artifact_id=denied_id,
                owner_user_id="u_owner",
            )
            is None
        )


@pytest.mark.asyncio
async def test_standard_template_denial_is_terminal(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """W0-B: standard path must not fall back after an unauthorized template."""

    class _Catalog:
        def get(self, _name: str) -> object:
            return object()

    class _Runtime:
        def __init__(self) -> None:
            self.called = False

        async def run(self, *_args: object, **_kwargs: object) -> object:
            self.called = True
            raise AssertionError("renderer must not run for a forbidden template")

    monkeypatch.setattr(canvas_export, "_db_path", None)
    art_store = _FakeArtStore()
    store = _FakeStore(
        _FakeCanvas(id="cv1", title="test", user_id="u_owner"),
        [_FakeReportSection(section_id="s1", title="x")],
    )
    runtime = _Runtime()
    job = ExportJob(id="exp_standard_denied", canvas_id="cv1", fmt="pptx", status="queued")

    await _run_export(
        job=job,
        store=store,
        catalog=_Catalog(),
        runtime=runtime,
        artifact_store=art_store,
        template_artifact_id="art_forbidden",
    )

    assert job.status == "failed"
    assert job.error == "template artifact missing_or_forbidden"
    assert runtime.called is False
    assert art_store.get_calls == [
        {
            "artifact_id": "art_forbidden",
            "check_created_by": "u_owner",
            "reject_null_owner": True,
        }
    ]
