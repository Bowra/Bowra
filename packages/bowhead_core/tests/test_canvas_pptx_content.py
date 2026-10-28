"""Standard PPT downloads must follow the complete Canvas, not demo slide slots."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.presentation import Presentation as PresentationType

SCRIPT = Path(__file__).resolve().parents[3] / "packages/bowhead_skill/catalog/canvas-pptx/scripts/render_from_canvas.py"
pytestmark = pytest.mark.integration


def _render(tmp_path: Path, sections: list[dict]) -> tuple[PresentationType, dict]:
    source = tmp_path / "input.json"
    source.write_text(json.dumps({"title": "Content coverage", "locale": "en-US", "sections": sections}), encoding="utf-8")
    result = subprocess.run(
        [sys.executable, str(SCRIPT), "--input", str(source), "--workdir", str(tmp_path)],
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    return Presentation(str(tmp_path / "output.pptx")), json.loads((tmp_path / "result.json").read_text())


def _texts(deck: PresentationType) -> list[str]:
    return [
        "\n".join(
            shape.text if shape.has_text_frame else "\n".join(cell.text for row in shape.table.rows for cell in row.cells) if shape.has_table else ""
            for shape in slide.shapes
        )
        for slide in deck.slides
    ]


def test_standard_download_follows_section_order_beyond_seven_slides(tmp_path: Path) -> None:
    sections = [{"kind": "h1", "text": "Synthetic review"}]
    sections.extend({"kind": "narrative", "markdown": f"Finding {index:02d}: full evidence for this chapter."} for index in range(12))
    deck, result = _render(tmp_path, sections)
    texts = _texts(deck)
    positions = [next(i for i, text in enumerate(texts) if f"Finding {index:02d}:" in text) for index in range(12)]
    assert positions == sorted(positions)
    assert result["slides"] > 7


def test_standard_download_keeps_long_text_and_all_recommendations(tmp_path: Path) -> None:
    paragraphs = [f"Evidence {index:03d}: " + "Detailed supporting text. " * 5 for index in range(40)]
    sections = [
        {"kind": "narrative", "markdown": "\n\n".join(paragraphs)},
        {"kind": "recommendations", "items": [{"action": f"Action {i:02d}", "rationale": f"Reason {i:02d}"} for i in range(9)]},
    ]
    deck, _ = _render(tmp_path, sections)
    text = " ".join(" ".join(_texts(deck)).split())
    for paragraph in paragraphs:
        assert paragraph.strip() in text
    for i in range(9):
        assert f"Action {i:02d}" in text
        assert f"Reason {i:02d}" in text


def test_standard_download_keeps_multiple_kpi_rows_and_table_cells(tmp_path: Path) -> None:
    sections = [
        {"kind": "kpi_row", "cells": [{"lbl": f"Metric {i}", "val": str(101 + i)} for i in range(7)]},
        {"kind": "kpi_row", "cells": [{"lbl": "Later metric", "val": "999"}]},
        {
            "kind": "data_table",
            "columns": [{"key": "name", "header": "Region"}, {"key": "amount", "header": "Revenue"}],
            "rows": [{"name": f"Region {i:02d}", "amount": 200 + i} for i in range(25)],
        },
    ]
    deck, _ = _render(tmp_path, sections)
    text = "\n".join(_texts(deck))
    for i in range(7):
        assert f"Metric {i}" in text
        assert str(101 + i) in text
    assert "Later metric" in text
    for i in range(25):
        assert f"Region {i:02d}" in text
        assert str(200 + i) in text


def test_standard_download_keeps_every_chart_and_series(tmp_path: Path) -> None:
    sections = [
        {
            "kind": "chart",
            "title": f"Chart {i}",
            "spec": {
                "kind": "bar",
                "columns": [{"key": "region", "type": "categorical"}, {"key": "actual", "type": "numeric"}, {"key": "plan", "type": "numeric"}],
                "rows": [{"region": "East", "actual": 100 + i, "plan": 200 + i}],
            },
        }
        for i in range(5)
    ]
    deck, _ = _render(tmp_path, sections)
    charts = [shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart]
    assert len(charts) == 5
    assert [list(chart.series[1].values) for chart in charts] == [[200 + i] for i in range(5)]


def test_title_only_canvas_is_one_slide(tmp_path: Path) -> None:
    deck, result = _render(tmp_path, [{"kind": "h1", "text": "A short note"}])
    assert result["slides"] == 1
    assert "A short note" in _texts(deck)[0]


def test_nested_sections_keep_chapters_and_safe_source_labels(tmp_path: Path) -> None:
    from bowhead_core.canvas_export import _static_safe_legacy_sections

    sections = _static_safe_legacy_sections(
        [
            {
                "kind": "section",
                "label": "01",
                "h2": "Regional delivery",
                "body": [
                    {"kind": "narrative", "markdown": "Nested evidence is part of the report."},
                    {"kind": "sources", "entries": [{"artifact_id": "art_private", "label": "Private board plan.xlsx"}]},
                ],
            }
        ],
        locale="en-US",
    )
    deck, _ = _render(tmp_path, sections)
    text = "\n".join(_texts(deck))
    assert "Regional delivery" in text
    assert "Nested evidence is part of the report." in text
    assert "[1] Generated analysis data" in text
    assert "Private board plan" not in text
    assert "art_private" not in text


def test_wide_table_and_long_cells_are_paginated_without_loss(tmp_path: Path) -> None:
    row = {f"c{i}": f"column-{i}-value" for i in range(8)}
    row["c7"] = " ".join(f"token-{i:03d}" for i in range(100))
    deck, _ = _render(tmp_path, [{"kind": "data_table", "rows": [row], "totalRow": {"c0": "Total", "c1": 0}}])
    text = "\n".join(_texts(deck))
    for i in range(7):
        assert f"column-{i}-value" in text
    for i in range(100):
        assert f"token-{i:03d}" in text
    assert "Total" in text


def test_unsupported_chart_keeps_table_and_percentage_semantics(tmp_path: Path) -> None:
    deck, result = _render(
        tmp_path,
        [
            {
                "kind": "chart",
                "spec": {
                    "kind": "dual_axis",
                    "title": "Actual vs conversion",
                    "columns": [
                        {"key": "region", "header": "Region", "type": "categorical"},
                        {"key": "conversion", "header": "Conversion", "type": "numeric", "format": "percent"},
                    ],
                    "rows": [{"region": "East", "conversion": 0.275}],
                },
            }
        ],
    )
    text = "\n".join(_texts(deck))
    assert "exported as its data table" in text
    assert "27.5%" in text
    assert result["warnings"] == ["chart_exported_as_table"]


def test_chart_rows_after_fifty_and_missing_values_are_preserved(tmp_path: Path) -> None:
    deck, _ = _render(
        tmp_path,
        [
            {
                "kind": "chart",
                "spec": {
                    "kind": "bar",
                    "columns": [{"key": "category", "type": "categorical"}, {"key": "amount", "type": "numeric"}],
                    "rows": [{"category": f"Item {i}", "amount": None if i == 3 else i} for i in range(55)],
                },
            }
        ],
    )
    charts = [shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart]
    assert [str(category) for chart in charts for category in chart.plots[0].categories] == [f"Item {i}" for i in range(55)]
    assert [value for chart in charts for value in chart.series[0].values] == [None if i == 3 else i for i in range(55)]


def test_recommendations_do_not_split_short_items_across_pages(tmp_path: Path) -> None:
    items = [
        {"priority": "P1", "action": f"Action {i}", "rationale": f"Reason {i}", "owner": f"Owner {i}", "timeline": f"Month {i}"} for i in range(5)
    ]
    deck, _ = _render(tmp_path, [{"kind": "recommendations", "items": items}])
    for item in items:
        page = next(text for text in _texts(deck) if item["action"] in text)
        assert all(value in page for value in item.values())


async def test_skill_runtime_publishes_the_content_driven_deck(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from bowhead_skill.loader import load_skill
    from bowhead_skill.runtime import SkillRuntime

    monkeypatch.setenv("BOWHEAD_SKILL_WORKDIR_ROOT", str(tmp_path))
    spec = load_skill(SCRIPT.parent.parent)
    result = await SkillRuntime().run(
        spec,
        entry_script="render_from_canvas.py",
        input_data={"sections": [{"kind": "narrative", "markdown": f"Runtime section {i:02d}"} for i in range(10)]},
    )
    assert result.ok, result.stderr
    output = next(path for path in result.outputs if path.suffix == ".pptx")
    text = "\n".join(_texts(Presentation(str(output))))
    assert result.result_json["slides"] == 10
    assert "Runtime section 09" in text


@pytest.mark.parametrize("kind", ["bar", "dual_axis"])
def test_percentage_points_override_generic_percent_format(tmp_path: Path, kind: str) -> None:
    deck, _ = _render(
        tmp_path,
        [
            {
                "kind": "chart",
                "spec": {
                    "kind": kind,
                    "columns": [
                        {"key": "region", "type": "categorical"},
                        {"key": "rate", "type": "numeric", "format": "percent", "percent_scale": "points", "unit": "%"},
                    ],
                    "rows": [{"region": "East", "rate": 27.5}],
                },
            }
        ],
    )
    if kind == "bar":
        chart = next(shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart)
        assert chart.value_axis.tick_labels.number_format == '0.0"%"'
        assert list(chart.series[0].values) == [27.5]
    else:
        text = "\n".join(_texts(deck))
        assert "27.5%" in text
        assert "2750%" not in text


def test_long_table_headers_keep_percentage_formatting(tmp_path: Path) -> None:
    deck, _ = _render(
        tmp_path,
        [
            {
                "kind": "data_table",
                "columns": [{"key": "rate", "header": "Conversion by eligible customer cohort " * 12, "percent_scale": "fraction"}],
                "rows": [{"rate": 0.275}],
            }
        ],
    )
    assert "27.5%" in "\n".join(_texts(deck))


@pytest.mark.parametrize("nested", [True, False], ids=["nested-chapter", "flat-heading"])
@pytest.mark.parametrize(
    "child",
    [
        {"kind": "kpi_row", "cells": [{"lbl": "Revenue", "val": "100"}]},
        {"kind": "data_table", "caption": "Revenue", "rows": [{"region": "East", "value": 100}]},
        {
            "kind": "chart",
            "spec": {
                "kind": "bar",
                "title": "Revenue",
                "columns": [{"key": "region", "type": "categorical"}, {"key": "value", "type": "numeric"}],
                "rows": [{"region": "East", "value": 100}],
            },
        },
    ],
    ids=["kpi", "table", "chart"],
)
def test_visual_only_chapters_keep_scope_headings(tmp_path: Path, nested: bool, child: dict) -> None:
    heading = "North America only - USD millions"
    sections = [{"kind": "section", "label": "01", "h2": heading, "body": [child]}] if nested else [{"kind": "h2", "text": heading}, child]
    deck, _ = _render(tmp_path, sections)
    texts = _texts(deck)
    assert heading in "\n".join(texts)
    assert next(i for i, text in enumerate(texts) if heading in text) <= next(i for i, text in enumerate(texts) if "Revenue" in text)


@pytest.mark.parametrize(
    ("kind", "key", "value"),
    [
        ("code_block", "code", "result = revenue ** 2\nvalue = record.__dict__\n# Keep this comment"),
        ("sql_block", "sql", "SELECT `__metric` FROM `facts`"),
        ("narrative", "text", "Literal __metric and 2 ** 3"),
    ],
    ids=["python", "sql", "plain-text"],
)
def test_literal_content_is_not_parsed_as_markdown(tmp_path: Path, kind: str, key: str, value: str) -> None:
    deck, _ = _render(tmp_path, [{"kind": kind, key: value}])
    assert value in "\n".join(_texts(deck))


def test_table_cells_keep_literal_identifiers(tmp_path: Path) -> None:
    deck, _ = _render(tmp_path, [{"kind": "data_table", "rows": [{"key": "__metric", "expression": "2 ** 3"}]}])
    text = "\n".join(_texts(deck))
    assert "__metric" in text
    assert "2 ** 3" in text


def test_markdown_preserves_code_literals_and_unpaired_operators(tmp_path: Path) -> None:
    deck, _ = _render(
        tmp_path,
        [
            {
                "kind": "narrative",
                "markdown": (
                    "**Evidence** uses `__metric` and the expression 2 ** 3.\n\n"
                    "```python\nresult = revenue ** 2\n# Keep this comment\n```\n\n"
                    "```sql\nSELECT `__metric` FROM `facts`\n```"
                ),
            }
        ],
    )
    text = "\n".join(_texts(deck))
    assert "Evidence uses __metric and the expression 2 ** 3." in text
    assert "result = revenue ** 2\n# Keep this comment" in text
    assert "SELECT `__metric` FROM `facts`" in text
