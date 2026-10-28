"""Exercise Community document creation through the authorized tool and Artifact flow."""

from __future__ import annotations

import hashlib
import io
import json
import sys
from collections.abc import AsyncIterator
from pathlib import Path
from types import SimpleNamespace
from typing import Any
from unittest.mock import AsyncMock

import pytest
from bowhead_core.artifact_store import ArtifactStore, LocalBackend
from bowhead_core.deck_render_inputs import prepare_deck_render_inputs
from bowhead_core.events import TurnEventBus
from bowhead_core.subagent import SubagentProfile, SubagentRunner, _collect_canvas_export_deliverables
from bowhead_core.subagent_return_tool import register_subagent_return_tool
from bowhead_core.tool_registry import ToolExecutionContext, ToolRegistry, ToolSpec
from bowhead_router.types import LLMChunk, LLMRequest, ToolSchema
from bowhead_skill.runtime import SkillRuntime
from bowhead_skill.tools import build_skill_tools
from pptx import Presentation

from bowhead_core import document_runtime
from bowhead_router import AgentBinding, LLMResponse, MockAdapter, ProviderConfig, ProviderRegistry, Router, ToolCall
from bowhead_skill.catalog import SkillCatalog, builtin_catalog_dir

pytestmark = pytest.mark.integration


class DocumentMockAdapter(MockAdapter):
    """Keep scripted tool calls visible when the peer switches to streaming."""

    async def stream(self, req: LLMRequest) -> AsyncIterator[LLMChunk]:
        response = await self.invoke(req)
        for call in response.tool_calls:
            yield LLMChunk(tool_call_delta=call)
        yield LLMChunk(delta=response.content, finish_reason=response.finish_reason, usage=response.usage)


def context(**kwargs: Any) -> ToolExecutionContext:
    return ToolExecutionContext(
        datarc=None,
        workspace_id="w_documents",
        member_uid="member_documents",
        user_id="u_documents",
        session_id="s_documents",
        turn_id="t_documents",
        **kwargs,
    )


def request(kind: str) -> dict[str, Any]:
    return json.loads((builtin_catalog_dir() / f"bowra-{kind}" / "examples/create.json").read_text())


def invoke_tool() -> Any:
    catalog = SkillCatalog(roots=[builtin_catalog_dir() / f"bowra-{kind}" for kind in ("docx", "pdf", "pptx", "xlsx")])
    return next(tool for tool in build_skill_tools(catalog, SkillRuntime(python_bin=sys.executable)) if tool.name == "skill_invoke")


async def collect(invocations: list[Any], store: Any, **kwargs: Any) -> list[dict[str, Any]]:
    fields = {"session_id": "s_documents", "user_id": "u_documents", "turn_id": "t_documents"}
    fields.update(kwargs)
    return await _collect_canvas_export_deliverables(invocations, store, bowra_documents=True, **fields)


async def test_real_four_format_creation_becomes_owned_current_turn_artifacts(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("BOWHEAD_SKILL_WORKDIR_ROOT", str(tmp_path))
    store = ArtifactStore(backend=LocalBackend(root_dir=tmp_path / "artifacts"), db_path=tmp_path / "artifacts.db")
    await store.initialize()
    tool = invoke_tool()
    invocations = []
    for kind in ("docx", "pdf", "pptx", "xlsx"):
        args = {"name": f"bowra-{kind}", "script": "run.py", "input": request(kind)}
        result = await tool.handler(args, context())
        assert result["ok"], result
        assert result["executed_document_skill"] == {"name": f"bowra-{kind}", "origin": "catalog", "script": "run.py"}
        invocations.append(SimpleNamespace(tool_name="skill_invoke", arguments=args, result=result))
    delivered = await collect(invocations, store)
    assert {item["kind"] for item in delivered} == {"docx", "pdf", "pptx", "xlsx"}
    rows = await store.list_deliverables_for_turn("t_documents")
    assert {row.id for row in rows} == {item["artifact_id"] for item in delivered}
    assert all(row.created_by == "u_documents" and row.kind == "deliverable" for row in rows)
    assert all(row.source == "community_document_skill" for row in rows)
    assert await store.list_deliverables_for_turn("t_previous") == []
    assert await store.get(rows[0].id, check_created_by="u_other") is None
    assert await store.get(rows[0].id, check_created_by="u_documents") is not None
    deck = next(item for item in delivered if item["kind"] == "pptx")
    assert deck["slide_count"] == 2
    source = Path(invocations[2].result["workdir"]) / "deck.pptx"
    presentation = Presentation(io.BytesIO(source.read_bytes()))
    assert len(presentation.slides) == 2
    assert presentation.slides[0].shapes[0].text == "Pilot results"


def invocation(tmp_path: Path) -> Any:
    path = tmp_path / "deck.pptx"
    content = b"bounded output fixture"
    path.write_bytes(content)
    return SimpleNamespace(
        tool_name="skill_invoke",
        arguments={"name": "bowra-pptx", "script": "run.py", "input": {"operation": "create"}},
        result={
            "ok": True,
            "workdir": str(tmp_path),
            "outputs": [str(path)],
            "executed_document_skill": {"name": "bowra-pptx", "origin": "catalog", "script": "run.py"},
            "result_json": {
                "operation": "create",
                "path": "deck.pptx",
                "bytes": len(content),
                "sha256": hashlib.sha256(content).hexdigest(),
                "slides": 1,
            },
        },
    )


@pytest.mark.parametrize(
    "case",
    [
        "failed",
        "inspect",
        "wrong_skill",
        "wrong_script",
        "unproven",
        "uploaded",
        "outside",
        "symlink",
        "wrong_digest",
        "wrong_size",
        "missing_owner",
        "missing_session",
        "missing_turn",
    ],
)
async def test_unproven_or_unowned_output_never_registers(case: str, tmp_path: Path) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    inv = invocation(workdir)
    extra = {}
    if case == "failed":
        inv.result["ok"] = False
    elif case == "inspect":
        inv.arguments["input"]["operation"] = "inspect"
    elif case == "wrong_skill":
        inv.arguments["name"] = "custom-pptx"
    elif case == "wrong_script":
        inv.arguments["script"] = "other.py"
    elif case == "unproven":
        inv.result.pop("executed_document_skill")
    elif case == "uploaded":
        inv.result["executed_document_skill"]["origin"] = "user"
    elif case in {"outside", "symlink"}:
        outside = tmp_path / "deck.pptx"
        outside.write_bytes((workdir / "deck.pptx").read_bytes())
        if case == "outside":
            inv.result["outputs"] = [str(outside)]
        else:
            (workdir / "deck.pptx").unlink()
            (workdir / "deck.pptx").symlink_to(outside)
    elif case == "wrong_digest":
        inv.result["result_json"]["sha256"] = "0" * 64
    elif case == "wrong_size":
        inv.result["result_json"]["bytes"] += 1
    elif case.startswith("missing_"):
        extra[{"missing_owner": "user_id", "missing_session": "session_id", "missing_turn": "turn_id"}[case]] = None
    store = SimpleNamespace(put_bytes=AsyncMock())
    assert await collect([inv], store, **extra) == []
    store.put_bytes.assert_not_awaited()


async def test_template_limitation_precedes_materialization_and_artifact_writes(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(document_runtime, "read_document_profile", lambda: "bowra-v1")
    uploaded = SimpleNamespace(id="art_template", filename="company.pptx", kind="upload", source="upload")
    store = SimpleNamespace(list_for_session=AsyncMock(return_value=[uploaded]), put_bytes=AsyncMock(), materialize_to_temp=AsyncMock())
    canvas = SimpleNamespace(get_canvas_by_session=AsyncMock())
    block = await prepare_deck_render_inputs(
        canvas_store=canvas, artifact_store=store, session_id="s_documents", user_id="u_documents", mode="consultant"
    )
    assert "needs_user_input" in block
    assert "do not generate any file" in block
    store.put_bytes.assert_not_awaited()
    store.materialize_to_temp.assert_not_awaited()
    canvas.get_canvas_by_session.assert_not_awaited()


async def test_fresh_deck_guidance_uses_real_contract_without_canvas(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(document_runtime, "read_document_profile", lambda: "bowra-v1")
    store = SimpleNamespace(list_for_session=AsyncMock(return_value=[]), put_bytes=AsyncMock())
    block = await prepare_deck_render_inputs(
        canvas_store=None, artifact_store=store, session_id="s_documents", user_id="u_documents", mode="consultant"
    )
    assert 'name="bowra-pptx",script="run.py"' in block
    assert '"operation":"create"' in block
    store.put_bytes.assert_not_awaited()


@pytest.mark.parametrize(
    ("produce", "document_profile", "report_claim"),
    [(True, "bowra-v1", False), (True, "bowra-v1", True), (False, "bowra-v1", True), (True, "legacy", False)],
)
async def test_runner_uses_real_output_and_discards_model_file_claims(
    produce: bool, document_profile: str, report_claim: bool, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import bowhead_core.subagent as module

    monkeypatch.setattr(module, "read_document_profile", lambda: document_profile)
    monkeypatch.setenv("BOWHEAD_SKILL_WORKDIR_ROOT", str(tmp_path))
    monkeypatch.setenv("BOWHEAD_SUBAGENT_PREFLIGHT", "0")
    tool = invoke_tool()
    monkeypatch.setitem(
        module.PROFILES,
        "consultant",
        SubagentProfile(
            name="consultant",
            llm_agent_type="analyst",
            tool_whitelist=("skill_invoke", "subagent_return"),
            system_prompt="Create verified documents.",
            skill_invoke_allowlist=("bowra-pptx",),
            deck_inputs=True,
        ),
    )
    registry = ToolRegistry()
    registry.register(ToolSpec(schema=ToolSchema(name=tool.name, description=tool.description, input_schema=tool.input_schema), handler=tool.handler))
    register_subagent_return_tool(registry)
    responses = []
    if produce:
        responses.append(
            LLMResponse(
                content="",
                tool_calls=[
                    ToolCall(id="create", name="skill_invoke", arguments={"name": "bowra-pptx", "script": "run.py", "input": request("pptx")})
                ],
            )
        )
    responses.append(
        LLMResponse(
            content="",
            tool_calls=[
                ToolCall(
                    id="return",
                    name="subagent_return",
                    arguments={
                        "payload": {
                            "title": "Pilot results",
                            "delivery_surface": "chat",
                            "exec_summary": "A basic text and table deck was requested for the pilot review.",
                            **({"deliverables": [{"kind": "pptx", "artifact_id": "art_model_claim"}]} if report_claim else {}),
                        }
                    },
                )
            ],
        )
    )
    adapter = DocumentMockAdapter(provider_name="documents")
    adapter.scripted_responses = responses
    providers = ProviderRegistry()
    providers.load([ProviderConfig(name="documents", type="mock")])
    providers._adapters["documents"] = adapter
    router = Router(providers)
    router.bind(AgentBinding(agent_type="analyst", primary=("documents", "mock")))
    store = SimpleNamespace(put_bytes=AsyncMock(return_value=SimpleNamespace(id="art_backend", filename="verified.pptx")))
    result = await SubagentRunner(router=router, tool_registry=registry).run(
        profile_name="consultant",
        task="Create a new basic pilot deck.",
        parent_turn_id="t_documents",
        parent_bus=TurnEventBus(),
        parent_ctx=context(artifact_store=store),
    )
    assert result.payload is not None, [(item.tool_name, item.result) for item in result.tool_invocations]
    if produce:
        assert result.payload["deliverables"] == [{"kind": "pptx", "artifact_id": "art_backend", "filename": "verified.pptx", "slide_count": 2}]
        store.put_bytes.assert_awaited_once()
        kwargs = store.put_bytes.call_args.kwargs
        assert kwargs["created_by_turn"] == "t_documents"
        assert len(Presentation(io.BytesIO(kwargs["content"])).slides) == 2
    else:
        assert not result.payload.get("deliverables")
        store.put_bytes.assert_not_awaited()


async def test_disabled_document_skill_does_not_execute() -> None:
    catalog = SkillCatalog(roots=[builtin_catalog_dir() / "bowra-pptx"])
    runtime = SimpleNamespace(run=AsyncMock())
    tool = next(item for item in build_skill_tools(catalog, runtime) if item.name == "skill_invoke")
    preferences = SimpleNamespace(get=AsyncMock(return_value=SimpleNamespace(source_type="catalog", enabled=False, import_status="ready")))
    result = await tool.handler({"name": "bowra-pptx", "script": "run.py", "input": request("pptx")}, context(skill_instance_store=preferences))
    assert "disabled" in result["error"]
    runtime.run.assert_not_awaited()


@pytest.mark.parametrize(
    "value", [{"schema_version": 1, "profile": "unknown"}, {"schema_version": 1, "profile": []}, {"schema_version": True, "profile": "legacy"}, {}]
)
def test_invalid_packaged_document_profile_fails_closed(value: dict[str, Any], tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(document_runtime, "__file__", str(tmp_path / "document_runtime.py"))
    (tmp_path / "document-runtime.json").write_text(json.dumps(value))
    document_runtime.read_document_profile.cache_clear()
    try:
        with pytest.raises(RuntimeError, match="Invalid packaged"):
            document_runtime.read_document_profile()
    finally:
        document_runtime.read_document_profile.cache_clear()
