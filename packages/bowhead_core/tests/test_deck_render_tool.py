"""Real rendering, Artifact storage and bounded retry for new presentations."""

from __future__ import annotations

import io
from dataclasses import replace
from pathlib import Path

import pytest
from bowhead_core.agents._lead_context import reset_current_peer_name, set_current_peer_name
from bowhead_core.artifact_store import ArtifactStore, LocalBackend
from bowhead_core.deck_render_inputs import _materialize_default_template
from bowhead_core.deck_render_tool import DeckRenderBinding, deck_target, register_deck_render_tool
from bowhead_core.tool_registry import ToolExecutionContext, ToolRegistry
from bowhead_skill.runtime import SkillRuntime
from pptx import Presentation

CONTENT = {
    "sections": [
        {"kind": "eyebrow", "payload": {"text": "2026年10月第二周"}},
        {"kind": "h1", "payload": {"text": "服务台周报"}},
        {"kind": "insight_card", "payload": {"title": "本周进展", "bullets": ["完成20项需求", "关闭8个缺陷", "上线2项改进"]}},
        {"kind": "insight_card", "payload": {"title": "下周计划", "bullets": ["优化搜索", "完善帮助文档", "回访5位用户"]}},
    ],
    "expected_slides": 3,
    "style": {"accent": "2F6FCE", "ink": "16324F", "background": "FFFFFF"},
}


@pytest.fixture
async def deck_case(tmp_path):
    store = ArtifactStore(backend=LocalBackend(root_dir=tmp_path / "files"), db_path=tmp_path / "artifacts.db")
    from bowhead_core.session_store import SessionStore

    sessions = SessionStore(tmp_path / "artifacts.db")
    await sessions.initialize()
    await sessions.ensure_session(session_id="s_test", workspace_id="ws_test", user_id="u_test")
    await sessions.record_turn(turn_id="t_test", session_id="s_test", user_text="Create a new three-slide presentation.")
    await store.initialize()
    ctx = ToolExecutionContext(
        workspace_id="ws_test",
        member_uid="m_test",
        user_id="u_test",
        team_id="team_test",
        session_id="s_test",
        turn_id="t_test",
        role="analyst",
        datarc=None,
        artifact_store=store,
    )
    template = _materialize_default_template()
    assert template is not None
    ctx.deck_render_binding = DeckRenderBinding(deck_target(ctx), template)

    workdirs = []

    class CountingRuntime(SkillRuntime):
        calls = 0

        async def run(self, *args, **kwargs):
            self.calls += 1
            result = await super().run(*args, **kwargs)
            if result.workdir:
                workdirs.append(Path(result.workdir))
            return result

    runtime = CountingRuntime()
    registry = ToolRegistry()
    register_deck_render_tool(registry, runtime)
    token = set_current_peer_name("consultant")
    try:
        yield ctx, store, registry.get("render_deck").handler, runtime
    finally:
        reset_current_peer_name(token)
        import shutil

        directories = {Path(template).parent, *workdirs}
        if ctx.deck_render_binding is not None:
            directories.add(Path(ctx.deck_render_binding.template_path).parent)
        for directory in directories:
            shutil.rmtree(directory, ignore_errors=True)


async def test_confirmed_three_slides_are_editable_registered_and_replayable(deck_case):
    ctx, store, render, runtime = deck_case
    receipt = await render(CONTENT, ctx)
    assert receipt["ok"] and receipt["slide_count"] == 3, receipt
    art = await store.get(receipt["artifact_id"])
    assert art.created_by == ctx.user_id and art.created_by_turn == ctx.turn_id
    assert art.kind == "deliverable" and art.source == "consultant_deck"
    assert any(a.id == art.id for a in await store.list_for_session(session_id=ctx.session_id, user_id=ctx.user_id))
    deck = Presentation(io.BytesIO(await store.open_read(art)))
    assert len(deck.slides) == 3
    text = ["\n".join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in deck.slides]
    for index, values in enumerate(
        [
            ["服务台周报", "2026年10月第二周"],
            ["本周进展", "完成20项需求", "关闭8个缺陷", "上线2项改进"],
            ["下周计划", "优化搜索", "完善帮助文档", "回访5位用户"],
        ]
    ):
        assert all(value in text[index] for value in values)
    for slide, section in zip(list(deck.slides)[1:], CONTENT["sections"][2:], strict=True):
        paragraphs = [p for shape in slide.shapes if shape.has_text_frame for p in shape.text_frame.paragraphs if p._p.xpath("./a:pPr/a:buChar")]
        assert [p.text for p in paragraphs] == section["payload"]["bullets"]
        assert all(p._p.xpath("./a:pPr/a:buChar")[0].get("char") == "•" for p in paragraphs)
    assert all(str(slide.background.fill.fore_color.rgb) == "FFFFFF" for slide in deck.slides)
    colors = {
        str(run.font.color.rgb)
        for slide in deck.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.font.color.type == 1
    }
    assert {"16324F", "2F6FCE"} <= colors
    assert await render(CONTENT, replace(ctx)) == receipt
    assert runtime.calls == 1


async def test_failed_registration_retries_original_bytes_without_rerender(deck_case, monkeypatch):
    ctx, store, render, runtime = deck_case
    original = store.open_read
    reads = []

    async def fail_once(art):
        reads.append(art.id)
        if len(reads) == 1:
            raise OSError("storage temporarily unavailable")
        return await original(art)

    monkeypatch.setattr(store, "open_read", fail_once)
    failed = await render(CONTENT, ctx)
    assert not failed["ok"] and failed["error"] == "deck_registration_failed"
    assert not await store.list_for_session(session_id=ctx.session_id, user_id=ctx.user_id)
    assert await store.get(reads[0]) is None
    recovered = await render(CONTENT, ctx)
    assert recovered["ok"] and runtime.calls == 1
    assert reads == [recovered["artifact_id"], recovered["artifact_id"]]
    assert len(await store.list_for_session(session_id=ctx.session_id, user_id=ctx.user_id)) == 1


@pytest.mark.parametrize(
    "change",
    [
        {"turn_id": "t_other"},
        {"session_id": "s_other"},
        {"user_id": "u_other"},
        {"team_id": "team_other"},
        {"workspace_id": "ws_other"},
        {"deck_render_binding": None},
    ],
)
async def test_render_authority_stays_with_original_assignment(deck_case, change):
    ctx, store, render, runtime = deck_case
    result = await render(CONTENT, replace(ctx, **change))
    assert not result["ok"] and result["error"] == "deck_render_not_admitted"
    assert runtime.calls == 0
    assert not await store.list_for_session(session_id=ctx.session_id, user_id=ctx.user_id)


async def test_read_only_peer_cannot_use_new_deck_authority(deck_case):
    ctx, store, render, runtime = deck_case
    token = set_current_peer_name("compute")
    try:
        assert (await render(CONTENT, ctx))["error"] == "deck_render_not_admitted"
    finally:
        reset_current_peer_name(token)
    assert runtime.calls == 0


async def test_slide_count_mismatch_has_no_download_claim(deck_case):
    ctx, store, render, runtime = deck_case
    result = await render({**CONTENT, "expected_slides": 4}, ctx)
    assert not result["ok"] and result["error"] == "deck_validation_failed"
    assert "3" in result["message"] and "4" in result["message"]
    assert not await store.list_for_session(session_id=ctx.session_id, user_id=ctx.user_id)


async def test_model_cannot_supply_template_or_output_paths(deck_case):
    ctx, store, render, runtime = deck_case
    result = await render({**CONTENT, "template_path": "/tmp/another-user.pptx"}, ctx)
    assert not result["ok"] and runtime.calls == 0


@pytest.mark.parametrize("render_first", [True, False])
async def test_actual_consultant_loop_requires_verified_download_before_success(deck_case, monkeypatch, render_first):
    from bowhead_core.agents import AgentRegistry
    from bowhead_core.agents.agent_runner_adapter import manifest_to_profile
    from bowhead_core.deck_render_inputs import prepare_deck_render_inputs
    from bowhead_core.events import TurnEventBus
    from bowhead_core.subagent import PROFILES, SubagentRunner
    from bowhead_core.subagent_return_tool import register_subagent_return_tool
    from test_peer_empty_closeout import ScriptedRouter

    from bowhead_router import LLMResponse, ToolCall

    ctx, store, _render, runtime = deck_case
    ctx.deck_render_binding = None
    guidance = await prepare_deck_render_inputs(
        canvas_store=None, artifact_store=store, session_id=ctx.session_id, user_id=ctx.user_id, mode="consultant", context=ctx
    )
    assert "render_deck" in guidance and ctx.deck_render_binding is not None
    profile = manifest_to_profile(AgentRegistry.bootstrap().get("consultant"))
    monkeypatch.setitem(PROFILES, "consultant", profile)
    registry = ToolRegistry()
    register_deck_render_tool(registry, runtime)
    register_subagent_return_tool(registry)
    responses = [LLMResponse(content="", tool_calls=[ToolCall(id="render", name="render_deck", arguments=CONTENT)])] if render_first else []
    responses.append(
        LLMResponse(
            content="",
            tool_calls=[
                ToolCall(
                    id="return",
                    name="subagent_return",
                    arguments={
                        "payload": {
                            "delivery_surface": "chat",
                            "exec_summary": "The requested deck is ready.",
                        }
                    },
                )
            ],
        )
    )
    router = ScriptedRouter(responses)
    result = await SubagentRunner(router=router, tool_registry=registry).run(
        profile_name="consultant",
        task="Create the confirmed three-slide presentation.",
        extra_instructions=guidance,
        parent_ctx=ctx,
        parent_bus=TurnEventBus(),
        parent_turn_id=ctx.turn_id,
    )
    assert result.ok is render_first, [(inv.tool_name, inv.result, inv.error) for inv in result.tool_invocations]
    if render_first:
        assert len(router.requests) == 2 and runtime.calls == 1
        download = result.payload["deliverables"][0]
        assert download["artifact_id"] != "art_forged" and download["slide_count"] == 3
        assert (await store.get(download["artifact_id"])).created_by_turn == ctx.turn_id
    else:
        assert result.error == "deck_delivery_missing" and not result.final_answer
        assert runtime.calls == 0


@pytest.mark.parametrize("attempt_render", [False, True])
async def test_uploaded_deck_keeps_inspection_possible_until_rendering_starts(deck_case, monkeypatch, attempt_render):
    """With an uploaded deck the assignment may be a read-only inspection: authority is
    available but a file becomes mandatory only once the consultant attempts rendering."""
    from bowhead_core.agents import AgentRegistry
    from bowhead_core.agents.agent_runner_adapter import manifest_to_profile
    from bowhead_core.deck_render_inputs import prepare_deck_render_inputs
    from bowhead_core.deck_render_tool import MIME
    from bowhead_core.events import TurnEventBus
    from bowhead_core.subagent import PROFILES, SubagentRunner
    from bowhead_core.subagent_return_tool import register_subagent_return_tool
    from test_peer_empty_closeout import ScriptedRouter

    from bowhead_router import LLMResponse, ToolCall

    ctx, store, _render, runtime = deck_case
    await store.put_bytes(
        content=Path(_materialize_default_template()).read_bytes(),
        filename="brand.pptx",
        mime=MIME,
        source="upload",
        session_id=ctx.session_id,
        created_by=ctx.user_id,
    )
    ctx.deck_render_binding = None
    guidance = await prepare_deck_render_inputs(
        canvas_store=None, artifact_store=store, session_id=ctx.session_id, user_id=ctx.user_id, mode="consultant", context=ctx
    )
    assert "render_deck" in guidance
    assert isinstance(ctx.deck_render_binding, DeckRenderBinding) and ctx.deck_render_binding.required is False
    monkeypatch.setitem(PROFILES, "consultant", manifest_to_profile(AgentRegistry.bootstrap().get("consultant")))
    registry = ToolRegistry()
    register_deck_render_tool(registry, runtime)
    register_subagent_return_tool(registry)
    responses = (
        [LLMResponse(content="", tool_calls=[ToolCall(id="render", name="render_deck", arguments={"sections": []})])] if attempt_render else []
    )
    responses.append(
        LLMResponse(
            content="",
            tool_calls=[
                ToolCall(
                    id="return",
                    name="subagent_return",
                    arguments={
                        "payload": {"delivery_surface": "chat", "exec_summary": "The uploaded deck has three slides about the weekly report."}
                    },
                )
            ],
        )
    )
    result = await SubagentRunner(router=ScriptedRouter(responses), tool_registry=registry).run(
        profile_name="consultant",
        task="Summarize the uploaded presentation.",
        extra_instructions=guidance,
        parent_ctx=ctx,
        parent_bus=TurnEventBus(),
        parent_turn_id=ctx.turn_id,
    )
    if attempt_render:
        assert result.ok is False and result.error == "deck_delivery_missing"
    else:
        assert result.ok is True, [(inv.tool_name, inv.result, inv.error) for inv in result.tool_invocations]
        assert not (result.payload or {}).get("deliverables")
    assert runtime.calls == 0


@pytest.mark.parametrize("require_delivery", [False, True])
async def test_no_uploaded_deck_makes_the_file_mandatory_without_a_report_assertion(deck_case, require_delivery):
    from bowhead_core.deck_render_inputs import prepare_deck_render_inputs

    ctx, store, _render, _runtime = deck_case
    ctx.deck_render_binding = None
    await prepare_deck_render_inputs(
        canvas_store=None,
        artifact_store=store,
        session_id=ctx.session_id,
        user_id=ctx.user_id,
        mode="consultant",
        context=ctx,
        require_delivery=require_delivery,
    )
    assert isinstance(ctx.deck_render_binding, DeckRenderBinding) and ctx.deck_render_binding.required is True


async def test_parallel_identical_render_calls_reuse_one_file(deck_case):
    import asyncio

    ctx, store, render, runtime = deck_case
    first, second = await asyncio.gather(render(CONTENT, ctx), render(CONTENT, replace(ctx)))
    assert first["ok"] and first == second and runtime.calls == 1
    assert len(await store.list_for_session(session_id=ctx.session_id, user_id=ctx.user_id)) == 1


@pytest.mark.parametrize("missing_output", [False, True])
async def test_render_failure_or_missing_file_never_produces_a_download(deck_case, monkeypatch, missing_output):
    from bowhead_skill.runtime import RunResult

    ctx, store, render, runtime = deck_case

    async def fail(*args, **kwargs):
        return RunResult(
            skill="pptx",
            ok=missing_output,
            exit_code=0 if missing_output else 1,
            stdout="",
            stderr="",
            result_json={"ok": missing_output},
            outputs=[],
        )

    monkeypatch.setattr(runtime, "run", fail)
    result = await render(CONTENT, ctx)
    assert not result["ok"] and result["error"] == ("deck_output_missing" if missing_output else "deck_render_failed")
    assert not await store.list_for_session(session_id=ctx.session_id, user_id=ctx.user_id)


async def test_lost_registration_ack_reconciles_original_artifact(deck_case, monkeypatch):
    ctx, store, render, runtime = deck_case
    original = store.commit_staged_artifacts
    commits = []

    async def lose_ack(staged):
        commits.append(staged[0].artifact.id)
        await original(staged)
        raise OSError("commit acknowledgement lost")

    monkeypatch.setattr(store, "commit_staged_artifacts", lose_ack)
    failed = await render(CONTENT, ctx)
    assert failed["error"] == "deck_registration_failed"
    receipt = await render(CONTENT, ctx)
    assert receipt["ok"] and commits == [receipt["artifact_id"]] and runtime.calls == 1
    assert len(await store.list_for_session(session_id=ctx.session_id, user_id=ctx.user_id)) == 1


@pytest.mark.parametrize(
    "payload",
    [
        {"title": "Progress", "bullets": []},
        {"title": "Progress", "bullets": "one\ntwo"},
        {"title": "Progress", "bullets": ["one", " "]},
        {"title": "Progress", "bullets": ["one", 2]},
        {"title": "Progress", "bullets": ["one"], "narrative": "two"},
    ],
)
async def test_invalid_list_structure_is_rejected_before_rendering(deck_case, payload):
    ctx, store, render, runtime = deck_case
    result = await render({"sections": [{"kind": "h1", "payload": {"text": "Weekly"}}, {"kind": "insight_card", "payload": payload}]}, ctx)
    assert result["error"] == "invalid_deck_bullets"
    assert runtime.calls == 0
    assert not await store.list_for_session(session_id=ctx.session_id, user_id=ctx.user_id)


@pytest.mark.parametrize("words", ["标题与副标题左对齐，行距放宽", "Leave a placeholder for the photo credit", "请输入你的旅行目的地"])
async def test_ordinary_wording_is_never_mistaken_for_a_leftover_placeholder(deck_case, words):
    """116 t_eb5a1e40bc18: the word 副标题 inside a bullet failed three renders in a row."""
    ctx, store, render, _runtime = deck_case
    content = {
        **CONTENT,
        "sections": [*CONTENT["sections"][:2], {"kind": "insight_card", "payload": {"title": "海岸与留白", "bullets": [words, "配色：海青"]}}],
    }
    result = await render({**content, "expected_slides": 2}, ctx)
    assert result["ok"], result


def test_unused_layout_placeholders_are_removed_from_a_user_template(tmp_path):
    """A user template's layouts carry placeholders; leaving them empty shows "Click to add" prompts."""
    import importlib.util

    from pptx import Presentation as _Presentation

    template = tmp_path / "user-template.pptx"
    _Presentation().save(template)  # python-pptx default layouts include title/body placeholders
    script = Path(__file__).resolve().parents[2] / "bowhead_skill/catalog/pptx/scripts/render_deck.py"
    spec = importlib.util.spec_from_file_location("render_deck_script", script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    out = tmp_path / "out.pptx"
    result = module.render_deck(str(template), CONTENT["sections"], str(out), {})
    assert result["placeholder_leftovers"] == 0
    assert not any(shape.is_placeholder for slide in _Presentation(out).slides for shape in slide.shapes)
