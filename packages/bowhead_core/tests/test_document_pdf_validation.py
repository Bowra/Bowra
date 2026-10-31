"""PDF text placement respects ordinary text overflow without losing identity checks."""

from __future__ import annotations

import hashlib
import io

import pdfplumber
import pytest
from bowhead_core.documents.package_io import OfficePackageError
from bowhead_core.documents.render import validate_pdf
from pptx import Presentation
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Pt
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen.canvas import Canvas

pytestmark = pytest.mark.unit


def _source(*, width=150, height=18, horizontal=None, vertical=None, table=False, left=36, autofit=False, alignment=None, anchor=None):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(240), Pt(180)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    if table:
        shape = slide.shapes.add_table(1, 1, Pt(left), Pt(36), Pt(width), Pt(height))
        frame = shape.table.cell(0, 0).text_frame
    else:
        shape = slide.shapes.add_textbox(Pt(left), Pt(36), Pt(width), Pt(height))
        frame = shape.text_frame
    frame.text = "Quarter review"
    if alignment is not None:
        frame.paragraphs[0].alignment = alignment
    if anchor is not None:
        frame.vertical_anchor = anchor
    body = frame._txBody.bodyPr
    for name, value in (("horzOverflow", horizontal), ("vertOverflow", vertical)):
        if value is not None:
            body.set(name, value)
    if autofit:
        from pptx.enum.text import MSO_AUTO_SIZE

        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    result = io.BytesIO()
    deck.save(result)
    return result.getvalue()


def _pdf(*, text="Quarter review", left=36, top=36, clip=None, page_width=240):
    result = io.BytesIO()
    canvas = Canvas(result, pagesize=(page_width, 180))
    if clip is not None:
        width, height = clip
        path = canvas.beginPath()
        path.rect(left, 180 - top - height, width, height)
        canvas.clipPath(path, stroke=0)
    canvas.setFont("Helvetica", 18)
    canvas.drawString(left, 180 - top - 18, text)
    canvas.showPage()
    canvas.save()
    return result.getvalue()


@pytest.mark.parametrize("overflow", [None, "overflow"], ids=["implicit-overflow", "explicit-overflow"])
def test_visible_text_with_font_descent_outside_frame_is_valid(overflow):
    source, pdf = _source(vertical=overflow), _pdf()
    with pdfplumber.open(io.BytesIO(pdf)) as document:
        assert max(char["bottom"] for char in document.pages[0].chars) > 36 + 18 + 2
    receipt = validate_pdf(source, pdf)
    assert receipt["source_sha256"] == hashlib.sha256(source).hexdigest()
    assert receipt["pdf_sha256"] == hashlib.sha256(pdf).hexdigest()
    assert receipt["visual_review"] == "user_confirmation_required"


@pytest.mark.parametrize("overflow", [None, "overflow"], ids=["implicit-overflow", "explicit-overflow"])
def test_visible_horizontal_overflow_is_not_clipping(overflow):
    assert validate_pdf(_source(width=35, height=40, horizontal=overflow), _pdf())["page_count"] == 1


@pytest.mark.parametrize("overflow", [None, "overflow"], ids=["implicit-overflow", "explicit-overflow"])
@pytest.mark.parametrize("direction", ["left", "above"])
def test_visible_overflow_before_the_frame_is_valid(overflow, direction):
    if direction == "left":
        source = _source(left=120, width=50, height=40, horizontal=overflow, alignment=PP_ALIGN.RIGHT)
        pdf = _pdf(left=170 - stringWidth("Quarter review", "Helvetica", 18))
    else:
        source = _source(height=8, vertical=overflow, anchor=MSO_VERTICAL_ANCHOR.BOTTOM)
        pdf = _pdf(top=20)
    assert validate_pdf(source, pdf)["text_and_geometry"] == "passed"


@pytest.mark.parametrize("direction", ["left", "above"])
def test_explicit_clipping_before_the_frame_is_still_rejected(direction):
    if direction == "left":
        source = _source(left=120, width=50, height=40, horizontal="clip", alignment=PP_ALIGN.RIGHT)
        pdf = _pdf(left=170 - stringWidth("Quarter review", "Helvetica", 18))
    else:
        source = _source(height=8, vertical="clip", anchor=MSO_VERTICAL_ANCHOR.BOTTOM)
        pdf = _pdf(top=20)
    with pytest.raises(OfficePackageError, match="document_pdf_text_missing_or_clipped"):
        validate_pdf(source, pdf)


@pytest.mark.parametrize("direction", ["left", "above"])
def test_overflow_does_not_match_text_wholly_before_the_frame(direction):
    source, pdf = (_source(left=180, width=20), _pdf(left=0)) if direction == "left" else (_source(), _pdf(top=0))
    with pytest.raises(OfficePackageError, match="document_pdf_text_missing_or_clipped"):
        validate_pdf(source, pdf)


@pytest.mark.parametrize(
    "options,clip",
    [
        ({"vertical": "clip"}, (150, 18)),
        ({"vertical": "ellipsis"}, (150, 18)),
        ({"width": 35, "height": 40, "horizontal": "clip"}, (35, 40)),
        ({"vertical": "clip", "autofit": True}, (150, 18)),
    ],
    ids=["vertical-clip", "vertical-ellipsis", "horizontal-clip", "autofit-does-not-disable-clip"],
)
def test_extractable_text_cut_by_explicit_clip_is_rejected(options, clip):
    pdf = _pdf(clip=clip)
    with pdfplumber.open(io.BytesIO(pdf)) as document:
        assert document.pages[0].extract_text() == "Quarter review"
    with pytest.raises(OfficePackageError, match="document_pdf_text_missing_or_clipped"):
        validate_pdf(_source(**options), pdf)


def test_table_text_keeps_its_cell_bounds():
    with pytest.raises(OfficePackageError, match="document_pdf_text_missing_or_clipped"):
        validate_pdf(_source(width=35, height=40, table=True), _pdf())


@pytest.mark.parametrize(
    "end,pdf_position",
    [((1, 0), {"left": 40, "top": 85}), ((0, 1), {"left": 120, "top": 40}), ((1, 1), {"left": 120, "top": 85})],
    ids=["vertical", "horizontal", "rectangle"],
)
@pytest.mark.parametrize("merged", [True, False], ids=["merged", "unmerged-clipping"])
def test_table_text_uses_its_actual_merged_extent(end, pdf_position, merged):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(240), Pt(180)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Pt(36), Pt(36), Pt(150), Pt(80)).table
    origin = table.cell(0, 0)
    if merged:
        origin.merge(table.cell(*end))
    origin.text = "Merged"
    source = io.BytesIO()
    deck.save(source)
    pdf = _pdf(text="Merged", **pdf_position)

    if merged:
        assert validate_pdf(source.getvalue(), pdf)["text_and_geometry"] == "passed"
    else:
        with pytest.raises(OfficePackageError, match="document_pdf_text_missing_or_clipped"):
            validate_pdf(source.getvalue(), pdf)


@pytest.mark.parametrize(
    "pdf",
    [{"text": "Different content"}, {"left": 190, "top": 100}],
    ids=["missing-text", "same-text-at-another-position"],
)
def test_missing_or_moved_text_does_not_match_the_source_frame(pdf):
    with pytest.raises(OfficePackageError, match="document_pdf_text_missing_or_clipped"):
        validate_pdf(_source(), _pdf(**pdf))


def test_overflow_must_remain_inside_the_page():
    with pytest.raises(OfficePackageError, match="document_pdf_text_missing_or_clipped"):
        validate_pdf(_source(left=216, width=20, height=40), _pdf(left=216))


def test_page_dimensions_still_match_the_source():
    with pytest.raises(OfficePackageError, match="document_pdf_page_mismatch"):
        validate_pdf(_source(), _pdf(page_width=241))
