"""Real PPTX value updates and independent package verification; not visual QA."""

from __future__ import annotations

import hashlib
import io
import json
from xml.dom import minidom
from zipfile import ZipFile

import pytest
from bowhead_core.documents.package_io import OfficePackageError
from bowhead_core.documents.pptx_text import (
    P,
    TextEditScope,
    TextReplacement,
    apply_pptx_text_update,
    inventory_pptx_text,
    parse_text_replacements,
    verify_pptx_text_update,
)
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

pytestmark = pytest.mark.unit


def _deck(*, mixed=False, merge=False, chart=False) -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    title.name = "Quarter heading"
    paragraph = title.text_frame.paragraphs[0]
    for text in ("Quarter ", "one"):
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(24)
    if mixed:
        paragraph.runs[1].font.bold = True
    group = slide.shapes.add_group_shape()
    caption = group.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    caption.text = "Stable grouped caption"
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(8), Inches(2)).table
    for row, values in enumerate((("Region", "Open cases"), ("North", "42"), ("South", "17"))):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    if merge:
        table.cell(1, 0).merge(table.cell(2, 0))
    if chart:
        data = CategoryChartData()
        data.categories = ["North", "South"]
        data.add_series("Cases", [42, 17])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4), Inches(8), Inches(2), data)
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _parts(source: bytes) -> dict[str, bytes]:
    with ZipFile(io.BytesIO(source)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _rewrite(source: bytes, changes: dict[str, bytes | None]) -> bytes:
    parts = _parts(source)
    for name, value in changes.items():
        if value is None:
            del parts[name]
        else:
            parts[name] = value
    output = io.BytesIO()
    with ZipFile(output, "w") as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    return output.getvalue()


def _plan(source: bytes, values: dict[str, str]):
    inventory = inventory_pptx_text(source)
    replacements = tuple(TextReplacement(slot.slot_id, values[slot.text]) for slot in inventory.slots if slot.text in values)
    assert len(replacements) == len(values)
    return TextEditScope(inventory.source_sha256, tuple(item.slot_id for item in replacements)), replacements


def test_inventory_exposes_context_not_xml_and_follows_presentation_order():
    source = _deck()
    parts = _parts(source)
    presentation = parts["ppt/presentation.xml"]
    # python-pptx creates these first two slide relationships after the masters.
    root = minidom.parseString(presentation)  # noqa: S318 — locally generated test fixture
    ids = root.getElementsByTagNameNS(P, "sldIdLst")[0]
    ids.appendChild(ids.firstChild)
    source = _rewrite(source, {"ppt/presentation.xml": root.toxml(encoding="UTF-8")})
    inventory = inventory_pptx_text(source)
    assert inventory.slide_count == 2
    title = next(slot for slot in inventory.slots if slot.text == "Quarter one")
    assert title.slide == 2 and title.shape_name == "Quarter heading" and title.cell is None
    north_value = next(slot for slot in inventory.slots if slot.text == "42")
    assert north_value.slide == 1 and north_value.cell == (1, 1)
    assert all(slot.blocked_reason is None for slot in inventory.slots)


def test_text_group_and_table_update_preserves_every_other_part_and_run_skeleton():
    source = _deck(chart=True)
    scope, updates = _plan(source, {"Quarter one": "Quarter two", "Stable grouped caption": "Updated grouped caption", "42": "53", "17": "21"})
    candidate, receipt = apply_pptx_text_update(source, scope, updates)
    assert receipt.source_sha256 == hashlib.sha256(source).hexdigest()
    assert receipt.candidate_sha256 == hashlib.sha256(candidate).hexdigest()
    assert receipt.changed_parts == ("ppt/slides/slide1.xml", "ppt/slides/slide2.xml")
    assert receipt.visual_verdict == "not_run"
    before, after = _parts(source), _parts(candidate)
    for part in before:
        if part not in receipt.changed_parts:
            assert after[part] == before[part]
    reopened = Presentation(io.BytesIO(candidate))
    runs = reopened.slides[0].shapes[0].text_frame.paragraphs[0].runs
    assert [run.text for run in runs] == ["Quarter two", ""]
    assert [run.font.size for run in runs] == [Pt(24), Pt(24)]
    assert reopened.slides[1].shapes[0].table.cell(1, 1).text == "53"
    # A preserved chart is explicitly unsupported, not claimed to be refreshed.
    assert inventory_pptx_text(candidate).unsupported_objects[0].reason == "non_text_graphic"


def test_next_period_reuses_slot_identity_but_requires_new_exact_source_hash():
    source = _deck()
    scope, updates = _plan(source, {"Quarter one": "Quarter two", "42": "53"})
    first, _ = apply_pptx_text_update(source, scope, updates)
    next_scope, next_updates = _plan(first, {"Quarter two": "Quarter three", "53": "61"})
    assert scope.slot_ids == next_scope.slot_ids
    with pytest.raises(OfficePackageError, match="source_hash_mismatch"):
        apply_pptx_text_update(first, scope, next_updates)
    second, _ = apply_pptx_text_update(first, next_scope, next_updates)
    assert "61" in {slot.text for slot in inventory_pptx_text(second).slots}
    assert source != first != second


def test_noop_returns_identical_source_bytes():
    source = _deck()
    scope, updates = _plan(source, {"Quarter one": "Quarter one"})
    candidate, receipt = apply_pptx_text_update(source, scope, updates)
    assert candidate == source
    assert receipt.changed_parts == ()


@pytest.mark.parametrize("replacement", ["", "中文 & <value> 😀", "A" * 4000])
def test_supported_values_roundtrip_without_inserting_markup(replacement):
    source = _deck()
    scope, updates = _plan(source, {"42": replacement})
    candidate, _ = apply_pptx_text_update(source, scope, updates)
    assert next(slot.text for slot in inventory_pptx_text(candidate).slots if slot.slot_id == updates[0].slot_id) == replacement


@pytest.mark.parametrize("replacement", ["two\nlines", "\x00", "\ud800", "\uffff", "A" * 4001, " padded "])
def test_values_requiring_new_structure_or_unbounded_capacity_are_rejected(replacement):
    source = _deck()
    scope, updates = _plan(source, {"42": replacement})
    with pytest.raises(OfficePackageError, match="unsupported_replacement|whitespace_requires_format_change"):
        apply_pptx_text_update(source, scope, updates)


@pytest.mark.parametrize("case", ["missing", "extra", "duplicate", "unknown"])
def test_model_proposal_cannot_expand_or_skip_trusted_scope(case):
    source = _deck()
    scope, updates = _plan(source, {"42": "53", "17": "21"})
    if case == "missing":
        updates = updates[:1]
    elif case == "extra":
        updates += _plan(source, {"North": "Changed"})[1]
    elif case == "duplicate":
        updates += updates[:1]
    else:
        scope = TextEditScope(scope.source_sha256, ("ps_unknown",))
        updates = (TextReplacement("ps_unknown", "x"),)
    with pytest.raises(OfficePackageError, match="text_scope_mismatch|overlapping_text_update|unsupported_text_target"):
        apply_pptx_text_update(source, scope, updates)


@pytest.mark.parametrize(
    "mixed,merge,value,reason", [(True, False, "Quarter one", "mixed_run_styles"), (False, True, "42", "merged_table_unsupported")]
)
def test_unsupported_targets_are_visible_but_cannot_be_edited(mixed, merge, value, reason):
    source = _deck(mixed=mixed, merge=merge)
    assert next(slot.blocked_reason for slot in inventory_pptx_text(source).slots if slot.text == value) == reason
    scope, updates = _plan(source, {value: "Changed"})
    with pytest.raises(OfficePackageError, match="unsupported_text_target"):
        apply_pptx_text_update(source, scope, updates)


@pytest.mark.parametrize("mutation", ["geometry", "style", "neighbor", "theme", "relationship", "missing_value", "extra_part", "run_removal"])
def test_independent_verifier_detects_unapproved_changes(mutation):
    source = _deck()
    scope, updates = _plan(source, {"Quarter one": "Quarter two"})
    candidate, _ = apply_pptx_text_update(source, scope, updates)
    parts = _parts(candidate)
    part = "ppt/slides/slide1.xml"
    xml = parts[part]
    if mutation == "geometry":
        changes = {part: xml.replace(b'x="914400"', b'x="914401"', 1)}
    elif mutation == "style":
        changes = {part: xml.replace(b'sz="2400"', b'sz="2500"', 1)}
    elif mutation == "neighbor":
        changes = {part: xml.replace(b"Stable grouped caption", b"Unapproved change")}
    elif mutation == "theme":
        changes = {"ppt/theme/theme1.xml": parts["ppt/theme/theme1.xml"].replace(b"Calibri", b"Arial")}
    elif mutation == "relationship":
        name = "ppt/slides/_rels/slide1.xml.rels"
        changes = {name: parts[name].replace(b"slideLayout7.xml", b"slideLayout1.xml")}
    elif mutation == "missing_value":
        changes = {part: xml.replace(b"Quarter two", b"Wrong value")}
    elif mutation == "extra_part":
        changes = {"unexpected.xml": b"<extra/>"}
    else:
        changes = {part: xml.replace(b'<a:r><a:rPr sz="2400"/><a:t/></a:r>', b"")}
    mutated = _rewrite(candidate, changes)
    assert mutated != candidate and any(parts.get(name) != value for name, value in changes.items())
    with pytest.raises(OfficePackageError):
        verify_pptx_text_update(source, mutated, scope, updates)


def test_noop_does_not_authorize_run_flattening_in_another_changed_paragraph():
    source = _deck()
    scope, updates = _plan(source, {"Quarter one": "Quarter one", "Stable grouped caption": "Updated caption"})
    candidate, _ = apply_pptx_text_update(source, scope, updates)
    part = "ppt/slides/slide1.xml"
    xml = _parts(candidate)[part].replace(b">Quarter </a:t>", b">Quarter one</a:t>").replace(b">one</a:t>", b"></a:t>")
    with pytest.raises(OfficePackageError, match="unapproved_node_change"):
        verify_pptx_text_update(source, _rewrite(candidate, {part: xml}), scope, updates)


def test_strict_model_payload_allows_only_slot_ids_and_text():
    response = json.dumps({"updates": [{"slot_id": "ps_example", "text": "New content"}]})
    assert parse_text_replacements(response) == (TextReplacement("ps_example", "New content"),)


@pytest.mark.parametrize("kind", ["field", "hyperlink", "extension"])
def test_field_hyperlink_and_maintainer_extensions_cannot_be_flattened(kind):
    source = _deck()
    part = "ppt/slides/slide1.xml"
    xml = _parts(source)[part]
    if kind == "field":
        xml = xml.replace(b'<a:r><a:rPr sz="2400"/><a:t>one</a:t></a:r>', b'<a:fld id="test" type="datetime"><a:t>one</a:t></a:fld>')
    elif kind == "hyperlink":
        xml = xml.replace(b'<a:rPr sz="2400"/>', b'<a:rPr sz="2400"><a:hlinkClick r:id="rId1"/></a:rPr>', 1)
    else:
        xml = xml.replace(b"</p:sp>", b'<p:extLst><p:ext uri="urn:test:maintainer"/></p:extLst></p:sp>', 1)
    source = _rewrite(source, {part: xml})
    title = next(slot for slot in inventory_pptx_text(source).slots if slot.text == "Quarter one")
    assert title.blocked_reason is not None
    scope, updates = _plan(source, {"Quarter one": "Quarter two"})
    with pytest.raises(OfficePackageError, match="unsupported_text_target"):
        apply_pptx_text_update(source, scope, updates)


def test_duplicate_shape_identity_is_not_a_valid_slot_address():
    source = _deck()
    part = "ppt/slides/slide1.xml"
    root = minidom.parseString(_parts(source)[part])  # noqa: S318 — locally generated fixture
    properties = root.getElementsByTagNameNS(P, "cNvPr")
    properties[-1].setAttribute("id", properties[1].getAttribute("id"))
    source = _rewrite(source, {part: root.toxml(encoding="UTF-8")})
    with pytest.raises(OfficePackageError, match="pptx_shape_identity_invalid"):
        inventory_pptx_text(source)


def test_password_protected_editing_is_not_bypassed():
    source = _deck()
    part = "ppt/presentation.xml"
    xml = _parts(source)[part].replace(b"</p:presentation>", b'<p:modifyVerifier hashData="test"/></p:presentation>')
    with pytest.raises(OfficePackageError, match="pptx_protected_unsupported"):
        inventory_pptx_text(_rewrite(source, {part: xml}))


def test_opaque_choice_and_fallback_are_preserved_without_indexing_duplicate_ids():
    source = _deck()
    part = "ppt/slides/slide1.xml"
    root = minidom.parseString(_parts(source)[part])  # noqa: S318 — locally generated fixture
    tree = root.getElementsByTagNameNS(P, "spTree")[0]
    shape = root.getElementsByTagNameNS(P, "sp")[1].cloneNode(deep=True)
    alternate = root.createElementNS("http://schemas.openxmlformats.org/markup-compatibility/2006", "mc:AlternateContent")
    alternate.setAttribute("xmlns:mc", "http://schemas.openxmlformats.org/markup-compatibility/2006")
    for name in ("Choice", "Fallback"):
        branch = root.createElementNS("http://schemas.openxmlformats.org/markup-compatibility/2006", "mc:" + name)
        branch.appendChild(shape.cloneNode(deep=True))
        alternate.appendChild(branch)
    tree.appendChild(alternate)
    source = _rewrite(source, {part: root.toxml(encoding="UTF-8")})
    inventory = inventory_pptx_text(source)
    assert [item.reason for item in inventory.unsupported_objects] == ["alternate_content"]
    assert len([slot for slot in inventory.slots if slot.text == "Stable grouped caption"]) == 1
    scope, updates = _plan(source, {"Quarter one": "Quarter two"})
    candidate, _ = apply_pptx_text_update(source, scope, updates)
    result = minidom.parseString(_parts(candidate)[part])  # noqa: S318 — locally generated fixture
    assert result.getElementsByTagName("mc:AlternateContent")[0].toxml() == alternate.toxml()


@pytest.mark.parametrize(
    "response",
    [
        '{"updates":[],"part":"ppt/slides/slide1.xml"}',
        '{"updates":[],"updates":[]}',
        '{"updates":[{"slot_id":"a","text":"x","text":"y"}]}',
        '{"updates":[{"slot_id":"a","text":"x","skip_validation":true}]}',
        '{"updates":[{"slot_id":"a","text":42}]}',
        '{"updates":{}}',
        "[]",
        "not JSON",
    ],
)
def test_model_payload_rejects_extra_keys_duplicates_and_invalid_shapes(response):
    with pytest.raises(OfficePackageError, match="text_proposal_invalid"):
        parse_text_replacements(response)
