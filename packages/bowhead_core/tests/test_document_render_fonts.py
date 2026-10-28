"""Font fallback reporting names the slides that declare an unavailable face."""

from __future__ import annotations

import io

import pytest
from bowhead_core.documents import render
from pptx import Presentation
from pptx.util import Inches

pytestmark = pytest.mark.unit


def _deck() -> bytes:
    deck = Presentation()
    for face in ("Arial", "Missing Display Face", "Arial", "Missing Display Face"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "Quarter"
        run.font.name = face
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def test_font_fallbacks_list_the_slides_that_use_the_missing_face(monkeypatch: pytest.MonkeyPatch) -> None:
    def fake_binary(argv: list[str], **_kwargs: object) -> str:
        return "Arial\nCalibri\n" if argv[0] == "fc-list" else "Noto Sans"

    monkeypatch.setattr(render, "run_binary", fake_binary)
    fallbacks = render.font_fallbacks(_deck())
    missing = [item for item in fallbacks if item["declared_font"] == "Missing Display Face"]
    assert missing == [{"declared_font": "Missing Display Face", "environment_fallback": "Noto Sans", "slides": [2, 4]}]
