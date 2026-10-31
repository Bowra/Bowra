"""Source identity, ordinary Word editing and untouched-format preservation."""

from __future__ import annotations

import io
from pathlib import Path
from zipfile import ZipFile

import pytest
from bowhead_core.documents.package_io import OfficePackageError
from bowhead_core.documents.word_template import DocumentTemplate, W
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

pytestmark = pytest.mark.unit


def _bytes(document: Document) -> bytes:
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _parts(content: bytes) -> dict[str, bytes]:
    with ZipFile(io.BytesIO(content)) as archive:
        return {item.filename: archive.read(item) for item in archive.infolist()}


def _repack(source: bytes, changes: dict[str, bytes]) -> bytes:
    output = io.BytesIO()
    with ZipFile(io.BytesIO(source)) as original, ZipFile(output, "w") as destination:
        for item in original.infolist():
            destination.writestr(item, changes.get(item.filename, original.read(item)))
        destination.comment = original.comment
    return output.getvalue()


@pytest.fixture
def source() -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.header.paragraphs[0].text = "SYNTHETIC BRAND"
    footer = section.footer.paragraphs[0]
    footer.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    footer._p.append(field)
    paragraph = document.add_paragraph()
    paragraph.add_run("Ac").bold = True
    company = paragraph.add_run("me")
    company.italic = True
    company.font.size = Pt(15)
    company.font.color.rgb = RGBColor.from_string("204060")
    paragraph.add_run(" / 2030 Q4 / unchanged")
    document.add_paragraph("Same phrase")
    document.add_paragraph("Same phrase")
    document.add_paragraph("Keep this numbered item", style="List Number")
    table = document.add_table(rows=2, cols=2)
    table.style = "Light Shading Accent 1"
    table.cell(0, 0).text, table.cell(0, 1).text = "Metric", "Amount"
    table.cell(1, 0).text = "Revenue"
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_paragraph("Second section remains unchanged")
    return _bytes(document)


def _block(document: DocumentTemplate, text: str, *, occurrence: int = 0) -> dict:
    matches = [block for block in document.inspect(detail="content")["blocks"] if block["text"] == text]
    return matches[occurrence]


def test_no_edit_and_same_value_preserve_exact_source_bytes(source: bytes) -> None:
    document = DocumentTemplate.open(source)
    assert document.to_bytes() is source
    block = _block(document, "Same phrase")
    assert document.replace_text(block["id"], start=0, end=4, text="Same") is None
    assert document.to_bytes() == source
    assert document.changes == ()


def test_split_run_replacement_preserves_other_xml_and_selected_run_properties(source: bytes) -> None:
    document = DocumentTemplate.open(source)
    block = _block(document, "Acme / 2030 Q4 / unchanged")
    document.replace_text(block["id"], start=0, end=4, text="New & <Company>", style_run=block["runs"][1]["id"])
    before, after = _parts(source), _parts(document.to_bytes())
    assert before.keys() == after.keys()
    assert [part for part in before if before[part] != after[part]] == ["word/document.xml"]
    expected = (
        before["word/document.xml"].replace(b"<w:t>Ac</w:t>", b"<w:t></w:t>").replace(b"<w:t>me</w:t>", b"<w:t>New &amp; &lt;Company&gt;</w:t>")
    )
    assert after["word/document.xml"] == expected
    restored = Document(io.BytesIO(document.to_bytes())).paragraphs[0]
    assert restored.text == "New & <Company> / 2030 Q4 / unchanged"
    assert restored.runs[0].bold
    assert restored.runs[1].italic
    assert restored.runs[1].font.size.pt == 15
    assert str(restored.runs[1].font.color.rgb) == "204060"
    assert len(document.changes) == 1


@pytest.mark.parametrize("reverse", [False, True])
def test_multiple_edits_use_original_unicode_offsets_in_either_order(source: bytes, reverse: bool) -> None:
    document = DocumentTemplate.open(source)
    block = _block(document, "Acme / 2030 Q4 / unchanged")
    operations = [(0, 4, "Regional company"), (7, 14, "2031 Q1")]
    for start, end, text in operations[::-1] if reverse else operations:
        document.replace_text(block["id"], start=start, end=end, text=text)
    assert Document(io.BytesIO(document.to_bytes())).paragraphs[0].text == "Regional company / 2031 Q1 / unchanged"
    assert _block(document, "Acme / 2030 Q4 / unchanged")["id"] == block["id"]


def test_repeated_text_does_not_select_an_unrelated_paragraph(source: bytes) -> None:
    document = DocumentTemplate.open(source)
    block = _block(document, "Same phrase", occurrence=1)
    document.replace_text(block["id"], start=0, end=11, text="Only the second match")
    paragraphs = [paragraph.text for paragraph in Document(io.BytesIO(document.to_bytes())).paragraphs]
    assert paragraphs[1:3] == ["Same phrase", "Only the second match"]


def test_blank_fixed_table_cell_keeps_table_and_page_format(source: bytes) -> None:
    document = DocumentTemplate.open(source)
    table = document.inspect()["tables"][0]
    assert document.set_table_cell(table["id"], row=1, column=1, text="207.75") is not None
    candidate = document.to_bytes()
    restored = Document(io.BytesIO(candidate))
    assert restored.tables[0].cell(1, 1).text == "207.75"
    original = _parts(source)["word/document.xml"]
    edited = _parts(candidate)["word/document.xml"]
    root_before, root_after = etree.fromstring(original), etree.fromstring(edited)
    for name in ("tblPr", "tblGrid", "tcPr", "sectPr", "numPr"):
        assert [etree.tostring(node) for node in root_before.iter(W + name)] == [etree.tostring(node) for node in root_after.iter(W + name)]
    assert all(_parts(source)[part] == _parts(candidate)[part] for part in _parts(source) if part != "word/document.xml")


def test_shared_header_reports_inherited_section_scope_and_edits_one_part(source: bytes) -> None:
    document = DocumentTemplate.open(source)
    block = _block(document, "SYNTHETIC BRAND")
    story = next(story for story in document.inspect()["stories"] if story["part"] == block["part"])
    assert story["kind"] == "header" and story["sections"] == [1, 2]
    assert story["references"] == [{"section": 1, "type": "default", "inherited": False}, {"section": 2, "type": "default", "inherited": True}]
    document.replace_text(block["id"], start=0, end=15, text="REVIEWED BRAND")
    before, after = _parts(source), _parts(document.to_bytes())
    assert [part for part in before if before[part] != after[part]] == [block["part"]]
    assert Document(io.BytesIO(document.to_bytes())).sections[1].header.paragraphs[0].text == "REVIEWED BRAND"


def test_field_is_readable_but_not_rewritten_and_does_not_block_other_paragraphs(source: bytes) -> None:
    document = DocumentTemplate.open(source)
    field = _block(document, "Page 1")
    assert field["limitation"] == "word_field_unsupported" and not field["editable"]
    with pytest.raises(OfficePackageError, match="word_field_unsupported"):
        document.replace_text(field["id"], start=0, end=4, text="Page")
    block = _block(document, "Same phrase")
    document.replace_text(block["id"], start=0, end=4, text="Other")
    for part in _parts(source):
        if part.startswith("word/footer"):
            assert _parts(document.to_bytes())[part] == _parts(source)[part]


@pytest.mark.parametrize("kind", ["merged", "nested"])
def test_unsupported_table_cannot_be_edited_through_a_paragraph_selector(kind: str) -> None:
    source = Document()
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Keep this cell"
    if kind == "merged":
        table.cell(0, 0).merge(table.cell(0, 1))
    else:
        table.cell(1, 0).add_table(rows=1, cols=1)
    document = DocumentTemplate.open(_bytes(source))
    block = _block(document, "Keep this cell")
    code = f"word_{kind}_table_unsupported"
    with pytest.raises(OfficePackageError, match=code):
        document.replace_text(block["id"], start=0, end=4, text="Change")
    with pytest.raises(OfficePackageError, match=code):
        document.set_table_cell(block["table_id"], row=0, column=0, text="Change")


def test_target_and_style_run_from_another_source_are_not_interchangeable(source: bytes) -> None:
    old = DocumentTemplate.open(source)
    previous = _block(old, "Same phrase")
    old.replace_text(previous["id"], start=0, end=4, text="New")
    document = DocumentTemplate.open(old.to_bytes())
    with pytest.raises(OfficePackageError, match="word_target_not_found"):
        document.replace_text(previous["id"], start=0, end=4, text="Next")
    block = _block(document, "New phrase")
    with pytest.raises(OfficePackageError, match="word_style_run_invalid"):
        document.replace_text(block["id"], start=0, end=3, text="Next", style_run=previous["runs"][0]["id"])


def test_overlapping_edits_fail_before_mutating_the_existing_edit(source: bytes) -> None:
    document = DocumentTemplate.open(source)
    block = _block(document, "Same phrase")
    document.replace_text(block["id"], start=0, end=4, text="Edited")
    candidate = document.to_bytes()
    with pytest.raises(OfficePackageError, match="word_text_ranges_overlap"):
        document.replace_text(block["id"], start=2, end=8, text="Wrong range")
    assert document.to_bytes() == candidate
    assert len(document.changes) == 1


@pytest.mark.parametrize("value", ["a\x00b", "a\nb", "a\tb", "\ud800"])
def test_unrepresentable_single_line_text_does_not_change_the_template(source: bytes, value: str) -> None:
    document = DocumentTemplate.open(source)
    block = _block(document, "Same phrase")
    with pytest.raises(OfficePackageError, match="word_replacement_text_invalid"):
        document.replace_text(block["id"], start=0, end=4, text=value)
    assert document.to_bytes() == source


@pytest.mark.parametrize(
    "paragraph_xml", [b"<w:p/>", b"<w:p><w:r/></w:p>", b"<w:p><w:r><w:rPr><w:b/></w:rPr></w:r></w:p>", b"<w:p><w:r><w:t/></w:r></w:p>"]
)
def test_empty_paragraph_and_run_forms_are_fillable_without_losing_properties(source: bytes, paragraph_xml: bytes) -> None:
    parts = _parts(source)
    changed = parts["word/document.xml"].replace(b"<w:p><w:r><w:t>Same phrase</w:t></w:r></w:p>", paragraph_xml, 1)
    document = DocumentTemplate.open(_repack(source, {"word/document.xml": changed}))
    block = next(
        block
        for block in document.inspect(detail="content")["blocks"]
        if block["part"] == "word/document.xml" and block["text"] == "" and block["table_id"] is None
    )
    document.replace_text(block["id"], start=0, end=0, text=" Filled ")
    restored = Document(io.BytesIO(document.to_bytes()))
    assert restored.paragraphs[1].text == " Filled "
    assert restored.paragraphs[1].runs[0].bold == (True if b"<w:b/>" in paragraph_xml else None)


def test_new_content_uses_unicode_offsets_and_preserves_leading_trailing_spaces() -> None:
    source = Document()
    source.add_paragraph("区域📈合计")
    document = DocumentTemplate.open(_bytes(source))
    block = _block(document, "区域📈合计")
    document.replace_text(block["id"], start=2, end=3, text=" 207.75 ")
    assert Document(io.BytesIO(document.to_bytes())).paragraphs[0].text == "区域 207.75 合计"


def test_save_retains_source_and_refuses_source_aliases(source: bytes, tmp_path: Path) -> None:
    original, output = tmp_path / "source.docx", tmp_path / "candidate.docx"
    original.write_bytes(source)
    document = DocumentTemplate.open(original)
    block = _block(document, "Same phrase")
    document.replace_text(block["id"], start=0, end=4, text="Changed")
    document.save(output)
    assert original.read_bytes() == source
    assert Document(io.BytesIO(output.read_bytes())).paragraphs[1].text == "Changed phrase"
    alias = tmp_path / "alias.docx"
    alias.symlink_to(original)
    with pytest.raises(OfficePackageError, match="word_source_overwrite_forbidden"):
        document.save(alias)


def test_inspection_is_paginated_and_not_a_mutable_candidate_view(source: bytes) -> None:
    document = DocumentTemplate.open(source)
    first = document.inspect(limit=1)
    second = document.inspect(offset=first["next_offset"], limit=1)
    assert first["basis"] == "original_source"
    assert first["blocks"][0]["id"] != second["blocks"][0]["id"]
    assert "runs" not in first["blocks"][0]
    assert "runs" in document.inspect(detail="content")["blocks"][0]


def test_pretty_printed_xml_whitespace_is_not_part_of_visible_paragraph_text(source: bytes) -> None:
    original = _parts(source)["word/document.xml"]
    pretty = etree.tostring(etree.fromstring(original), encoding="UTF-8", xml_declaration=True, pretty_print=True)
    document = DocumentTemplate.open(_repack(source, {"word/document.xml": pretty}))
    block = _block(document, "Acme / 2030 Q4 / unchanged")
    document.replace_text(block["id"], start=0, end=4, text="Reviewed")
    edited = _parts(document.to_bytes())["word/document.xml"]
    assert edited == pretty.replace(b"<w:t>Ac</w:t>", b"<w:t>Reviewed</w:t>").replace(b"<w:t>me</w:t>", b"<w:t></w:t>")


@pytest.mark.parametrize("default_namespace", [False, True])
def test_empty_paragraph_fill_uses_the_original_namespace(source: bytes, default_namespace: bool) -> None:
    if default_namespace:
        xml = f'<document xmlns="{W[1:-1]}"><body><p/></body></document>'.encode()
    else:
        xml = f'<word:document xmlns:word="{W[1:-1]}"><word:body><word:p/></word:body></word:document>'.encode()
    document = DocumentTemplate.open(_repack(source, {"word/document.xml": xml}))
    block = document.inspect()["blocks"][0]
    document.replace_text(block["id"], start=0, end=0, text="Filled")
    assert Document(io.BytesIO(document.to_bytes())).paragraphs[0].text == "Filled"


def test_unqualified_story_encoding_is_read_only_while_other_stories_remain_editable(source: bytes) -> None:
    original = _parts(source)
    utf16 = etree.tostring(etree.fromstring(original["word/document.xml"]), encoding="UTF-16", xml_declaration=True)
    document = DocumentTemplate.open(_repack(source, {"word/document.xml": utf16}))
    body = _block(document, "Same phrase")
    assert body["limitation"] == "word_story_encoding_unsupported"
    with pytest.raises(OfficePackageError, match="word_story_encoding_unsupported"):
        document.replace_text(body["id"], start=0, end=4, text="New")
    header = _block(document, "SYNTHETIC BRAND")
    document.replace_text(header["id"], start=0, end=15, text="REVIEWED BRAND")
    assert _parts(document.to_bytes())["word/document.xml"] == utf16


def test_bookmark_and_proofing_markers_are_preserved_around_ordinary_text(source: bytes) -> None:
    original = _parts(source)["word/document.xml"]
    markers = b'<w:bookmarkStart w:id="42" w:name="Company"/><w:proofErr w:type="spellStart"/><w:r><w:t>Same phrase</w:t></w:r><w:proofErr w:type="spellEnd"/><w:bookmarkEnd w:id="42"/>'
    marked = original.replace(b"<w:r><w:t>Same phrase</w:t></w:r>", markers, 1)
    document = DocumentTemplate.open(_repack(source, {"word/document.xml": marked}))
    block = _block(document, "Same phrase")
    assert block["editable"]
    document.replace_text(block["id"], start=0, end=4, text="Reviewed")
    assert _parts(document.to_bytes())["word/document.xml"] == marked.replace(b"<w:t>Same phrase</w:t>", b"<w:t>Reviewed phrase</w:t>", 1)


@pytest.mark.parametrize("kind", ["revision", "control"])
def test_revision_and_control_wrappers_are_readable_without_becoming_edit_targets(source: bytes, kind: str) -> None:
    original = _parts(source)["word/document.xml"]
    run = b"<w:r><w:t>Same phrase</w:t></w:r>"
    replacement = (
        b'<w:ins w:id="1" w:author="Synthetic QA">' + run + b"</w:ins>"
        if kind == "revision"
        else b"<w:sdt><w:sdtContent>" + run + b"</w:sdtContent></w:sdt>"
    )
    document = DocumentTemplate.open(_repack(source, {"word/document.xml": original.replace(run, replacement, 1)}))
    block = _block(document, "Same phrase")
    code = "word_tracked_change_unsupported" if kind == "revision" else "word_content_control_unsupported"
    assert block["limitation"] == code
    with pytest.raises(OfficePackageError, match=code):
        document.replace_text(block["id"], start=0, end=4, text="New")
    editable = _block(document, "Same phrase", occurrence=1)
    assert editable["editable"]


@pytest.mark.parametrize("original_attribute", [b"", b" xml:space='default'"])
def test_replacement_edge_spaces_have_explicit_word_preservation(source: bytes, original_attribute: bytes) -> None:
    original = _parts(source)["word/document.xml"]
    xml = original.replace(b"<w:t>Same phrase</w:t>", b"<w:t" + original_attribute + b">Same phrase</w:t>", 1)
    document = DocumentTemplate.open(_repack(source, {"word/document.xml": xml}))
    block = _block(document, "Same phrase")
    document.replace_text(block["id"], start=0, end=11, text=" reviewed ")
    root = etree.fromstring(_parts(document.to_bytes())["word/document.xml"])
    replacement = next(node for node in root.iter(W + "t") if node.text == " reviewed ")
    assert replacement.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"


def test_custom_relationship_type_is_not_reported_as_a_native_word_header(source: bytes) -> None:
    relationships = _parts(source)["word/_rels/document.xml.rels"]
    changed = relationships.replace(b"http://schemas.openxmlformats.org/officeDocument/2006/relationships/header", b"urn:synthetic-opaque/header")
    assert changed != relationships
    with pytest.raises(OfficePackageError, match="word_story_reference_invalid"):
        DocumentTemplate.open(_repack(source, {"word/_rels/document.xml.rels": changed}))
