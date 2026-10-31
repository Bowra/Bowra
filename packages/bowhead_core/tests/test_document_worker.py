"""Fixed worker inputs, evidence fidelity and real-source render checks."""

from __future__ import annotations

import hashlib
import io
import json

import pytest
from bowhead_core.documents.evidence import prepare_evidence
from bowhead_core.documents.package_io import OfficePackageError
from bowhead_core.documents.worker import WORKER_PROFILE, execute
from test_document_pptx_plan import case as case

pytestmark = pytest.mark.unit


def _inputs(tmp_path, case, operation="inventory", **changes):
    inputs, outputs = tmp_path / "input", tmp_path / "output"
    inputs.mkdir(exist_ok=True)
    outputs.mkdir(exist_ok=True)
    entries = []
    for index, (name, role, content) in enumerate((("template.pptx", "template", case[0]), ("counts.json", "data", case[2][0].content))):
        (inputs / f"input-{index}").write_bytes(content)
        entries.append({"name": name, "role": role, "sha256": hashlib.sha256(content).hexdigest()})
    control = {
        "profile": WORKER_PROFILE,
        "operation": operation,
        "instruction": case[1],
        "inputs": entries,
        "response": None if operation == "inventory" else json.dumps(case[4]),
        "context_digest": None if operation == "inventory" else case[3].digest,
        **changes,
    }
    raw = json.dumps(control).encode()
    (inputs / "control.json").write_bytes(raw)
    return inputs, outputs, hashlib.sha256(raw).hexdigest()


def test_inventory_reproduces_planning_identity_without_rendering(tmp_path, case):
    args = _inputs(tmp_path, case)
    result = execute(*args)
    assert result["kind"] == "context"
    assert result["context"]["digest"] == case[3].digest
    assert list(args[1].iterdir()) == []


def test_v5_worker_inventory_and_question_use_the_requested_profile(tmp_path, case):
    from bowhead_core.documents.pptx_plan import CURRENT_PROFILE, prepare_planning_context

    context = prepare_planning_context(case[0], case[1], case[2], profile=CURRENT_PROFILE)
    result = execute(*_inputs(tmp_path, case, planner_profile=CURRENT_PROFILE))
    assert result["context"] == {"payload": context.payload, "digest": context.digest}
    question = {"action": "ask", "reason": "scope", "slots": ["s0"], "question": "Which heading?"}
    args = _inputs(tmp_path, case, "apply_verify", planner_profile=CURRENT_PROFILE, context_digest=context.digest, response=json.dumps(question))
    assert execute(*args)["kind"] == "question"
    assert list(args[1].iterdir()) == []


@pytest.mark.parametrize("profile", ["pptx-grounded-plan-v4", "pptx-grounded-plan-v5"])
def test_inventory_reader_requires_the_frozen_profile(tmp_path, case, profile):
    from bowhead_core.documents.executor import read_result

    inputs, output, digest = _inputs(tmp_path, case, planner_profile=profile)
    record = execute(inputs, output, digest)
    (output / "result.json").write_text(json.dumps(record))
    assert read_result(output, image_id="fixture", operation="inventory", context_digest=None, planner_profile=profile).record == record
    other = "pptx-grounded-plan-v5" if profile.endswith("v4") else "pptx-grounded-plan-v4"
    with pytest.raises(OfficePackageError, match="planning_context_changed"):
        read_result(output, image_id="fixture", operation="inventory", context_digest=None, planner_profile=other)


@pytest.mark.parametrize("payload", ["null", "[]", '"text"', "1"])
def test_inventory_reader_rejects_non_object_payload(tmp_path, payload):
    from bowhead_core.documents.executor import read_result

    record = {"profile": WORKER_PROFILE, "kind": "context", "context": {"payload": payload, "digest": "0" * 64}}
    (tmp_path / "result.json").write_text(json.dumps(record))
    with pytest.raises(OfficePackageError, match="document_worker_output_invalid"):
        read_result(tmp_path, image_id="fixture", operation="inventory", context_digest=None)


@pytest.mark.parametrize("profile", [None, "pptx-grounded-plan-v99", {}, []])
def test_worker_refuses_unrecognized_explicit_planner_without_output(tmp_path, case, profile):
    args = _inputs(tmp_path, case, planner_profile=profile)
    with pytest.raises(OfficePackageError, match="planning_profile_unsupported"):
        execute(*args)
    assert list(args[1].iterdir()) == []


@pytest.mark.parametrize("changed", ["control", "input", "context"])
def test_worker_rejects_changed_identity(tmp_path, case, changed):
    args = _inputs(tmp_path, case, "apply_verify", **({"context_digest": "0" * 64} if changed == "context" else {}))
    if changed == "control":
        (args[0] / "control.json").write_bytes(b"{}")
    if changed == "input":
        (args[0] / "input-1").write_bytes(b"{}")
    with pytest.raises(OfficePackageError):
        execute(*args)
    assert list(args[1].iterdir()) == []


def test_question_produces_no_files(tmp_path, case):
    question = {"action": "ask", "reason": "comparison", "slots": ["s0"], "question": "Which comparison period?"}
    args = _inputs(tmp_path, case, "apply_verify", response=json.dumps(question))
    result = execute(*args)
    assert result["kind"] == "question"
    assert result["review"]["semantic_status"] == "unapproved"
    assert list(args[1].iterdir()) == []


def test_csv_preserves_ids_and_exact_decimal_values():
    data = prepare_evidence("data.csv", b'id,value,negative,formatted\n001,0.10,-4,"1,234"\n')
    assert data.content == b'[["id","value","negative","formatted"],["001",0.10,-4,"1,234"]]'


def _workbook(value):
    from openpyxl import Workbook

    book = Workbook()
    book.active.append(["Metric", "Value"])
    book.active.append(["Cases", value])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def test_xlsx_values_are_named_by_sheet_without_recalculation():
    data = prepare_evidence("data.xlsx", _workbook(42))
    assert json.loads(data.content) == {"Sheet": [["Metric", "Value"], ["Cases", 42]]}


def test_xlsx_formulas_require_explicit_values():
    with pytest.raises(OfficePackageError, match="document_data_formula_unsupported"):
        prepare_evidence("data.xlsx", _workbook("=SUM(1,2)"))


@pytest.mark.parametrize("name,content", [("data.json", b" " * 128_001), ("data.csv", b"0," * 513), ("data.txt", b"x" * 128_001)])
def test_evidence_never_silently_truncates(name, content):
    with pytest.raises(OfficePackageError, match="document_data_limit"):
        prepare_evidence(name, content)


def test_bounded_binary_output_rejects_truncation():
    import sys

    from bowhead_core.documents.render import run_binary

    with pytest.raises(OfficePackageError, match="document_renderer_output_limit"):
        run_binary([sys.executable, "-c", "print('x' * 100)"], output_limit=10)


def test_font_inventory_output_does_not_include_diagnostic_warnings():
    import sys

    from bowhead_core.documents.render import run_binary

    assert (
        run_binary([sys.executable, "-c", "import sys; print('Font cache warning', file=sys.stderr); print('/fonts/font.ttf')"], merge_stderr=False)
        == "/fonts/font.ttf"
    )


@pytest.mark.parametrize("transform", ["flipH", "vertical"])
def test_pdf_geometry_rejects_unsupported_text_transforms(case, transform):
    from bowhead_core.documents.render import _text_boxes
    from pptx import Presentation

    deck = Presentation(io.BytesIO(case[0]))
    shape = next(shape for shape in deck.slides[0].shapes if shape.has_text_frame)
    if transform == "flipH":
        shape._element.spPr.xfrm.set("flipH", "1")
    else:
        shape.text_frame._txBody.bodyPr.set("vert", "vert")
    output = io.BytesIO()
    deck.save(output)
    with pytest.raises(OfficePackageError, match="document_layout_unsupported"):
        _text_boxes(output.getvalue())


@pytest.mark.parametrize("transform", ["group_rotation", "group_flip", "shape_rotation", "vertical"])
def test_pdf_geometry_preserves_text_positions_with_transformed_decoration(case, transform):
    from bowhead_core.documents.render import _text_boxes
    from pptx import Presentation
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.util import Inches

    deck = Presentation(io.BytesIO(case[0]))
    group = deck.slides[0].shapes.add_group_shape()
    nested = group.shapes.add_group_shape()
    decoration = nested.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
    decoration.text = " \t "
    if transform == "group_rotation":
        group.rotation = 270
    elif transform == "group_flip":
        group._element.grpSpPr.xfrm.set("flipH", "1")
    elif transform == "shape_rotation":
        decoration.rotation = 270
    else:
        decoration.text_frame._txBody.bodyPr.set("vert", "vert")
    output = io.BytesIO()
    deck.save(output)

    assert _text_boxes(output.getvalue()) == _text_boxes(case[0])


@pytest.mark.parametrize("transform", ["rotation", "flip"])
def test_pdf_geometry_checks_new_text_inside_transformed_groups(case, transform):
    from bowhead_core.documents.render import _text_boxes
    from pptx import Presentation
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.util import Inches

    deck = Presentation(io.BytesIO(case[0]))
    group = deck.slides[0].shapes.add_group_shape()
    nested = group.shapes.add_group_shape()
    shape = nested.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
    if transform == "rotation":
        group.rotation = 270
    else:
        group._element.grpSpPr.xfrm.set("flipV", "1")
    empty = io.BytesIO()
    deck.save(empty)
    assert _text_boxes(empty.getvalue()) == _text_boxes(case[0])

    shape.text = "Visible note"
    filled = io.BytesIO()
    deck.save(filled)
    with pytest.raises(OfficePackageError, match="document_layout_unsupported"):
        _text_boxes(filled.getvalue())


async def test_executor_reconciles_lost_dispatch_without_repeating_post(tmp_path, monkeypatch):
    import httpx2
    from bowhead_core.documents.executor import PROFILE, DocumentExecutor, WorkerInput

    run_id = "a" * 64
    calls = []

    def transport(request):
        calls.append(request.method)
        if request.method == "POST":
            (tmp_path / ("document-" + run_id) / "output" / "result.json").write_text(
                json.dumps({"profile": PROFILE, "kind": "error", "error_code": "draft_template_unsupported"})
            )
            raise httpx2.ReadTimeout("Lost response", request=request)
        raw = (tmp_path / ("document-" + run_id) / "input" / "control.json").read_bytes()
        return httpx2.Response(
            200, json={"run_id": run_id, "control_sha256": hashlib.sha256(raw).hexdigest(), "status": "completed", "image_id": "sha256:" + "b" * 64}
        )

    client_type = httpx2.AsyncClient
    monkeypatch.setattr(httpx2, "AsyncClient", lambda **kwargs: client_type(transport=httpx2.MockTransport(transport), **kwargs))
    result = await DocumentExecutor(root=tmp_path, url="http://executor", token="test").run(
        run_id=run_id,
        inputs=(WorkerInput("template.pptx", "template", b"unsupported"),),
        instruction="Use this template",
        operation="inventory",
    )
    assert result.record["error_code"] == "draft_template_unsupported"
    assert calls == ["POST", "GET"]
