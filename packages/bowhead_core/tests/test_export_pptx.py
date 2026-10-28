"""PPTX V2 clean-shell rendering and closed-file parity contracts."""

from __future__ import annotations

import copy
import hashlib
import io
import json
import shutil
import zipfile
from dataclasses import dataclass, replace
from pathlib import Path

import pytest
from bowhead_core.export_assets import (
    ResolvedSnapshotAsset,
    SnapshotAssetResolution,
)
from bowhead_core.export_ir import (
    ExportIR,
    ExportIRBlock,
    normalize_export_snapshot,
)
from bowhead_core.export_materializer import (
    BrowserRenderedVisual,
    MaterializedVisual,
    materialize_export_visuals,
)
from bowhead_core.export_native_palette import native_export_palette_from_tokens
from bowhead_core.export_pptx import (
    PptxExportError,
    PptxRenderResult,
    extract_pptx_evidence,
    render_export_pptx,
)
from bowhead_core.export_snapshot import ExportSnapshot
from bowhead_core.export_validation import validate_export_output
from bowhead_core.export_visual_identity import EXPORT_VISUAL_RENDERER_VERSION
from bowhead_core.print_theme import current_print_theme_snapshot
from bowhead_core.theme_validation import contrast_ratio
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from bowhead_core import export_pptx as export_pptx_module

_FIXTURE = Path(__file__).resolve().parents[3] / "eval/fixtures/canvas_export/surface_projection_v1.json"
_SOFFICE = shutil.which("soffice")
_PDFTOPPM = shutil.which("pdftoppm")


def test_pptx_blank_gate_ignores_slide_chrome_but_detects_body_content() -> None:
    slide = Image.new("RGB", (1280, 720), "#f6f1e7")
    draw = ImageDraw.Draw(slide)
    draw.rectangle((50, 40, 70, 90), fill="#c9674b")
    draw.text((1150, 690), "01 / 01", fill="#65726d")
    assert export_pptx_module._pptx_body_is_blank(slide)  # noqa: SLF001

    draw.rectangle((200, 250, 600, 450), fill="#173d32")
    assert not export_pptx_module._pptx_body_is_blank(slide)  # noqa: SLF001


@dataclass
class _FakeBrowserRenderer:
    name: str = "canvas-browser-visual-v1"
    version: str = EXPORT_VISUAL_RENDERER_VERSION

    async def render(
        self,
        block: ExportIRBlock,
        *,
        asset_paths: dict[str, Path],
    ) -> BrowserRenderedVisual:
        del asset_paths
        color = (
            int(
                hashlib.sha256(block.block_key.encode()).hexdigest()[:2],
                16,
            ),
            90,
            150,
            255,
        )
        image = Image.new("RGBA", (1600, 900), "#fffaf1")
        draw = ImageDraw.Draw(image)
        draw.line((150, 720, 1470, 720), fill="#173d32", width=5)
        draw.line((150, 120, 150, 720), fill="#173d32", width=5)
        draw.text(
            (150, 55),
            f"{block.source.section_kind} · APAC / EMEA",
            fill="#173d32",
        )
        for index, value in enumerate((420, 650, 510, 730)):
            left = 260 + index * 280
            draw.rounded_rectangle(
                (left, 720 - value, left + 135, 720),
                radius=18,
                fill=color if index % 2 == 0 else "#c9674b",
            )
            draw.text(
                (left + 35, 750),
                f"Q{index + 1}",
                fill="#173d32",
            )
        output = io.BytesIO()
        image.save(output, format="PNG")
        return BrowserRenderedVisual(
            content=output.getvalue(),
            mime="image/png",
            visual_input_sha256=str(block.visual_input_sha256),
        )


def _snapshot(profile: str) -> ExportSnapshot:
    fixture = json.loads(_FIXTURE.read_text(encoding="utf-8"))
    by_id = {section["id"]: section for section in fixture["sections"]}
    entries = copy.deepcopy(fixture["expected_entries"])
    sections = [copy.deepcopy(by_id[entry["section_id"]]) for entry in entries]
    section_by_id = {section["id"]: section for section in sections}
    section_by_id["sec_h1"]["payload"]["text"] = "季度经营复盘 · Quarterly Review"
    section_by_id["sec_deck"]["payload"]["markdown"] = "管理层摘要：收入保持增长，区域表现出现分化。"
    section_by_id["sec_narrative"]["payload"]["markdown"] = "本季度收入同比增长 12%，APAC 是主要增长来源。"
    sections[1]["source_refs"] = ["tl_pptx_source"]
    parity_contract = "full_content_v1" if profile == "full_report" else "summary_policy_v1"
    payload = {
        "request": {
            "format": "pptx",
            "profile": profile,
            "parity_contract": parity_contract,
            "client_surface_receipt_sha256": f"sha256:{'d' * 64}",
            "locale": "zh-CN",
            "timezone": "Asia/Singapore",
            "theme": {"kind": "default"},
            "print_theme": current_print_theme_snapshot(),
            "options": {"include_appendix": True},
        },
        "surface_projection": {
            "contract_version": 1,
            "token": fixture["surface_token"],
            "entries": entries,
            "exclusions": fixture["expected_exclusions"],
        },
        "sections": sections,
    }
    return ExportSnapshot(
        id=f"exs_pptx_{profile}",
        schema_version=1,
        canvas_id="cv_pptx",
        session_id="s_pptx",
        team_id="team_pptx",
        workspace_id="ws_pptx",
        canvas_owner="u_owner",
        requested_by="u_owner",
        fmt="pptx",
        profile=profile,
        title="PPTX 完整性测试",
        locale="zh-CN",
        timezone="Asia/Singapore",
        projection_token=f"sha256:{'a' * 64}",
        surface_contract_version=1,
        surface_token=fixture["surface_token"],
        parity_contract=parity_contract,
        payload=payload,
        payload_sha256=f"sha256:{'c' * 64}",
        size_bytes=4096,
        section_count=len(sections),
        created_at="2026-07-17T00:00:00+00:00",
        expires_at="2026-07-24T00:00:00+00:00",
    )


async def _render_fixture(
    tmp_path: Path,
    *,
    profile: str,
) -> tuple[
    ExportSnapshot,
    ExportIR,
    tuple[MaterializedVisual, ...],
    Path,
    PptxRenderResult,
]:
    attempt_dir = tmp_path / f"attempt-{profile}"
    image = _illustration_png((640, 480))
    background = _illustration_png((1080, 1080))
    assets_dir = attempt_dir / "assets"
    assets_dir.mkdir(parents=True)
    assets = (
        _stage_asset(
            assets_dir=assets_dir,
            artifact_id="art_image",
            content=image,
        ),
        _stage_asset(
            assets_dir=assets_dir,
            artifact_id="art_bg",
            content=background,
        ),
    )
    snapshot = _snapshot(profile)
    export_ir = normalize_export_snapshot(
        snapshot,
        asset_content_hashes={asset.artifact_id: asset.content_sha256 for asset in assets},
    )
    visuals = await materialize_export_visuals(
        export_ir=export_ir,
        asset_resolution=SnapshotAssetResolution(
            resolved=assets,
            warnings=(),
        ),
        attempt_dir=attempt_dir,
        browser_renderer=_FakeBrowserRenderer(),
    )
    result = render_export_pptx(
        print_theme=current_print_theme_snapshot(),
        export_ir=export_ir,
        materialized_visuals=visuals,
        attempt_dir=attempt_dir,
        profile=profile,  # type: ignore[arg-type]
        include_appendix=True,
    )
    return (
        snapshot,
        export_ir,
        visuals,
        attempt_dir / result.relative_path,
        result,
    )


def _isolated_export(
    *,
    section_id: str,
    payload_override: dict[str, object] | None = None,
) -> tuple[ExportSnapshot, ExportIR]:
    snapshot = _snapshot("full_report")
    payload = copy.deepcopy(snapshot.payload)
    sections = payload["sections"]
    surface = payload["surface_projection"]
    assert isinstance(sections, list)
    assert isinstance(surface, dict)
    entries = surface["entries"]
    assert isinstance(entries, list)
    section = next(item for item in sections if item["id"] == section_id)
    entry = next(item for item in entries if item["section_id"] == section_id)
    if payload_override is not None:
        section["payload"] = payload_override
    payload["sections"] = [section]
    surface["entries"] = [entry]
    surface["exclusions"] = []
    isolated_snapshot = replace(
        snapshot,
        payload=payload,
        section_count=1,
    )
    return (
        isolated_snapshot,
        normalize_export_snapshot(
            isolated_snapshot,
            asset_content_hashes={
                "art_image": f"sha256:{'1' * 64}",
                "art_bg": f"sha256:{'2' * 64}",
            },
        ),
    )


def _disable_slide_raster_gate(monkeypatch: pytest.MonkeyPatch) -> None:
    def _no_render_failures(**_: object) -> set[str]:
        return set()

    monkeypatch.setattr(
        export_pptx_module,
        "_render_all_slides",
        _no_render_failures,
    )


@pytest.mark.skipif(
    _SOFFICE is None or _PDFTOPPM is None,
    reason="LibreOffice and Poppler are PPTX V2 runtime gates",
)
async def test_pptx_full_closed_file_matches_all_kinds(
    tmp_path: Path,
) -> None:
    snapshot, export_ir, visuals, pptx_path, result = await _render_fixture(
        tmp_path,
        profile="full_report",
    )

    assert result.slide_count < 24
    assert result.excluded_section_count == 0
    with zipfile.ZipFile(pptx_path) as package:
        names = tuple(package.namelist())
    assert not [
        name
        for name in names
        if (
            name.startswith(
                (
                    "ppt/charts/",
                    "ppt/embeddings/",
                )
            )
            or "vbaProject" in name
            or "/oleObjects/" in name
        )
    ]

    evidence = extract_pptx_evidence(
        pptx_path=pptx_path,
        export_ir=export_ir,
        materialized_visuals=visuals,
        profile="full_report",
        include_appendix=True,
        validation_dir=tmp_path / "validation-full",
        soffice_path=Path(str(_SOFFICE)),
        pdftoppm_path=Path(str(_PDFTOPPM)),
    )
    manifest = validate_export_output(
        job_id="exp_pptx_full",
        snapshot=snapshot,
        export_ir=export_ir,
        target="pptx_full",
        include_appendix=True,
        renderer_name="canvas-clean-pptx-v2",
        renderer_version=EXPORT_VISUAL_RENDERER_VERSION,
        evidence=evidence,
    )

    assert evidence.parser_opened
    assert evidence.hard_check_failures == ()
    assert evidence.demo_residue_count == 0
    assert len(evidence.sections) == 22
    assert manifest.validation.result == "warnings"
    assert manifest.summary.rendered == 19
    assert manifest.summary.fallback == 3
    assert manifest.summary.excluded_by_policy == 0
    assert manifest.summary.semantic_parity_rate == 1.0
    assert manifest.summary.visual_identity_rate == 1.0
    assert manifest.summary.typed_data_parity_rate == 1.0
    appendix = {section.section_kind for section in manifest.sections if any(target.kind == "appendix" for target in section.targets)}
    assert appendix == {
        "data_table",
        "code_block",
    }

    forged = (
        replace(
            visuals[0],
            pixel_sha256=f"sha256:{'f' * 64}",
        ),
        *visuals[1:],
    )
    forged_evidence = extract_pptx_evidence(
        pptx_path=pptx_path,
        export_ir=export_ir,
        materialized_visuals=forged,
        profile="full_report",
        include_appendix=True,
        validation_dir=tmp_path / "validation-forged",
        soffice_path=Path(str(_SOFFICE)),
        pdftoppm_path=Path(str(_PDFTOPPM)),
    )
    forged_manifest = validate_export_output(
        job_id="exp_pptx_forged",
        snapshot=snapshot,
        export_ir=export_ir,
        target="pptx_full",
        include_appendix=True,
        renderer_name="canvas-clean-pptx-v2",
        renderer_version=EXPORT_VISUAL_RENDERER_VERSION,
        evidence=forged_evidence,
    )

    assert "pptx_visual_pixel_mismatch" in (forged_evidence.hard_check_failures)
    assert forged_manifest.validation.result == "rejected"
    assert forged_manifest.summary.visual_identity_rate is not None
    assert forged_manifest.summary.visual_identity_rate < 1.0


@pytest.mark.skipif(
    _SOFFICE is None or _PDFTOPPM is None,
    reason="LibreOffice and Poppler are PPTX V2 runtime gates",
)
async def test_pptx_summary_has_explicit_selection_and_appendix(
    tmp_path: Path,
) -> None:
    snapshot, export_ir, visuals, pptx_path, result = await _render_fixture(
        tmp_path,
        profile="executive_summary",
    )

    assert 5 <= result.core_slide_count <= 7
    assert result.excluded_section_count > 0
    evidence = extract_pptx_evidence(
        pptx_path=pptx_path,
        export_ir=export_ir,
        materialized_visuals=visuals,
        profile="executive_summary",
        include_appendix=True,
        validation_dir=tmp_path / "validation-summary",
        soffice_path=Path(str(_SOFFICE)),
        pdftoppm_path=Path(str(_PDFTOPPM)),
    )
    manifest = validate_export_output(
        job_id="exp_pptx_summary",
        snapshot=snapshot,
        export_ir=export_ir,
        target="pptx_summary",
        include_appendix=True,
        renderer_name="canvas-clean-pptx-v2",
        renderer_version=EXPORT_VISUAL_RENDERER_VERSION,
        evidence=evidence,
    )

    assert evidence.hard_check_failures == ()
    assert len(evidence.sections) == 22
    assert manifest.validation.result == "warnings"
    assert manifest.summary.accounted_rate == 1.0
    assert manifest.summary.excluded_by_policy > 0
    assert manifest.summary.semantic_parity_rate == 1.0
    assert manifest.summary.visual_identity_rate == 1.0
    assert manifest.summary.typed_data_parity_rate == 1.0
    excluded = [section for section in manifest.sections if section.disposition == "excluded_by_policy"]
    assert all(not section.targets for section in excluded)
    assert {section.reason_code for section in excluded} <= {
        "summary_detail_omitted",
        "summary_visual_limit",
    }
    appendix = {section.section_kind for section in manifest.sections if any(target.kind == "appendix" for target in section.targets)}
    assert appendix == {
        "data_table",
        "code_block",
        "sources",
    }


@pytest.mark.parametrize(
    ("cell_count", "expected_slide_count"),
    [
        (7, 2),
        (12, 2),
        (13, 3),
    ],
)
def test_pptx_kpi_paginates_without_closed_file_parity_loss(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    cell_count: int,
    expected_slide_count: int,
) -> None:
    cells = [
        {
            "label": f"KPI {index + 1}",
            "value": index + 1,
            "unit": "%",
        }
        for index in range(cell_count)
    ]
    snapshot, export_ir = _isolated_export(
        section_id="sec_kpi",
        payload_override={"cells": cells},
    )
    attempt_dir = tmp_path / f"attempt-kpi-{cell_count}"
    result = render_export_pptx(
        print_theme=current_print_theme_snapshot(),
        export_ir=export_ir,
        materialized_visuals=(),
        attempt_dir=attempt_dir,
        profile="full_report",
        include_appendix=True,
    )
    pptx_path = attempt_dir / result.relative_path

    closed = Presentation(str(pptx_path))
    tagged_shapes = [shape for slide in closed.slides for shape in slide.shapes if shape.name.startswith("BH1B_")]
    assert result.slide_count == expected_slide_count
    assert len(closed.slides) == expected_slide_count
    assert len(tagged_shapes) == cell_count

    _disable_slide_raster_gate(monkeypatch)
    evidence = extract_pptx_evidence(
        pptx_path=pptx_path,
        export_ir=export_ir,
        materialized_visuals=(),
        profile="full_report",
        include_appendix=True,
        validation_dir=tmp_path / f"validation-kpi-{cell_count}",
        soffice_path=tmp_path / "unused-soffice",
        pdftoppm_path=tmp_path / "unused-pdftoppm",
    )
    manifest = validate_export_output(
        job_id=f"exp_pptx_kpi_{cell_count}",
        snapshot=snapshot,
        export_ir=export_ir,
        target="pptx_full",
        include_appendix=True,
        renderer_name="canvas-clean-pptx-v2",
        renderer_version=EXPORT_VISUAL_RENDERER_VERSION,
        evidence=evidence,
    )

    assert evidence.hard_check_failures == ()
    assert len(evidence.sections) == 1
    assert len(evidence.sections[0].targets) == expected_slide_count
    assert evidence.sections[0].typed_data_sha256 == export_ir.blocks[0].typed_data_sha256
    assert manifest.validation.result == "clean"
    assert manifest.summary.accounted_rate == 1.0
    assert manifest.summary.typed_data_parity_rate == 1.0


def test_pptx_large_table_slide_budget_fails_before_presentation_creation(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _, export_ir = _isolated_export(
        section_id="sec_data",
        payload_override={
            "caption": "Large table",
            "columns": [{"key": "value", "label": "Value"}],
            "rows": [{"value": index} for index in range(6_001)],
        },
    )

    def _presentation_must_not_be_created(
        *args: object,
        **kwargs: object,
    ) -> None:
        del args, kwargs
        raise AssertionError("Presentation must not be created over the slide budget")

    monkeypatch.setattr(
        export_pptx_module,
        "Presentation",
        _presentation_must_not_be_created,
    )

    with pytest.raises(
        PptxExportError,
        match="pptx_slide_count_invalid",
    ):
        render_export_pptx(
            print_theme=current_print_theme_snapshot(),
            export_ir=export_ir,
            materialized_visuals=(),
            attempt_dir=tmp_path / "attempt-large-table",
            profile="full_report",
            include_appendix=True,
        )


@pytest.mark.parametrize(
    ("mutation", "expected_failure"),
    [
        ("duplicate", "pptx_tagged_shape_duplicate"),
        ("gap", "pptx_tagged_shape_sequence_invalid"),
        ("extra", "pptx_tagged_shape_type_invalid"),
    ],
)
def test_pptx_rejects_invalid_tagged_shape_sequences(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    mutation: str,
    expected_failure: str,
) -> None:
    snapshot, export_ir = _isolated_export(
        section_id="sec_kpi",
        payload_override={"cells": [{"label": "Revenue", "value": 120}]},
    )
    attempt_dir = tmp_path / f"attempt-tagged-{mutation}"
    result = render_export_pptx(
        print_theme=current_print_theme_snapshot(),
        export_ir=export_ir,
        materialized_visuals=(),
        attempt_dir=attempt_dir,
        profile="full_report",
        include_appendix=True,
    )
    closed = Presentation(str(attempt_dir / result.relative_path))
    slide = closed.slides[0]
    tagged = next(shape for shape in slide.shapes if shape.name.startswith("BH1B_"))
    assert tagged.name.endswith("0000")
    if mutation == "duplicate":
        extra = slide.shapes.add_textbox(
            Inches(0.8),
            Inches(6.2),
            Inches(2.0),
            Inches(0.4),
        )
        extra.name = tagged.name
        extra.text = "duplicate"
    elif mutation == "gap":
        tagged.name = f"{tagged.name[:-4]}0001"
    else:
        extra = slide.shapes.add_textbox(
            Inches(0.8),
            Inches(6.2),
            Inches(2.0),
            Inches(0.4),
        )
        extra.name = f"{tagged.name[:-4]}0001"
        extra.text = "forged extra"
    mutated_path = tmp_path / f"mutated-{mutation}.pptx"
    closed.save(str(mutated_path))

    _disable_slide_raster_gate(monkeypatch)
    evidence = extract_pptx_evidence(
        pptx_path=mutated_path,
        export_ir=export_ir,
        materialized_visuals=(),
        profile="full_report",
        include_appendix=True,
        validation_dir=tmp_path / f"validation-tagged-{mutation}",
        soffice_path=tmp_path / "unused-soffice",
        pdftoppm_path=tmp_path / "unused-pdftoppm",
    )
    manifest = validate_export_output(
        job_id=f"exp_pptx_tagged_{mutation}",
        snapshot=snapshot,
        export_ir=export_ir,
        target="pptx_full",
        include_appendix=True,
        renderer_name="canvas-clean-pptx-v2",
        renderer_version=EXPORT_VISUAL_RENDERER_VERSION,
        evidence=evidence,
    )

    assert expected_failure in evidence.hard_check_failures
    assert manifest.validation.result == "rejected"


def test_pptx_text_capacity_rejects_the_previous_1400_character_budget() -> None:
    presentation = Presentation()
    presentation.slide_width = export_pptx_module.SLIDE_WIDTH
    presentation.slide_height = export_pptx_module.SLIDE_HEIGHT
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(
        Inches(0.85),
        Inches(1.3),
        Inches(11.55),
        Inches(5.25),
    )
    export_pptx_module._set_shape_text(  # noqa: SLF001
        shape,
        text="经营质量与风险整改" * 140,
        size=17,
        color=export_pptx_module._rgb("#061223"),  # noqa: SLF001
    )

    assert not export_pptx_module._shape_text_within_capacity(shape)  # noqa: SLF001


def test_pptx_chunker_output_always_satisfies_the_capacity_gate() -> None:
    # Every chunk the geometry-aware paginator emits must satisfy the same
    # line-budget the closed-file overflow check enforces, at the box the block
    # actually renders into -- otherwise a slide would be produced that the
    # validator rejects with pptx_text_overflow.
    paragraphs = [
        "本季度公司整体经营呈现稳健增长态势，收入端与利润端均实现同比改善，但区域之间的分化在持续加深，需要在下一阶段的资源配置中重点关注。"
        for _ in range(24)
    ]
    for fallback in (False, True):
        width_in, height_in = export_pptx_module._FALLBACK_CARD_BOX_IN[2:] if fallback else export_pptx_module._BODY_BOX_IN[2:]  # noqa: SLF001
        font_pt = 17.0 if not fallback else 19.0
        units_per_line, max_lines = export_pptx_module._box_line_budget(  # noqa: SLF001
            width_in=width_in,
            height_in=height_in,
            font_pt=font_pt,
        )
        chunks = export_pptx_module._chunk_tokens(  # noqa: SLF001
            list(paragraphs),
            units_per_line=units_per_line,
            max_lines=max_lines,
        )
        assert len(chunks) > 1
        for chunk in chunks:
            assert export_pptx_module._required_lines_for_text(chunk, units_per_line) <= max_lines  # noqa: SLF001
    # A single section that flows across many PDF pages arrives as one very long
    # token; it must still be split so each piece fits, without dropping content.
    long_token = "数据洞察" * 1200
    units_per_line, max_lines = export_pptx_module._box_line_budget(  # noqa: SLF001
        width_in=export_pptx_module._BODY_BOX_IN[2],  # noqa: SLF001
        height_in=export_pptx_module._BODY_BOX_IN[3],  # noqa: SLF001
        font_pt=17.0,
    )
    split = export_pptx_module._chunk_tokens(  # noqa: SLF001
        [long_token],
        units_per_line=units_per_line,
        max_lines=max_lines,
    )
    assert len(split) > 1
    assert all(export_pptx_module._required_lines_for_text(piece, units_per_line) <= max_lines for piece in split)  # noqa: SLF001
    rejoined = "".join(export_pptx_module._normalize_visible(piece) for piece in split)  # noqa: SLF001
    assert rejoined == export_pptx_module._normalize_visible(long_token)  # noqa: SLF001


def test_pptx_long_cjk_narrative_paginates_without_text_overflow(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Regression for pptx_text_overflow: a multi-thousand-character CJK narrative
    # (the kind that flows across pages in the PDF) must paginate across slides
    # at a readable font instead of overflowing one fixed text box.
    narrative = "\n\n".join(
        f"第{index + 1}节：本季度公司整体经营呈现稳健增长态势，"
        "收入端与利润端均实现同比改善，区域之间的分化在持续加深，"
        "需要在下一阶段的资源配置和风险整改中重点关注并给出明确的行动路径。"
        for index in range(18)
    )
    snapshot, export_ir = _isolated_export(
        section_id="sec_narrative",
        payload_override={"markdown": narrative},
    )
    attempt_dir = tmp_path / "attempt-long-narrative"
    result = render_export_pptx(
        print_theme=current_print_theme_snapshot(),
        export_ir=export_ir,
        materialized_visuals=(),
        attempt_dir=attempt_dir,
        profile="full_report",
        include_appendix=True,
    )
    # Rendering succeeding at all proves the planner's slide prediction matched
    # the paginated output (otherwise pptx_slide_count_invalid would raise).
    assert result.slide_count > 1

    closed = Presentation(str(attempt_dir / result.relative_path))
    business_shapes = [shape for slide in closed.slides for shape in slide.shapes if shape.name.startswith("BH1B_")]
    assert len(business_shapes) > 1
    for shape in business_shapes:
        assert export_pptx_module._shape_text_within_capacity(shape)  # noqa: SLF001

    _disable_slide_raster_gate(monkeypatch)
    evidence = extract_pptx_evidence(
        pptx_path=attempt_dir / result.relative_path,
        export_ir=export_ir,
        materialized_visuals=(),
        profile="full_report",
        include_appendix=True,
        validation_dir=tmp_path / "validation-long-narrative",
        soffice_path=tmp_path / "unused-soffice",
        pdftoppm_path=tmp_path / "unused-pdftoppm",
    )
    assert "pptx_text_overflow" not in evidence.hard_check_failures
    manifest = validate_export_output(
        job_id="exp_pptx_long_narrative",
        snapshot=snapshot,
        export_ir=export_ir,
        target="pptx_full",
        include_appendix=True,
        renderer_name="canvas-clean-pptx-v2",
        renderer_version=EXPORT_VISUAL_RENDERER_VERSION,
        evidence=evidence,
    )
    # The paginated narrative must still read back in order with full semantic parity.
    assert manifest.summary.semantic_parity_rate == 1.0


def test_pptx_cover_metadata_box_grows_to_fit_multiline_meta() -> None:
    # Regression for a real-report pptx_text_overflow: the cover stacks
    # fixed-height meta boxes; a 3-line meta block (更新于… / N 个数据系统 /
    # Bowhead 生成) overflowed the 0.45in box at 12pt and got the whole deck
    # rejected even though every parity rate was 1.0. The box must grow to fit.
    text = "更新于 2026-07-16\n1 个数据系统\nBowhead 生成"
    height = export_pptx_module._fit_box_height_in(  # noqa: SLF001
        text, width_in=10.9, font_pt=12, min_height_in=0.45
    )
    assert height > 0.45  # grew past the fixed design height to hold three lines

    presentation = Presentation()
    presentation.slide_width = export_pptx_module.SLIDE_WIDTH
    presentation.slide_height = export_pptx_module.SLIDE_HEIGHT
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.95), Inches(3.05), Inches(10.9), Inches(height))
    export_pptx_module._set_shape_text(  # noqa: SLF001
        box,
        text=text,
        size=12,
        color=export_pptx_module._rgb("#E5EFFF"),  # noqa: SLF001
    )
    assert export_pptx_module._shape_text_within_capacity(box)  # noqa: SLF001
    # A single-line meta block keeps the compact design height (no needless growth).
    assert (
        export_pptx_module._fit_box_height_in(  # noqa: SLF001
            "更新于 2026-07-16", width_in=10.9, font_pt=12, min_height_in=0.45
        )
        == 0.45
    )


def test_pptx_cover_keeps_only_leading_masthead_and_stays_on_slide() -> None:
    # Regression for pptx_shape_overflow on a 75-section report: every cover-kind
    # block (incl. later h1 chapter headings) piled onto the cover and stacked
    # off the 7.5in slide. Only the leading masthead run belongs on the cover;
    # chapter headings flow to the body, and the run is capped to the slide.
    from types import SimpleNamespace

    def blk(kind: str, key: str, tokens: list[str]) -> object:
        return SimpleNamespace(source=SimpleNamespace(section_kind=kind), block_key=key, semantic_tokens=tuple(tokens))

    def render_entry() -> object:
        return SimpleNamespace(action="render")

    blocks = (
        blk("eyebrow", "k0", ["季度经营复盘"]),
        blk("h1", "k1", ["完整报告标题"]),
        blk("narrative", "k2", ["正文段落。"]),  # ends the leading masthead run
        blk("h1", "k3", ["第二章：区域表现"]),  # chapter heading — must NOT be on the cover
    )
    plan = {b.block_key: render_entry() for b in blocks}
    cover = export_pptx_module._select_cover_blocks(  # noqa: SLF001
        blocks,  # type: ignore[arg-type]
        rendered=frozenset(b.block_key for b in blocks),  # type: ignore[attr-defined]
        plan_by_key=plan,  # type: ignore[arg-type]
    )
    assert {b.block_key for b in cover} == {"k0", "k1"}  # type: ignore[attr-defined]

    # A masthead run long enough to overflow is capped so no box leaves the slide.
    tall = tuple(blk("h1", f"m{index}", ["很长的章节标题占位" * 3]) for index in range(12))
    cover2 = export_pptx_module._select_cover_blocks(  # noqa: SLF001
        tall,  # type: ignore[arg-type]
        rendered=frozenset(b.block_key for b in tall),  # type: ignore[attr-defined]
        plan_by_key={b.block_key: render_entry() for b in tall},  # type: ignore[arg-type]
    )
    assert 0 < len(cover2) < 12  # capped, not every block
    top = export_pptx_module._COVER_TOP_IN  # noqa: SLF001
    for block in cover2:
        _, _, height = export_pptx_module._cover_block_layout(block)  # type: ignore[arg-type]  # noqa: SLF001
        assert top + height <= export_pptx_module._COVER_BOTTOM_IN  # noqa: SLF001  # bottom stays on slide
        top += height + export_pptx_module._COVER_GAP_IN  # noqa: SLF001


def test_pptx_fixed_layout_visual_title_autofits_to_not_overflow() -> None:
    # Regression for a real-report pptx_text_overflow: a long chart/visual title
    # ("品类贡献瀑布图：…") overflowed the fixed 0.8in title box at font 20 (needs
    # 2 lines, box holds 1). Fixed-layout title/card boxes must auto-fit the font.
    text = "品类贡献瀑布图：生鲜（-13.00pp）和日用（-12.65pp）是整体降幅的最大贡献者，合计贡献 25.65pp，占整体降幅的 60.6%"
    size = export_pptx_module._fit_font_pt(  # noqa: SLF001
        text, width_in=11.65, height_in=0.8, max_pt=20, min_pt=11
    )
    assert size < 20  # shrank from the design size to fit

    presentation = Presentation()
    presentation.slide_width = export_pptx_module.SLIDE_WIDTH
    presentation.slide_height = export_pptx_module.SLIDE_HEIGHT
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.85), Inches(0.82), Inches(11.65), Inches(0.8))
    export_pptx_module._set_shape_text(  # noqa: SLF001
        box,
        text=text,
        size=size,
        color=export_pptx_module._rgb("#061223"),  # noqa: SLF001
        bold=True,
    )
    assert export_pptx_module._shape_text_within_capacity(box)  # noqa: SLF001


async def test_pptx_rejects_tagged_shape_type_mismatch(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    snapshot, export_ir = _isolated_export(section_id="sec_chart")
    attempt_dir = tmp_path / "attempt-tagged-type"
    visuals = await materialize_export_visuals(
        export_ir=export_ir,
        asset_resolution=SnapshotAssetResolution(
            resolved=(),
            warnings=(),
        ),
        attempt_dir=attempt_dir,
        browser_renderer=_FakeBrowserRenderer(),
    )
    result = render_export_pptx(
        print_theme=current_print_theme_snapshot(),
        export_ir=export_ir,
        materialized_visuals=visuals,
        attempt_dir=attempt_dir,
        profile="full_report",
        include_appendix=True,
    )
    closed = Presentation(str(attempt_dir / result.relative_path))
    visual = next(shape for slide in closed.slides for shape in slide.shapes if shape.name.startswith("BH1V_"))
    visual_match = export_pptx_module._VISUAL_RE.fullmatch(  # noqa: SLF001
        visual.name
    )
    assert visual_match is not None
    marker, source_digest, _, part = visual_match.groups()
    visual.name = f"BH1B_{marker}_{source_digest}_{part}"
    mutated_path = tmp_path / "mutated-tagged-type.pptx"
    closed.save(str(mutated_path))

    _disable_slide_raster_gate(monkeypatch)
    evidence = extract_pptx_evidence(
        pptx_path=mutated_path,
        export_ir=export_ir,
        materialized_visuals=visuals,
        profile="full_report",
        include_appendix=True,
        validation_dir=tmp_path / "validation-tagged-type",
        soffice_path=tmp_path / "unused-soffice",
        pdftoppm_path=tmp_path / "unused-pdftoppm",
    )
    manifest = validate_export_output(
        job_id="exp_pptx_tagged_type",
        snapshot=snapshot,
        export_ir=export_ir,
        target="pptx_full",
        include_appendix=True,
        renderer_name="canvas-clean-pptx-v2",
        renderer_version=EXPORT_VISUAL_RENDERER_VERSION,
        evidence=evidence,
    )

    assert "pptx_tagged_shape_type_invalid" in evidence.hard_check_failures
    assert "pptx_tagged_shape_duplicate" not in evidence.hard_check_failures
    assert "pptx_tagged_shape_sequence_invalid" not in evidence.hard_check_failures
    assert manifest.validation.result == "rejected"


async def test_pptx_output_name_cannot_escape_attempt_directory(
    tmp_path: Path,
) -> None:
    export_ir = normalize_export_snapshot(
        _snapshot("full_report"),
        asset_content_hashes={
            "art_image": f"sha256:{'1' * 64}",
            "art_bg": f"sha256:{'2' * 64}",
        },
    )

    with pytest.raises(
        PptxExportError,
        match="pptx_staging_invalid",
    ):
        render_export_pptx(
            print_theme=current_print_theme_snapshot(),
            export_ir=export_ir,
            materialized_visuals=(),
            attempt_dir=tmp_path / "other-attempt",
            profile="full_report",
            include_appendix=True,
            output_name="../escaped.pptx",
        )


_PALE_ACCENT = {"accent": "#A9C8FF", "accentHover": "#C4DAFF", "accentInk": "#FFFFFF", "accentSoft": "#EEF4FF", "accentTint": "#F8FBFF"}
_DARK_ACCENT = {"accent": "#0B1A3C", "accentHover": "#050C1E", "accentInk": "#FFFFFF", "accentSoft": "#1B2A4A", "accentTint": "#E9EDF5"}


def _solid_fill(shape: object) -> str | None:
    fill = getattr(shape, "fill", None)
    if fill is None or fill.type != MSO_FILL.SOLID:
        return None
    return f"#{fill.fore_color.rgb}"


def _run_colors(text_frame: object) -> list[str]:
    return [f"#{run.font.color.rgb}" for paragraph in text_frame.paragraphs for run in paragraph.runs if run.text.strip()]  # type: ignore[attr-defined]


@pytest.mark.parametrize("accent", [None, _PALE_ACCENT, _DARK_ACCENT], ids=["product_default", "pale_accent", "dark_accent"])
async def test_rendered_slides_meet_contrast_floors(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    accent: dict[str, str] | None,
) -> None:
    # Walk the closed deck itself, not the palette: every visible run against the fill it sits
    # on (its own card/cell fill, else the slide background), and every bar/border against the
    # slide background. A pale and a dark accent stand in for future team brands.
    if accent is not None:
        tokens = {**current_print_theme_snapshot()["tokens"], **accent}
        monkeypatch.setattr(export_pptx_module, "native_export_palette", lambda _snapshot: native_export_palette_from_tokens(tokens))
    *_, pptx_path, _result = await _render_fixture(tmp_path, profile="full_report")

    failures: list[str] = []
    checked = 0
    for slide_number, slide in enumerate(Presentation(str(pptx_path)).slides, start=1):
        shapes = list(slide.shapes)
        slide_fill = _solid_fill(shapes[0])
        assert slide_fill is not None
        for shape in shapes[1:]:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                continue
            fill = _solid_fill(shape)
            if shape.has_text_frame:
                for color in _run_colors(shape.text_frame):
                    checked += 1
                    background = fill or slide_fill
                    if contrast_ratio(color, background) < 4.5:
                        failures.append(f"slide {slide_number} {shape.name}: text {color} on {background}")
                if fill is not None and not shape.text_frame.text.strip() and contrast_ratio(fill, slide_fill) < 3.0:
                    failures.append(f"slide {slide_number} {shape.name}: mark {fill} on {slide_fill}")
            if fill is not None and shape.line.fill.type == MSO_FILL.SOLID:
                border = f"#{shape.line.color.rgb}"
                checked += 1
                if contrast_ratio(border, slide_fill) < 3.0:
                    failures.append(f"slide {slide_number} {shape.name}: border {border} on {slide_fill}")
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        cell_fill = f"#{cell.fill.fore_color.rgb}"
                        for color in _run_colors(cell.text_frame):
                            checked += 1
                            if contrast_ratio(color, cell_fill) < 4.5:
                                failures.append(f"slide {slide_number} table: text {color} on {cell_fill}")
    assert checked > 50
    assert failures == []


@pytest.mark.parametrize(
    ("member_name", "member_content", "expected_failure"),
    [
        (
            "../escaped.xml",
            b"<root/>",
            "pptx_zip_path_invalid",
        ),
        (
            "ppt/embeddings/demo.xlsx",
            b"PK\x03\x04demo",
            "pptx_forbidden_package_member",
        ),
        (
            "ppt/slides/_rels/slide1.xml.rels",
            (
                b'<?xml version="1.0"?>'
                b'<Relationships xmlns="http://schemas.openxmlformats.org/'
                b'package/2006/relationships">'
                b'<Relationship Id="rId1" TargetMode="External" '
                b'Target="https://example.invalid/asset"/>'
                b"</Relationships>"
            ),
            "pptx_external_relationship",
        ),
    ],
)
def test_pptx_package_scanner_rejects_untrusted_members(
    tmp_path: Path,
    member_name: str,
    member_content: bytes,
    expected_failure: str,
) -> None:
    package_path = tmp_path / "malicious.pptx"
    with zipfile.ZipFile(
        package_path,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as package:
        package.writestr(
            "[Content_Types].xml",
            b"<Types/>",
        )
        package.writestr(
            "ppt/presentation.xml",
            b"<p:presentation/>",
        )
        package.writestr(member_name, member_content)

    failures = export_pptx_module._inspect_package(  # noqa: SLF001
        package_path
    )

    assert expected_failure in failures


def _stage_asset(
    *,
    assets_dir: Path,
    artifact_id: str,
    content: bytes,
) -> ResolvedSnapshotAsset:
    path = assets_dir / f"{artifact_id}.png"
    path.write_bytes(content)
    return ResolvedSnapshotAsset(
        artifact_id=artifact_id,
        content_sha256=_sha256(content),
        size_bytes=len(content),
        mime="image/png",
        relative_path=f"assets/{path.name}",
    )


def _png_bytes(
    size: tuple[int, int],
    color: tuple[int, int, int, int],
) -> bytes:
    output = io.BytesIO()
    Image.new("RGBA", size, color).save(
        output,
        format="PNG",
    )
    return output.getvalue()


def _illustration_png(size: tuple[int, int]) -> bytes:
    image = Image.new("RGBA", size, "#f4efe5")
    draw = ImageDraw.Draw(image)
    width, height = size
    draw.rounded_rectangle(
        (
            width * 0.1,
            height * 0.1,
            width * 0.9,
            height * 0.9,
        ),
        radius=max(8, int(min(size) * 0.06)),
        fill="#173d32",
    )
    draw.ellipse(
        (
            width * 0.23,
            height * 0.22,
            width * 0.77,
            height * 0.78,
        ),
        fill="#c9674b",
    )
    draw.rectangle(
        (
            width * 0.32,
            height * 0.34,
            width * 0.68,
            height * 0.66,
        ),
        fill="#f6f1e7",
    )
    draw.text(
        (width * 0.37, height * 0.48),
        "CAMPAIGN",
        fill="#173d32",
    )
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _sha256(content: bytes) -> str:
    return f"sha256:{hashlib.sha256(content).hexdigest()}"


def test_normalize_visible_strips_wrap_artifact_whitespace() -> None:
    """export_pptx has its own _normalize_visible used for closed-file semantic
    parity (mirror of the PDF regression). LibreOffice/python-pptx text
    extraction inserts stray spaces at wrap boundaries that, in a CJK-dominant
    report, land anywhere -- between ideographs, mid-number ("29.6%" ->
    "2 9.6%"), mid-citation ("[10]" -> "[1 0]"), inside an ASCII acronym
    ("SOP" -> "S OP"). Strip-all whitespace normalization, applied symmetrically
    to expected token and observed text, must collapse them all while never
    fabricating a false match."""
    n = export_pptx_module._normalize_visible  # noqa: SLF001
    for wrapped, source in (
        ("侵蚀利润高达88.5 %，退款率", "侵蚀利润高达88.5%，退款率"),  # digit-percent
        ("同比下滑4.22% ，降幅", "同比下滑4.22%，降幅"),  # percent-comma
        ("高度集中特征— —这种", "高度集中特征——这种"),  # split em-dash
        ('这种" 尖峰', '这种"尖峰'),  # quote-CJK
        ("东湖街单店退款2, 886.8万", "东湖街单店退款2,886.8万"),  # comma-digit
        ("退款率2 9.6%", "退款率29.6%"),  # digit-digit (mid-number)
        ("区域差异高达1.75倍[1 0]", "区域差异高达1.75倍[10]"),  # digit-digit (citation)
        ("重点审查S OP进行", "重点审查SOP进行"),  # letter-letter (acronym)
    ):
        assert n(wrapped) == n(source), f"wrap artifact not collapsed: {wrapped!r}"
    assert n("hello world") == n("helloworld") == "helloworld"
    assert export_pptx_module._contains_tokens_in_order(  # noqa: SLF001
        "近30天营业额10.64亿，同比下滑4.28 %和4.22% ，退款率升至1 5.31%[1 0]。S OP流程",
        ("近30天营业额10.64亿，同比下滑4.28%和4.22%，退款率升至15.31%[10]。SOP流程",),
    )
