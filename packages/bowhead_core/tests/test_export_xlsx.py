"""XLSX V2 typed-workbook and closed-file parity contracts."""

from __future__ import annotations

import copy
import hashlib
import io
import json
import shutil
import zipfile
from dataclasses import dataclass, replace
from pathlib import Path

import pytest
from bowhead_core.export_assets import (
    ResolvedSnapshotAsset,
    SnapshotAssetResolution,
)
from bowhead_core.export_ir import ExportIR, ExportIRBlock, normalize_export_snapshot
from bowhead_core.export_materializer import (
    BrowserRenderedVisual,
    MaterializedVisual,
    materialize_export_visuals,
)
from bowhead_core.export_snapshot import ExportSnapshot
from bowhead_core.export_validation import validate_export_output
from bowhead_core.export_visual_identity import EXPORT_VISUAL_RENDERER_VERSION
from bowhead_core.export_xlsx import (
    XlsxExportError,
    extract_xlsx_evidence,
    render_export_xlsx,
)
from bowhead_core.print_theme import current_print_theme_snapshot
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw

from bowhead_core import export_xlsx as export_xlsx_module

_FIXTURE = Path(__file__).resolve().parents[3] / "eval/fixtures/canvas_export/surface_projection_v1.json"
_SOFFICE = shutil.which("soffice")


@dataclass
class _FakeBrowserRenderer:
    name: str = "canvas-browser-visual-v1"
    version: str = EXPORT_VISUAL_RENDERER_VERSION

    async def render(
        self,
        block: ExportIRBlock,
        *,
        asset_paths: dict[str, Path],
    ) -> BrowserRenderedVisual:
        del asset_paths
        color = (
            int(hashlib.sha256(block.block_key.encode()).hexdigest()[:2], 16),
            90,
            150,
            255,
        )
        image = Image.new("RGBA", (1600, 900), "#fffaf1")
        draw = ImageDraw.Draw(image)
        draw.line((150, 720, 1470, 720), fill="#173d32", width=5)
        draw.line((150, 120, 150, 720), fill="#173d32", width=5)
        draw.text(
            (150, 55),
            f"{block.source.section_kind} · APAC / EMEA",
            fill="#173d32",
        )
        for index, value in enumerate((420, 650, 510, 730)):
            left = 260 + index * 280
            draw.rounded_rectangle(
                (left, 720 - value, left + 135, 720),
                radius=18,
                fill=color if index % 2 == 0 else "#c9674b",
            )
        output = io.BytesIO()
        image.save(output, format="PNG")
        return BrowserRenderedVisual(
            content=output.getvalue(),
            mime="image/png",
            visual_input_sha256=str(block.visual_input_sha256),
        )


def _snapshot(
    *,
    long_text: str | None = None,
    revenue: float = 120.5,
    region: str | None = None,
) -> ExportSnapshot:
    fixture = json.loads(_FIXTURE.read_text(encoding="utf-8"))
    by_id = {section["id"]: section for section in fixture["sections"]}
    entries = copy.deepcopy(fixture["expected_entries"])
    sections = [copy.deepcopy(by_id[entry["section_id"]]) for entry in entries]
    section_by_id = {section["id"]: section for section in sections}
    section_by_id["sec_h1"]["payload"]["text"] = "=2+2 is ordinary report text"
    section_by_id["sec_narrative"]["payload"]["markdown"] = long_text if long_text is not None else "+cmd|' /C calc'!A0 is ordinary narrative text"
    section_by_id["sec_data"]["payload"]["columns"] = [
        {"key": "region", "label": "Region", "value_type": "string"},
        {
            "key": "revenue",
            "label": "Revenue",
            "value_type": "number",
            "currency_code": "CAD",
            "display": {"decimal_places": 3, "decimal_mode": "fixed"},
        },
        {"key": "active", "label": "Active", "value_type": "boolean"},
        {"key": "as_of", "label": "As of", "value_type": "date"},
        {"key": "fraction", "label": "Fraction", "value_type": "number", "percent_scale": "fraction"},
        {"key": "points", "label": "Points", "value_type": "number", "percent_scale": "points", "unit": "%"},
        {"key": "identifier", "label": "Identifier", "value_type": "string"},
    ]
    section_by_id["sec_data"]["payload"]["rows"] = [
        {
            "region": (long_text if long_text is not None else region if region is not None else "@APAC"),
            "revenue": revenue,
            "active": True,
            "as_of": "2026-07-17",
            "fraction": 0.0035,
            "points": 0.35,
            "identifier": "9007199254740993",
        }
    ]
    sections[1]["source_refs"] = ["tl_xlsx_source"]
    payload = {
        "request": {
            "format": "xlsx",
            "profile": "data_workbook",
            "parity_contract": "data_parity_v1",
            "client_surface_receipt_sha256": f"sha256:{'d' * 64}",
            "locale": "zh-CN",
            "timezone": "Asia/Singapore",
            "theme": {"kind": "default"},
            "print_theme": current_print_theme_snapshot(),
            "options": {"include_appendix": True},
        },
        "surface_projection": {
            "contract_version": 1,
            "token": fixture["surface_token"],
            "entries": entries,
            "exclusions": fixture["expected_exclusions"],
        },
        "sections": sections,
    }
    return ExportSnapshot(
        id="exs_xlsx_data_workbook",
        schema_version=1,
        canvas_id="cv_xlsx",
        session_id="s_xlsx",
        team_id="team_xlsx",
        workspace_id="ws_xlsx",
        canvas_owner="u_owner",
        requested_by="u_owner",
        fmt="xlsx",
        profile="data_workbook",
        title="XLSX 完整性测试",
        locale="zh-CN",
        timezone="Asia/Singapore",
        projection_token=f"sha256:{'a' * 64}",
        surface_contract_version=1,
        surface_token=fixture["surface_token"],
        parity_contract="data_parity_v1",
        payload=payload,
        payload_sha256=f"sha256:{'c' * 64}",
        size_bytes=4096,
        section_count=len(sections),
        created_at="2026-07-17T00:00:00+00:00",
        expires_at="2026-07-24T00:00:00+00:00",
    )


async def _render_fixture(
    tmp_path: Path,
    *,
    long_text: str | None = None,
    revenue: float = 120.5,
    region: str | None = None,
) -> tuple[
    ExportSnapshot,
    ExportIR,
    tuple[MaterializedVisual, ...],
    Path,
]:
    attempt_dir = tmp_path / "attempt"
    assets_dir = attempt_dir / "assets"
    assets_dir.mkdir(parents=True)
    assets = (
        _stage_asset(
            assets_dir=assets_dir,
            artifact_id="art_image",
            content=_illustration_png((640, 480)),
        ),
        _stage_asset(
            assets_dir=assets_dir,
            artifact_id="art_bg",
            content=_illustration_png((1080, 1080)),
        ),
    )
    snapshot = _snapshot(
        long_text=long_text,
        revenue=revenue,
        region=region,
    )
    export_ir = normalize_export_snapshot(
        snapshot,
        asset_content_hashes={asset.artifact_id: asset.content_sha256 for asset in assets},
    )
    visuals = await materialize_export_visuals(
        export_ir=export_ir,
        asset_resolution=SnapshotAssetResolution(
            resolved=assets,
            warnings=(),
        ),
        attempt_dir=attempt_dir,
        browser_renderer=_FakeBrowserRenderer(),
    )
    result = render_export_xlsx(
        export_ir=export_ir,
        materialized_visuals=visuals,
        attempt_dir=attempt_dir,
    )
    return (
        snapshot,
        export_ir,
        visuals,
        attempt_dir / result.relative_path,
    )


@pytest.mark.skipif(
    _SOFFICE is None,
    reason="LibreOffice is the XLSX V2 recalc gate",
)
async def test_xlsx_closed_file_matches_all_kinds_and_cell_types(
    tmp_path: Path,
) -> None:
    snapshot, export_ir, visuals, xlsx_path = await _render_fixture(tmp_path)

    workbook = load_workbook(xlsx_path, data_only=False)
    assert workbook.sheetnames[0] == "Summary"
    data_sheet_count = len([name for name in workbook.sheetnames if name.startswith("Data - ")])
    assert workbook.sheetnames[1 : data_sheet_count + 1] == [f"Data - {index}" for index in range(1, data_sheet_count + 1)]
    assert workbook.sheetnames[data_sheet_count + 1 :] == [
        "Insights",
        "Recommendations",
        "Visual Index",
        "SQL & Code",
        "Sources",
        "Manifest",
    ]
    assert all(len(name) <= 31 for name in workbook.sheetnames)
    assert not [cell.coordinate for worksheet in workbook.worksheets for row in worksheet.iter_rows() for cell in row if cell.data_type == "f"]
    typed_cells = {
        (cell.value, cell.data_type, cell.number_format)
        for worksheet in workbook.worksheets
        if worksheet.title.startswith("Data - ")
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value in {120.5, True} or str(cell.value).startswith("2026-07-17")
    }
    workbook.close()
    assert any(value == 120.5 and data_type == "n" for value, data_type, _ in typed_cells)
    assert any(value == 120.5 and fmt == '#,##0.000 "CAD"' for value, _, fmt in typed_cells)
    assert any(value is True and data_type == "b" for value, data_type, _ in typed_cells)
    assert any(str(value).startswith("2026-07-17") and number_format == "yyyy-mm-dd" for value, _data_type, number_format in typed_cells)

    evidence = extract_xlsx_evidence(
        xlsx_path=xlsx_path,
        export_ir=export_ir,
        materialized_visuals=visuals,
        validation_dir=tmp_path / "validation",
        soffice_path=Path(str(_SOFFICE)),
    )
    manifest = validate_export_output(
        job_id="exp_xlsx_data",
        snapshot=snapshot,
        export_ir=export_ir,
        target="xlsx_data",
        include_appendix=True,
        renderer_name="canvas-data-workbook-v2",
        renderer_version=EXPORT_VISUAL_RENDERER_VERSION,
        evidence=evidence,
    )

    assert evidence.parser_opened
    assert evidence.hard_check_failures == ()
    assert evidence.demo_residue_count == 0
    assert len(evidence.sections) == 22
    assert manifest.validation.result == "warnings"
    assert manifest.summary.rendered == 19
    assert manifest.summary.fallback == 3
    assert manifest.summary.excluded_by_policy == 0
    assert manifest.summary.accounted_rate == 1.0
    assert manifest.summary.semantic_parity_rate == 1.0
    assert manifest.summary.visual_identity_rate == 1.0
    assert manifest.summary.typed_data_parity_rate == 1.0


@pytest.mark.skipif(
    _SOFFICE is None,
    reason="LibreOffice is the XLSX V2 recalc gate",
)
async def test_xlsx_accepts_excel_float_rounding_and_literal_error_text(
    tmp_path: Path,
) -> None:
    snapshot, export_ir, visuals, xlsx_path = await _render_fixture(
        tmp_path,
        revenue=10 / 3,
        region="#N/A",
    )

    evidence = extract_xlsx_evidence(
        xlsx_path=xlsx_path,
        export_ir=export_ir,
        materialized_visuals=visuals,
        validation_dir=tmp_path / "validation-float",
        soffice_path=Path(str(_SOFFICE)),
    )
    manifest = validate_export_output(
        job_id="exp_xlsx_float",
        snapshot=snapshot,
        export_ir=export_ir,
        target="xlsx_data",
        include_appendix=True,
        renderer_name="canvas-data-workbook-v2",
        renderer_version=EXPORT_VISUAL_RENDERER_VERSION,
        evidence=evidence,
    )

    assert evidence.hard_check_failures == ()
    assert manifest.summary.typed_data_parity_rate == 1.0


@pytest.mark.skipif(
    _SOFFICE is None,
    reason="LibreOffice is the XLSX V2 recalc gate",
)
async def test_xlsx_forged_visual_pixel_is_rejected(
    tmp_path: Path,
) -> None:
    snapshot, export_ir, visuals, xlsx_path = await _render_fixture(tmp_path)
    forged = (
        replace(
            visuals[0],
            pixel_sha256=f"sha256:{'f' * 64}",
        ),
        *visuals[1:],
    )

    evidence = extract_xlsx_evidence(
        xlsx_path=xlsx_path,
        export_ir=export_ir,
        materialized_visuals=forged,
        validation_dir=tmp_path / "validation-forged",
        soffice_path=Path(str(_SOFFICE)),
    )
    manifest = validate_export_output(
        job_id="exp_xlsx_forged",
        snapshot=snapshot,
        export_ir=export_ir,
        target="xlsx_data",
        include_appendix=True,
        renderer_name="canvas-data-workbook-v2",
        renderer_version=EXPORT_VISUAL_RENDERER_VERSION,
        evidence=evidence,
    )

    assert manifest.validation.result == "rejected"
    assert manifest.summary.visual_identity_rate is not None
    assert manifest.summary.visual_identity_rate < 1.0


@pytest.mark.skipif(
    _SOFFICE is None,
    reason="LibreOffice is the XLSX V2 recalc gate",
)
async def test_xlsx_tampered_public_manifest_is_rejected(
    tmp_path: Path,
) -> None:
    _snapshot_value, export_ir, visuals, xlsx_path = await _render_fixture(tmp_path)
    workbook = load_workbook(xlsx_path)
    workbook["Manifest"].cell(6, 2, "sec_forged")
    tampered_path = tmp_path / "tampered.xlsx"
    workbook.save(tampered_path)
    workbook.close()

    evidence = extract_xlsx_evidence(
        xlsx_path=tampered_path,
        export_ir=export_ir,
        materialized_visuals=visuals,
        validation_dir=tmp_path / "validation-tampered",
        soffice_path=Path(str(_SOFFICE)),
    )

    assert "xlsx_manifest_entry_mismatch" in (evidence.hard_check_failures)


@pytest.mark.skipif(
    _SOFFICE is None,
    reason="LibreOffice is the XLSX V2 recalc gate",
)
async def test_xlsx_long_semantic_and_typed_text_is_chunked_without_loss(
    tmp_path: Path,
) -> None:
    long_text = "".join(
        hashlib.sha256(
            f"长期经营风险与整改建议-{index}".encode(),
        ).hexdigest()
        for index in range(1_000)
    )
    snapshot, export_ir, visuals, xlsx_path = await _render_fixture(
        tmp_path,
        long_text=long_text,
    )
    workbook = load_workbook(xlsx_path)
    text_cells = [cell for worksheet in workbook.worksheets for row in worksheet.iter_rows() for cell in row if isinstance(cell.value, str)]
    assert max(len(cell.value) for cell in text_cells) <= (export_xlsx_module.MAX_CELL_TEXT_CHARS)
    assert sum(cell.comment is not None and cell.comment.text.startswith("BH-TEXT1:") for cell in text_cells) >= 4
    workbook.close()

    evidence = extract_xlsx_evidence(
        xlsx_path=xlsx_path,
        export_ir=export_ir,
        materialized_visuals=visuals,
        validation_dir=tmp_path / "validation-long-text",
        soffice_path=Path(str(_SOFFICE)),
    )
    manifest = validate_export_output(
        job_id="exp_xlsx_long_text",
        snapshot=snapshot,
        export_ir=export_ir,
        target="xlsx_data",
        include_appendix=True,
        renderer_name="canvas-data-workbook-v2",
        renderer_version=EXPORT_VISUAL_RENDERER_VERSION,
        evidence=evidence,
    )

    assert evidence.hard_check_failures == ()
    assert manifest.summary.semantic_parity_rate == 1.0
    assert manifest.summary.typed_data_parity_rate == 1.0


@pytest.mark.skipif(
    _SOFFICE is None,
    reason="LibreOffice is the XLSX V2 recalc gate",
)
@pytest.mark.parametrize("hidden_target", ["sheet", "row", "column"])
async def test_xlsx_hidden_business_content_is_rejected(
    tmp_path: Path,
    hidden_target: str,
) -> None:
    _snapshot_value, export_ir, visuals, xlsx_path = await _render_fixture(
        tmp_path,
    )
    workbook = load_workbook(xlsx_path)
    worksheet = workbook["Data - 1"]
    locator_cell = next(cell for row in worksheet.iter_rows() for cell in row if cell.comment is not None and cell.comment.text.startswith("BH1:"))
    if hidden_target == "sheet":
        worksheet.sheet_state = "hidden"
    elif hidden_target == "row":
        worksheet.row_dimensions[locator_cell.row].hidden = True
    else:
        worksheet.column_dimensions[locator_cell.column_letter].hidden = True
    tampered_path = tmp_path / f"hidden-{hidden_target}.xlsx"
    workbook.save(tampered_path)
    workbook.close()

    evidence = extract_xlsx_evidence(
        xlsx_path=tampered_path,
        export_ir=export_ir,
        materialized_visuals=visuals,
        validation_dir=tmp_path / f"validation-hidden-{hidden_target}",
        soffice_path=Path(str(_SOFFICE)),
    )

    expected_failure = "xlsx_sheet_hidden" if hidden_target == "sheet" else "xlsx_locator_range_hidden"
    assert expected_failure in evidence.hard_check_failures


def test_xlsx_demo_fingerprint_allows_expected_business_text() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "客户 CMP-2604-001 的整体合规率"

    assert (
        export_xlsx_module._demo_residue_count(  # noqa: SLF001
            worksheet,
            expected_fingerprints={
                "CMP-2604-001",
                "整体合规率",
            },
        )
        == 0
    )
    assert (
        export_xlsx_module._demo_residue_count(  # noqa: SLF001
            worksheet,
            expected_fingerprints=set(),
        )
        == 2
    )
    workbook.close()


async def test_xlsx_typed_cell_limit_fails_instead_of_truncating(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(export_xlsx_module, "MAX_TYPED_CELLS", 1)

    with pytest.raises(XlsxExportError, match="xlsx_typed_cell_limit"):
        await _render_fixture(tmp_path)


async def test_xlsx_output_name_cannot_escape_attempt_directory(
    tmp_path: Path,
) -> None:
    export_ir = normalize_export_snapshot(
        _snapshot(),
        asset_content_hashes={
            "art_image": f"sha256:{'1' * 64}",
            "art_bg": f"sha256:{'2' * 64}",
        },
    )

    with pytest.raises(XlsxExportError, match="xlsx_output_path_invalid"):
        render_export_xlsx(
            export_ir=export_ir,
            materialized_visuals=(),
            attempt_dir=tmp_path / "attempt",
            output_name="../escaped.xlsx",
        )


@pytest.mark.parametrize(
    ("member_name", "member_content", "expected_failure"),
    [
        (
            "../escaped.xml",
            b"<root/>",
            "xlsx_zip_path_invalid",
        ),
        (
            "xl/embeddings/demo.xlsm",
            b"PK\x03\x04demo",
            "xlsx_forbidden_package_member",
        ),
        (
            "xl/_rels/workbook.xml.rels",
            (
                b'<?xml version="1.0"?>'
                b'<Relationships xmlns="http://schemas.openxmlformats.org/'
                b'package/2006/relationships">'
                b'<Relationship Id="rId1" TargetMode="External" '
                b'Target="https://example.invalid/data"/>'
                b"</Relationships>"
            ),
            "xlsx_external_relationship_forbidden",
        ),
    ],
)
def test_xlsx_package_scanner_rejects_untrusted_members(
    tmp_path: Path,
    member_name: str,
    member_content: bytes,
    expected_failure: str,
) -> None:
    package_path = tmp_path / "malicious.xlsx"
    with zipfile.ZipFile(
        package_path,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as package:
        package.writestr("[Content_Types].xml", b"<Types/>")
        package.writestr("xl/workbook.xml", b"<workbook/>")
        package.writestr(member_name, member_content)

    failures = export_xlsx_module._inspect_package(package_path)  # noqa: SLF001

    assert expected_failure in failures


def _stage_asset(
    *,
    assets_dir: Path,
    artifact_id: str,
    content: bytes,
) -> ResolvedSnapshotAsset:
    path = assets_dir / f"{artifact_id}.png"
    path.write_bytes(content)
    return ResolvedSnapshotAsset(
        artifact_id=artifact_id,
        content_sha256=_sha256(content),
        size_bytes=len(content),
        mime="image/png",
        relative_path=f"assets/{path.name}",
    )


def _illustration_png(size: tuple[int, int]) -> bytes:
    image = Image.new("RGBA", size, "#f4efe5")
    draw = ImageDraw.Draw(image)
    width, height = size
    draw.rounded_rectangle(
        (
            width * 0.1,
            height * 0.1,
            width * 0.9,
            height * 0.9,
        ),
        radius=max(8, int(min(size) * 0.06)),
        fill="#173d32",
    )
    draw.ellipse(
        (
            width * 0.23,
            height * 0.22,
            width * 0.77,
            height * 0.78,
        ),
        fill="#c9674b",
    )
    draw.rectangle(
        (
            width * 0.32,
            height * 0.34,
            width * 0.68,
            height * 0.66,
        ),
        fill="#f6f1e7",
    )
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _sha256(content: bytes) -> str:
    return f"sha256:{hashlib.sha256(content).hexdigest()}"


async def test_numeric_display_survives_closed_workbook_without_changing_source_values(tmp_path):
    _, _, _, path = await _render_fixture(tmp_path)
    workbook = load_workbook(path)
    cells = [cell for sheet in workbook if sheet.title.startswith("Data - ") for row in sheet for cell in row]
    assert any(cell.value == 0.0035 and cell.data_type == "n" and cell.number_format == "#,##0.##%" for cell in cells)
    assert any(cell.value == 0.35 and cell.data_type == "n" and cell.number_format == '#,##0.## "%"' for cell in cells)
    assert any(cell.value == "9007199254740993" and cell.data_type == "s" for cell in cells)
    workbook.close()
