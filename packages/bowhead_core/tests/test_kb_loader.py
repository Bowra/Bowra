"""KB loader tests.

Covers:
- dispatch: mime → loader (含 P1-S1 text/plain → MarkdownLoader)
- MarkdownLoader: 直读 .md / .txt 文件 → KbDocument
- DocxLoader / PdfLoader: import OK + 简单 sanity (实际 PDF/DOCX 重 dep, e2e 测试覆盖)
- PptxLoader: 真 deck 端到端 (python-pptx 是纯 python dep, 测试内造 deck 不重)
"""

from __future__ import annotations

import io
import zipfile

import pytest
from bowhead_core.kb.errors import UnsupportedMimeError
from bowhead_core.kb.loader import get_loader
from bowhead_core.kb.loader.docx_loader import DocxLoader
from bowhead_core.kb.loader.md_loader import MarkdownLoader
from bowhead_core.kb.loader.pdf_loader import PdfLoader
from bowhead_core.kb.loader.pptx_loader import PPTX_MIME, PptxLoader

# ─── dispatch ─────────────────────────────────────────


def test_dispatch_pdf_returns_pdf_loader() -> None:
    loader = get_loader("application/pdf")
    assert isinstance(loader, PdfLoader)


def test_dispatch_docx_returns_docx_loader() -> None:
    loader = get_loader("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert isinstance(loader, DocxLoader)


def test_dispatch_markdown_returns_md_loader() -> None:
    assert isinstance(get_loader("text/markdown"), MarkdownLoader)
    assert isinstance(get_loader("text/x-markdown"), MarkdownLoader)


def test_dispatch_text_plain_returns_md_loader() -> None:
    """[P1-S1] python-magic 多把 .md detect 为 text/plain, 走 MarkdownLoader."""
    assert isinstance(get_loader("text/plain"), MarkdownLoader)


def test_dispatch_pptx_returns_pptx_loader() -> None:
    """pptx-loader sprint: 上传闸早放行 pptx, dispatch 补上 loader (原 UnsupportedMime)."""
    assert isinstance(get_loader(PPTX_MIME), PptxLoader)


def test_dispatch_unknown_raises() -> None:
    with pytest.raises(UnsupportedMimeError) as exc:
        get_loader("application/x-evil")
    assert "application/x-evil" in str(exc.value)


# ─── MarkdownLoader ───────────────────────────────────


@pytest.mark.asyncio
async def test_md_loader_returns_single_doc(tmp_path) -> None:
    md_file = tmp_path / "test.md"
    md_file.write_text("# Header\n\nContent paragraph.\n", encoding="utf-8")

    loader = MarkdownLoader()
    result = await loader.load(str(md_file))

    assert result.layer_used == "L0"
    assert len(result.documents) == 1
    assert "# Header" in result.documents[0].page_content
    assert result.metadata["char_count"] > 0


@pytest.mark.asyncio
async def test_md_loader_handles_chinese(tmp_path) -> None:
    md_file = tmp_path / "中文.md"
    md_file.write_text("# 业务定义\n\nGMV 是商品交易总额.\n", encoding="utf-8")

    loader = MarkdownLoader()
    result = await loader.load(str(md_file))

    assert "GMV" in result.documents[0].page_content
    assert "业务定义" in result.documents[0].page_content


@pytest.mark.asyncio
async def test_md_loader_empty_file(tmp_path) -> None:
    md_file = tmp_path / "empty.md"
    md_file.write_text("", encoding="utf-8")

    loader = MarkdownLoader()
    result = await loader.load(str(md_file))

    assert len(result.documents) == 1
    assert result.documents[0].page_content == ""
    assert result.metadata["char_count"] == 0


# ─── PdfLoader / DocxLoader basic instantiation ───────


def test_pdf_loader_instantiates() -> None:
    """PdfLoader 可 instantiate (import marker via 延迟 import 不阻塞 module load)."""
    loader = PdfLoader()
    assert loader is not None


def test_docx_loader_instantiates() -> None:
    loader = DocxLoader()
    assert loader is not None


# ─── PptxLoader ───────────────────────────────────────
#
# Build real deck, real extraction (python-pptx pure Python, lightweight); each case pins observed spike behavior.

_PNG_1X1 = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01"
    b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)


def _build_deck(path, *, with_notes: bool = True):
    """标题页 / 要点页 / 表格页 / 组合形状页 / 纯图页 / 备注页 / 空页."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "2026 Q1 竞品分析"
    s1.placeholders[1].text = "Marketing 团队"

    s2 = prs.slides.add_slide(blank)
    tf = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3)).text_frame
    tf.text = "核心发现"
    p = tf.add_paragraph()
    p.text = "星舰 Pro 定价上探高端"

    s3 = prs.slides.add_slide(blank)
    tbl = s3.shapes.add_table(2, 3, Inches(0.5), Inches(1), Inches(8), Inches(2)).table
    for r, row in enumerate([["型号", "价格", "情感"], ["星舰 Pro", "¥6999", "正面"]]):
        for c, val in enumerate(row):
            tbl.cell(r, c).text = val

    # Group shapes: slide.shapes nonrecursive; only recursion retrieves text.
    s4 = prs.slides.add_slide(blank)
    grp = s4.shapes.add_group_shape()
    grp.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "组合内文本 A"
    grp.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1)).text_frame.text = "组合内文本 B"

    s5 = prs.slides.add_slide(blank)
    s5.shapes.add_picture(io.BytesIO(_PNG_1X1), Inches(1), Inches(1), Inches(2), Inches(2))

    s6 = prs.slides.add_slide(blank)
    s6.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "总结页"
    if with_notes:
        s6.notes_slide.notes_text_frame.text = "下季度重点跟进星舰 Pro 高端线"

    prs.slides.add_slide(blank)  # Empty slide.
    prs.save(str(path))
    return path


@pytest.mark.asyncio
async def test_pptx_loader_extracts_slides(tmp_path) -> None:
    """真 deck: 一 slide 一 doc, page=slide 序号 (溯源锚)."""
    deck = _build_deck(tmp_path / "deck.pptx")
    result = await PptxLoader().load(str(deck))

    assert result.layer_used == "L0"
    assert result.metadata["slide_count"] == 7
    # Image-only(5) + empty(7) lack text → skip; 5 remain.
    assert len(result.documents) == 5
    pages = [d.metadata["page"] for d in result.documents]
    assert pages == [1, 2, 3, 4, 6]
    assert result.documents[0].page_content.startswith("## Slide 1")


@pytest.mark.asyncio
async def test_pptx_loader_extracts_title_bullets_and_table(tmp_path) -> None:
    deck = _build_deck(tmp_path / "deck.pptx")
    result = await PptxLoader().load(str(deck))
    by_page = {d.metadata["page"]: d.page_content for d in result.documents}

    assert "2026 Q1 竞品分析" in by_page[1]  # title placeholder
    assert "Marketing 团队" in by_page[1]  # subtitle placeholder
    assert "星舰 Pro 定价上探高端" in by_page[2]  # Bullet paragraph.
    assert "| 型号 | 价格 | 情感 |" in by_page[3]  # Table → pipe row.
    assert "| 星舰 Pro | ¥6999 | 正面 |" in by_page[3]


@pytest.mark.asyncio
async def test_pptx_loader_recurses_group_shapes(tmp_path) -> None:
    """⚠ 回归闸: slide.shapes 不自动递归 GROUP —— 不手动递归会静默丢组合内全部文本."""
    deck = _build_deck(tmp_path / "deck.pptx")
    result = await PptxLoader().load(str(deck))
    by_page = {d.metadata["page"]: d.page_content for d in result.documents}

    assert "组合内文本 A" in by_page[4]
    assert "组合内文本 B" in by_page[4]


@pytest.mark.asyncio
async def test_pptx_loader_includes_speaker_notes(tmp_path) -> None:
    deck = _build_deck(tmp_path / "deck.pptx")
    result = await PptxLoader().load(str(deck))
    by_page = {d.metadata["page"]: d.page_content for d in result.documents}

    assert "Notes: 下季度重点跟进星舰 Pro 高端线" in by_page[6]


@pytest.mark.asyncio
async def test_pptx_loader_no_notes_deck_has_no_notes_marker(tmp_path) -> None:
    """无备注的 deck 不该冒出 Notes: 空串.

    ⚠ 这条只守 `if note:` 那道判空, **守不住 has_notes_slide 门控**(把门控换成
    `if True:` 本测试仍绿 —— 实验验证过). 门控由下面那条守.
    """
    deck = _build_deck(tmp_path / "deck.pptx", with_notes=False)
    result = await PptxLoader().load(str(deck))

    assert all("Notes:" not in d.page_content for d in result.documents)


def test_pptx_loader_does_not_create_phantom_notes_slide(tmp_path) -> None:
    """⚠ 回归闸(坑#2): 直接碰 `slide.notes_slide` 会惰性造出空备注页.

    必须先 `has_notes_slide` 门控. 断言**副作用**而非输出文本 —— 后者被 `if note:`
    判空掩盖, 门控删了也看不出来(自嗨测试的反例).
    """
    from bowhead_core.kb.loader.pptx_loader import _slide_to_text
    from pptx import Presentation

    deck = _build_deck(tmp_path / "deck.pptx", with_notes=False)
    prs = Presentation(str(deck))
    slide = list(prs.slides)[1]  # No-notes slide.

    assert slide.has_notes_slide is False  # Precondition.
    _slide_to_text(slide, 2)
    assert slide.has_notes_slide is False, "门控失守: 抽取过程凭空造出了备注页"


@pytest.mark.asyncio
async def test_pptx_loader_extracts_chart_text(tmp_path) -> None:
    """⚠ 回归闸: 图表是 GraphicFrame 且 has_text_frame=False —— 光判 text_frame 会把
    标题(常是该页核心结论)/分类/系列**静默丢光**, 文件却标 done(marketing deck 图表密集).
    """
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cd = CategoryChartData()
    cd.categories = ["星舰 Pro", "星舰 Air", "竞品X"]
    cd.add_series("2026 Q1 营收(亿元)", (12.5, 8.3, 6.1))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), cd)
    gf.chart.has_title = True
    gf.chart.chart_title.text_frame.text = "星舰 Pro 领跑高端市场"
    deck = tmp_path / "chart.pptx"
    prs.save(str(deck))

    result = await PptxLoader().load(str(deck))

    assert result.metadata.get("parse_failed") is not True
    body = result.documents[0].page_content
    assert "Chart: 星舰 Pro 领跑高端市场" in body  # Title = slide key conclusion.
    assert "2026 Q1 营收(亿元)" in body  # Series name.
    assert "竞品X" in body  # Category.
    assert "12.5" in body  # Value for matrix extraction.


@pytest.mark.asyncio
async def test_pptx_loader_recurses_nested_groups(tmp_path) -> None:
    """组合套组合: 递归必须是任意深度, 不能只展开一层."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    outer = s.shapes.add_group_shape()
    outer.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "外层文本"
    inner = outer.shapes.add_group_shape()
    inner.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1)).text_frame.text = "内层嵌套文本"
    deck = tmp_path / "nested.pptx"
    prs.save(str(deck))

    result = await PptxLoader().load(str(deck))

    body = result.documents[0].page_content
    assert "外层文本" in body
    assert "内层嵌套文本" in body, "嵌套组合(深度 2)的文本丢了"


@pytest.mark.asyncio
async def test_pptx_loader_escapes_pipe_and_newline_in_table(tmp_path) -> None:
    """表格 cell 里的 | 和换行会破坏 pipe 行格式 → 必须转义/压平."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tbl = s.shapes.add_table(1, 2, Inches(0.5), Inches(1), Inches(8), Inches(1)).table
    tbl.cell(0, 0).text = "含|竖线"
    tbl.cell(0, 1).text = "含\n换行"
    deck = tmp_path / "tbl.pptx"
    prs.save(str(deck))

    result = await PptxLoader().load(str(deck))
    body = result.documents[0].page_content

    assert r"含\|竖线" in body  # Escape pipe, preserve columns.
    assert "含 换行" in body  # Flatten newlines, preserve row per line.
    tbl_lines = [ln for ln in body.splitlines() if ln.startswith("| ")]
    assert len(tbl_lines) == 1, f"cell 内换行撑出了多行: {tbl_lines}"


@pytest.mark.asyncio
async def test_pptx_loader_corrupt_file_marks_parse_failed(tmp_path) -> None:
    """损坏字节 → PackageNotFoundError → parse_failed (不 raise; worker 据此标 failed)."""
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a real pptx at all")

    result = await PptxLoader().load(str(bad))

    assert result.metadata["parse_failed"] is True
    assert result.metadata["reason"]
    assert len(result.documents) == 1  # Placeholder doc, caller safe.


@pytest.mark.asyncio
async def test_pptx_loader_zip_but_not_pptx_marks_parse_failed(tmp_path) -> None:
    """合法 zip 但非 OOXML → KeyError([Content_Types].xml) → 同样 parse_failed."""
    fake = tmp_path / "fake.pptx"
    with zipfile.ZipFile(fake, "w") as z:
        z.writestr("hello.txt", "i am not office xml")

    result = await PptxLoader().load(str(fake))

    assert result.metadata["parse_failed"] is True


@pytest.mark.asyncio
async def test_pptx_loader_image_only_deck_marks_parse_failed(tmp_path) -> None:
    """纯图 deck 抽不出文本 → parse_failed(诚实): 存成 done 会让用户以为可查, 实则搜不到."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(io.BytesIO(_PNG_1X1), Inches(1), Inches(1), Inches(2), Inches(2))
    deck = tmp_path / "images.pptx"
    prs.save(str(deck))

    result = await PptxLoader().load(str(deck))

    assert result.metadata["parse_failed"] is True
    assert "no extractable text" in result.metadata["reason"]


@pytest.mark.asyncio
async def test_pptx_loader_empty_deck_marks_parse_failed(tmp_path) -> None:
    from pptx import Presentation

    deck = tmp_path / "empty.pptx"
    Presentation().save(str(deck))

    result = await PptxLoader().load(str(deck))

    assert result.metadata["parse_failed"] is True
    assert result.metadata["slide_count"] == 0


@pytest.mark.asyncio
async def test_pptx_loader_caps_slides(tmp_path, monkeypatch) -> None:
    """slide cap: 超限截断但**记 metadata 不静默** (no-silent-caps)."""
    from bowhead_core.kb.loader import pptx_loader as mod
    from pptx import Presentation
    from pptx.util import Inches

    monkeypatch.setattr(mod, "MAX_SLIDES", 2)
    prs = Presentation()
    for i in range(4):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = f"页 {i + 1}"
    deck = tmp_path / "big.pptx"
    prs.save(str(deck))

    result = await PptxLoader().load(str(deck))

    assert result.metadata["slide_count"] == 4
    assert result.metadata["truncated"] is True
    assert result.metadata["max_slides"] == 2
    assert len(result.documents) == 2
