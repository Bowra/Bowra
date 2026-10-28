"""KB worker tests with mocks (no real dashscope/marker).

Covers:
- end-to-end parse flow: pending → parsing → chunking → embedding → done
- P0-D: count_mismatch > 5% → status=failed
- P1-S2: semaphore exists
- P1-Q6: parse_progress_pct state map
"""

from __future__ import annotations

import asyncio
from unittest.mock import AsyncMock, MagicMock

import pytest
from bowhead_core.kb.embedder import EMBED_DIM
from bowhead_core.kb.files_store import KbFilesStore
from bowhead_core.kb.worker import (
    _PARSE_SEMAPHORE,
    PARSE_PROGRESS_MAP,
    KbWorker,
)


@pytest.fixture
async def files_store(initialize_migrated_kb_files_database, tmp_path):
    initialize_migrated_kb_files_database(tmp_path / "kb_files.db")
    s = KbFilesStore(tmp_path / "kb_files.db")
    await s.initialize()
    return s


def test_parse_progress_map_state_values() -> None:
    """[P1-Q6] state machine map: pending 0 → parsing 20 → chunking 50 → embedding 75 → done 100."""
    assert PARSE_PROGRESS_MAP["pending"] == 0
    assert PARSE_PROGRESS_MAP["parsing"] == 20
    assert PARSE_PROGRESS_MAP["chunking"] == 50
    assert PARSE_PROGRESS_MAP["embedding"] == 75
    assert PARSE_PROGRESS_MAP["done"] == 100
    assert PARSE_PROGRESS_MAP["failed"] is None


def test_semaphore_exists() -> None:
    """[P1-S2] in-process semaphore 限并发 5 个."""
    assert _PARSE_SEMAPHORE._value == 5


@pytest.mark.asyncio
async def test_parse_and_index_many_concurrent_and_failure_isolated() -> None:
    """parse_and_index_many fans out CONCURRENTLY (not serial) and one file's
    failure does not cancel the rest (return_exceptions=True).

    Pins the bundle-ingest contract: confirm hands ALL children to this single
    method so they parse under _PARSE_SEMAPHORE, instead of per-child
    background_tasks.add_task (which FastAPI awaits sequentially → serial).
    """
    worker = KbWorker(files_store=MagicMock(), store=MagicMock(), embedder=MagicMock())

    started: list[str] = []
    finished: list[str] = []
    active = 0
    max_concurrent = 0

    async def fake_parse(fid: str) -> None:
        nonlocal active, max_concurrent
        active += 1
        max_concurrent = max(max_concurrent, active)
        started.append(fid)
        try:
            await asyncio.sleep(0.02)
            if fid == "boom":
                raise RuntimeError("boom")
            finished.append(fid)
        finally:
            active -= 1

    worker.parse_and_index = fake_parse  # type: ignore[assignment]
    ids = [f"f{i}" for i in range(6)] + ["boom"]
    await worker.parse_and_index_many(ids)

    assert set(started) == set(ids)  # all dispatched
    assert max_concurrent >= 2  # ran concurrently, not one-at-a-time
    assert "boom" not in finished  # the raising file did not "finish"
    assert sorted(finished) == sorted(f"f{i}" for i in range(6))  # others survived


@pytest.mark.asyncio
async def test_worker_md_file_end_to_end(files_store, tmp_path) -> None:
    """E2E: markdown file → load → split → embed → chromadb add → done."""
    md_file = tmp_path / "test.md"
    md_file.write_text("# Section A\n\nContent A1.\n\n## Section B\n\nContent B1.\n", encoding="utf-8")

    file_id = await files_store.create(
        artifact_id=str(md_file),  # Direct path.
        filename="test.md",
        mime_type="text/markdown",
        size_bytes=md_file.stat().st_size,
        scope="workspace",
        scope_id="ws_42",
        workspace_id="ws_42",
        team_id="team_5",
        uploaded_by="user_a",
    )

    embedder = MagicMock()
    embedder.embed_batch = AsyncMock(return_value=[[0.1] * EMBED_DIM, [0.2] * EMBED_DIM])

    store = MagicMock()
    store.hard_delete_file = AsyncMock(return_value=0)
    store.add_chunks = AsyncMock()
    store.count_chunks = AsyncMock(return_value=2)

    worker = KbWorker(
        files_store=files_store,
        store=store,
        embedder=embedder,
        framework_inferrer=None,
        artifact_fetcher=None,
    )
    await worker.parse_and_index(file_id)

    dto = await files_store.get(file_id)
    assert dto is not None
    assert dto.parse_status == "done"
    assert dto.chunks_count == 2
    assert dto.parse_progress_pct == 100


@pytest.mark.asyncio
async def test_worker_pptx_file_end_to_end(files_store, tmp_path) -> None:
    """E2E: pptx → load → split → embed → chromadb add → done.

    pptx-loader sprint 的**缺口闭合证**: 上传闸 (routes/files.py ALLOWED_MIME) 早放行
    pptx, 但 dispatch 无 loader → 此前必 failed("unsupported mime"), 文件不可查.
    本测试锁死 pptx 真能走完摄入链路, 且**页码溯源 (page=slide 序号) 一路传到 chunk**.
    """
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    for text in ["星舰 Pro 定价上探高端", "星舰 Air 主打性价比"]:
        s = prs.slides.add_slide(blank)
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text_frame.text = text
    deck = tmp_path / "competitor.pptx"
    prs.save(str(deck))

    file_id = await files_store.create(
        artifact_id=str(deck),  # Direct path.
        filename="competitor.pptx",
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        size_bytes=deck.stat().st_size,
        scope="workspace",
        scope_id="ws_42",
        workspace_id="ws_42",
        team_id="team_5",
        uploaded_by="user_a",
    )

    embedder = MagicMock()
    embedder.embed_batch = AsyncMock(side_effect=lambda texts: [[0.1] * EMBED_DIM] * len(texts))

    added: dict[str, list] = {}

    async def _fake_add(file, chunks) -> None:
        added["chunks"] = list(chunks)

    store = MagicMock()
    store.hard_delete_file = AsyncMock(return_value=0)
    store.add_chunks = AsyncMock(side_effect=_fake_add)
    store.count_chunks = AsyncMock(side_effect=lambda file: len(added.get("chunks", [])))

    worker = KbWorker(files_store=files_store, store=store, embedder=embedder)
    await worker.parse_and_index(file_id)

    dto = await files_store.get(file_id)
    assert dto is not None
    assert dto.parse_status == "done", f"pptx 未走通摄入: {dto.parse_error}"
    assert dto.parse_layer == "L0"
    assert dto.chunks_count == 2

    # Actual content enters chunks and page provenance present (worker._build_kb_chunks reads metadata['page']).
    chunks = added["chunks"]
    assert any("星舰 Pro 定价上探高端" in c.content for c in chunks)
    assert sorted(c.metadata.page for c in chunks) == [1, 2]


@pytest.mark.asyncio
async def test_worker_count_mismatch_5pct_marks_failed(files_store, tmp_path) -> None:
    """[P0-D] expected 10, actual 4 → mismatch 60% > 5% → status=failed."""
    md_file = tmp_path / "test.md"
    md_file.write_text("\n\n".join(f"## H{i}\n\nContent {i}." for i in range(10)), encoding="utf-8")

    file_id = await files_store.create(
        artifact_id=str(md_file),
        filename="test.md",
        mime_type="text/markdown",
        size_bytes=100,
        scope="workspace",
        scope_id="ws_42",
        workspace_id="ws_42",
        team_id="team_5",
        uploaded_by="u",
    )

    embedder = MagicMock()
    # split 10 chunks → 10 embeddings
    embedder.embed_batch = AsyncMock(return_value=[[0.0] * EMBED_DIM] * 10)

    store = MagicMock()
    store.hard_delete_file = AsyncMock(return_value=0)
    store.add_chunks = AsyncMock()
    # Fake chromadb adds only 4 successfully (60% mismatch).
    store.count_chunks = AsyncMock(return_value=4)

    worker = KbWorker(
        files_store=files_store,
        store=store,
        embedder=embedder,
    )
    await worker.parse_and_index(file_id)

    dto = await files_store.get(file_id)
    assert dto is not None
    assert dto.parse_status == "failed"
    assert "mismatch" in (dto.parse_error or "").lower()


@pytest.mark.asyncio
async def test_worker_unsupported_mime_marks_failed(files_store, tmp_path) -> None:
    """Unknown mime → UnsupportedMimeError → status=failed."""
    file = tmp_path / "test.weird"
    file.write_text("data", encoding="utf-8")

    file_id = await files_store.create(
        artifact_id=str(file),
        filename="test.weird",
        mime_type="application/x-evil",
        size_bytes=4,
        scope="personal",
        scope_id="user_a",
        workspace_id="ws_42",
        team_id="team_5",
        uploaded_by="user_a",
    )

    worker = KbWorker(
        files_store=files_store,
        store=MagicMock(),
        embedder=MagicMock(),
    )
    await worker.parse_and_index(file_id)

    dto = await files_store.get(file_id)
    assert dto is not None
    assert dto.parse_status == "failed"
    assert "unsupported mime" in (dto.parse_error or "")


@pytest.mark.asyncio
async def test_worker_nonexistent_file_id_returns_silently(files_store) -> None:
    """file_id not in DB → log warn but no exception."""
    worker = KbWorker(
        files_store=files_store,
        store=MagicMock(),
        embedder=MagicMock(),
    )
    # Must not raise.
    await worker.parse_and_index("kbf_nonexistent_xxxxx")


@pytest.mark.asyncio
async def test_worker_parse_failed_placeholder_marks_failed(files_store, tmp_path, monkeypatch) -> None:
    """Finding B: a loader returning a parse_failed LoadResult (e.g. a PDF that
    PDFium can't open / scan-only with failed vision) must mark the file
    `failed` with the real reason — NOT silently `done` with a placeholder
    chunk (which hides an unsearchable file from the user)."""
    import bowhead_core.kb.worker as worker_mod
    from bowhead_core.kb.loader.base import LoadResult

    file_id = await files_store.create(
        artifact_id=str(tmp_path / "scan.pdf"),
        filename="scan.pdf",
        mime_type="application/pdf",
        size_bytes=123,
        scope="workspace",
        scope_id="ws_42",
        workspace_id="ws_42",
        team_id="team_5",
        uploaded_by="u",
    )

    class _FailLoader:
        async def load(self, file_path, *, use_llm=False):
            return LoadResult(
                documents=[],
                layer_used="L0",
                duration_ms=1,
                metadata={"parse_failed": True, "reason": "vision_l2_failed: PDFium Data format error"},
            )

    monkeypatch.setattr(worker_mod, "get_loader", lambda *a, **k: _FailLoader())

    worker = KbWorker(
        files_store=files_store,
        store=MagicMock(),
        embedder=MagicMock(),
        artifact_fetcher=None,
    )
    await worker.parse_and_index(file_id)

    dto = await files_store.get(file_id)
    assert dto is not None
    assert dto.parse_status == "failed"
    assert "parse_failed" in (dto.parse_error or "")
    assert not dto.chunks_count


@pytest.mark.asyncio
async def test_load_with_retry_recovers_transient_failure() -> None:
    """#2a: a transient PDFium open failure ('Data format error') retries and
    recovers on a fresh attempt (the batch casualties' actual bytes were valid)."""
    from types import SimpleNamespace

    from bowhead_core.kb.document import KbDocument
    from bowhead_core.kb.loader.base import LoadResult

    calls = []

    class _FlakyLoader:
        async def load(self, file_path, *, use_llm=False):
            calls.append(1)
            if len(calls) == 1:
                return LoadResult(
                    documents=[],
                    layer_used="L0",
                    duration_ms=1,
                    metadata={"parse_failed": True, "reason": "Failed to load document (PDFium: Data format error)."},
                )
            return LoadResult(
                documents=[KbDocument(page_content="recovered content", metadata={})],
                layer_used="L0",
                duration_ms=1,
                metadata={},
            )

    worker = KbWorker(
        files_store=MagicMock(),
        store=MagicMock(),
        embedder=MagicMock(),
        artifact_fetcher=None,
    )
    file = SimpleNamespace(artifact_id="/tmp/x.pdf", use_llm_marker=False, filename="x.pdf")
    lr = await worker._load_with_retry(file, "kbf_x", _FlakyLoader())

    assert len(calls) == 2  # retried once
    assert not lr.metadata.get("parse_failed")  # recovered


@pytest.mark.asyncio
async def test_load_with_retry_no_retry_on_nontransient() -> None:
    """A non-transient failure (e.g. scan-only, no vision client) must NOT retry
    — don't burn a second vision/parse pass on a genuinely unparseable doc."""
    from types import SimpleNamespace

    from bowhead_core.kb.loader.base import LoadResult

    calls = []

    class _ScanOnlyLoader:
        async def load(self, file_path, *, use_llm=False):
            calls.append(1)
            return LoadResult(
                documents=[],
                layer_used="L0",
                duration_ms=1,
                metadata={"parse_failed": True, "reason": "scan_only_no_vision_client"},
            )

    worker = KbWorker(
        files_store=MagicMock(),
        store=MagicMock(),
        embedder=MagicMock(),
        artifact_fetcher=None,
    )
    file = SimpleNamespace(artifact_id="/tmp/y.pdf", use_llm_marker=False, filename="y.pdf")
    lr = await worker._load_with_retry(file, "kbf_y", _ScanOnlyLoader())

    assert len(calls) == 1  # no wasted retry
    assert lr.metadata.get("parse_failed")


@pytest.mark.asyncio
async def test_cleanup_temp_unlinks_only_when_fetcher_set(tmp_path) -> None:
    """#2: _fetch_artifact_path writes NamedTemporaryFile(delete=False) → the
    worker must unlink it (a big bundle otherwise leaks dozens of temp files).
    In test mode (artifact_fetcher=None) file_path is the caller's own file —
    must NOT be deleted."""
    real = tmp_path / "blob.pdf"
    real.write_bytes(b"%PDF-1.4 data")

    # fetcher set → delete
    w1 = KbWorker(files_store=MagicMock(), store=MagicMock(), embedder=MagicMock(), artifact_fetcher=MagicMock())
    f = tmp_path / "scratch.tmp"
    f.write_bytes(b"x")
    w1._cleanup_temp(str(f))
    assert not f.exists()

    # fetcher None → keep
    w2 = KbWorker(files_store=MagicMock(), store=MagicMock(), embedder=MagicMock(), artifact_fetcher=None)
    w2._cleanup_temp(str(real))
    assert real.exists()  # caller's own file untouched


@pytest.mark.asyncio
async def test_load_with_retry_uses_structured_transient_flag() -> None:
    """Retry fires on the structured metadata['transient'] flag even when the
    reason text carries no known marker — decouples retry from exact wording."""
    from types import SimpleNamespace

    from bowhead_core.kb.document import KbDocument
    from bowhead_core.kb.loader.base import LoadResult

    calls = []

    class _L:
        async def load(self, file_path, *, use_llm=False):
            calls.append(1)
            if len(calls) == 1:
                return LoadResult(
                    documents=[],
                    layer_used="L0",
                    duration_ms=1,
                    metadata={"parse_failed": True, "reason": "opaque", "transient": True},
                )
            return LoadResult(
                documents=[KbDocument(page_content="ok", metadata={})],
                layer_used="L0",
                duration_ms=1,
                metadata={},
            )

    worker = KbWorker(files_store=MagicMock(), store=MagicMock(), embedder=MagicMock(), artifact_fetcher=None)
    file = SimpleNamespace(artifact_id="/tmp/z.pdf", use_llm_marker=False, filename="z.pdf")
    lr = await worker._load_with_retry(file, "kbf_z", _L())
    assert len(calls) == 2
    assert not lr.metadata.get("parse_failed")


@pytest.mark.asyncio
async def test_reindex_file_chunks_is_version_aware_and_does_not_update_parse_status(
    monkeypatch,
) -> None:
    from types import SimpleNamespace

    file = SimpleNamespace(
        file_id="kbf_a",
        version=3,
        soft_deleted_at=None,
        mime_type="text/plain",
    )
    files_store = MagicMock()
    files_store.get = AsyncMock(side_effect=[file, file, file])
    store = MagicMock()
    store.hard_delete_file = AsyncMock(return_value=2)
    store.add_chunks = AsyncMock()
    store.count_chunks = AsyncMock(return_value=2)
    worker = KbWorker(files_store=files_store, store=store, embedder=MagicMock())
    load_result = SimpleNamespace(metadata={})
    docs = [SimpleNamespace(page_content="a"), SimpleNamespace(page_content="b")]
    embeddings = [[1.0], [2.0]]
    chunks = [object(), object()]
    worker._load_with_retry = AsyncMock(return_value=load_result)  # type: ignore[method-assign]
    worker._split = AsyncMock(return_value=docs)  # type: ignore[method-assign]
    worker._embed_all = AsyncMock(return_value=embeddings)  # type: ignore[method-assign]
    worker._build_kb_chunks = MagicMock(return_value=chunks)  # type: ignore[method-assign]
    monkeypatch.setattr("bowhead_core.kb.worker.get_loader", lambda *_args, **_kwargs: object())

    indexed = await worker.reindex_file_chunks("kbf_a", 3)

    assert indexed is True
    store.hard_delete_file.assert_awaited_once_with(file)
    store.add_chunks.assert_awaited_once_with(file, chunks)
    assert {call[0] for call in files_store.mock_calls} == {"get"}


def _load_result(text: str):
    from types import SimpleNamespace

    return SimpleNamespace(documents=[SimpleNamespace(page_content=text, metadata={})], metadata={})


@pytest.mark.asyncio
async def test_split_caps_plain_text_without_headings() -> None:
    """没有标题的 .txt 不能整份成为一个 chunk（116 实测：3MB 文件只得 1 个 chunk，rerank 整次被拒）。"""
    from bowhead_core.kb.splitter.paragraph_aware import DEFAULT_CHUNK_SIZE, count_tokens

    text = "\n\n".join(f"第{i}段：华东区门店营业额同比增长，会员复购率在第四季度回升，生鲜促销贡献最大。" * 6 for i in range(200))
    worker = KbWorker(files_store=MagicMock(), store=MagicMock(), embedder=MagicMock())
    file = MagicMock(mime_type="text/plain")

    docs = await worker._split(file, _load_result(text))

    assert len(docs) > 10
    assert max(count_tokens(d.page_content) for d in docs) <= DEFAULT_CHUNK_SIZE
    assert "第199段" in docs[-1].page_content


@pytest.mark.asyncio
async def test_split_keeps_short_markdown_sections_whole_and_carries_headings_into_long_ones() -> None:
    long_body = "\n\n".join("会员复购率在第四季度出现明显回升，主要来自生鲜品类的促销活动。" * 8 for _ in range(60))
    text = f"# 总览\n\n短内容。\n\n## 会员\n\n{long_body}\n"
    worker = KbWorker(files_store=MagicMock(), store=MagicMock(), embedder=MagicMock())
    file = MagicMock(mime_type="text/markdown")

    docs = await worker._split(file, _load_result(text))

    assert docs[0].page_content.startswith("# 总览")
    assert "短内容。" in docs[0].page_content
    assert "会员复购率" not in docs[0].page_content
    member = docs[1:]
    assert len(member) > 1
    assert all(d.metadata.get("h2") == "会员" for d in member)


@pytest.mark.asyncio
async def test_missing_original_file_fails_as_source_missing(files_store) -> None:
    """原文件不在存储里：归为 source_missing（前端不给重试），而不是通用解析失败。"""
    from types import SimpleNamespace

    from bowhead_core.kb.parse_recovery import parse_error_code

    file_id = await files_store.create(
        artifact_id="art_gone",
        filename="deck.pptx",
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        size_bytes=100,
        scope="personal",
        scope_id="user_a",
        workspace_id="ws_1",
        team_id="team_1",
        uploaded_by="user_a",
    )
    backend = MagicMock()
    backend.read = AsyncMock(side_effect=FileNotFoundError("gone"))
    fetcher = MagicMock(backend=backend)
    fetcher.get = AsyncMock(return_value=SimpleNamespace(path="artifacts-116/art_gone_deck.pptx"))
    worker = KbWorker(files_store=files_store, store=MagicMock(), embedder=MagicMock(), artifact_fetcher=fetcher)

    await worker.parse_and_index(file_id)

    dto = await files_store.get(file_id)
    assert (dto.parse_status, dto.parse_error_code) == ("failed", "source_missing")
    # Other read errors remain generic failures.
    assert parse_error_code("failed", "[Errno 13] Permission denied") == "parse_failed"
