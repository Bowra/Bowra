"""Confirmed Office upload, explicit read, original-card resume without rereading."""

import io
import json
import sqlite3

import pytest
import respx
from bowhead_core.artifact_store import ArtifactStore, LocalBackend
from bowhead_core.documents.package_io import OFFICE_MIMES
from bowhead_gateway.routes import files_presign
from test_reference_http_start import answer_response, rows, settled
from test_reference_lead_prompt import http_start as http_start
from test_reference_lead_prompt import lead_prompt as lead_prompt
from test_reference_lead_prompt import prohibit_unmocked_network as prohibit_unmocked_network
from test_reference_lead_prompt import prompt_http as prompt_http
from test_reference_lead_prompt import receiver_case as receiver_case
from test_reference_lead_prompt import start_case as start_case
from test_skill_extension_runtime import tool_response

pytestmark = [
    pytest.mark.integration,
    pytest.mark.asyncio,
    pytest.mark.parametrize("start_case", ["team"], indirect=True),
    pytest.mark.parametrize("receiver_case", ["admission_office_unused_sop"], indirect=True),
]


def office_bytes(kind):
    output = io.BytesIO()
    if kind == "docx":
        from docx import Document

        document = Document()
        document.add_paragraph("Original Office profit is 42.")
    else:
        from pptx import Presentation

        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[1])
        slide.shapes.title.text = "Original Office profit is 42."
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize("kind,changed", [("docx", False), ("pptx", False), ("docx", True), ("pptx", True)])
async def test_office_read_survives_original_question(lead_prompt, monkeypatch, kind, changed):
    p = lead_prompt
    case = p.case
    backend = LocalBackend(
        case.path.parent / "office-uploads",
        allow_local_receiver=True,
        gateway_base_url="http://test",
        local_presign_secret="office-native-test-secret-32-characters",
    )
    uploads = ArtifactStore(str(case.path), backend)
    case.app.state.artifact_store = uploads
    case.app.state.artifact_backend = backend
    case.app.state.orchestrator.artifact_store = uploads
    case.app.include_router(files_presign.router)
    content = office_bytes(kind)
    space = {"X-Space-Id": case.fixture.scope.workspace_id, "Content-Type": "application/json"}
    presign = await case.client.post(
        "/api/v1/files/presign",
        headers=space,
        json={"filename": "profit." + kind, "mime": OFFICE_MIMES["." + kind], "size_bytes": len(content)},
    )
    assert presign.status_code == 200, presign.text
    artifact_id = presign.json()["file_id"]
    await backend.put(aid=artifact_id, filename="profit." + kind, content=content, mime=OFFICE_MIMES["." + kind])
    confirmed = await case.client.post(presign.json()["confirm_url"], headers=space, json={})
    assert confirmed.status_code == 200, confirmed.text
    art = await uploads.get(artifact_id)
    blob_reads = []
    read_bounded = backend.read_bounded

    async def read_blob(path, *, limit):
        blob_reads.append(path)
        return await read_bounded(path, limit=limit)

    monkeypatch.setattr(backend, "read_bounded", read_blob)
    body = {
        **p.body,
        "attachments": [artifact_id],
        "attachment_refs": [{"fileId": art.id, "filename": art.filename, "mime": art.mime, "size": art.size, "referenceKind": "upload"}],
    }
    calls = []
    parses = []
    from bowhead_core import attachment_read

    extract = attachment_read.extract_attachment_bounded

    async def parsed(*args):
        parses.append(args[0].id)
        return await extract(*args)

    monkeypatch.setattr(attachment_read, "extract_attachment_bounded", parsed)

    def respond(request):
        model = json.loads(request.content)
        names = {tool["function"]["name"] for tool in model.get("tools", [])}
        if "subagent_spawn" in names:
            return tool_response(
                "subagent_spawn",
                {"agent_id": p.agent.id, "task": "Read the attached source document and analyze its profit.", "delegation_role": "primary"},
                stream=model.get("stream", False),
            )
        if "attachment_read" not in names:
            return answer_response(request)
        calls.append(model)
        if len(calls) == 1:
            assert "Original Office profit" not in json.dumps(model)
            assert blob_reads == []
            return tool_response("attachment_read", {"artifact_id": artifact_id}, stream=model.get("stream", False))
        assert "Original Office profit is 42." in json.dumps(model), model
        if len(calls) == 2:
            return tool_response(
                "subagent_return",
                {
                    "payload": {
                        "error": {
                            "kind": "needs_user_input",
                            "message": "Choose period.",
                            "clarification_payload": {"questions": [{"question": "Which period?", "type": "free_text", "options": []}]},
                        }
                    }
                },
                stream=model.get("stream", False),
            )
        assert "FY2026" in json.dumps(model)
        return answer_response(request)

    with respx.mock() as mock:
        mock.post(url__regex=r"https://api\..*").mock(side_effect=respond)
        response = await case.client.post("/api/v1/chat/send", json=body)
        assert response.status_code == 200, response.text
        turn_id = response.json()["turn_id"]
        await settled(turn_id)
        turn = await case.fixture.case.store.get_turn(turn_id)
        assert turn.status == "clarification_pending", turn
        cards = [row for row in rows(case, "clarification_requests") if row["turn_id"] == turn_id]
        assert len(cards) == 1
        if changed:
            original = await backend.read_bounded(art.path, limit=art.size)
            from pathlib import Path

            Path(art.path).write_bytes(original[:-1] + bytes([original[-1] ^ 1]))
        resolved = await case.client.post(
            f"/api/v1/clarifications/{cards[0]['request_id']}/resolve",
            json={"chosen_values": ["FY2026"], "answers": [{"question": "Which period?", "source": "custom", "answer": "FY2026"}]},
        )
        assert resolved.status_code == 200, resolved.text
        if changed:
            with pytest.raises(RuntimeError, match="legacy_chat_durability_recovery_handoff"):
                await settled(turn_id)
            assert len(calls) == 2
        else:
            await settled(turn_id)
            assert (await case.fixture.case.store.get_turn(turn_id)).status == "completed"
            assert len(calls) == 3
            facts = await case.client.get(f"/api/v1/turns/{turn_id}/reference-evidence")
            assert facts.status_code == 200, facts.text
            assert facts.json()["continuations"][0]["resume_generation"] == 2
    assert parses == [artifact_id]
    from types import SimpleNamespace

    current = SimpleNamespace(
        session_store=case.fixture.case.store,
        artifact_store=uploads,
        turn_id=turn_id,
        session_id=response.json()["session_id"],
        user_id=case.fixture.scope.user_id,
        workspace_id=case.fixture.scope.workspace_id,
    )
    for field, value in (("user_id", "u_other"), ("workspace_id", "ws_other"), ("session_id", "s_other")):
        with pytest.raises(ValueError, match="attachment_read_unavailable"):
            await attachment_read._original_attachment(SimpleNamespace(**{**vars(current), field: value}), artifact_id)
    unrelated = await uploads.put_bytes(
        content=content,
        filename="unrelated." + kind,
        mime=art.mime,
        created_by=current.user_id,
        workspace_id=current.workspace_id,
    )
    with pytest.raises(ValueError, match="attachment_read_unavailable"):
        await attachment_read._original_attachment(current, unrelated.id)
    with sqlite3.connect(case.path) as db:
        assert db.execute("SELECT count(DISTINCT execution_id) FROM chat_reference_model_invocations WHERE turn_id=?", (turn_id,)).fetchone() == (1,)
