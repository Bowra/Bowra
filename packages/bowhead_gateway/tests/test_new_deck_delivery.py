"""New-deck dispatch reaches a real registered PPTX through the normal Turn writer."""

from __future__ import annotations

import io
import json

from pptx import Presentation
from test_lead_llm_failure_persistence import _read_turn, _wait_done_event
from test_lead_llm_failure_persistence import r4_app as r4_app

from bowhead_router import LLMChunk, LLMResponse, ToolCall


async def test_confirmed_deck_is_registered_and_projected_in_live_and_replay(r4_app):
    app, db_file = r4_app
    orch = app.state.orchestrator
    user_text = "Create a new editable 3-slide PPTX: cover Service Desk Weekly; progress: Completed 20 requests, Closed 8 defects, Released 2 improvements; next week: Improve search, Expand help, Contact 5 users. Blue and white."
    sections = [
        {"kind": "h1", "payload": {"text": "Service Desk Weekly"}},
        {
            "kind": "insight_card",
            "payload": {"title": "Progress", "bullets": ["Completed 20 requests", "Closed 8 defects", "Released 2 improvements"]},
        },
        {"kind": "insight_card", "payload": {"title": "Next week", "bullets": ["Improve search", "Expand help", "Contact 5 users"]}},
    ]
    calls = []
    original_invoke = orch.router.invoke

    async def invoke(agent_type, request, **kwargs):
        names = {tool.name for tool in request.tools or []}
        if "render_deck" in names and "subagent_spawn" not in names:
            phase = "consultant"
        elif "subagent_spawn" in names:
            phase = "lead"
        else:
            return await original_invoke(agent_type, request, **kwargs)
        previous = calls.count(phase)
        calls.append(phase)
        if phase == "lead" and previous == 0:
            return LLMResponse(
                content="",
                tool_calls=[
                    ToolCall(
                        id="spawn",
                        name="subagent_spawn",
                        arguments={
                            "profile": "consultant",
                            "task": user_text,
                            "delivery": "report",
                            "delivery_quote": user_text,
                            "canvas_mode": "new_report",
                        },
                    )
                ],
            )
        if phase == "consultant":
            if previous == 0:
                return LLMResponse(
                    content="",
                    tool_calls=[
                        ToolCall(
                            id="render",
                            name="render_deck",
                            arguments={"sections": sections, "expected_slides": 3, "style": {"accent": "2F6FCE", "background": "FFFFFF"}},
                        )
                    ],
                )
            assert previous == 1, "Native deck closeout must finish after its original render"
            return LLMResponse(
                content="",
                tool_calls=[
                    ToolCall(
                        id="return",
                        name="subagent_return",
                        arguments={"payload": {"delivery_surface": "chat", "exec_summary": "The editable three-slide presentation is ready."}},
                    )
                ],
            )
        spawn_messages = [json.loads(message.content) for message in request.messages if message.role == "tool" and message.name == "subagent_spawn"]
        assert spawn_messages[-1]["delivered"] == "files"
        assert spawn_messages[-1]["deliverables"][0]["artifact_id"].startswith("art_")
        assert "downloadable PPTX" in spawn_messages[-1]["directive"]
        assert "composer" not in spawn_messages[-1]["final_answer"]
        return LLMResponse(content="The editable three-slide presentation is ready.")

    async def stream(agent_type, request, **kwargs):
        response = await invoke(agent_type, request, **kwargs)
        if response.content:
            yield LLMChunk(delta=response.content)
        for call in response.tool_calls:
            yield LLMChunk(tool_call_delta=call)
        yield LLMChunk(finish_reason=response.finish_reason, usage=response.usage)

    orch.router.invoke = invoke
    orch.router.stream = stream
    live = []
    from bowhead_core.types import TurnRequest
    from bowhead_gateway.turn_service import admit_legacy_chat_execution

    users = app.state.user_store
    team = await users.create_team(name="Presentation test", max_seats=2)
    member = await users.create_user(team_id=team.id, email="deck-test@example.com")
    workspace_id = await users.create_local_workspace(team_id=team.id, name="Presentation test", creator_user_id=member.id)
    tid, sid = "t_new_deck", "s_new_deck"
    request = TurnRequest(
        session_id=sid,
        workspace_id=workspace_id,
        user_id=member.id,
        member_uid=member.id,
        text=user_text,
        turn_id=tid,
        team_id=team.id,
        team_router=orch.router,
    )
    bus = await admit_legacy_chat_execution(orch, request)
    try:
        await orch.session_store.ensure_session(session_id=sid, workspace_id=workspace_id, user_id=member.id)
        await orch.session_store.record_turn(turn_id=tid, session_id=sid, user_text=user_text)
        await orch.run_turn(request)
        live.extend(bus.snapshot())
    finally:
        await orch.event_registry.close(tid, expected=bus)
    await _wait_done_event(db_file, tid)
    row, events = _read_turn(db_file, tid)
    assert row["status"] == "completed", [(e["event_type"], e["payload_json"]) for e in events if e["event_type"] in {"turn_failed", "tool_result"}]
    payloads = [(e["event_type"], json.loads(e["payload_json"])) for e in events]
    renders = [p for kind, p in payloads if kind == "tool_result" and p.get("tool_name") == "render_deck"]
    assert len(renders) == 1 and renders[0]["ok"]
    receipt = renders[0]["result_inline"]
    spawn = next(p["result_inline"] for kind, p in payloads if kind == "tool_result" and p.get("tool_name") == "subagent_spawn")
    assert spawn["delivered"] == "files"
    assert spawn["deliverables"][0]["artifact_id"] == receipt["artifact_id"]
    art = await orch.artifact_store.get(receipt["artifact_id"])
    assert art.created_by_turn == tid and art.created_by == member.id
    deck = Presentation(io.BytesIO(await orch.artifact_store.open_read(art)))
    assert len(deck.slides) == 3
    for slide, section in zip(list(deck.slides)[1:], sections[1:], strict=True):
        points = [
            paragraph.text
            for shape in slide.shapes
            if shape.has_text_frame
            for paragraph in shape.text_frame.paragraphs
            if paragraph._p.xpath("./a:pPr/a:buChar")
        ]
        assert points == section["payload"]["bullets"]
    assert calls.count("consultant") == 2
    assert not any(
        kind == "step_started" and p.get("tool_name") in {"bi_query", "skill_read", "skill_list", "python_sandbox"} for kind, p in payloads
    )
    finals = [p for kind, p in payloads if kind == "lead.final.payload"]
    assert receipt["artifact_id"] in json.dumps(finals)
    assert [(e.event_type, e.payload) for e in live if e.event_type in {"lead.final.payload", "done"}] == [
        (kind, p) for kind, p in payloads if kind in {"lead.final.payload", "done"}
    ]
