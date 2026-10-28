# SPDX-License-Identifier: MIT
"""Standalone document CLI using the Bowhead Skill input/workdir protocol."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import sys
import zipfile
from pathlib import Path
from typing import Any

from jsonschema import Draft202012Validator, FormatChecker, ValidationError

# JSON schemas constrain these heterogeneous input/output mappings at the CLI boundary.
JsonObject = dict[str, Any]
MAX_INPUT_BYTES = 2 * 1024 * 1024
MAX_SOURCE_BYTES = 16 * 1024 * 1024
MAX_OUTPUT_BYTES = 16 * 1024 * 1024


def reject_constant(value: str) -> None:
    raise ValueError("Nonfinite JSON numbers are unsupported")


def load_request(path: Path) -> JsonObject:
    if path.stat().st_size > MAX_INPUT_BYTES:
        raise ValueError("Request exceeds the 2 MiB limit")
    raw = path.read_bytes()
    if len(raw) > MAX_INPUT_BYTES:
        raise ValueError("Request exceeds the 2 MiB limit")
    request = json.loads(raw.decode("utf-8"), parse_constant=reject_constant)
    schema = json.loads((Path(__file__).resolve().parents[1] / "input-schema.json").read_text(encoding="utf-8"))
    try:
        Draft202012Validator(schema, format_checker=FormatChecker()).validate(request)
    except ValidationError as exc:
        raise ValueError(f"Request violates the input schema ({exc.validator})") from None
    return request


def source_bytes(workdir: Path, source: str, suffix: str) -> bytes:
    parts = source.split("/")
    if any(part in {"", ".", ".."} for part in parts) or "\\" in source or ":" in source:
        raise ValueError("Source must be a canonical path relative to the workdir")
    candidate = workdir
    for part in parts:
        candidate = candidate / part
        if candidate.is_symlink():
            raise ValueError("Source symlinks are unsupported")
    if candidate.suffix.lower() != suffix or not candidate.is_file():
        raise ValueError(f"Source must be an existing {suffix} file in the workdir")
    if candidate.stat().st_size > MAX_SOURCE_BYTES:
        raise ValueError("Source exceeds the 16 MiB limit")
    raw = candidate.read_bytes()
    if len(raw) > MAX_SOURCE_BYTES:
        raise ValueError("Source exceeds the 16 MiB limit")
    return raw


def check_office_archive(raw: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 4096 or sum(entry.file_size for entry in entries) > 64 * 1024 * 1024:
            raise ValueError("Office package exceeds the expanded-size or entry limit")
        names = [entry.filename for entry in entries]
        if len(set(names)) != len(names):
            raise ValueError("Duplicate Office package entries are unsupported")


def write_output(workdir: Path, name: str, raw: bytes) -> JsonObject:
    if len(raw) > MAX_OUTPUT_BYTES:
        raise ValueError("Output exceeds the 16 MiB limit")
    with (workdir / name).open("xb") as output:
        output.write(raw)
    return {"path": name, "bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}


def encode_json(value: JsonObject) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode("utf-8")


def run(request: JsonObject, workdir: Path) -> JsonObject:
    """Create a document or inspect staged source bytes."""
    if request["operation"] == "create":
        return create(request, workdir)
    raw = source_bytes(workdir, request["source"], SUFFIX)
    if SUFFIX != ".pdf":
        check_office_archive(raw)
    report = inspect(raw)
    report["source_sha256"] = hashlib.sha256(raw).hexdigest()
    report["operation"] = "inspect"
    result = write_output(workdir, "inspection.json", encode_json(report))
    return {**result, "operation": "inspect", "scope": report["scope"]}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--workdir", type=Path, required=True)
    args = parser.parse_args()
    try:
        request = load_request(args.input)
        workdir = args.workdir.resolve()
        workdir.mkdir(parents=True, exist_ok=True)
        if (workdir / "result.json").exists():
            raise FileExistsError("Workdir already contains result.json; choose a fresh output directory")
        result = run(request, workdir)
        write_output(workdir, "result.json", encode_json(result))
        sys.stdout.write(json.dumps(result, ensure_ascii=False, allow_nan=False) + "\n")
    except (ValueError, OSError, zipfile.BadZipFile) as exc:
        sys.stderr.write(json.dumps({"error": type(exc).__name__, "message": str(exc)}, ensure_ascii=False) + "\n")
        return 2
    return 0


from docx import Document
from docx.shared import Inches, Pt

SUFFIX = ".docx"


def create(request: JsonObject, workdir: Path) -> JsonObject:
    document = Document()
    document.core_properties.title = request["title"]
    document.core_properties.author = "Bowra"
    document.core_properties.last_modified_by = "Bowra"
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(request["title"], level=0)
    cell_count = 0
    for block in request["blocks"]:
        kind = block["type"]
        if kind == "heading":
            document.add_heading(block["text"], level=block["level"])
        elif kind == "paragraph":
            document.add_paragraph(block["text"])
        elif kind == "bullets":
            for item in block["items"]:
                document.add_paragraph(item, style="List Bullet")
        else:
            rows = [block["columns"], *block["rows"]]
            if any(len(row) != len(block["columns"]) for row in rows):
                raise ValueError("Every table row must match its column count")
            cell_count += sum(len(row) for row in rows)
            if cell_count > 4000:
                raise ValueError("Document exceeds the 4,000 table-cell limit")
            table = document.add_table(rows=len(rows), cols=len(block["columns"]), style="Light Shading Accent 1")
            for index, row in enumerate(rows):
                for cell, value in zip(table.rows[index].cells, row, strict=True):
                    cell.text = value
                    if index == 0:
                        for paragraph in cell.paragraphs:
                            for text_run in paragraph.runs:
                                text_run.bold = True
    stream = io.BytesIO()
    document.save(stream)
    return {**write_output(workdir, "document.docx", stream.getvalue()), "operation": "create"}


def inspect(raw: bytes) -> JsonObject:
    from docx.table import Table

    document = Document(io.BytesIO(raw))
    blocks: list[JsonObject] = []
    characters = 0
    cells = 0
    for item in document.iter_inner_content():
        if len(blocks) >= 2000:
            raise ValueError("Document exceeds the 2,000 body-block limit")
        if isinstance(item, Table):
            rows = [[cell.text for cell in row.cells] for row in item.rows]
            cells += sum(len(row) for row in rows)
            if cells > 4000:
                raise ValueError("Document exceeds the 4,000 table-cell limit")
            characters += sum(len(cell) for row in rows for cell in row)
            blocks.append({"type": "table", "rows": rows})
        else:
            characters += len(item.text)
            blocks.append({"type": "paragraph", "text": item.text, "style": item.style.name if item.style else ""})
        if characters > 200_000:
            raise ValueError("Document exceeds the 200,000-character inspection limit")
    return {
        "scope": "main-body paragraphs and top-level table text in document order",
        "title": document.core_properties.title,
        "blocks": blocks,
        "limitations": ["No page layout, headers, footers, tracked revisions, comments, drawings or nested-table inventory."],
    }


if __name__ == "__main__":
    raise SystemExit(main())
