# SPDX-License-Identifier: MIT
"""Standalone document CLI using the Bowhead Skill input/workdir protocol."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import sys
import zipfile
from pathlib import Path
from typing import Any

from jsonschema import Draft202012Validator, FormatChecker, ValidationError

# JSON schemas constrain these heterogeneous input/output mappings at the CLI boundary.
JsonObject = dict[str, Any]
MAX_INPUT_BYTES = 2 * 1024 * 1024
MAX_SOURCE_BYTES = 16 * 1024 * 1024
MAX_OUTPUT_BYTES = 16 * 1024 * 1024


def reject_constant(value: str) -> None:
    raise ValueError("Nonfinite JSON numbers are unsupported")


def load_request(path: Path) -> JsonObject:
    if path.stat().st_size > MAX_INPUT_BYTES:
        raise ValueError("Request exceeds the 2 MiB limit")
    raw = path.read_bytes()
    if len(raw) > MAX_INPUT_BYTES:
        raise ValueError("Request exceeds the 2 MiB limit")
    request = json.loads(raw.decode("utf-8"), parse_constant=reject_constant)
    schema = json.loads((Path(__file__).resolve().parents[1] / "input-schema.json").read_text(encoding="utf-8"))
    try:
        Draft202012Validator(schema, format_checker=FormatChecker()).validate(request)
    except ValidationError as exc:
        raise ValueError(f"Request violates the input schema ({exc.validator})") from None
    return request


def source_bytes(workdir: Path, source: str, suffix: str) -> bytes:
    parts = source.split("/")
    if any(part in {"", ".", ".."} for part in parts) or "\\" in source or ":" in source:
        raise ValueError("Source must be a canonical path relative to the workdir")
    candidate = workdir
    for part in parts:
        candidate = candidate / part
        if candidate.is_symlink():
            raise ValueError("Source symlinks are unsupported")
    if candidate.suffix.lower() != suffix or not candidate.is_file():
        raise ValueError(f"Source must be an existing {suffix} file in the workdir")
    if candidate.stat().st_size > MAX_SOURCE_BYTES:
        raise ValueError("Source exceeds the 16 MiB limit")
    raw = candidate.read_bytes()
    if len(raw) > MAX_SOURCE_BYTES:
        raise ValueError("Source exceeds the 16 MiB limit")
    return raw


def check_office_archive(raw: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 4096 or sum(entry.file_size for entry in entries) > 64 * 1024 * 1024:
            raise ValueError("Office package exceeds the expanded-size or entry limit")
        names = [entry.filename for entry in entries]
        if len(set(names)) != len(names):
            raise ValueError("Duplicate Office package entries are unsupported")


def write_output(workdir: Path, name: str, raw: bytes) -> JsonObject:
    if len(raw) > MAX_OUTPUT_BYTES:
        raise ValueError("Output exceeds the 16 MiB limit")
    with (workdir / name).open("xb") as output:
        output.write(raw)
    return {"path": name, "bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}


def encode_json(value: JsonObject) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode("utf-8")


def run(request: JsonObject, workdir: Path) -> JsonObject:
    """Create a document or inspect staged source bytes."""
    if request["operation"] == "create":
        return create(request, workdir)
    raw = source_bytes(workdir, request["source"], SUFFIX)
    if SUFFIX != ".pdf":
        check_office_archive(raw)
    report = inspect(raw)
    report["source_sha256"] = hashlib.sha256(raw).hexdigest()
    report["operation"] = "inspect"
    result = write_output(workdir, "inspection.json", encode_json(report))
    return {**result, "operation": "inspect", "scope": report["scope"]}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--workdir", type=Path, required=True)
    args = parser.parse_args()
    try:
        request = load_request(args.input)
        workdir = args.workdir.resolve()
        workdir.mkdir(parents=True, exist_ok=True)
        if (workdir / "result.json").exists():
            raise FileExistsError("Workdir already contains result.json; choose a fresh output directory")
        result = run(request, workdir)
        write_output(workdir, "result.json", encode_json(result))
        sys.stdout.write(json.dumps(result, ensure_ascii=False, allow_nan=False) + "\n")
    except (ValueError, OSError, zipfile.BadZipFile) as exc:
        sys.stderr.write(json.dumps({"error": type(exc).__name__, "message": str(exc)}, ensure_ascii=False) + "\n")
        return 2
    return 0


from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

SUFFIX = ".pptx"


def create(request: JsonObject, workdir: Path) -> JsonObject:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = request["title"]
    presentation.core_properties.author = "Bowra"
    presentation.core_properties.last_modified_by = "Bowra"
    for slide_index, spec in enumerate(request["slides"], start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(11.9), Inches(1.15)).text_frame
        title.word_wrap = True
        title.text = spec["title"]
        for paragraph in title.paragraphs:
            paragraph.font.name = "Calibri"
            paragraph.font.size = Pt(30)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor.from_string("183153")
        if "bullets" in spec:
            body = slide.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(11.5), Inches(4.8)).text_frame
            body.word_wrap = True
            for index, value in enumerate(spec["bullets"]):
                paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
                paragraph.text = "• " + value
                paragraph.font.name = "Calibri"
                paragraph.font.size = Pt(21)
                paragraph.space_after = Pt(14)
        else:
            rows = [spec["columns"], *spec["rows"]]
            if any(len(row) != len(spec["columns"]) for row in rows):
                raise ValueError("Every slide table row must match its column count")
            table = slide.shapes.add_table(len(rows), len(spec["columns"]), Inches(0.7), Inches(1.9), Inches(11.9), Inches(0.55 * len(rows))).table
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    cell = table.cell(row_index, column_index)
                    cell.text = value
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string("EAF0F8" if row_index == 0 else "FFFFFF")
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.name = "Calibri"
                        paragraph.font.size = Pt(14)
                        paragraph.font.bold = row_index == 0
                        paragraph.font.color.rgb = RGBColor.from_string("183153")
        footer = slide.shapes.add_textbox(Inches(11.9), Inches(7.0), Inches(0.7), Inches(0.3)).text_frame
        footer.text = str(slide_index)
        footer.paragraphs[0].font.size = Pt(10)
    stream = io.BytesIO()
    presentation.save(stream)
    return {**write_output(workdir, "deck.pptx", stream.getvalue()), "operation": "create", "slides": len(presentation.slides)}


def inspect(raw: bytes) -> JsonObject:
    presentation = Presentation(io.BytesIO(raw))
    if len(presentation.slides) > 64:
        raise ValueError("Presentation exceeds the 64-slide inspection limit")
    slides: list[JsonObject] = []
    characters = 0
    shapes = 0
    for index, slide in enumerate(presentation.slides, start=1):
        items: list[JsonObject] = []
        for shape in slide.shapes:
            shapes += 1
            if shapes > 2048:
                raise ValueError("Presentation exceeds the 2,048-shape inspection limit")
            if shape.has_text_frame:
                characters += len(shape.text)
                items.append({"type": "text", "text": shape.text})
            elif shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                characters += sum(len(cell) for row in rows for cell in row)
                items.append({"type": "table", "rows": rows})
            else:
                items.append({"type": "uninspected", "shape_type": str(shape.shape_type)})
            if characters > 200_000:
                raise ValueError("Presentation exceeds the 200,000-character inspection limit")
        slides.append({"slide": index, "items": items})
    return {
        "scope": "top-level slide text and table text; other shapes are identified but not interpreted",
        "slides": slides,
        "limitations": ["No master/layout fidelity, notes, charts, images, animation or grouped-object text inspection."],
    }


if __name__ == "__main__":
    raise SystemExit(main())
