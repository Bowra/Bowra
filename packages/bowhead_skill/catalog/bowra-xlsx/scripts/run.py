# SPDX-License-Identifier: MIT
"""Standalone document CLI using the Bowhead Skill input/workdir protocol."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import sys
import zipfile
from pathlib import Path
from typing import Any

from jsonschema import Draft202012Validator, FormatChecker, ValidationError

# JSON schemas constrain these heterogeneous input/output mappings at the CLI boundary.
JsonObject = dict[str, Any]
MAX_INPUT_BYTES = 2 * 1024 * 1024
MAX_SOURCE_BYTES = 16 * 1024 * 1024
MAX_OUTPUT_BYTES = 16 * 1024 * 1024


def reject_constant(value: str) -> None:
    raise ValueError("Nonfinite JSON numbers are unsupported")


def load_request(path: Path) -> JsonObject:
    if path.stat().st_size > MAX_INPUT_BYTES:
        raise ValueError("Request exceeds the 2 MiB limit")
    raw = path.read_bytes()
    if len(raw) > MAX_INPUT_BYTES:
        raise ValueError("Request exceeds the 2 MiB limit")
    request = json.loads(raw.decode("utf-8"), parse_constant=reject_constant)
    schema = json.loads((Path(__file__).resolve().parents[1] / "input-schema.json").read_text(encoding="utf-8"))
    try:
        Draft202012Validator(schema, format_checker=FormatChecker()).validate(request)
    except ValidationError as exc:
        raise ValueError(f"Request violates the input schema ({exc.validator})") from None
    return request


def source_bytes(workdir: Path, source: str, suffix: str) -> bytes:
    parts = source.split("/")
    if any(part in {"", ".", ".."} for part in parts) or "\\" in source or ":" in source:
        raise ValueError("Source must be a canonical path relative to the workdir")
    candidate = workdir
    for part in parts:
        candidate = candidate / part
        if candidate.is_symlink():
            raise ValueError("Source symlinks are unsupported")
    if candidate.suffix.lower() != suffix or not candidate.is_file():
        raise ValueError(f"Source must be an existing {suffix} file in the workdir")
    if candidate.stat().st_size > MAX_SOURCE_BYTES:
        raise ValueError("Source exceeds the 16 MiB limit")
    raw = candidate.read_bytes()
    if len(raw) > MAX_SOURCE_BYTES:
        raise ValueError("Source exceeds the 16 MiB limit")
    return raw


def check_office_archive(raw: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 4096 or sum(entry.file_size for entry in entries) > 64 * 1024 * 1024:
            raise ValueError("Office package exceeds the expanded-size or entry limit")
        names = [entry.filename for entry in entries]
        if len(set(names)) != len(names):
            raise ValueError("Duplicate Office package entries are unsupported")


def write_output(workdir: Path, name: str, raw: bytes) -> JsonObject:
    if len(raw) > MAX_OUTPUT_BYTES:
        raise ValueError("Output exceeds the 16 MiB limit")
    with (workdir / name).open("xb") as output:
        output.write(raw)
    return {"path": name, "bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}


def encode_json(value: JsonObject) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode("utf-8")


def run(request: JsonObject, workdir: Path) -> JsonObject:
    """Create a document or inspect staged source bytes."""
    if request["operation"] == "create":
        return create(request, workdir)
    raw = source_bytes(workdir, request["source"], SUFFIX)
    if SUFFIX != ".pdf":
        check_office_archive(raw)
    report = inspect(raw)
    report["source_sha256"] = hashlib.sha256(raw).hexdigest()
    report["operation"] = "inspect"
    result = write_output(workdir, "inspection.json", encode_json(report))
    return {**result, "operation": "inspect", "scope": report["scope"]}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--workdir", type=Path, required=True)
    args = parser.parse_args()
    try:
        request = load_request(args.input)
        workdir = args.workdir.resolve()
        workdir.mkdir(parents=True, exist_ok=True)
        if (workdir / "result.json").exists():
            raise FileExistsError("Workdir already contains result.json; choose a fresh output directory")
        result = run(request, workdir)
        write_output(workdir, "result.json", encode_json(result))
        sys.stdout.write(json.dumps(result, ensure_ascii=False, allow_nan=False) + "\n")
    except (ValueError, OSError, zipfile.BadZipFile) as exc:
        sys.stderr.write(json.dumps({"error": type(exc).__name__, "message": str(exc)}, ensure_ascii=False) + "\n")
        return 2
    return 0


from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SUFFIX = ".xlsx"
FORMATS = {"general": "General", "integer": "0", "decimal": "0.00", "percent": "0.00%"}


def create(request: JsonObject, workdir: Path) -> JsonObject:
    workbook = Workbook()
    workbook.remove(workbook.active)
    names: set[str] = set()
    cells = 0
    for spec in request["sheets"]:
        name = spec["name"]
        if name.casefold() in names or any(character in name for character in "[]:*?/\\") or name.startswith("'") or name.endswith("'"):
            raise ValueError("Sheet names must be unique and valid Excel names")
        names.add(name.casefold())
        rows = [spec["columns"], *spec["rows"]]
        if any(len(row) != len(spec["columns"]) for row in rows):
            raise ValueError("Every worksheet row must match its column count")
        formats = spec.get("formats", ["general"] * len(spec["columns"]))
        if len(formats) != len(spec["columns"]):
            raise ValueError("Worksheet formats must match its column count")
        cells += sum(len(row) for row in rows)
        if cells > 50_000:
            raise ValueError("Workbook exceeds the 50,000-cell limit")
        sheet = workbook.create_sheet(name)
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                cell = sheet.cell(row_index, column_index)
                if isinstance(value, dict):
                    cell.value = date.fromisoformat(value["date"])
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(value, str):
                    cell.value = value
                    cell.data_type = "s"
                else:
                    if isinstance(value, float) and not math.isfinite(value):
                        raise ValueError("Nonfinite cell numbers are unsupported")
                    cell.value = value
                if row_index == 1:
                    cell.font = Font(bold=True, color="183153")
                    cell.fill = PatternFill("solid", fgColor="EAF0F8")
                elif not isinstance(value, dict):
                    cell.number_format = FORMATS[formats[column_index - 1]]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_index, heading in enumerate(spec["columns"], start=1):
            has_date = any(isinstance(row[column_index - 1], dict) for row in spec["rows"])
            sheet.column_dimensions[get_column_letter(column_index)].width = min(40, max(20 if has_date else 14, len(heading) + 4))
    stream = io.BytesIO()
    workbook.save(stream)
    return {**write_output(workdir, "workbook.xlsx", stream.getvalue()), "operation": "create", "sheets": len(workbook.worksheets)}


def inspect(raw: bytes) -> JsonObject:
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) > 8:
            raise ValueError("Workbook exceeds the eight-sheet inspection limit")
        sheets: list[JsonObject] = []
        total = 0
        for sheet in workbook.worksheets:
            if sheet.max_row is None or sheet.max_column is None or sheet.max_row > 2001 or sheet.max_column > 32:
                raise ValueError("Worksheet dimensions are missing or exceed 2,001 rows / 32 columns")
            total += sheet.max_row * sheet.max_column
            if total > 50_000:
                raise ValueError("Workbook exceeds the 50,000-cell inspection limit")
            rows: list[list[Any]] = []
            for row in sheet.iter_rows():
                values: list[Any] = []
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f":
                        value = {"formula": value, "calculated": False}
                    elif isinstance(value, (date, datetime)):
                        value = {"date": value.isoformat()}
                    values.append(value)
                rows.append(values)
            sheets.append({"name": sheet.title, "rows": rows})
        return {
            "scope": "worksheet cell values and formula text without calculation or external-link evaluation",
            "sheets": sheets,
            "limitations": ["No chart, drawing, macro, conditional-format or template preservation; formulas are never evaluated."],
        }
    finally:
        workbook.close()


if __name__ == "__main__":
    raise SystemExit(main())
