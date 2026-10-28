"""Build the independently authored, synthetic Community role template.

The only presentation input is python-pptx's bundled empty presentation.
No existing project presentation, artwork, font, or business data is read.
"""

from __future__ import annotations

import argparse
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.slide import Slide
from pptx.util import Inches, Pt

INK = "243247"
MUTED = "5C697B"
ACCENT = "167D8D"
PAPER = "FAFBFD"
CARD = "EDF2F7"
STAMP = datetime(2000, 1, 1)
OUTPUT = Path(__file__).resolve().parents[1] / "templates" / "community-neutral.pptx"


def text(slide: Slide, x: float, y: float, w: float, h: float, size: int = 18, value: str = "", color: str = INK) -> None:
    """Append one styled textbox; even empty slots retain their run style."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = f"Slot {len(slide.shapes) - 1}"
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    run = frame.paragraphs[0].add_run()
    run.text = value
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def block(slide: Slide, x: float, y: float, w: float, h: float, color: str = CARD) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()


def header(slide: Slide) -> None:
    block(slide, 0.5, 0.5, 0.08, 0.42, ACCENT)  # Slot 0.
    text(slide, 0.8, 0.54, 11.7, 0.3, 11, color=ACCENT)  # Section.
    text(slide, 0.8, 1.18, 11.7, 0.95, 30)  # Title.
    text(slide, 0.8, 2.12, 11.7, 0.68, 17, color=MUTED)  # Subtitle.


def cover(slide: Slide) -> None:
    text(slide, 0.8, 1.0, 11.7, 0.4, 12, color=ACCENT)
    text(slide, 0.8, 1.5, 11.7, 0.3, 10)
    text(slide, 0.8, 2.1, 11.7, 1.9, 42)
    text(slide, 0.8, 4.2, 10.8, 1.3, 22, color=MUTED)
    text(slide, 0.8, 6.3, 6, 0.3, 10)
    text(slide, 7.5, 6.3, 5, 0.45, 11, color=MUTED)
    block(slide, 0.8, 0.6, 1.2, 0.09, ACCENT)


def executive(slide: Slide) -> None:
    header(slide)
    for i in range(3):
        x = 0.8 + i * 4.0
        block(slide, x, 3.08, 3.72, 1.72)
        text(slide, x + 0.2, 3.3, 3.32, 0.3, 10, color=ACCENT)
        text(slide, x + 0.2, 3.84, 3.32, 0.75, 18)
    block(slide, 0.8, 5.25, 11.72, 1.35)
    text(slide, 1.0, 5.46, 11.32, 0.3, 10, "RECOMMENDATION", ACCENT)
    text(slide, 1.0, 5.98, 11.32, 0.42, 18)


def kpis(slide: Slide) -> None:
    header(slide)
    for i in range(4):
        x = 0.8 + i * 3.0
        block(slide, x, 3.1, 2.72, 2.25)
        text(slide, x + 0.16, 3.33, 2.4, 0.6, 10, color=MUTED)
        text(slide, x + 0.16, 4.0, 2.4, 0.7, 44)
        text(slide, x + 0.16, 4.86, 2.4, 0.3, 11, color=MUTED)
    text(slide, 0.8, 5.8, 11.7, 0.98, 19, color=MUTED)


def chart(slide: Slide, *, line: bool) -> None:
    header(slide)
    data = CategoryChartData()
    data.categories = ["Example A", "Example B"]
    data.add_series("Synthetic example", [1, 2])
    native = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE if line else XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.8),
        Inches(3.05),
        Inches(7.3),
        Inches(3.65),
        data,
    ).chart
    native.has_legend = False
    native.chart_style = 10
    native.category_axis.tick_labels.font.size = Pt(11)
    native.value_axis.tick_labels.font.size = Pt(10)
    native.value_axis.has_major_gridlines = True
    series = native.series[0]
    series.format.line.color.rgb = RGBColor.from_string(ACCENT)
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = RGBColor.from_string(ACCENT)
    if not line:
        native.plots[0].has_data_labels = True
        native.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    block(slide, 8.55, 3.05, 3.97, 3.65)
    text(slide, 8.78, 3.28, 3.51, 0.32, 10, "INSIGHT", ACCENT)
    text(slide, 8.78, 3.87, 3.51, 1.05, 21)
    text(slide, 8.78, 5.12, 3.51, 1.32, 14, color=MUTED)


def recommendations(slide: Slide) -> None:
    header(slide)
    for i in range(3):
        x = 0.8 + i * 4.0
        block(slide, x, 3.05, 3.72, 3.92)
        text(slide, x, 3.05, 0.1, 0.1)
        text(slide, x, 3.05, 0.1, 0.1)
        text(slide, x + 0.2, 3.27, 3.32, 0.32, 11, color=ACCENT)
        text(slide, x + 0.2, 3.85, 3.32, 0.82, 21)
        text(slide, x + 0.2, 4.82, 3.32, 0.8, 15, color=MUTED)
        text(slide, x + 0.2, 5.81, 0.7, 0.3, 10, "OWNER", ACCENT)
        text(slide, x + 0.99, 5.79, 2.52, 0.46, 12)
        text(slide, x + 0.2, 6.4, 0.7, 0.3, 10, "WHEN", ACCENT)
        text(slide, x + 0.99, 6.38, 2.52, 0.46, 12)


def closing(slide: Slide) -> None:
    text(slide, 0.8, 1.0, 11.7, 0.4, 12, color=ACCENT)
    text(slide, 0.8, 2.5, 11.7, 1.2, 48, "Thank you.")
    text(slide, 0.8, 4.0, 11.7, 0.6)
    text(slide, 0.8, 4.8, 11.7, 0.6)
    text(slide, 0.8, 6.3, 6, 0.3, 10)
    text(slide, 7.5, 6.3, 5, 0.45, 11, color=MUTED)
    block(slide, 0.8, 0.6, 1.2, 0.09, ACCENT)


def canonical_zip(blob: bytes) -> bytes:
    """Normalize ZIP ordering, timestamps and nested workbook metadata."""
    output = BytesIO()
    with ZipFile(BytesIO(blob)) as source, ZipFile(output, "w", ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            if name == "docProps/thumbnail.jpeg":
                continue
            data = source.read(name)
            if name.endswith(".xlsx"):
                data = canonical_zip(data)
            if name == "docProps/core.xml":
                root = etree.fromstring(data)
                for element in root:
                    tag = etree.QName(element).localname
                    if tag in {"created", "modified"}:
                        element.text = "2000-01-01T00:00:00Z"
                    elif tag in {"creator", "lastModifiedBy", "description", "keywords", "subject"}:
                        element.text = ""
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif name == "docProps/app.xml":
                root = etree.fromstring(data)
                for element in list(root):
                    tag = etree.QName(element).localname
                    if tag in {"Company", "Manager", "HyperlinkBase"}:
                        element.text = ""
                    elif tag in {"HeadingPairs", "TitlesOfParts"}:
                        root.remove(element)
                    elif tag == "Application":
                        element.text = "python-pptx"
                    elif tag == "PresentationFormat":
                        element.text = "Widescreen (16:9)"
                    elif tag == "Slides":
                        element.text = "7"
                    elif tag == "TotalTime":
                        element.text = "0"
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif name.endswith(".rels"):
                root = etree.fromstring(data)
                for element in list(root):
                    if element.get("Type", "").endswith("/metadata/thumbnail"):
                        root.remove(element)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            target.writestr(info, data)
    return output.getvalue()


def build() -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    # Empty all library placeholders so shape indices are defined by this file.
    for owner in [*prs.slide_masters, *prs.slide_layouts]:
        for shape in list(owner.shapes):
            shape.element.getparent().remove(shape.element)
    for layout, role in zip(prs.slide_layouts, ["COVER", "TITLE_ONLY", "KPI_GRID", "CHART_SLIDE", "CLOSING"], strict=False):
        layout.name = role
    for layout_index, draw in [(0, cover), (1, executive), (2, kpis), (3, None), (3, None), (1, recommendations), (4, closing)]:
        slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(PAPER)
        if draw is None:
            chart(slide, line=len(prs.slides) == 5)
        else:
            draw(slide)
    props = prs.core_properties
    props.title = "Community neutral role template"
    props.author = props.last_modified_by = props.subject = props.keywords = props.comments = ""
    props.created = props.modified = STAMP
    props.revision = 1
    output = BytesIO()
    prs.save(output)
    return canonical_zip(output.getvalue())


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--check", action="store_true", help="Check reproducibility without replacing the packaged file.")
    args = parser.parse_args()
    data = build()
    if args.check:
        if not args.output.is_file() or args.output.read_bytes() != data:
            parser.exit(1, "Community template differs from its generator output.\n")
    else:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_bytes(data)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
