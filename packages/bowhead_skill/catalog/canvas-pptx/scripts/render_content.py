"""Content-driven standard export, independent of sample-deck slide slots.

This module runs in the existing skill subprocess. It keeps native text, tables
and supported category charts editable without model calls or content summaries.
"""

from __future__ import annotations

import math
import re
import sys
import tempfile
import unicodedata
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

# python-pptx exposes dynamic shape/table/chart objects; Any stays at that boundary.
INK = RGBColor.from_string("1A2925")
ACCENT = RGBColor.from_string("2E5D4F")
MUTED = RGBColor.from_string("65726D")
PAPER = RGBColor.from_string("FFFCF6")
FONT = "Arial Unicode MS" if sys.platform == "darwin" else "Noto Sans CJK SC"
MAX_SLIDES = 500
BODY_LINES = 15
LINE_WIDTH = 76


def _text(value: object) -> str:
    """Keep literal values intact, including code operators and identifiers."""
    return (str(value) if value is not None else "").strip()


def _markdown_text(value: object) -> str:
    """Flatten only fields explicitly carrying Markdown, never arbitrary values."""
    parts = re.split(r"(```[\s\S]*?```|`[^`\n]+`)", _text(value))
    for index, text in enumerate(parts):
        if index % 2 and text.startswith("```"):
            parts[index] = re.sub(r"^```[^\n]*\n", "", text)[:-3]
        elif index % 2:
            parts[index] = text[1:-1]
        else:
            text = re.sub(r"\[figure:[^\]]+\]", "", text)
            text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)
            text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
            text = re.sub(r"^\s{0,3}#{1,6}\s+", "", text, flags=re.MULTILINE)
            parts[index] = re.sub(r"(?<!\w)(\*\*|__)(?=\S)(.+?)(?<=\S)\1(?!\w)", r"\2", text)
    return "".join(parts).strip()


def _width(text: str) -> int:
    return sum(0 if unicodedata.combining(char) else 2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1 for char in text)


def _lines(text: str, width: int) -> list[str]:
    """Wrap without dropping characters, preferring word boundaries."""
    lines: list[str] = []
    for paragraph in text.split("\n"):
        remaining = paragraph
        while _width(remaining) > width:
            used = 0
            end = 0
            for char in remaining:
                used += _width(char)
                if used > width:
                    break
                end += 1
            boundary = remaining.rfind(" ", 0, end + 1)
            if boundary > end // 2:
                end = boundary + 1
            lines.append(remaining[:end])
            remaining = remaining[end:]
        lines.append(remaining)
    return lines


def _style(frame: Any, text: str, size: int, *, color: RGBColor = INK, bold: bool = False) -> None:
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.04)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_before = paragraph.space_after = Pt(0)
        # Fixed leading avoids font-dependent CJK line metrics overflowing a page.
        paragraph.line_spacing = Pt(size * 1.2)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        props = run._r.get_or_add_rPr()
        east_asian = OxmlElement("a:ea")
        east_asian.set("typeface", FONT)
        props.append(east_asian)


def _box(slide: Any, text: str, x: float, y: float, w: float, h: float, size: int, *, color: RGBColor = INK, bold: bool = False) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _style(shape.text_frame, text, size, color=color, bold=bold)


class ContentDeck:
    def __init__(self, data: dict[str, Any], workdir: Path) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.workdir = workdir.resolve()
        self.zh = str(data.get("locale", "zh-CN")).startswith("zh")
        self.title = _text(next((s.get("text") for s in data.get("sections", []) if s.get("kind") == "h1"), None) or data.get("title"))
        self.warnings: list[str] = []

    def label(self, en: str, zh: str) -> str:
        return zh if self.zh else en

    def slide(self, title: str) -> Any:
        if len(self.prs.slides) >= MAX_SLIDES:
            raise ValueError("PPT export exceeds the 500-slide limit; split the report")
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PAPER
        _box(slide, title, 0.7, 0.45, 11.9, 1.08, 28, bold=True)
        return slide

    def heading(self, title: str) -> str:
        if len(_lines(title, 55)) <= 2:
            return title
        # A long heading is content too, never a character-count truncation.
        self.paragraphs(self.label("Report heading", "报告标题"), title)
        return self.label("Report content", "报告内容")

    def continued(self, slide: Any) -> None:
        _box(slide, self.label("Continued", "续页"), 0.75, 7.0, 2.0, 0.3, 10, color=MUTED)

    def paragraphs(self, title: str, text: str) -> None:
        if not text:
            return
        title = self.heading(title)
        pages: list[list[str]] = []
        page: list[str] = []
        for paragraph in text.split("\n\n"):
            wrapped = _lines(paragraph, LINE_WIDTH)
            groups = [wrapped] if len(wrapped) <= BODY_LINES else [_lines(line, LINE_WIDTH) for line in paragraph.split("\n")]
            for group in groups:
                for offset in range(0, len(group), BODY_LINES):
                    chunk = group[offset : offset + BODY_LINES]
                    if page and len(page) + len(chunk) + 1 > BODY_LINES:
                        pages.append(page)
                        page = []
                    if page:
                        page.append("")
                    page.extend(chunk)
        if page:
            pages.append(page)
        for index, lines in enumerate(pages):
            slide = self.slide(title)
            if index:
                self.continued(slide)
            _box(slide, "\n".join(lines), 0.72, 1.65, 11.85, 5.0, 18)

    def cover(self, sections: list[dict[str, Any]]) -> None:
        blocks = []
        for section in sections:
            if section.get("kind") == "meta":
                text = "\n".join(map(_text, section.get("pills") or []))
            else:
                text = _text(section["text"]) if section.get("text") else _markdown_text(section.get("markdown"))
            size = 32 if section.get("kind") == "h1" else 18
            lines = _lines(text, 48 if size == 32 else LINE_WIDTH)
            height = len(lines) * size * 1.2 / 72 + 0.15
            if text:
                blocks.append(("\n".join(lines), size, height))
        if sum(height + 0.25 for _, _, height in blocks) > 5.6:
            self.paragraphs(self.label("Report", "报告"), "\n\n".join(text for text, _, _ in blocks))
            return
        slide = self.slide("")
        top = 1.4
        for text, size, height in blocks:
            _box(slide, text, 0.75, top, 11.8, height, size, color=INK if size == 32 else MUTED, bold=size == 32)
            top += height + 0.25

    def kpis(self, section: dict[str, Any], title: str) -> None:
        cells = section.get("cells") or []
        title = self.heading(title)
        # Avoid an otherwise empty continuation page containing just a fifth KPI.
        page_count = max(1, math.ceil(len(cells) / 4))
        group_size = max(1, math.ceil(len(cells) / page_count))
        for offset in range(0, len(cells), group_size):
            group = cells[offset : offset + group_size]
            # Long labels/values use the paginated text layout at readable size.
            values = [
                "\n".join(_text(c.get(key, c.get(alias))) for key, alias in (("lbl", "label"), ("val", "value"), ("sub", "delta"))) for c in group
            ]
            if any(
                len(_lines(_text(cell.get("lbl", cell.get("label"))), 38)) > 2
                or _width(_text(cell.get("val", cell.get("value")))) > 21
                or len(_lines(_text(cell.get("sub", cell.get("delta"))), 40)) > 2
                for cell in group
            ):
                self.paragraphs(title, "\n\n".join(values))
                continue
            slide = self.slide(title)
            if offset:
                self.continued(slide)
            for index, cell in enumerate(group):
                x = 0.75 + (index % 2) * 6.0
                y = 1.75 + (index // 2) * 2.4
                _box(slide, _text(cell.get("lbl", cell.get("label"))), x, y, 5.55, 0.65, 17, color=MUTED)
                _box(slide, _text(cell.get("val", cell.get("value"))), x, y + 0.72, 5.55, 0.75, 30, color=ACCENT, bold=True)
                _box(slide, _text(cell.get("sub", cell.get("delta"))), x, y + 1.55, 5.55, 0.6, 16, color=MUTED)

    def table(self, section: dict[str, Any], title: str) -> None:
        rows = list(section.get("rows") or [])
        total = section.get("total_row", section.get("totalRow"))
        if isinstance(total, dict):
            rows.append(total)
        columns = section.get("columns") or [{"key": key, "header": key} for key in dict.fromkeys(key for row in rows for key in row)]
        if not columns:
            self.paragraphs(title, self.label("No data rows", "暂无数据"))
            return
        title = self.heading(title)
        for start in range(0, len(columns), 5):
            group = columns[start : start + 5]
            # Repeat the first column when a wide table spans multiple pages.
            if start:
                group = [columns[0], *group]
            keys = [str(col.get("key") or col.get("header") or "") for col in group]
            headers = [
                _text(col.get("header") or key) + (f" ({col['unit']})" if col.get("unit") else "") for col, key in zip(group, keys, strict=True)
            ]
            width = max(12, int(115 / len(group)))
            if any(len(_lines(header, width)) > 3 for header in headers):
                self.paragraphs(
                    title,
                    "\n\n".join(
                        "\n".join(f"{header}: {self.cell_text(row.get(key), col)}" for header, key, col in zip(headers, keys, group, strict=True))
                        for row in rows
                    )
                    or "\n".join(headers),
                )
                continue
            head_height = max(len(_lines(header, width)) for header in headers) * 0.24 + 0.14
            slide = None
            top = 0.0
            for row_index, row in enumerate(rows or [{}]):
                wrapped = [_lines(self.cell_text(row.get(key), col), width) for key, col in zip(keys, group, strict=True)]
                for line_start in range(0, max(map(len, wrapped)), 12):
                    values = ["\n".join(lines[line_start : line_start + 12]) for lines in wrapped]
                    height = max(max(len(lines[line_start : line_start + 12]), 1) for lines in wrapped) * 0.24 + 0.14
                    if slide is None or top + height > 6.7:
                        slide = self.slide(title)
                        if start or row_index or line_start:
                            self.continued(slide)
                        top = 1.65
                        self.table_row(slide, headers, top, head_height, header=True)
                        top += head_height
                    self.table_row(slide, values, top, height)
                    top += height

    @staticmethod
    def cell_text(value: object, column: dict[str, Any]) -> str:
        from render_from_canvas import _excel_number_format

        if isinstance(value, (int, float)) and not isinstance(value, bool):
            # Use the same storage-scale contract in charts and table fallbacks.
            number_format = _excel_number_format(column)
            if number_format == "0.0%":
                return f"{value * 100:g}%"
            if number_format == '0.0"%"':
                return f"{value:g}%"
        return _text(value)

    def table_row(self, slide: Any, values: list[str], top: float, height: float, *, header: bool = False) -> None:
        table = slide.shapes.add_table(1, len(values), Inches(0.75), Inches(top), Inches(11.8), Inches(height)).table
        for cell, value in zip(table.rows[0].cells, values, strict=True):
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if header else RGBColor.from_string("FFFFFF")
            _style(cell.text_frame, value, 12, color=RGBColor.from_string("FFFFFF") if header else INK, bold=header)

    def chart(self, section: dict[str, Any], title: str) -> None:
        # Reuse the existing unit contract, including percent fraction vs points.
        from render_from_canvas import _excel_number_format

        spec = section.get("spec") or {}
        rows = spec.get("rows") or []
        columns = spec.get("columns") or []
        numeric = [c for c in columns if c.get("type") == "numeric"]
        categories = [c for c in columns if c.get("type") != "numeric"]
        category = categories[0] if len(categories) == 1 else None
        types = {
            "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "compare": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "hbar": XL_CHART_TYPE.BAR_CLUSTERED,
            "rank": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE,
            "trend": XL_CHART_TYPE.LINE,
        }
        chart_type = types.get(spec.get("kind"))
        formats = {_excel_number_format(c) for c in numeric}
        units = {c.get("unit", "") for c in numeric}
        valid = rows and numeric and category and chart_type is not None and len(formats) == 1 and len(units) == 1 and len(numeric) <= 6
        if valid:
            for row in rows:
                for col in numeric:
                    value = row.get(col.get("key") or col.get("header"))
                    if value is not None and (isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value)):
                        valid = False
        if not valid:
            self.warnings.append("chart_exported_as_table")
            self.paragraphs(title, self.label("This chart is exported as its data table.", "此图表以数据表形式导出。"))
            self.table(spec, title)
            return
        if any(_width(str(row.get(category.get("key") or category.get("header"), ""))) > 25 for row in rows):
            self.warnings.append("chart_exported_as_table")
            self.paragraphs(
                title,
                self.label("This chart is exported as its data table to keep labels readable.", "此图表以数据表形式导出，以完整保留较长的标签。"),
            )
            self.table(spec, title)
            return
        title = self.heading(title)
        category_key = category.get("key") or category.get("header")
        for offset in range(0, len(rows), 12):
            data = CategoryChartData()
            batch = rows[offset : offset + 12]
            data.categories = [str(row.get(category_key, "")) for row in batch]
            for col in numeric:
                key = col.get("key") or col.get("header")
                data.add_series(str(col.get("header") or key), [row.get(key) for row in batch], number_format=_excel_number_format(col))
            slide = self.slide(title)
            if offset:
                self.continued(slide)
            chart = slide.shapes.add_chart(chart_type, Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.9), data).chart
            chart.has_title = False
            chart.has_legend = len(numeric) > 1
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.include_in_layout = False
                chart.legend.font.size = Pt(12)
                chart.legend.font.name = FONT
            for axis in (chart.category_axis, chart.value_axis):
                axis.tick_labels.font.name = FONT
                axis.tick_labels.font.size = Pt(12)
            chart.value_axis.tick_labels.number_format = next(iter(formats))
            chart.value_axis.tick_labels.number_format_is_linked = False
            if len(numeric) == 1:
                chart.value_axis.has_title = True
                col = numeric[0]
                _style(chart.value_axis.axis_title.text_frame, _text(col.get("header") or col.get("key")), 12, color=MUTED)

    def image(self, section: dict[str, Any], title: str) -> None:
        path = Path(section.get("_local_path") or "").resolve()
        roots = (self.workdir, Path(tempfile.gettempdir()).resolve())
        if not path.is_file() or not any(path.is_relative_to(root) for root in roots):
            self.warnings.append("image_unavailable")
            self.paragraphs(title, self.label("Image unavailable in this export.", "此图片暂无法导出。"))
            return
        title = self.heading(title)
        slide = self.slide(title)
        picture = slide.shapes.add_picture(str(path), Inches(0.8), Inches(1.65))
        scale = min(Inches(11.7) / picture.width, Inches(4.9) / picture.height)
        picture.width = round(picture.width * scale)
        picture.height = round(picture.height * scale)
        picture.left = round((self.prs.slide_width - picture.width) / 2)
        picture.top = Inches(1.65) + round((Inches(4.9) - picture.height) / 2)

    def sections(self, sections: list[dict[str, Any]], *, heading: str = "") -> None:
        heading = heading or self.title or self.label("Report", "报告")
        index = 0
        while index < len(sections):
            section = sections[index]
            kind = section.get("kind")
            index += 1
            title = _text(section.get("title") or section.get("caption")) or heading
            if kind in {"eyebrow", "h1", "deck", "meta"}:
                masthead = [section]
                while index < len(sections) and sections[index].get("kind") in {"eyebrow", "h1", "deck", "meta"}:
                    masthead.append(sections[index])
                    index += 1
                heading = _text(next((item.get("text") for item in masthead if item.get("kind") == "h1"), None)) or heading
                self.cover(masthead)
            elif kind in {"h2", "h3"}:
                heading = _text(section.get("text")) or heading
                self.cover([{"kind": "h1", "text": heading}])
            elif kind == "section":
                chapter = " ".join(_text(section.get(key)) for key in ("label", "h2") if section.get(key)) or heading
                # Child captions cannot substitute for a chapter's scope or units.
                self.cover([{"kind": "h1", "text": chapter}])
                if section.get("body"):
                    self.sections(section["body"], heading=chapter)
            elif kind == "kpi_row":
                self.kpis(section, _text(section.get("title")) or self.label("Key metrics", "关键指标"))
            elif kind == "chart":
                self.chart(section, _text(section.get("title") or (section.get("spec") or {}).get("title")) or title)
            elif kind == "data_table":
                self.table(section, title)
            elif kind == "image":
                self.image(section, title)
            elif kind == "recommendations":
                self.paragraphs(
                    title,
                    "\n\n".join(
                        "\n".join(
                            _text(item[key])
                            for key in ("priority", "action", "rationale", "expected_impact", "owner", "timeline")
                            if item.get(key) is not None
                        )
                        for item in section.get("items") or []
                    ),
                )
            elif kind == "insight_card":
                paragraphs = [_text(section[key]) for key in ("thesis", "narrative", "confidence") if section.get(key) is not None]
                impact = section.get("impact")
                if isinstance(impact, dict):
                    paragraphs.append(
                        " ".join(_text(impact[key]) for key in ("direction", "quantity", "unit", "timeframe", "note") if impact.get(key) is not None)
                    )
                elif impact:
                    paragraphs.append(_text(impact))
                for evidence in section.get("evidence") or []:
                    paragraphs.append("\n".join(_text(evidence[key]) for key in ("caption", "subcaption", "callout_text") if evidence.get(key)))
                self.paragraphs(title, "\n\n".join(paragraphs))
            elif kind == "sources":
                self.paragraphs(
                    _text(section.get("label")) or title,
                    "\n".join(f"[{entry.get('ordinal', i)}] {_text(entry.get('label'))}" for i, entry in enumerate(section.get("entries") or [], 1)),
                )
            elif kind == "next_list":
                self.paragraphs(
                    title,
                    "\n\n".join(
                        "\n".join(_text(item[key]) for key in ("text", "label", "title", "detail", "owner", "timeline") if item.get(key))
                        for item in section.get("items") or []
                    ),
                )
            elif kind == "design":
                self.warnings.append("design_exported_as_text")
                self.paragraphs(
                    title,
                    "\n\n".join(
                        [
                            self.label("Design text only; view the original layout in the report.", "仅导出设计稿文字，原版式请查看在线报告。"),
                            *[_text(layer.get("text")) for layer in section.get("layers") or []],
                        ]
                    ),
                )
            else:
                text = "\n\n".join(
                    _markdown_text(section[key]) if key == "markdown" else _text(section[key])
                    for key in ("label", "markdown", "text", "sql", "code", "caption", "alt", "text_content", "embedded_data_copy")
                    if section.get(key)
                )
                if not text:
                    self.warnings.append(f"unsupported_section:{kind}")
                    text = self.label("This section requires the interactive report.", "此区块需在在线报告中查看。")
                self.paragraphs(title, text)


def render_content(data: dict[str, Any], workdir: Path) -> dict[str, Any]:
    deck = ContentDeck(data, workdir)
    sections = data.get("sections") or []
    deck.sections(sections)
    if not deck.prs.slides:
        deck.paragraphs(deck.title or deck.label("Report", "报告"), deck.label("No content to export.", "暂无可导出的内容。"))
    for index, slide in enumerate(deck.prs.slides, 1):
        _box(slide, f"{index} / {len(deck.prs.slides)}", 11.5, 7.0, 1.1, 0.3, 10, color=MUTED)
    props = deck.prs.core_properties
    props.title = deck.title
    props.author = props.last_modified_by = props.subject = props.comments = ""
    # Remove the library's default editing instruction from all master layouts.
    for layout in deck.prs.slide_layouts:
        for shape in layout.shapes:
            if shape.has_text_frame:
                shape.text_frame.clear()
    output = workdir / "output.pptx"
    deck.prs.save(str(output))
    return {"path": str(output), "slides": len(deck.prs.slides), "sections": len(sections), "demo_residue_count": 0, "warnings": deck.warnings}
