"""脏 excel 读取 helper —— python_sandbox 沙箱内可 import(零 bowhead 依赖,仅 openpyxl+pandas)。

excel RC-6(docs/sprints/kb-optimization/rc6-dirty-table-design.md):compute peer 读
脏银行表(多级表头/合并单元格/标题行/单位行)时手写 pandas 崩溃。本模块提供
`read_dirty_table()` 让 peer 一行读成干净 DataFrame。规整表也能读(退化成普通 read)。

沙箱隔离:此文件由 runner.py 运行时 copy 进 workdir,user 代码 `from dirty_table import
read_dirty_table` 即用。**不 import bowhead_core**(沙箱够不到),逻辑自足。
"""

from __future__ import annotations

import re
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    import pandas as pd

# Unit rows: "单位:万元" / "（金额单位：元）" / a whole-cell parenthetical note without numeric values.
_UNIT_ROW_RE = re.compile(r"单位\s*[:：]|[（(][^）)]*(?:万元|亿元|千元|元|万|亿)[）)]")

_MAX_HEADER_ROWS = 3


def read_dirty_table(path: str, sheet: str | int | None = None, *, max_header_rows: int = _MAX_HEADER_ROWS) -> pd.DataFrame:
    """读脏 excel → 干净 DataFrame。

    处理:合并单元格 forward-fill · 标题行跳过 · 多级表头(≤max_header_rows 行)合成列名 ·
    单位行跳过 · 空行分隔的多 island 取最大表。规整表退化成普通 read。

    Args:
        path: xlsx 路径(通常 INPUTS[key]["path"])。
        sheet: sheet 名 / 0-based 序号 / None(取非空行最多的 sheet)。
        max_header_rows: 多级表头最多几行(默认 3,中国式 group+sub+unit)。

    任何步骤异常 → fallback `pandas.read_excel`(绝不比裸 pandas 更糟)。
    """
    import pandas as pd

    try:
        return _read_dirty_impl(path, sheet, max_header_rows)
    except Exception:
        # Fallback to a normal read (degraded path); if it raises again, let the caller see the actual error.
        return pd.read_excel(path, sheet_name=sheet if sheet is not None else 0)


def _read_dirty_impl(path: str, sheet: str | int | None, max_header_rows: int) -> pd.DataFrame:
    import openpyxl
    import pandas as pd

    # Avoid read_only: ReadOnlyWorksheet has no merged_cells (observed AttributeError),
    # and merged cells cause dirty-table failures. data_only reads computed values instead of formulas.
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = _pick_sheet(wb, sheet)
    _forward_fill_merged(ws)
    rows = [[c.value for c in row] for row in ws.iter_rows()]

    islands = _detect_islands(rows)
    if not islands:
        return pd.DataFrame()
    island = max(islands, key=len)  # The main table for compute analysis is the largest island.
    n_other = len(islands) - 1  # Report how many tables were dropped (review P2-3: the original incident left this requirement partly addressed).

    h_start, h_end = _detect_header_span(island, max_header_rows)
    header_rows = island[h_start:h_end]
    n_cols = max((len(r) for r in header_rows), default=0) or max((len(r) for r in island), default=0)
    columns = _synth_columns(header_rows, n_cols)

    data: list[list[Any]] = []
    for row in island[h_end:]:
        if _is_unit_row(row):
            continue
        padded = (list(row) + [None] * n_cols)[:n_cols]
        data.append(padded)

    df = pd.DataFrame(data, columns=columns)
    df = _coerce_numeric(df)
    # Report n_other additional blank-row-separated tables dropped from this sheet after choosing the largest. For all tables/sheets,
    # user code can switch sheets or split islands itself. Most pandas operations discard df.attrs,
    # so treat it only as an immediate signal.
    if n_other > 0:
        df.attrs["dirty_table_other_islands"] = n_other
    return df


# ── Helpers (pure functions, unit-testable) ───────────────────────────────────────


def _is_empty_cell(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _nonempty(row: list[Any]) -> list[Any]:
    return [c for c in row if not _is_empty_cell(c)]


def _row_is_empty(row: list[Any]) -> bool:
    return not _nonempty(row)


def _looks_numeric(v: Any) -> bool:
    """cell 是数值 或 可解析成数值的字符串(如 '100'/'1,234.5'/'12%')。"""
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    if isinstance(v, str):
        s = v.strip().replace(",", "").replace("%", "").replace("￥", "").replace("$", "")
        if not s:
            return False
        try:
            float(s)
            return True
        except ValueError:
            return False
    return False


def _is_data_row(row: list[Any]) -> bool:
    """多数非空格可解析为数值 → 数据行(不是表头)。用于给 header span 收尾。"""
    cells = _nonempty(row)
    if not cells:
        return False
    return sum(1 for c in cells if _looks_numeric(c)) / len(cells) >= 0.5


def _is_text_heavy(row: list[Any]) -> bool:
    """非空格里字符串占比 ≥0.7 → 文字密集行(用于探多级表头)。数值型字符串('100')不算文字。"""
    cells = _nonempty(row)
    if not cells:
        return False
    text = sum(1 for c in cells if isinstance(c, str) and not _looks_numeric(c))
    return text / len(cells) >= 0.7


def _detect_islands(rows: list[list[Any]]) -> list[list[list[Any]]]:
    """按全空行切成多张表(中国式空行分隔的 island)。"""
    islands: list[list[list[Any]]] = []
    current: list[list[Any]] = []
    for row in rows:
        if _row_is_empty(row):
            if current:
                islands.append(current)
                current = []
            continue
        current.append(row)
    if current:
        islands.append(current)
    return islands


def _is_title_row(row: list[Any]) -> bool:
    """标题行:非空格是单一 distinct 值。覆盖 ①单 cell 标题 ②满宽合并居中标题
    banner(forward-fill 后变 [报表,报表,报表],review P1-1)。"""
    vals = _nonempty(row)
    return bool(vals) and len({str(v).strip() for v in vals}) == 1


def _detect_header_span(island: list[list[Any]], max_header_rows: int) -> tuple[int, int]:
    """(header_start, header_end_exclusive)。第 0 行是标题行(单值/合并 banner) → 表头
    从行 1 起;表头 = 从起点连续 text-heavy 行(≤max_header_rows)。"""
    if not island:
        return (0, 0)
    # Skip row0 only if it is a title and the next row looks like a real header (text-heavy, not data); avoid skipping the sole
    # header as a title (review P1-1: detect merged banners with set-size==1, no longer len==1).
    start = 1 if (len(island) >= 2 and _is_title_row(island[0]) and not _is_data_row(island[0]) and _is_text_heavy(island[1])) else 0
    end = start
    while end < len(island) and end - start < max_header_rows:
        row = island[end]
        # Stop headers at unit/data rows (mostly numeric cells), or text-heavy detection would consume them as headers.
        if _is_unit_row(row) or _is_data_row(row) or not _is_text_heavy(row):
            break
        end += 1
    if end == start:
        end = start + 1  # At least 1 header row.
    return (start, end)


def _is_unit_row(row: list[Any]) -> bool:
    """单位行:含单位关键词且无实数值(单位行不是数据行)。"""
    cells = _nonempty(row)
    if not cells:
        return True  # Skip empty rows too.
    has_unit = any(isinstance(c, str) and _UNIT_ROW_RE.search(c) for c in cells)
    has_number = any(isinstance(c, (int, float)) and not isinstance(c, bool) for c in cells)
    return has_unit and not has_number


def _synth_columns(header_rows: list[list[Any]], n_cols: int) -> list[str]:
    """多行表头 → 单层列名(`本期 / 金额` join),去重(pandas 拒重复列名)。"""
    raw: list[str] = []
    for j in range(n_cols):
        parts: list[str] = []
        for hr in header_rows:
            v = hr[j] if j < len(hr) else None
            s = "" if _is_empty_cell(v) else str(v).strip()
            if s and (not parts or parts[-1] != s):
                parts.append(s)
        raw.append(" / ".join(parts) if parts else f"col{j}")
    seen: dict[str, int] = {}
    out: list[str] = []
    for c in raw:
        if c in seen:
            seen[c] += 1
            out.append(f"{c}.{seen[c]}")
        else:
            seen[c] = 0
            out.append(c)
    return out


def _forward_fill_merged(ws: Any) -> None:
    """合并区左上值填满整个区域(openpyxl 合并区只左上有值)。先 unmerge 再写。"""
    for rng in list(ws.merged_cells.ranges):
        top_left = ws.cell(rng.min_row, rng.min_col).value
        ws.unmerge_cells(str(rng))
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                ws.cell(r, c).value = top_left


def _pick_sheet(wb: Any, sheet: str | int | None) -> Any:
    if isinstance(sheet, str):
        return wb[sheet]
    if isinstance(sheet, int):
        return wb.worksheets[sheet]
    best, best_n = wb.worksheets[0], -1
    for ws in wb.worksheets:
        n = sum(1 for row in ws.iter_rows(values_only=True) if _nonempty(list(row)))
        if n > best_n:
            best, best_n = ws, n
    return best


def _coerce_numeric(df: pd.DataFrame) -> pd.DataFrame:
    """列尽量转数值(pandas 2.2 起 to_numeric(errors='ignore') 弃用,逐列 try)。"""
    import pandas as pd

    for col in df.columns:
        converted = pd.to_numeric(df[col], errors="coerce")
        # Use numeric columns only after successful conversion (not all NaN, and not originally all empty); avoid clearing text columns to NaN.
        if converted.notna().any() and converted.notna().sum() >= df[col].notna().sum() * 0.8:
            df[col] = converted
    return df
