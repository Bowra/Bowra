"""python_sandbox runner — 在 gateway 本机准备输入并委派 executor 执行 agent Python 代码。

调用：
    python3 runner.py --input <workdir>/input.json --workdir <workdir>

输入（input.json 由 SkillRuntime 按 skill.invoke 参数写）:
    {
        "code": "<agent 生成的 python>",
        "inputs": {"name1": {"artifact_id": "art_xxx", "filename": "data.csv"}, ...},
        "timeout_sec": 30,
        # Control keys (`_`-prefixed; set by tools.py, NOT LLM; ignored by
        # _materialize_inputs):
        #   _pi_k_enabled / _pi_k_redact_pii  — π §K INPUTS-schema injection.
        #   _session_id / _turn_id            — pptx-from-template M3w.2 (RD-2):
        #     thread session/turn so outputs are session-pinned + typed at upload
        #     (a .pptx → kind=deliverable, created_by_turn=_turn_id). python_sandbox
        #     only. created_by_turn == the owning Lead turn (peer tools run on
        #     parent_ctx), which is what the is_report_worthy keep-signal queries.
        #   _sandbox_requirements             — pptx-from-template M2 (wheelhouse):
        #     skill-declared long-tail deps (from the resolved SkillSpec, NOT the
        #     LLM). Regex+lockset-validated here; when present + a wheelhouse is
        #     available，executor 会挂载 /wheelhouse:ro 并运行离线 bootstrap。
    }

输出（写到 workdir/result.json）:
    {
        "ok": true, "exit_code": 0,
        "stdout": "...", "stderr": "...",
        "duration_ms": 842,
        "outputs": [{"name":"x.png","artifact_id":"art_new","mime":"image/png","size":1234}, ...]
    }

环境变量（runner 自身接）:
    BOWHEAD_SANDBOX_BACKEND   executor|subprocess 默认 executor；subprocess 仅开发/CI
    BOWHEAD_GATEWAY_URL       默认 http://127.0.0.1:8080 —— 用来 resolve artifact_id
    # pptx-from-template M2 (wheelhouse) — offline hash-pinned long-tail deps:
    BOWHEAD_SANDBOX_WHEELHOUSE   host wheelhouse 目录
                                 默认 UNSET; prod /var/lib/bowhead/wheelhouse.
                                 当 skill 声明 _sandbox_requirements 但此目录
                                 unset/缺失 → 返 ok:false（不静默 ModuleNotFound）。
"""

from __future__ import annotations

import argparse
import contextlib
import hashlib
import json
import mimetypes
import os
import re
import shutil
import socket
import subprocess
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any

# Truncation limit: inline stdout/stderr shown to the agent must not be too large.
_MAX_STDIO_CHARS = 8 * 1024
# Maximum 50MB per output file (needed for large datasets; the former 10 MB limit blocked real analyses).
_MAX_OUTPUT_BYTES = 50 * 1024 * 1024
# Sanitize filenames: strip path separators and control characters.
_UNSAFE_NAME_RE = re.compile(r"[/\\\x00-\x1f\x7f]+")

# ─────────────────────────────────────────────────────────
# pptx-from-template M2 (wheelhouse) — offline hash-pinned long-tail deps.
# ─────────────────────────────────────────────────────────
# Strict requirement grammar — IDENTICAL to sandbox_bootstrap._REQ_RE. Only a
# bare name or `name==version`; the charset blocks argv injection / extra pip
# flags even though deps originate from a SKILL.md on the trusted gateway FS (a
# SKILL.md can be a user-uploaded skill). The wheelhouse install (--no-index
# --require-hashes) + RO root-owned wheelhouse dir + lock-set membership contain
# this: an attacker can at most name an unlocked package (fails closed).
_REQ_RE = re.compile(r"^[A-Za-z0-9._-]+(==[A-Za-z0-9._-]+)?$")


def _norm_req(name: str) -> str:
    """PEP 503 normalize a project name for wheelhouse lock-set membership."""
    return re.sub(r"[-_.]+", "-", name).lower()


def _wheelhouse_lockset(wheelhouse: str | None) -> set[str]:
    """Project-name set from <wheelhouse>/requirements.lock — the reviewed,
    hash-locked names a skill is allowed to request. Absent/empty lock → empty
    set (everything fails closed). Same parse as sandbox_bootstrap._lockset so
    the runner pre-check and the in-container validate agree.
    """
    if not wheelhouse:
        return set()
    p = Path(wheelhouse) / "requirements.lock"
    if not p.is_file():
        return set()
    names: set[str] = set()
    try:
        for raw in p.read_text(encoding="utf-8").splitlines():
            line = raw.strip()
            if not line or line.startswith(("#", "--")):
                continue
            head = line.split("\\", 1)[0].split(";", 1)[0].strip()
            m = re.match(r"^([A-Za-z0-9._-]+)", head)
            if m:
                names.add(_norm_req(m.group(1)))
    except Exception:
        return set()
    return names


def _validate_sandbox_requirements(
    reqs: list,
    lockset: set[str],
) -> tuple[list[str], list[str]]:
    """(valid, errors). IDENTICAL contract to sandbox_bootstrap._validate: each
    requirement must match the strict regex AND its project name must be in the
    wheelhouse lock-set, else it is DROPPED + an error recorded (FAIL CLOSED —
    never install an unvalidated/unlocked dep, never reach the network).
    """
    valid: list[str] = []
    errors: list[str] = []
    for req in reqs or []:
        r = str(req).strip()
        if not r:
            continue
        if not _REQ_RE.match(r):
            errors.append(f"sandbox_requirement {r!r} fails strict regex — dropped")
            continue
        proj = _norm_req(r.split("==", 1)[0])
        if proj not in lockset:
            errors.append(f"sandbox_requirement {r!r} not in wheelhouse lock-set — dropped (fail closed)")
            continue
        valid.append(r)
    return valid, errors


# ─────────────────────────────────────────────────────────
# Sprint π §K — INPUTS schema injection (Layer 1)
# ─────────────────────────────────────────────────────────
# Runtime constants (per feedback_design_no_hard_caps — tweakable, no magic
# numbers buried in code). Defaults match tech-design-r2 §K.2.
_K_MAX_INPUT_PREVIEW_ROWS = 5
_K_MAX_INPUT_PREVIEW_COLS = 20
_K_MAX_INPUT_PREVIEW_CELL_CHARS = 200
_K_MAX_INPUT_PREVIEW_BYTES = 2048
# JSON path: read first 64 KB raw bytes; never load full file
_K_JSON_HEAD_BYTES = 64 * 1024


# K.3 PII column-name regex (tech-design-r2 expanded).
_K_PII_COLUMN_RE = re.compile(
    r"(?:phone|mobile|tel|手机|电话|"
    r"id_?card|身份证|证件|"
    r"email|邮箱|"
    r"address|地址|"
    r"name|姓名|客户名|客户姓名|"
    r"passport|护照|"
    r"bank_?card|银行卡|account|账号|"
    r"wechat|微信|qq|"
    r"salary|薪资|工资|"
    r"patient|病人|医保|"
    r"dob|birthday|出生|"
    r"compound|drug|trial|药品|临床|"
    r"ip_?address|ip(?![a-z])|"
    r"diagnosis|诊断)",
    re.IGNORECASE,
)

# K.3 PII cell-value regex (when column NAME doesn't match, fall back to
# value-shape detection).
_K_PII_VALUE_PATTERNS = {
    "cn_mobile": re.compile(r"^1[3-9]\d{9}$"),
    "cn_id_card": re.compile(r"^\d{17}[\dXx]$"),
    "email": re.compile(r"^[\w.+-]+@[\w-]+\.[\w.-]+$"),
    "cn_passport": re.compile(r"^[GE]\d{8}$"),
    "ipv4": re.compile(r"^\d{1,3}(?:\.\d{1,3}){3}$"),
}


_K_SUFFIX_TO_MIME = {
    ".csv": "text/csv",
    ".tsv": "text/tab-separated-values",
    ".json": "application/json",
    ".jsonl": "application/jsonl",
    ".parquet": "application/x-parquet",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _k_guess_mime(path: Path) -> str | None:
    return _K_SUFFIX_TO_MIME.get(path.suffix.lower())


def _k_redact_cell(value, column_name: str) -> tuple[object, str | None]:
    """Per K.3: (redacted_value, reason or None).

    Column NAMES are preserved — LLM still needs the schema. Only cell VALUES
    are masked.
    """
    if column_name and _K_PII_COLUMN_RE.search(str(column_name)):
        return "<redacted_col>", f"col_match:{column_name}"
    if isinstance(value, str):
        for kind, pat in _K_PII_VALUE_PATTERNS.items():
            if pat.match(value):
                return "<redacted_val>", f"val_match:{kind}"
    return value, None


def _k_clip_cell(v, max_chars: int = _K_MAX_INPUT_PREVIEW_CELL_CHARS) -> str:
    """Stringify cell with hard char cap + ellipsis suffix."""
    s = str(v if v is not None else "")
    if len(s) > max_chars:
        return s[:max_chars] + "…"
    return s


def _k_guess_dtype(values: list) -> str:
    """Best-effort dtype label for LLM hint. Skip masked sentinels so they
    don't poison the guess.
    """
    if not values:
        return "str"
    ints = floats = 0
    for v in values:
        if isinstance(v, str) and v in ("<redacted_col>", "<redacted_val>"):
            continue
        s = str(v if v is not None else "").strip()
        if not s:
            continue
        try:
            int(s)
            ints += 1
        except (ValueError, TypeError):
            try:
                float(s)
                floats += 1
            except (ValueError, TypeError):
                return "str"
    if ints and not floats:
        return "int64"
    if ints or floats:
        return "float64"
    return "str"


#: Helpers copied into workdir on each run (available inside the sandbox via `from <module_name> import ...`).
#: dirty_table = excel RC-6 dirty-table reading; period_growth = comparable periods for growth / YoY rankings.
_SANDBOX_HELPERS = ("dirty_table.py", "period_growth.py")


def _stage_sandbox_helpers(workdir: Path) -> None:
    """把 `_SANDBOX_HELPERS` stage 进 workdir，让沙箱 user 代码直接 import
    （如 `from dirty_table import read_dirty_table`、`from period_growth import compare_growth`）。
    best-effort：copy 失败只 log，user 代码仍可裸 pandas(不回归)。
    随 app 镜像 ship(runner 是 gateway 侧编排器)，无需重 build 沙箱镜像。"""
    here = Path(__file__).resolve().parent
    for name in _SANDBOX_HELPERS:
        src = here / name
        try:
            if src.is_file():
                shutil.copy(src, workdir / name)
        except OSError as e:
            sys.stderr.write(f"[runner] stage {name} failed (non-fatal): {e}\n")


def _rc6_detect_header(buf: list[list[Any]]) -> tuple[int, int, list[str]]:
    """excel RC-6: 用 dirty_table 的纯表头检测给 pi-K schema 产出对的列名。

    复用 sibling dirty_table 模块(runner 是 gateway 侧,能 import 同目录)。
    返回 (h_start, h_end, columns);检测不到/import 失败 → (0, 0, []) 让 caller
    退回原 naive(第 0 行当表头)。best-effort,绝不抛。
    """
    try:
        import sys as _sys
        from pathlib import Path as _P

        _d = str(_P(__file__).resolve().parent)
        if _d not in _sys.path:
            _sys.path.insert(0, _d)
        import dirty_table as _dt  # sibling module

        if not buf:
            return (0, 0, [])
        h_start, h_end = _dt._detect_header_span(buf, _dt._MAX_HEADER_ROWS)
        n_cols = max((len(r) for r in buf[h_start:h_end]), default=0) or max((len(r) for r in buf), default=0)
        cols = _dt._synth_columns(buf[h_start:h_end], n_cols)
        return (h_start, h_end, cols)
    except Exception:
        return (0, 0, [])


def _rc6_is_unit_row(row: list[Any]) -> bool:
    """excel RC-6 P2-4: 复用 dirty_table 的单位行判定(sample 跳单位行)。失败→False。"""
    try:
        import sys as _sys
        from pathlib import Path as _P

        _d = str(_P(__file__).resolve().parent)
        if _d not in _sys.path:
            _sys.path.insert(0, _d)
        import dirty_table as _dt

        return bool(_dt._is_unit_row(row))
    except Exception:
        return False


def _preview_structured_streaming(
    path: Path,
    *,
    mime: str | None = None,
    max_rows: int = _K_MAX_INPUT_PREVIEW_ROWS,
    max_cols: int = _K_MAX_INPUT_PREVIEW_COLS,
    cell_chars: int = _K_MAX_INPUT_PREVIEW_CELL_CHARS,
    byte_cap: int = _K_MAX_INPUT_PREVIEW_BYTES,
    redact_pii: bool = True,
) -> tuple[dict | None, str, list[dict]]:
    """π §K Layer 1 — extract columns + sample + dtypes WITHOUT full-file load.

    Returns (preview_dict, truncate_reason, pii_records).

    preview_dict shape:
        {
          "columns": [...],                       # ≤ max_cols
          "dtypes": {col: "str"/"int64"/...},
          "sample_first_5": [{...}, ...],         # ≤ max_rows, PII-redacted
          "row_count": int | None,
          "preview_truncated": bool,
          "_more_columns_count": int              # only when truncated by col cap
        }

    truncate_reason ∈ {"", "size_cap", "col_cap", "unsupported_mime",
                        "read_error"}.

    pii_records: list of {"column_pattern": ..., "redact_kind": ...} for the
    sidecar metric file (one per masked cell).
    """
    import csv as _csv
    import itertools as _it

    suffix = path.suffix.lower()
    if suffix not in (".csv", ".tsv", ".json", ".jsonl", ".parquet", ".xlsx"):
        # Mime hint can override suffix if explicit
        if not (mime and mime in _K_SUFFIX_TO_MIME.values()):
            return None, "unsupported_mime", []

    header_all: list[str] = []
    header: list[str] = []
    sample_rows_raw: list[list] = []
    row_count: int | None = None

    try:
        if suffix in (".csv", ".tsv"):
            delim = "\t" if suffix == ".tsv" else ","
            with path.open(newline="", encoding="utf-8", errors="replace") as f:
                reader = _csv.reader(f, delimiter=delim)
                header_row = next(reader, None)
                if header_row is None:
                    return None, "read_error", []
                header_all = [str(c) for c in header_row]
                header = header_all[:max_cols]
                # streaming — never `list(reader)`
                for row in _it.islice(reader, max_rows):
                    sample_rows_raw.append(list(row))

        elif suffix in (".json", ".jsonl"):
            head_bytes = path.read_bytes()[:_K_JSON_HEAD_BYTES]
            items_raw: list = []
            try:
                parsed = json.loads(head_bytes.decode("utf-8", errors="replace"))
                if isinstance(parsed, dict) and isinstance(parsed.get("rows"), list):
                    parsed = parsed["rows"]
                if isinstance(parsed, list):
                    items_raw = parsed
            except Exception:
                # JSONL fallback: line-by-line over the head buffer
                for line in head_bytes.decode("utf-8", errors="replace").splitlines():
                    s = line.strip()
                    if not s:
                        continue
                    with contextlib.suppress(Exception):
                        items_raw.append(json.loads(s))
                    if len(items_raw) >= max_rows * 4:
                        break
            if not items_raw or not isinstance(items_raw[0], dict):
                return None, "unsupported_mime", []
            header_set: dict[str, None] = {}
            for it in items_raw[: max_rows * 4]:
                if isinstance(it, dict):
                    for k in it:
                        header_set.setdefault(str(k), None)
            header_all = list(header_set.keys())
            header = header_all[:max_cols]
            sample_rows_raw = [[it.get(h, "") for h in header] for it in items_raw[:max_rows]]

        elif suffix == ".parquet":
            try:
                import pyarrow.parquet as _pq  # type: ignore[import-not-found]
            except ImportError:
                return None, "read_error", []
            try:
                pf = _pq.ParquetFile(str(path))
                schema = pf.schema_arrow
                header_all = [f.name for f in schema]
                header = header_all[:max_cols]
                # SRE P0-4: MUST iter_batches, NOT read_row_group (which loads
                # the entire row group ~128 MB into RAM).
                batch_iter = pf.iter_batches(batch_size=max_rows, columns=header)
                first_batch = next(batch_iter, None)
                if first_batch is not None:
                    sample_dicts = first_batch.to_pylist()[:max_rows]
                    sample_rows_raw = [[d.get(h, "") for h in header] for d in sample_dicts]
                row_count = int(pf.metadata.num_rows) if pf.metadata else None
            except Exception:
                return None, "read_error", []

        elif suffix == ".xlsx":
            try:
                import openpyxl  # type: ignore[import-not-found]
            except ImportError:
                return None, "read_error", []
            try:
                wb = openpyxl.load_workbook(
                    filename=str(path),
                    read_only=True,
                    data_only=True,
                )
                ws = wb.active
                rows_it = ws.iter_rows(values_only=True)
                # excel RC-6: buffer initial rows and use dirty_table header detection (titles/multilevel
                # headers synthesized into column names) to produce the right schema. Otherwise a title header or many "" names
                # leads peers to write brittle pandas code. Keep read_only streaming to avoid memory regressions on large files;
                # if detection fails, fall back to the original naive approach (row 0 as the header).
                _buf = [list(r) for r in _it.islice(rows_it, max(max_rows, 10))]
                _h_start, _h_end, _cols = _rc6_detect_header(_buf)
                if _h_end and _cols:
                    header_all = _cols
                    # Skip unit rows in samples (review P2-4: otherwise sample_first_5[0] may be "单位:万元").
                    _data = [r for r in _buf[_h_end:] if not _rc6_is_unit_row(r)]
                    sample_rows_raw = _data[:max_rows]
                else:
                    header_row = _buf[0] if _buf else None
                    if header_row is None:
                        return None, "read_error", []
                    header_all = [str(c) if c is not None else "" for c in header_row]
                    sample_rows_raw = _buf[1 : 1 + max_rows]
                header = header_all[:max_cols]
                # Continue sampling remaining rows beyond the buffer until max_rows.
                if len(sample_rows_raw) < max_rows:
                    for row in _it.islice(rows_it, max_rows - len(sample_rows_raw)):
                        sample_rows_raw.append(list(row))
                row_count = ws.max_row if ws.max_row else None
                wb.close()
            except Exception:
                return None, "read_error", []
        else:
            return None, "unsupported_mime", []
    except Exception:
        return None, "read_error", []

    more_cols = max(0, len(header_all) - max_cols)

    # PII + clip cells
    pii_records: list[dict] = []
    sample_dicts_out: list[dict] = []
    for row in sample_rows_raw:
        d: dict[str, object] = {}
        for i, h in enumerate(header):
            raw_val = row[i] if i < len(row) else ""
            redacted, reason = _k_redact_cell(raw_val, h) if redact_pii else (raw_val, None)
            if reason is not None:
                kind = "col" if reason.startswith("col_match:") else "val"
                pii_records.append(
                    {
                        "column_pattern": reason.split(":", 1)[1],
                        "redact_kind": kind,
                    }
                )
                d[h] = redacted
            else:
                d[h] = _k_clip_cell(raw_val, max_chars=cell_chars)
        sample_dicts_out.append(d)

    # Dtype guess — per-column over raw sample (skip masked sentinels)
    dtypes: dict[str, str] = {}
    for i, h in enumerate(header):
        col_was_masked = redact_pii and bool(_K_PII_COLUMN_RE.search(str(h)))
        if col_was_masked:
            dtypes[h] = "str"
            continue
        column_vals = [(r[i] if i < len(r) else "") for r in sample_rows_raw]
        dtypes[h] = _k_guess_dtype(column_vals)

    out: dict = {
        "columns": header,
        "dtypes": dtypes,
        "sample_first_5": sample_dicts_out,
        "row_count": row_count,
        "preview_truncated": False,
    }
    if more_cols > 0:
        out["_more_columns_count"] = more_cols
        out["preview_truncated"] = True

    # Hard byte cap (K.2) — if stringified > byte_cap, drop sample rows until
    # under cap. If even with no sample we're over, still return columns
    # (the LLM needs at minimum the schema).
    reason = "col_cap" if more_cols > 0 else ""
    encoded = json.dumps(out, ensure_ascii=False).encode("utf-8")
    if len(encoded) > byte_cap:
        while out["sample_first_5"] and len(json.dumps(out, ensure_ascii=False).encode("utf-8")) > byte_cap:
            out["sample_first_5"].pop()
        out["preview_truncated"] = True
        reason = "size_cap"

    return out, reason, pii_records


def _k_record_metric(workdir: Path, rec: dict) -> None:
    """Append one JSONL line to <workdir>/.k_metrics.json. tools.py drains
    after runner exits and emits Prom counters.

    Best-effort: any IO failure is swallowed (metrics are non-critical).
    """
    try:
        with (workdir / ".k_metrics.json").open("a", encoding="utf-8") as f:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    except Exception:  # noqa: S110
        pass


def _safe_name(name: str, max_len: int = 200) -> str:
    cleaned = _UNSAFE_NAME_RE.sub("_", (name or "").strip()).strip(".") or "file"
    return cleaned[:max_len]


class _NoAuthOnRedirect(urllib.request.HTTPRedirectHandler):
    """cca P3 (F-B review fix): strip Authorization when following a redirect.

    The artifact-fetch Bearer is scoped to the gateway's owner-check route. On the S3/MinIO
    backend the gateway answers the download with a 302 to a SELF-AUTHENTICATING presigned URL
    on a DIFFERENT host. CPython's default redirect handler copies every non-content header
    (incl. Authorization) to the redirect target, so the scoped token would otherwise leak to
    the object-store endpoint (and its access logs / any TLS-terminating proxy). The presigned
    URL needs no credential, so drop Authorization on every redirect hop.
    """

    def redirect_request(self, req, fp, code, msg, headers, newurl):  # type: ignore[no-untyped-def]
        new = super().redirect_request(req, fp, code, msg, headers, newurl)
        if new is not None:
            # Request stores add_header() items capitalized ("Authorization"); drop case-insensitively.
            new.headers = {k: v for k, v in new.headers.items() if k.lower() != "authorization"}
        return new


def _fetch_artifact(artifact_id: str, gateway: str, token: str | None = None) -> bytes:
    """从 gateway 下载 artifact bytes. Gateway `/api/v1/files/{id}` 要么 FileResponse
    要么 302 到 S3 presigned URL —— 用自定义 opener 跟随 redirect，但在跳转时剥掉 Authorization.

    cca P3 F-B: when ``token`` is set (the gateway-minted artifact-fetch token), send it as
    ``Authorization: Bearer`` so the gateway download route runs the owner-check as the real
    caller (not a dropped header / world-readable fetch). Host-side only — never in-container.
    The token is dropped before following the gateway's 302 to the presigned object-store URL
    (_NoAuthOnRedirect, review fix), so the scoped credential never reaches a non-gateway host.
    """
    base = gateway.rstrip("/")
    url = f"{base}/api/v1/files/{artifact_id}"
    req = urllib.request.Request(url, method="GET")  # noqa: S310
    if token:
        req.add_header("Authorization", f"Bearer {token}")
    opener = urllib.request.build_opener(_NoAuthOnRedirect())
    try:
        with opener.open(req, timeout=30) as resp:
            return resp.read()
    except urllib.error.HTTPError as e:
        # This route exists in every build, so a 404 almost certainly means BOWHEAD_GATEWAY_URL is misconfigured.
        # Usually the deployment has a subpath (e.g. https://host/bowhead/) but only the domain was configured here,
        # reaching the frontend SPA 404 page. A customer observed this on 2026-08-18: the same
        # artifact ID returned 404 without the prefix and 200 with it. A bare "HTTP Error 404" hides the cause.
        if e.code == 404:
            raise RuntimeError(
                f"HTTP Error 404: Not Found (gateway={base}) —— 该 artifact 存在但这个 URL 取不到；"
                "多半是 BOWHEAD_GATEWAY_URL 少了部署的 base path（子路径部署要写成 https://<host>/<prefix>），"
                "或指向了前端而不是 app。",
            ) from e
        raise


# Authored-document formats an agent produces AS the user-facing deliverable —
# a deck / spreadsheet / doc / pdf / web page is something it MEANT to hand back.
# Data formats (.csv / .tsv / .json / .jsonl) are deliberately excluded: the
# runner already renders those inline as Data-tab tables (see _preview_structured
# below) and they're routinely intermediate (fetched/scratch data, the run
# manifest) — tagging them deliverable would both double-surface them and flip
# ordinary data-dump turns to report-worthy. Unknown types still upload as normal
# sandbox outputs (visible in the Data tab); they just don't get a chat chip.
_DELIVERABLE_EXTS = (
    ".pptx",
    ".xlsx",
    ".xls",
    ".docx",
    ".doc",
    ".pdf",
    ".html",
    ".htm",
)


def _kind_for_output(
    name: str,
    *,
    purpose: str = "analysis",
    deliverable_outputs: frozenset[str] = frozenset(),
) -> str | None:
    """pptx-from-template M3w.2 (RD-2) / subagent-unification P5a: per-output
    asset kind by filename.

    A produced document — deck / spreadsheet / doc / pdf / html (see
    ``_DELIVERABLE_EXTS``) — keeps its backward-compatible downloadable
    ``deliverable`` classification. Other formats are promoted only when the
    direct python_sandbox call carried explicit ``purpose=deliverable`` *and*
    named this exact top-level output in ``deliverable_outputs``; MIME or extension
    alone never promotes media such as MP4/PNG, and helper files from the same run
    cannot hitchhike into the delivery group. Default analysis files remain normal
    sandbox outputs (None → gateway defaults kind from ``source``).
    The set is about document *format*, not the producing agent — any agent's
    .xlsx/.docx/.pdf downloads, not just the consultant's deck.

    Template-refresh (M4 preview): ``deck-slide-NN.png`` → ``deck_preview`` (the
    ``deck-`` prefix scopes this to verify_deck's per-slide renders so an unrelated
    run's stray ``slide-*.png`` is NOT mis-tagged / made report-worthy).
    """
    n = name.lower()
    if n.startswith("deck-slide-") and n.endswith(".png"):
        return "deck_preview"
    if purpose == "deliverable" and name in deliverable_outputs:
        return "deliverable"
    if n.endswith(_DELIVERABLE_EXTS):
        return "deliverable"
    return None


def _missing_declared_deliverables(
    outputs: list[dict],
    *,
    purpose: str,
    deliverable_outputs: frozenset[str],
) -> list[str]:
    """Return declared final files that were not captured and persisted.

    The declaration is a delivery contract, not merely a classification hint.
    An uploaded artifact record is the only success evidence accepted here; a
    filename that never appeared, was oversized, or failed upload must fail the
    sandbox result instead of letting the model claim a nonexistent download.
    """
    if purpose != "deliverable":
        return []
    produced = {str(output.get("name")) for output in outputs if isinstance(output, dict) and output.get("artifact_id") and output.get("name")}
    return sorted(deliverable_outputs - produced)


# Sandbox/app images both use python:3.12-slim: no /etc/mime.types, and CPython's built-in table **omits** OOXML
# (`mimetypes.MimeTypes(filenames=()).guess_type("a.pptx")` → None). Without an explicit mapping:
# `.pptx/.docx/.xlsx` all resolve to None → upload as octet-stream → Gateway 415 file_type_unsupported.
# Try the explicit table first, then mimetypes.
_OOXML_MIME_BY_SUFFIX: dict[str, str] = {
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _guess_upload_mime(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    explicit = _OOXML_MIME_BY_SUFFIX.get(suffix)
    if explicit:
        return explicit
    mime, _ = mimetypes.guess_type(filename)
    return mime or "application/octet-stream"


def _upload_artifact(
    path: Path,
    gateway: str,
    source: str = "python_sandbox",
    *,
    session_id: str | None = None,
    kind: str | None = None,
    created_by_turn: str | None = None,
    derived_from: str | None = None,
    token: str | None = None,
) -> dict | None:
    """把 workdir 里产出的文件 POST /api/v1/files/upload → 拿 artifact_id.

    失败返 ``None`` —— **本函数自己不做终止决定，由调用方定**。别把这句读成「上传失败
    不影响 turn」：现存两个调用方（`_materialize_data_outputs` 与 outputs 收集处）**都**
    会调 `_fail_data_materialization` 判整轮失败，这是有意的 ——

      · stdout / structured 走的是「artifact 即真相」：result 里内联的 ``stdout`` 只是
        **有界的模型预览**（``_stdout_full`` 会被 pop 掉），完整数据只以 artifact 形式存在。
        上传失败还报 ok=true，等于把「数据被静默截断」的成功交给用户。
      · 同 D3.5 deliverable-presence 守卫的立场（见本文件 `_command_mode` 分支）：
        **绝不返 ok:true/outputs:[]**，宁可判失败也不报静默假成功。

    （2026-08-08：原文「runner 不 fail turn」措辞有误导性，实测把人带偏过 —— 依它推断
    「成功的计算被误判失败」并差点改掉这条正确行为。）

    Sprint ι: ``source`` 透传给 gateway，写进 artifacts.source；Data tab
    origin 映射用（python_sandbox → 计算）。

    pptx-from-template M3w.2 (RD-2): ``session_id``/``kind``/``created_by_turn``/
    ``derived_from`` 作为 query params 透传给 ``/files/upload`` (gateway 已在 M1
    支持这些 param 并写进 put_bytes)，让 sandbox 产出能在创建时即 session-pin + 定型
    (deck → kind=deliverable, created_by_turn=本 turn). 全 optional，老路径不变.
    """
    mime = _guess_upload_mime(path.name)
    try:
        # Construct multipart/form-data manually (no dependencies).
        boundary = "----bowheadsandboxboundary"
        body_parts = [
            f"--{boundary}\r\n".encode(),
            (f'Content-Disposition: form-data; name="file"; filename="{path.name}"\r\nContent-Type: {mime}\r\n\r\n').encode(),
            path.read_bytes(),
            f"\r\n--{boundary}--\r\n".encode(),
        ]
        body = b"".join(body_parts)
        # Build query string: source + any of the optional asset-model params.
        from urllib.parse import urlencode

        _qs = {"source": source}
        if session_id:
            _qs["session_id"] = session_id
        if kind:
            _qs["kind"] = kind
        if created_by_turn:
            _qs["created_by_turn"] = created_by_turn
        if derived_from:
            _qs["derived_from"] = derived_from
        _headers = {
            "Content-Type": f"multipart/form-data; boundary={boundary}",
            "X-Bowhead-Upload-Source": source,
        }
        # cca P3 F-B: carry the artifact-fetch token so the sandbox OUTPUT is owned by the
        # REAL caller (sub), symmetric with the fetch side — otherwise a chained
        # output→input read 404s (output owned by dev_default_user, fetched as the real user).
        if token:
            _headers["Authorization"] = f"Bearer {token}"
        req = urllib.request.Request(  # noqa: S310
            f"{gateway.rstrip('/')}/api/v1/files/upload?{urlencode(_qs)}",
            data=body,
            method="POST",
            headers=_headers,
        )
        with urllib.request.urlopen(req, timeout=60) as resp:  # noqa: S310
            data = json.loads(resp.read().decode())
        return {
            "name": path.name,
            "artifact_id": data.get("file_id"),
            "mime": data.get("mime") or mime,
            "size": int(data.get("size") or path.stat().st_size),
        }
    except Exception as e:
        # Some MIME types are outside the Gateway allowlist (e.g. .npy / .parquet); do not crash, fall back to the filename.
        sys.stderr.write(f"[runner] upload failed for {path.name}: {e}\n")
        return None


# Structured previews for the frontend Data tab; generate only for structured tabular files.
_MAX_PREVIEW_ROWS = 50
_MAX_PREVIEW_CELL_CHARS = 200


def _preview_structured(path: Path) -> dict | None:
    """对 .csv / .tsv / .json / .jsonl 读头部转 {columns, rows}. 失败返 None.
    产物足够前端 `data_table` section 渲染；同时保留原 artifact 链接供下载.
    """
    suffix = path.suffix.lower()
    try:
        if suffix in (".csv", ".tsv"):
            import csv as _csv

            delim = "\t" if suffix == ".tsv" else ","
            with path.open(newline="", encoding="utf-8") as f:
                reader = _csv.reader(f, delimiter=delim)
                rows_raw = list(reader)
            if not rows_raw:
                return None
            header = [str(c) for c in rows_raw[0]]
            data_rows = rows_raw[1 : 1 + _MAX_PREVIEW_ROWS]
            truncated = len(rows_raw) - 1 > _MAX_PREVIEW_ROWS
        elif suffix in (".json", ".jsonl"):
            if suffix == ".jsonl":
                items = []
                with path.open(encoding="utf-8") as f:
                    for line in f:
                        s = line.strip()
                        if not s:
                            continue
                        items.append(json.loads(s))
            else:
                payload = json.loads(path.read_text(encoding="utf-8"))
                if isinstance(payload, list):
                    items = payload
                elif isinstance(payload, dict) and isinstance(payload.get("rows"), list):
                    items = payload["rows"]
                else:
                    # Not a list / no rows field: not tabular, so skip it.
                    return None
            if not items or not isinstance(items[0], dict):
                return None
            # union keys as header, stable order
            header_set: dict[str, None] = {}
            for it in items[:1000]:
                if isinstance(it, dict):
                    for k in it:
                        header_set.setdefault(str(k), None)
            header = list(header_set.keys())
            data_rows = [[str(it.get(k, "")) for k in header] for it in items[:_MAX_PREVIEW_ROWS]]
            truncated = len(items) > _MAX_PREVIEW_ROWS
        else:
            return None
    except Exception as e:
        sys.stderr.write(f"[runner] preview_structured failed for {path.name}: {e}\n")
        return None

    # Truncate overly long cells that Canvas data_table cannot display.
    def _clip(v: str) -> str:
        v = str(v)
        return v if len(v) <= _MAX_PREVIEW_CELL_CHARS else v[:_MAX_PREVIEW_CELL_CHARS] + "…"

    return {
        "columns": header,
        "rows": [[_clip(c) for c in row] for row in data_rows],
        "truncated": truncated,
        "row_count_shown": len(data_rows),
    }


_INPUTS_DIRNAME = "inputs"


def _input_target(workdir: Path, name: str, filename: str, taken: set[str]) -> Path:
    """输入文件的落盘位置：`<workdir>/inputs/<key><ext>`.

    key 直接当文件名（清洗后），扩展名沿用原文件的——模型写 pandas 时靠扩展名判读法。
    两个 key 清洗后撞名时加序号，绝不互相覆盖（覆盖 = 一份数据被另一份悄悄顶掉）。
    """
    d = workdir / _INPUTS_DIRNAME
    d.mkdir(parents=True, exist_ok=True)
    stem = _safe_name(name) or "input"
    ext = Path(filename).suffix[:16]
    candidate = f"{stem}{ext}"
    i = 2
    while candidate in taken:
        candidate = f"{stem}_{i}{ext}"
        i += 1
    taken.add(candidate)
    return d / candidate


def _materialize_inputs(
    inputs: dict,
    workdir: Path,
    gateway: str,
    *,
    k_enabled: bool | None = None,
    k_redact_pii: bool | None = None,
    artifact_fetch_token: str | None = None,
) -> tuple[dict, list[str]]:
    """把 inputs 里每个 artifact_id 下载到 workdir. 返新的 inputs manifest + 错误列表.

    π §K (Layer 1): when ``k_enabled`` (env BOWHEAD_PI_K_ENABLED, default ON since 批1;
    tools.py gate decides), each resolved entry is extended with
    ``columns / dtypes / sample_first_5 / row_count / preview_truncated`` so
    the LLM can write pandas code against real column names instead of
    hallucinating them.

    The flag itself is read in tools.py (since allowlist needs ctx.workspace_id
    which the runner doesn't have access to). Runner only checks env defaults
    for the no-context case.
    """
    if k_enabled is None:
        k_enabled = os.environ.get("BOWHEAD_PI_K_ENABLED", "1") == "1"
    if k_redact_pii is None:
        k_redact_pii = os.environ.get("BOWHEAD_PI_K_PII_REDACT", "1") == "1"

    resolved: dict[str, dict] = {}
    errors: list[str] = []
    taken: set[str] = set()  # Deduplicate filenames on disk (two keys may collide after sanitization).
    # K.4: per-materialize-call (= per skill_invoke) cache so same artifact_id
    # mapped to multiple input names only previews once. Cross-invoke caching
    # lives in tools.py (ctx-level, scoped to turn_id).
    preview_cache: dict[str, tuple[dict | None, str, list[dict]]] = {}

    for name, spec in (inputs or {}).items():
        if not isinstance(spec, dict):
            errors.append(f"input '{name}' must be object, got {type(spec).__name__}")
            continue
        aid = spec.get("artifact_id")
        if not aid:
            errors.append(f"input '{name}' missing artifact_id")
            continue
        fn = _safe_name(str(spec.get("filename") or aid))
        # Place files in the fixed `inputs/` subdirectory, named after **the caller-declared key**, rather than
        # the artifact's long, unwieldy original name. Observed on 2026-08-18: the previous path was
        # `/workdir/sales_onc产品线2026年1-7月销售金额大区经理辖区_e094413.csv`;
        # the model could neither remember nor copy it, inventing `INPUTS/sales/xxx.csv` and `/data/inputs/xxx.csv` instead.
        # It expected a directory convention. Now `inputs/sales.csv` is short, predictable, isolated, and easy to guess;
        # `INPUTS[key]['path']` remains the sole authoritative access method; the contract shape is unchanged.
        target = _input_target(workdir, name, fn, taken)
        try:
            data = _fetch_artifact(str(aid), gateway, token=artifact_fetch_token)
            target.write_bytes(data)
            entry: dict = {
                "path": f"/workdir/{_INPUTS_DIRNAME}/{target.name}",
                "artifact_id": aid,
                "filename": spec.get("filename") or fn,
                "size": len(data),
                "sha256": hashlib.sha256(data).hexdigest(),
            }
            if k_enabled:
                mime = spec.get("mime") or _k_guess_mime(target)
                aid_str = str(aid)
                if aid_str in preview_cache:
                    preview, reason, pii_recs = preview_cache[aid_str]
                else:
                    preview, reason, pii_recs = _preview_structured_streaming(
                        target,
                        mime=mime,
                        redact_pii=k_redact_pii,
                    )
                    preview_cache[aid_str] = (preview, reason, pii_recs)
                # Sidecar metric record (tools.py drains).
                mime_label = mime or "unknown"
                if preview is None:
                    _k_record_metric(
                        workdir,
                        {
                            "mime": mime_label,
                            "reason": reason or "unsupported_mime",
                            "schema_present": False,
                            "bytes": 0,
                            "pii": [],
                        },
                    )
                else:
                    # Merge preview fields into INPUTS entry — this is the
                    # whole point of K: LLM sees columns + sample on its
                    # next sandbox call.
                    entry["columns"] = preview.get("columns", [])
                    entry["dtypes"] = preview.get("dtypes", {})
                    entry["sample_first_5"] = preview.get("sample_first_5", [])
                    entry["row_count"] = preview.get("row_count")
                    entry["preview_truncated"] = preview.get(
                        "preview_truncated",
                        False,
                    )
                    if preview.get("_more_columns_count"):
                        entry["_more_columns_count"] = preview["_more_columns_count"]
                    encoded_bytes = len(
                        json.dumps(preview, ensure_ascii=False).encode("utf-8"),
                    )
                    _k_record_metric(
                        workdir,
                        {
                            "mime": mime_label,
                            "reason": reason,
                            "schema_present": True,
                            "bytes": encoded_bytes,
                            "pii": pii_recs,
                        },
                    )
            resolved[name] = entry
        except Exception as e:
            errors.append(f"input '{name}' fetch failed: {e}")
            if k_enabled:
                _k_record_metric(
                    workdir,
                    {
                        "mime": "fetch_failed",
                        "reason": "read_error",
                        "schema_present": False,
                        "bytes": 0,
                        "pii": [],
                    },
                )
    return resolved, errors


# ─────────────────────────────────────────────────────────
# Backend: docker
# ─────────────────────────────────────────────────────────
def _sandbox_not_ready_cause(stderr_lower: str) -> str:
    """把「sandbox 未就绪」细分成运维能直接照做的三类（只回枚举，不回原文）.

    - ``daemon_unreachable``  : docker daemon 没起 / socket 不可达 → 起 daemon、查挂载
    - ``image_missing``       : 沙箱镜像本地没有且拉不到      → 拉镜像 / 校对 tag
    - ``runtime_unregistered``: --runtime 的值 docker 不认识  → 装 gVisor 或改回 runc
    """
    if "unknown runtime" in stderr_lower or "invalid runtime" in stderr_lower:
        return "runtime_unregistered"
    if "unable to find image" in stderr_lower:
        return "image_missing"
    if "cannot connect" in stderr_lower or "daemon" in stderr_lower:
        return "daemon_unreachable"
    return "unknown"


def _friendly_runtime_error(stdout: str, stderr: str, dur: int, exit_code: int, command_mode: bool = False) -> dict | None:
    """Classify a docker-level failure (daemon down / image missing / runsc not
    registered) into the user-facing friendly ok:false. Returns None if it is NOT
    such an error (so the caller surfaces the real exit). Shared by the local
    `docker run` path and the executor path (OQ-1). prod scrubs stdio (host-path
    leak risk via daemon-missing stderr); dev keeps it for diagnosis.

    follow-up② (SRE P2): in command mode the container's aggregated stderr is the
    SKILL's own arbitrary output — a skill that prints the literal "docker" + a
    daemon keyword would otherwise be misclassified (and prod would blank its real
    error, defeating self-correction). Docker-level launch failures always return
    a docker-runner exit code (125 daemon/image/runtime, 126/127 exec) — require it
    in command mode so a skill's own non-zero exit (typically 1/2) surfaces verbatim.
    """
    if command_mode and exit_code not in (125, 126, 127):
        return None
    sl = stderr.lower()
    if not (
        "docker" in sl
        and (
            "daemon" in sl
            or "not found" in sl
            or "cannot connect" in sl
            or "unable to find image" in sl
            or "unknown runtime" in sl
            or "invalid runtime" in sl
        )
    ):
        return None
    _is_prod = os.environ.get("BOWHEAD_ENV") == "prod"
    # cause is **an enum derived from matched keywords**, without host paths/image tags/environment details,
    # so it can be exposed in prod too. A customer observed 7 sandbox failures in one morning on 2026-08-18:
    # all exited immediately with 125, but after prod cleared stdout/stderr, only "code execution currently unsupported" remained.
    # A dead daemon / missing image / unregistered runtime require different operations, yet shared one message;
    # staff had to infer the cause from exit codes and timing. Preserve the cause enum; redact only details.
    return {
        "ok": False,
        "exit_code": exit_code,
        "stdout": "" if _is_prod else _trunc(stdout),
        "stderr": "" if _is_prod else _trunc(stderr),
        "error": "暂不支持代码执行 (sandbox 未就绪), 该问题需手动计算",
        "cause": _sandbox_not_ready_cause(sl),
        "duration_ms": dur,
    }


def _run_via_executor(
    executor_url: str,
    executor_token: str,
    *,
    workdir: Path,
    timeout: int,
    requirements: list[str] | None,
    command_mode: bool = False,
) -> dict:
    """sandbox-hardening 2026-06-17 (Phase 3 de-DooD): POST the run to the
    sandbox-executor sibling, which owns the Docker socket, image, runtime and limits.
    We only hand it the validated workdir + requirements +
    the per-invoke timeout. Executor-unreachable → friendly ok:false (sandbox not
    ready), same fail-closed shape as a missing docker daemon.
    """
    body = json.dumps(
        {
            "workdir": str(workdir.resolve()),
            "requirements": requirements or [],
            "timeout": timeout,
        }
    ).encode("utf-8")
    req = urllib.request.Request(  # noqa: S310
        f"{executor_url.rstrip('/')}/run",
        data=body,
        method="POST",
        headers={"Content-Type": "application/json", "X-Executor-Token": executor_token},
    )
    t0 = time.perf_counter()
    try:
        # executor must outlive the container run; give it the timeout + slack
        with urllib.request.urlopen(req, timeout=timeout + 30) as resp:  # noqa: S310
            result = json.loads(resp.read().decode())
    except Exception as e:
        _is_prod = os.environ.get("BOWHEAD_ENV") == "prod"
        reason = e.reason if isinstance(e, urllib.error.URLError) else e
        if isinstance(e, urllib.error.HTTPError):
            cause = "executor_http_error"
            message = "The sandbox executor returned an HTTP error. Its service is reachable; ask the administrator to inspect its diagnostics."
        elif isinstance(reason, TimeoutError):
            cause = "executor_response_timeout"
            message = "The sandbox executor did not reply before the transport deadline. Execution may still be running; check the existing run before retrying."
        elif isinstance(reason, (ConnectionRefusedError, socket.gaierror)):
            cause = "executor_unreachable"
            message = "The sandbox executor could not be reached. Ask the administrator to check its service and connection."
        elif isinstance(e, (json.JSONDecodeError, UnicodeDecodeError)):
            cause = "executor_invalid_response"
            message = "The sandbox executor returned an unreadable response. Check the existing run and executor diagnostics before retrying."
        else:
            cause = "executor_transport_error"
            message = "The sandbox connection failed without a confirmed execution result. Check the existing run before retrying."
        return {
            "ok": False,
            "exit_code": -1,
            "stdout": "",
            "stderr": "" if _is_prod else f"{cause}: {e}",
            "error": message,
            "cause": cause,
            "execution_state": "not_started" if cause == "executor_unreachable" else "unknown",
            "retry_safe": cause == "executor_unreachable",
            "duration_ms": int((time.perf_counter() - t0) * 1000),
        }
    # Map a docker-level failure surfaced by the executor to the friendly message
    # (OQ-1: runsc-not-registered etc.); otherwise pass the result through, with
    # stdio truncated for inline display.
    friendly = _friendly_runtime_error(
        result.get("stdout") or "",
        result.get("stderr") or "",
        int(result.get("duration_ms") or 0),
        int(result.get("exit_code") or -1),
        command_mode=command_mode,
    )
    if friendly is not None and not result.get("ok"):
        return {**friendly, "execution_state": result.get("execution_state", "unknown"), "timeout_sec": result.get("timeout_sec", timeout)}
    stdout = result.get("stdout") or ""
    result["_stdout_full"] = stdout
    result["stdout"] = _trunc(stdout)
    result["stderr"] = _trunc(result.get("stderr") or "")
    return result


# ─────────────────────────────────────────────────────────
# Backend: subprocess (INSECURE, dev fallback)
# ─────────────────────────────────────────────────────────
def _run_subprocess(*, workdir: Path, timeout: int) -> dict:
    """纯 subprocess 跑 user.py —— **无沙箱隔离**. 仅 dev 急救或 CI 用.
    `BOWHEAD_SANDBOX_BACKEND=subprocess` 显式启用；默认 executor.
    """
    sys.stderr.write("[runner] WARNING: using subprocess backend (NO isolation)\n")
    cmd = [sys.executable, str(workdir / "user.py")]
    # dogfood-0708 F6 follow-up (self-review): docker backend container env is naturally clean (docker run
    # without -e), but subprocess inherits runner env, exposing BOWHEAD_SANDBOX_* orchestration settings
    # allowed through by F6 to user code. They are not secrets (EXECUTOR_TOKEN uses payload, not env,
    # and was not allowlisted), but _build_env intends to hide our config from user code; F6 must not weaken that.
    # Remove BOWHEAD_-prefixed variables to restore this backend's previous invariant (dev/CI only, without isolation).
    _user_env = {k: v for k, v in os.environ.items() if not k.startswith("BOWHEAD_")}
    t0 = time.perf_counter()
    try:
        proc = subprocess.run(  # noqa: S603
            cmd,
            capture_output=True,
            timeout=timeout,
            text=True,
            check=False,
            cwd=str(workdir),
            env=_user_env,
        )
    except subprocess.TimeoutExpired as e:
        stdout = e.stdout or ""
        if isinstance(stdout, bytes):
            stdout = stdout.decode(errors="replace")
        stderr = e.stderr or ""
        if isinstance(stderr, bytes):
            stderr = stderr.decode(errors="replace")
        return {
            "ok": False,
            "exit_code": -1,
            "stdout": _trunc(stdout),
            "_stdout_full": stdout,
            "stderr": _trunc(stderr),
            "error": f"timeout after {timeout}s",
            "duration_ms": int((time.perf_counter() - t0) * 1000),
        }
    dur = int((time.perf_counter() - t0) * 1000)
    stdout = proc.stdout or ""
    return {
        "ok": proc.returncode == 0,
        "exit_code": proc.returncode,
        "stdout": _trunc(stdout),
        "_stdout_full": stdout,
        "stderr": _trunc(proc.stderr or ""),
        "error": None if proc.returncode == 0 else f"exit_code={proc.returncode}",
        "duration_ms": dur,
    }


def _trunc(s: str) -> str:
    if len(s) <= _MAX_STDIO_CHARS:
        return s
    return s[:_MAX_STDIO_CHARS] + f"\n[... {_MAX_STDIO_CHARS} chars truncated ...]"


# ─────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────
# Locate input.json relative to **the script itself**, without hardcoding `/workdir`; valid for both backends:
#   executor    : user.py is in the executor-created container at /workdir/user.py → /workdir/input.json
#   subprocess  : user.py is in a temporary host workdir → input.json in the same directory
# Hardcoded `/workdir` always breaks subprocess (FileNotFoundError; the host has no /workdir),
# causing `BOWHEAD_SANDBOX_BACKEND=subprocess` to fail on **every** call with this scaffolding.
# Locate input metadata independently of the selected execution backend.
# sandbox-input-contract L1: `INPUTS` is a **self-explanatory** dict.
#
# Compare omitted inputs across the two supported code shapes:
#   Hardcoded `pd.read_csv('rows.csv')` → FileNotFoundError → `_hint_missing_file` explains
#     that no inputs were declared and how to declare them; the model fixes it on the next iteration.
#   Documented `INPUTS['rows']['path']` → KeyError → `error` contains only `exit_code=1`.
# Both failures need actionable input-declaration guidance.
#
# Why this belongs in the scaffold rather than `_hint_missing_file`:
#   ① Explain **on the failing line**, without requiring correlation across fields;
#   ② Use **actual runtime state**: separate empty/nonempty messages and list available keys when nonempty,
#      also covering mismatched keys invisible to `_hint_missing_file` (declared as d16,
#      referenced as rows in code);
#   ③ The scaffold is the only convergence point for both entry points and belongs to no agent.
#
# `tools.py` renders `_call_shape` for **the actual entry point** and passes it through a control key. runner is a separate
# process unable to import `bowhead_skill.sandbox_call_shape`, so pass the rendered string, not the entry point name,
# to avoid duplicating rendering logic. If unavailable, use fallback wording valid for both entry points.
_SCAFFOLD = """\
# --- bowhead python_sandbox scaffolding (auto-injected) ---
import difflib as _bh_difflib
import json as _bh_json
import sys as _bh_sys
from pathlib import Path as _bh_Path
with open(_bh_Path(__file__).resolve().parent / 'input.json') as _bh_f:
    _bh_payload = _bh_json.load(_bh_f)
_bh_shape = _bh_payload.get('_call_shape') or "inputs={'<key>': {'artifact_id': 'art_...'}}（与 code 同级）"
_bh_unresolved = _bh_payload.get('_unresolved_inputs') or []


class _BhInputs(dict):
    def __missing__(self, _bh_k):
        # 声明过但一个都没落地 / 部分没落地 —— 病因在平台侧（取不到、形状不合法），
        # **不是**调用方没声明。这条必须排在 `not self` 前面：全部取失败时 INPUTS 同样
        # 为空，若先判空就会把平台的锅说成「你没声明」，模型于是原地重试改代码。
        # A fetch failure must not be described as an absent input declaration.
        if _bh_unresolved:
            _bh_left = ("本次仍可用的 key 是: %s 。" % ", ".join(sorted(self))) if self else "本次没有任何可用输入。"
            raise KeyError(
                "%r —— 声明过的输入 %s **没能取到**（原因见 stderr 末尾的 "
                "[runner] input errors）。这不是调用写法的问题，重写代码没用；"
                "%s" % (_bh_k, ", ".join(_bh_unresolved), _bh_left)
            )
        if not self:
            raise KeyError(
                "%r —— 本次调用没有声明任何输入，INPUTS 是空的。声明方式：%s ；"
                "文件会落在 inputs/<key>.csv，代码里用 INPUTS['<key>']['path'] 读。"
                % (_bh_k, _bh_shape)
            )
        raise KeyError(
            "%r —— 本次可用的 key 是: %s 。key 是**调用时你自己起的名字**，"
            "不是 artifact 的原文件名。" % (_bh_k, ", ".join(sorted(self)))
        )


INPUTS = _BhInputs(_bh_payload.get('inputs') or {})
# L3：小写别名。`inputs` 正是**工具参数名**，模型写 `inputs['d3']` 是被我们自己的
# 命名不一致诱导的（116 近 30 天 5 次 NameError）。别名让两种写法都对，代价是零；
# 用户代码若自己 `inputs = ...` 会自然覆盖它（scaffold 在前，用户代码在后）。
inputs = INPUTS


# ── L4：列名不存在时，把**真实列名**附在异常后面 ──
# 116 近 30 天：修完 INPUTS 那层（裸 KeyError 09-07 止）之后，头号沙箱失败变成
# 「模型编了个不存在的业务列名」——30/142，缺的键是 '营业额' '品牌' '退款额_25'
# '6月环比%' 这种。其中 12 次发生在 groupby/merge 的**派生帧**上，所以只看入参 CSV
# 表头的主机侧做法覆盖不到，必须在容器内按运行时真实对象说话。
# 不接管 traceback：先让原始异常照常打印，再把提示追加到 stderr 末尾——`_hint_missing_file`
# 的 `[runner] ` 前缀也在那儿，模型已经习惯从末尾读。
def _bh_live_frames(_bh_tb):
    _bh_pd = _bh_sys.modules.get('pandas')
    if _bh_pd is None:
        return []
    _bh_seen = []
    _bh_found = []
    while _bh_tb is not None:
        _bh_fr = _bh_tb.tb_frame
        if _bh_fr.f_code.co_filename.endswith('user.py'):
            for _bh_n, _bh_val in list(_bh_fr.f_locals.items()):
                if _bh_n.startswith('_bh_') or not isinstance(_bh_val, _bh_pd.DataFrame):
                    continue
                if id(_bh_val) in _bh_seen:
                    continue
                _bh_seen.append(id(_bh_val))
                _bh_found.append((_bh_n, [str(_bh_c) for _bh_c in _bh_val.columns]))
        _bh_tb = _bh_tb.tb_next
    return _bh_found[:3]


def _bh_near(_bh_want, _bh_cols):
    # 先看包含关系再看编辑距离：实测最常见的形状是 '营业额' vs '营业额 / 元'（带单位后缀），
    # 而 difflib 对这种长度悬殊的配对打分只有 0.55，默认 cutoff=0.6 正好漏掉。
    _bh_hits = [_bh_c for _bh_c in _bh_cols if _bh_want and (_bh_want in _bh_c or _bh_c in _bh_want)]
    for _bh_c in _bh_difflib.get_close_matches(_bh_want or '', _bh_cols, n=3, cutoff=0.5):
        if _bh_c not in _bh_hits:
            _bh_hits.append(_bh_c)
    return _bh_hits[:3]


def _bh_excepthook(_bh_t, _bh_v, _bh_tb):
    _bh_prev_excepthook(_bh_t, _bh_v, _bh_tb)
    try:
        if not issubclass(_bh_t, KeyError):
            return
        # `_BhInputs.__missing__` 抛的那三档文案已经自解释，别再叠一层把它冲淡。
        _bh_last = _bh_tb
        while _bh_last.tb_next is not None:
            _bh_last = _bh_last.tb_next
        if _bh_last.tb_frame.f_code.co_name == '__missing__':
            return
        _bh_frames = _bh_live_frames(_bh_tb)
        if not _bh_frames:
            return
        _bh_want = str(_bh_v.args[0]) if _bh_v.args else ''
        for _bh_n, _bh_cols in _bh_frames:
            _bh_shown = ', '.join(_bh_c[:40] for _bh_c in _bh_cols[:40])
            if len(_bh_cols) > 40:
                _bh_shown += ' …(共 %d 列)' % len(_bh_cols)
            _bh_msg = '[runner] %s 的真实列名：%s' % (_bh_n, _bh_shown)
            _bh_hit = _bh_near(_bh_want, _bh_cols)
            if _bh_hit:
                _bh_msg += ' ——最接近 %r 的是：%s' % (_bh_want, ', '.join(repr(_bh_h) for _bh_h in _bh_hit))
            print(_bh_msg, file=_bh_sys.stderr)
        print(
            "[runner] 列名要以数据为准，不能照着业务习惯写。先 print(df.columns.tolist()) 看一眼再取。",
            file=_bh_sys.stderr,
        )
    except Exception:
        # 诊断增强绝不允许盖掉用户的真异常：原始 traceback 上面已经打过了。
        pass


_bh_prev_excepthook = _bh_sys.excepthook
_bh_sys.excepthook = _bh_excepthook
# --- user code below ---
"""


# ─────────────────────────────────────────────────────────────
# skill-multiruntime S1 — execute user skill scripts in multiple languages inside the container
# ─────────────────────────────────────────────────────────────
_MAX_SKILL_BUNDLE_BYTES = 16 * 1024 * 1024  # Match the upload quota.
# review-P2: skill_result.json is skill-authored (written in the container to a shared bind mount); limit size before reading
# to prevent malicious multi-GB results from exhausting host memory/disk. Structured returns should stay small.
_MAX_SKILL_RESULT_BYTES = 1 * 1024 * 1024
# Control keys + secrets never pass into container input.json; other user business keys pass to the script.
_SKILL_CONTROL_KEYS = frozenset(
    {
        "code",
        "timeout_sec",
        "_skill_dir",
        "_skill_entry",
        # follow-up② command-mode control keys (server-injected; the `_`-prefix shape guard already excludes them from the container);
        # list them here to preserve SSOT parity.
        "_skill_command",
        "_skill_command_enabled",
        "_skill_argv0_allowlist",
        "_sandbox_requirements",
        "_pi_k_enabled",
        "_pi_k_redact_pii",
        "_session_id",
        "_turn_id",
        "_output_purpose",
        "_deliverable_outputs",
        "_artifact_fetch_token",
        "_sandbox_executor_url",
        "_sandbox_executor_token",
        # L1: entry-specific rendered call skeleton, used only for scaffold KeyError messages.
        "_call_shape",
        # Declared input names that were not materialized; used only for scaffold KeyError messages.
        "_unresolved_inputs",
    }
)

# Trusted dispatch shim runs as user.py in the container. Choose the interpreter by extension (built-in constants, ignoring
# manifest interp), execute skill/scripts/<entry>, pass through stdout/stderr/exit, and rename the child script's
# result.json (host contract) to skill_result.json to prevent the runner envelope overwriting it.
_SKILL_DISPATCH_SHIM = r"""
# --- bowhead skill-multiruntime dispatch shim (auto-injected) ---
import json as _bh_j, os as _bh_os, subprocess as _bh_sp, sys as _bh_sys
from pathlib import Path as _BhPath

_BH_WD = "/workdir"
_BH_EXT_INTERP = {".py": "python3", ".sh": "bash", ".js": "node", ".mjs": "node"}


def _bh_relay_result() -> None:
    # 子脚本若写 /workdir/result.json → 改名 skill_result.json（避免被 runner envelope 覆写）
    child = _BhPath(_BH_WD, "result.json")
    if child.is_file():
        _bh_os.replace(str(child), str(_BhPath(_BH_WD, "skill_result.json")))


def _bh_dispatch() -> int:
    manifest = _bh_j.loads(_BhPath(_BH_WD, "skill_manifest.json").read_text())
    # follow-up② command 模式：跑 model-composed argv（已 host-side 替换 token + argv[0]
    # allowlist + 路径守卫；此处不信任、只执行）。cwd/argv 由 host 侧 runner.main 写死进
    # manifest；shim 只做 subprocess.run(list)（无 shell）。exit code 必透传（result.json.ok
    # 权威=容器 returncode；假成功类 s_3389df456f/t_2527d23b07f7）。
    cmd = manifest.get("command")
    if isinstance(cmd, list) and cmd:
        cwd = str(manifest.get("cwd") or _BH_WD)
        proc = _bh_sp.run([str(x) for x in cmd], cwd=cwd)
        _bh_relay_result()
        return proc.returncode
    # legacy 固定契约默认（`--input/--workdir` + result.json 的 skill 保底）
    entry = str(manifest.get("entry") or "")
    interp = _BH_EXT_INTERP.get(_BhPath(entry).suffix)
    if not interp:
        _bh_sys.stderr.write("[skill-shim] unsupported entry: %r\n" % entry)
        return 2
    script = _BhPath(_BH_WD, "skill", "scripts", entry)
    if not script.is_file():
        _bh_sys.stderr.write("[skill-shim] entry not found: skill/scripts/%s\n" % entry)
        return 2
    proc = _bh_sp.run([interp, str(script), "--input", _BH_WD + "/input.json", "--workdir", _BH_WD], cwd=_BH_WD)
    _bh_relay_result()
    return proc.returncode


_bh_sys.exit(_bh_dispatch())
"""


def _stage_user_skill(workdir: Path, skill_dir: Path, entry: str) -> str | None:
    """把用户 skill bundle 拷进 workdir/skill/（host 侧；workdir 是共享 mount → 容器见
    /workdir/skill/）。返回错误串或 None。

    权威路径守卫在 tools.py（_skill_dir 出自 resolved spec.source_dir，只可能是 .data
    的 user/team skill 根）；此处只做存在性 + SKILL.md 在位 + 非符号链的防御性复核
    （runner cwd=workdir，无法可靠做根前缀判定）。全 bundle 拷贝——脚本靠 __file__ 相对
    读 references/assets 才不碎（与 host 语义一致）。
    """
    if skill_dir.is_symlink() or not skill_dir.is_dir() or not (skill_dir / "SKILL.md").is_file():
        return f"skill dir invalid or missing SKILL.md: {skill_dir}"
    # review-P2 defense in depth: tools.py owns the guard, but runner also asserts _skill_dir is under
    # .data/{user_skills,team_skills}, preventing future runner callers from pointing it at /app or elsewhere.
    _parts = skill_dir.resolve().parts
    if not any(part == ".data" and _parts[i + 1] in {"user_skills", "team_skills", "org_skills"} for i, part in enumerate(_parts[:-3])):
        return f"skill dir not under .data/{{user,team}}_skills: {skill_dir}"
    dst = workdir / "skill"

    def _ignore_symlinks(directory: str, names: list[str]) -> set[str]:
        return {n for n in names if (Path(directory) / n).is_symlink()}

    try:
        shutil.copytree(skill_dir, dst, symlinks=False, ignore=_ignore_symlinks)
    except OSError as e:
        return f"skill copy failed: {e}"
    total = 0
    for p in dst.rglob("*"):
        if not p.is_file():
            continue
        try:
            total += p.stat().st_size
            p.chmod(0o644)  # Remove exec bits (the environment enforces noexec semantics).
        except OSError:
            continue
    if total > _MAX_SKILL_BUNDLE_BYTES:
        shutil.rmtree(dst, ignore_errors=True)
        return "skill bundle exceeds size cap"
    (workdir / "skill_manifest.json").write_text(
        json.dumps({"entry": entry}),
        encoding="utf-8",
    )
    return None


# follow-up② command mode: replace tokens and run guards host-side (runner.main, inside the trusted app container).
# Recognize delimited tokens `{inputs.<name>}` / `{output_dir}`; replacements are always server-generated /workdir paths, never
# LLM/user content, so even a `bash -c` element has no injection. argv is a vector; subprocess.run(list) uses no shell.
_BH_CMD_TOKEN_RE = re.compile(r"\{(inputs\.[A-Za-z0-9_.\-]+|output_dir)\}")


def _escapes_workdir(el: str) -> bool:
    """argv 参数路径守卫：含 `..` 段、或绝对路径不在 /workdir 下 → True（拒）。

    相对路径（如 `scripts/run_pipeline.py`）放行（cwd=/workdir/skill ⊂ /workdir）。
    """
    parts = el.replace("\\", "/").split("/")
    if ".." in parts:
        return True
    if el.startswith("/"):
        norm = os.path.normpath(el)
        return not (norm == "/workdir" or norm.startswith("/workdir/"))
    return False


def _resolve_command(
    command: Any,
    resolved: dict,
    allowlist: set[str],
) -> tuple[list[str] | None, str | None]:
    """把 model-composed command argv 的 token 替换成沙箱真实路径 + 守卫。返 (argv, err)。

    - `{inputs.<名>}` → resolved[名]['path']（/workdir/...）；`{output_dir}` → /workdir。
    - 未知 / 未解析 `{...}` → loud err 且回显合法词表（模型可一轮自纠）。
    - argv[0] ∈ allowlist（解释器 + guaranteed_cli；basename 或全名）。
    - argv[1:] 路径守卫（_escapes_workdir）。
    """
    if not isinstance(command, list) or not command:
        return None, "command 必须是非空 argv 数组"
    declared = sorted(resolved.keys())
    out: list[str] = []
    for el in command:
        if not isinstance(el, (str, int, float, bool)):
            return None, f"command 元素必须是标量（string），得到 {type(el).__name__}"
        s = str(el)
        if "\x00" in s:
            return None, "command 元素含 NUL 字符"
        # Replace each match without closures (avoid loop-variable binding; values are always server-generated /workdir paths).
        parts: list[str] = []
        last = 0
        bad_token: str | None = None
        for m in _BH_CMD_TOKEN_RE.finditer(s):
            parts.append(s[last : m.start()])
            tok = m.group(1)
            if tok == "output_dir":
                parts.append("/workdir")
            else:
                ent = resolved.get(tok[len("inputs.") :])
                if isinstance(ent, dict) and ent.get("path"):
                    parts.append(str(ent["path"]))
                else:
                    bad_token = tok
                    break
            last = m.end()
        if bad_token is not None:
            return None, (f"未知 token {{{bad_token}}}；可引用的 inputs: {declared}；输出目录 token: {{output_dir}}")
        parts.append(s[last:])
        s2 = "".join(parts)
        leftover = re.findall(r"\{[^{}]*\}", s2)
        if leftover:
            return None, (f"未解析 token {leftover}；可引用的 inputs: {declared}；输出目录 token: {{output_dir}}")
        out.append(s2)
    argv0 = out[0]
    base0 = argv0.rsplit("/", 1)[-1]
    if allowlist and base0 not in allowlist and argv0 not in allowlist:
        return None, (f"argv[0] {argv0!r} 不在允许的解释器/CLI 集 {sorted(allowlist)}（跑 skill 脚本用 python3/node/bash + 脚本相对路径）")
    for el in out[1:]:
        if _escapes_workdir(el):
            return None, f"命令路径 {el!r} 跳出 /workdir（脚本用相对路径，输出用 {{output_dir}}）"
    return out, None


def _relay_skill_result(workdir: Path, result: dict) -> None:
    """把子脚本 result.json（shim 已改名 skill_result.json）嵌进 runner envelope."""
    sr = workdir / "skill_result.json"
    if not sr.is_file():
        return
    try:
        if sr.stat().st_size > _MAX_SKILL_RESULT_BYTES:
            result["skill_result"] = {"ok": False, "error": "skill_result.json exceeds size cap"}
            return
    except OSError:
        return
    with contextlib.suppress(OSError, json.JSONDecodeError):
        result["skill_result"] = json.loads(sr.read_text(encoding="utf-8"))


def _unique_data_output_path(workdir: Path, preferred_name: str) -> Path:
    """Reserve a deterministic top-level name without replacing user output."""

    candidate = workdir / preferred_name
    if not candidate.exists():
        return candidate
    stem = candidate.stem
    suffix = candidate.suffix
    for ordinal in range(1, 10_000):
        candidate = workdir / f"{stem}-data-{ordinal}{suffix}"
        if not candidate.exists():
            return candidate
    raise RuntimeError("sandbox data output namespace exhausted")


def _fail_data_materialization(result: dict, reason: str) -> None:
    result["ok"] = False
    if not result.get("exit_code"):
        result["exit_code"] = 1
    prior_error = str(result.get("error") or "").strip()
    result["error"] = f"{prior_error}; {reason}" if prior_error else reason
    result["stderr"] = (result.get("stderr") or "") + f"\n[runner] {reason}"


def _materialize_data_outputs(
    *,
    workdir: Path,
    result: dict,
    outputs: list[dict],
    gateway: str,
    session_id: str | None,
    created_by_turn: str | None,
    token: str | None,
) -> None:
    """Persist stdout and structured results through Artifact truth.

    The inline ``stdout`` field remains a bounded model preview. ``_stdout_full``
    is runner-private and is always removed before result.json is written.
    """

    stdout = result.pop("_stdout_full", result.get("stdout") or "")
    skill_result = result.get("skill_result")
    candidates: list[tuple[str, str, str]] = []
    materialization_errors: list[str] = []
    if stdout:
        candidates.append(("sandbox-stdout.txt", stdout, "stdout"))
    if skill_result is not None:
        try:
            structured = (
                json.dumps(
                    skill_result,
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                    allow_nan=False,
                )
                + "\n"
            )
        except (TypeError, ValueError) as exc:
            materialization_errors.append(
                f"sandbox structured result cannot be materialized: {type(exc).__name__}",
            )
        else:
            candidates.append(("sandbox-result.json", structured, "structured"))

    for preferred_name, content, data_role in candidates:
        encoded = content.encode("utf-8", errors="replace")
        if len(encoded) > _MAX_OUTPUT_BYTES:
            materialization_errors.append(
                f"sandbox {data_role} exceeds Data artifact limit",
            )
            continue
        path = _unique_data_output_path(workdir, preferred_name)
        path.write_bytes(encoded)
        record = _upload_artifact(
            path,
            gateway,
            session_id=session_id,
            created_by_turn=created_by_turn,
            token=token,
        )
        if record is None or not record.get("artifact_id"):
            materialization_errors.append(
                f"sandbox {data_role} Data artifact upload failed",
            )
            continue
        record["data_role"] = data_role
        if data_role == "structured":
            preview = _preview_structured(path)
            if preview:
                record["preview"] = preview
        outputs.append(record)
    for materialization_error in materialization_errors:
        _fail_data_materialization(result, materialization_error)


def _is_collectable_output(path: Path) -> bool:
    """Accept only regular files, without following container-created links."""
    # `/workdir/result -> /etc/passwd` resolves against the host when the trusted
    # runner later calls read_bytes(), because /workdir is a shared bind mount.
    return not path.is_symlink() and path.is_file()


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--workdir", required=True)
    args, _ = ap.parse_known_args()

    workdir = Path(args.workdir).resolve()
    workdir.mkdir(parents=True, exist_ok=True)

    payload = json.loads(Path(args.input).read_text(encoding="utf-8"))

    # skill-multiruntime S1: user skill script mode for multiple languages; _skill_dir/_skill_entry are
    # injected server-side by tools.py, never LLM values. A trusted dispatch shim supplies code in this mode (no
    # LLM code required); copy the scripts into workdir/skill/ and dispatch by extension inside the container.
    _skill_dir = str(payload.get("_skill_dir") or "").strip()
    _skill_entry = str(payload.get("_skill_entry") or "").strip()
    # follow-up② command mode (server marker + argv; see design §4.2 D2/D3). tools.py sets the marker
    # only when the flag is ON; _skill_command alone (inserted into input by an LLM) cannot trigger it or bypass rollout.
    _skill_command = payload.get("_skill_command")
    _command_mode = bool(payload.get("_skill_command_enabled") and isinstance(_skill_command, list) and _skill_command)
    _skill_argv0_allowlist = {str(x) for x in (payload.get("_skill_argv0_allowlist") or [])}
    _skill_mode = bool(_skill_dir and (_skill_entry or _command_mode))
    if _skill_mode:
        _stage_err = _stage_user_skill(workdir, Path(_skill_dir), _skill_entry)
        if _stage_err:
            _write_result(
                workdir,
                {"ok": False, "exit_code": -1, "stdout": "", "stderr": "", "error": _stage_err, "duration_ms": 0, "outputs": []},
            )
            return 0
        code = _SKILL_DISPATCH_SHIM
    else:
        code = str(payload.get("code") or "").strip()
        if not code:
            _write_result(
                workdir,
                {
                    "ok": False,
                    "exit_code": -1,
                    "stdout": "",
                    "stderr": "",
                    "error": "input.code is empty",
                    "duration_ms": 0,
                    "outputs": [],
                },
            )
            return 0

    timeout = int(payload.get("timeout_sec") or 120)
    timeout = max(5, timeout)

    gateway = os.environ.get("BOWHEAD_GATEWAY_URL", "http://127.0.0.1:8080")
    backend = (os.environ.get("BOWHEAD_SANDBOX_BACKEND") or "executor").lower()
    # pptx-from-template M2 (wheelhouse) — offline hash-pinned long-tail deps.
    _wheelhouse = (os.environ.get("BOWHEAD_SANDBOX_WHEELHOUSE") or "").strip() or None

    # 1. Materialize inputs (download artifacts).
    # π §K: tools.py gate decides whether to populate INPUTS schema. The flag
    # is passed in via input.json (env BOWHEAD_PI_K_ENABLED is stripped by
    # SkillRuntime._build_env). When tools.py is bypassed (direct runner call
    # in tests / debug), fall back to env var.
    _k_payload_enabled = payload.get("_pi_k_enabled")
    _k_payload_redact = payload.get("_pi_k_redact_pii")
    _k_enabled = os.environ.get("BOWHEAD_PI_K_ENABLED", "1") == "1" if _k_payload_enabled is None else bool(_k_payload_enabled)
    _k_redact = os.environ.get("BOWHEAD_PI_K_PII_REDACT", "1") == "1" if _k_payload_redact is None else bool(_k_payload_redact)

    # pptx-from-template M3w.2 (RD-2): session/turn threaded from tools.py via
    # input.json control keys (like _pi_k_* above), so sandbox outputs are
    # session-pinned + typed at upload — a generated deck (.pptx) becomes
    # kind=deliverable with created_by_turn = the owning turn, which the
    # is_report_worthy keep-signal (turn_has_deliverable) then sees.
    _session_id = payload.get("_session_id") or None
    _created_by_turn = payload.get("_turn_id") or None
    _output_purpose = str(payload.get("_output_purpose") or "analysis").strip()
    if _output_purpose not in {"analysis", "deliverable"}:
        _output_purpose = "analysis"
    _raw_deliverable_outputs = payload.get("_deliverable_outputs")
    _deliverable_outputs = frozenset(
        str(name) for name in (_raw_deliverable_outputs if isinstance(_raw_deliverable_outputs, list) else []) if isinstance(name, str) and name
    )
    # cca P3 F-B: artifact-fetch token (set by tools.py from the gateway ctx minter). Read
    # HOST-SIDE only; threaded into _materialize_inputs → _fetch_artifact as a Bearer header
    # so the gateway download owner-check runs as the real caller. NEVER written into the
    # container-visible input.json (_container_input below), so it never enters the sandbox.
    _artifact_fetch_token = payload.get("_artifact_fetch_token") or None

    # sandbox-hardening 2026-06-17 (Phase 3 de-DooD): the sandbox-executor URL +
    # shared-secret token, threaded from tools.py (which reads the gateway's full
    # env). The env fallback below is a last resort: SkillRuntime._build_env strips
    # BOWHEAD_SANDBOX_EXECUTOR_* (secrets stay denied even after dogfood-0708 F6,
    # which only allow-listed the NON-secret sandbox config), so env-reading here
    # normally misses them — payload threading is the real path.
    # Like the
    # artifact-fetch token, these are read HOST-SIDE only and are NOT copied into
    # the container-visible _container_input below → never enter the sandbox.
    _executor_url = (payload.get("_sandbox_executor_url") or os.environ.get("BOWHEAD_SANDBOX_EXECUTOR_URL", "")).strip()
    _executor_token = payload.get("_sandbox_executor_token") or os.environ.get("BOWHEAD_SANDBOX_EXECUTOR_TOKEN", "")

    # pptx-from-template M2 (wheelhouse) — skill-declared long-tail deps (set by
    # tools.py from the resolved SkillSpec, NOT the LLM). Regex+lockset-validate
    # here (IDENTICAL contract to sandbox_bootstrap); only validated entries flow
    # into the container, where the bootstrap re-validates before the offline
    # install. FAIL CLOSED: if deps are declared but no wheelhouse is
    # available, return a clear ok:false (never a silent ModuleNotFoundError).
    _raw_reqs = payload.get("_sandbox_requirements") or []
    if not isinstance(_raw_reqs, list):
        _raw_reqs = []
    _valid_reqs: list[str] = []
    _req_errors: list[str] = []
    if _raw_reqs:
        if not (_wheelhouse and (Path(_wheelhouse) / "requirements.lock").is_file()):
            _write_result(
                workdir,
                {
                    "ok": False,
                    "exit_code": -1,
                    "stdout": "",
                    "stderr": "",
                    "error": (
                        "declared sandbox_requirements but no wheelhouse available "
                        "(BOWHEAD_SANDBOX_WHEELHOUSE unset/absent or requirements.lock "
                        "missing) — refusing to run without the requested deps"
                    ),
                    "duration_ms": 0,
                    "outputs": [],
                },
            )
            return 0
        _lockset = _wheelhouse_lockset(_wheelhouse)
        _valid_reqs, _req_errors = _validate_sandbox_requirements(_raw_reqs, _lockset)
        for _err in _req_errors:
            sys.stderr.write(f"[runner] {_err}\n")

    resolved, input_errors = _materialize_inputs(
        payload.get("inputs") or {},
        workdir,
        gateway,
        k_enabled=_k_enabled,
        k_redact_pii=_k_redact,
        artifact_fetch_token=_artifact_fetch_token,  # cca P3 F-B
    )

    # follow-up② command mode: replace tokens and guard argv[0]/paths on the HOST (resolved paths +
    # allowlist live here, inside the trusted app container). Pass resolved argv into the container via **extended skill_manifest.json**,
    # not the stripped input.json. The container shim only calls subprocess.run(argv).
    if _command_mode:
        _cmd_argv, _cmd_err = _resolve_command(_skill_command, resolved, _skill_argv0_allowlist)
        if _cmd_err:
            _write_result(
                workdir,
                {"ok": False, "exit_code": -1, "stdout": "", "stderr": _cmd_err, "error": _cmd_err, "duration_ms": 0, "outputs": []},
            )
            return 0
        (workdir / "skill_manifest.json").write_text(
            json.dumps({"command": _cmd_argv, "cwd": "/workdir/skill"}),
            encoding="utf-8",
        )

    # 2. Write user.py + input.json (for container to read)
    (workdir / "user.py").write_text(_SCAFFOLD + "\n" + code, encoding="utf-8")
    _stage_sandbox_helpers(workdir)
    # pptx-from-template M2 (wheelhouse): include the VALIDATED reqs in the
    # container-visible input.json so sandbox_bootstrap.py reads them (it
    # re-validates). Only written when present → common path input.json shape
    # is unchanged.
    _container_input: dict = {"inputs": resolved}
    # Declared but not materialized = declared names − resolved names. The 3 `_materialize_inputs` failure paths
    # (non-object / missing artifact_id / failed fetch) append an error and continue, so the difference
    # exactly matches unavailable inputs. Derive it from actual runtime state, not parsed error strings.
    _declared_inputs = payload.get("inputs")
    if isinstance(_declared_inputs, dict):
        _unresolved = [str(n) for n in _declared_inputs if n not in resolved]
        if _unresolved:
            _container_input["_unresolved_inputs"] = sorted(_unresolved)
    # L1: scaffold KeyError messages need the skeleton for **the actual entry point**. Render host-side
    # (`tools.py` → `sandbox_call_shape`); only pass through the rendered string here.
    _call_shape = payload.get("_call_shape")
    if isinstance(_call_shape, str) and _call_shape.strip():
        _container_input["_call_shape"] = _call_shape
    if _valid_reqs:
        _container_input["_sandbox_requirements"] = _valid_reqs
    if _skill_mode:
        # P1-2: user business payload (non-control keys such as title/options) passes into the container for scripts to read
        # parameters from /workdir/input.json; control keys + secrets never enter the container.
        # review-P2: beyond explicit control keys, **never pass through any `_`-prefixed key** (fail-closed by shape),
        # so future secret control keys added in tools.py cannot leak even if omitted from _SKILL_CONTROL_KEYS.
        for _k, _v in payload.items():
            if _k in _SKILL_CONTROL_KEYS or _k == "inputs" or _k.startswith("_"):
                continue
            _container_input.setdefault(_k, _v)
    (workdir / "input.json").write_text(
        json.dumps(_container_input, ensure_ascii=False),
        encoding="utf-8",
    )

    # Record pre-existing files to exclude when scanning for new outputs.
    pre_existing = {p.name for p in workdir.iterdir()}

    # 3. Run
    if backend == "subprocess":
        run_result = _run_subprocess(workdir=workdir, timeout=timeout)
    elif backend == "executor":
        if not _executor_url or not _executor_token:
            run_result = {
                "ok": False,
                "exit_code": -1,
                "stdout": "",
                "stderr": "",
                "error": "sandbox executor is not configured",
                "cause": "executor_not_configured",
                "duration_ms": 0,
            }
        else:
            run_result = _run_via_executor(
                _executor_url,
                _executor_token,
                workdir=workdir,
                timeout=timeout,
                requirements=_valid_reqs,
                command_mode=_command_mode,
            )
    else:
        run_result = {
            "ok": False,
            "exit_code": -1,
            "stdout": "",
            "stderr": "",
            "error": f"unsupported sandbox backend: {backend}",
            "cause": "unsupported_sandbox_backend",
            "duration_ms": 0,
        }

    # 4. Scan outputs + upload
    # Capture-on-partial-failure applies to every sandbox mode: stdout,
    # Markdown and files are user-visible Data even when a later statement
    # fails. Run success remains independent from artifact presence.
    outputs: list[dict] = []
    output_data_errors: list[str] = []
    for p in sorted(workdir.iterdir()):
        if p.name in pre_existing:
            continue
        if _skill_mode and p.name == "skill_result.json":
            continue  # Child-script structured return goes into the envelope via _relay_skill_result, not as an artifact.
        if not _is_collectable_output(p):
            continue
        try:
            size = p.stat().st_size
        except OSError:
            continue
        if size <= 0:
            sys.stderr.write(f"[runner] skip {p.name} size={size} (limit {_MAX_OUTPUT_BYTES})\n")
            continue
        if size > _MAX_OUTPUT_BYTES:
            output_data_errors.append(f"sandbox file {p.name} exceeds Data artifact limit")
            continue
        _out_kind = _kind_for_output(
            p.name,
            purpose=_output_purpose,
            deliverable_outputs=_deliverable_outputs,
        )
        rec = _upload_artifact(
            p,
            gateway,
            session_id=_session_id,
            kind=_out_kind,
            created_by_turn=_created_by_turn,
            token=_artifact_fetch_token,  # cca P3 F-B: own outputs as the real caller
        )
        if rec:
            # Structured tabular files (csv/tsv/json/jsonl) get previews so the canvas
            # decomposer can convert them into data_table sections in the Data tab.
            preview = _preview_structured(p)
            if preview:
                rec["preview"] = preview
            outputs.append(rec)
        else:
            output_data_errors.append(f"sandbox file {p.name} Data artifact upload failed")

    # 5. Stitch result.json
    result = {
        **run_result,
        "outputs": outputs,
        "input_artifacts": [
            {"key": key, "artifact_id": item["artifact_id"], "size": item["size"], "sha256": item["sha256"]} for key, item in resolved.items()
        ],
    }
    missing_deliverables = _missing_declared_deliverables(
        outputs,
        purpose=_output_purpose,
        deliverable_outputs=_deliverable_outputs,
    )
    if missing_deliverables:
        output_data_errors.append("declared deliverable outputs were not produced and persisted: " + ", ".join(missing_deliverables))
    for output_data_error in output_data_errors:
        _fail_data_materialization(result, output_data_error)
    if _skill_mode:
        _relay_skill_result(workdir, result)
    # follow-up② command-mode deliverable-presence guard (design D3.5) addresses LLM-behavior P0 silent false
    # success: exit0 without top-level outputs fails with a model-legible hint (often after copying the relative
    # `output/` directory from SKILL.md into the unscanned skill subtree). Never return ok:true/outputs:[] (s_3389df456f cases).
    if _command_mode and result.get("ok") and not outputs:
        result["ok"] = False
        _hint = (
            "命令 exit 0 但 /workdir 顶层无产物被捕获。产物必须写到 {output_dir}（→ 沙箱工作目录顶层），"
            "不要用 SKILL.md 里的 output/ 相对目录（写进 skill 子目录不会被捕获）。"
        )
        result["error"] = ((result.get("error") + "; ") if result.get("error") else "") + _hint
        result["stderr"] = (result.get("stderr") or "") + "\n[runner] " + _hint
    _mark_input_errors(result, input_errors)
    _hint_missing_file(result, resolved)
    _materialize_data_outputs(
        workdir=workdir,
        result=result,
        outputs=outputs,
        gateway=gateway,
        session_id=_session_id,
        created_by_turn=_created_by_turn,
        token=_artifact_fetch_token,
    )
    _write_result(workdir, result)
    # Skill runtime also reads this stdout segment as a simple preview.
    return 0


def _hint_missing_file(result: dict, resolved: dict) -> None:
    """代码报 FileNotFoundError 时，把**真实可用的输入清单**贴回去让模型自纠.

    2026-08-18 某客户实测的失败形状：compute agent 拿到 bi_query 的 artifact 后**没有声明
    `inputs`**，直接照着 bi_query 的文件名去猜路径——`pd.read_csv('INPUTS/sales/xxx.csv')`、
    `pd.read_csv('2026年1-7月onc销售数量中文省份_225e27a.csv')`。裸的 FileNotFoundError 不
    带任何可行动信息，模型于是换个目录再猜一次（同一文件被两个杜撰目录各试一遍）。
    对照组：另一轮它第二次改用 `INPUTS['df26']['path']` 就一次成功——**只要错误里有清单，
    它自己能纠正**。所以这里补的是判据不是散文：有输入就列出名字和真实路径；没输入就明说
    「一个都没声明」并给出声明方式。
    """
    err = (result.get("stderr") or "") + (result.get("error") or "")
    if "FileNotFoundError" not in err and "No such file or directory" not in err:
        return
    # compute-agent-optimize batch 1: on 116 over 30 days, 25% (28/112) of FileNotFoundError cases quoted
    # `INPUTS['x']['path']` as a **string literal**; the "filename" in the error was
    # `INPUTS[...]` itself. The generic hint below differs from that code only by a pair of quotes, which the model may miss,
    # so explicitly identify this shape.
    _literal_hint = ""
    if "No such file or directory: " in err and "INPUTS[" in err.split("No such file or directory: ", 1)[1][:80]:
        _literal_hint = (
            "你把 INPUTS['<key>']['path'] **加了引号写成字符串字面量**, pandas 去找一个叫这个名字的文件当然不存在。"
            "INPUTS 是 user.py 顶部已注入的 dict 变量: 写 pd.read_csv(INPUTS['<key>']['path'])（不加引号）。"
        )
        if resolved:
            # Review batch 1-R7: a literal is the sole cause only when inputs are declared; otherwise removing quotes still raises KeyError.
            # Also explain that inputs must be declared first; do not return early and omit that.
            result["stderr"] = (result.get("stderr") or "") + "\n[runner] " + _literal_hint
            result["error"] = ((result.get("error") + "; ") if result.get("error") else "") + _literal_hint
            return
    if resolved:
        # Report only keys, **never actual paths**: showing `/workdir/a.csv` teaches hardcoding,
        # precisely the problem being fixed. The model needs only the available keys and how to access them.
        listing = ", ".join(f"{name}({Path(str(spec.get('path') or '')).name})" for name, spec in list(resolved.items())[:12])
        hint = (
            f"输入都在 {_INPUTS_DIRNAME}/ 目录下，文件名就是你声明的 key。本次可用：{listing}。"
            f"取法：pd.read_csv(INPUTS['<key>']['path'])（等价于 '{_INPUTS_DIRNAME}/<key>.csv'）。别照抄 artifact 原文件名。"
        )
    else:
        hint = (
            f"本次调用**一个 input 都没声明**，所以 {_INPUTS_DIRNAME}/ 是空的（INPUTS 为空），"
            "凭空猜文件名不会有结果。要读 bi_query 的数据，请在调用 python_sandbox 时传 "
            "inputs={'<key>': {'artifact_id': '<bi_query 返回的 _ref_artifact_id>'}}，"
            f"文件就会出现在 {_INPUTS_DIRNAME}/<key>.csv，代码里用 INPUTS['<key>']['path'] 读。"
        )
        if _literal_hint:
            hint = _literal_hint + " 另外：" + hint
    result["stderr"] = (result.get("stderr") or "") + "\n[runner] " + hint
    result["error"] = ((result.get("error") + "; ") if result.get("error") else "") + hint


def _mark_input_errors(result: dict, input_errors: list) -> None:
    """followup-context-deck #5: 声明的 input 拉取失败 = run 没拿到该拿的数据, 必须
    fail loudly. 旧逻辑只拼 stderr 不翻 ok → 当 user.py catch 缺失输入异常后 exit 0 时
    result.ok 仍 True → tools.py outer_ok=True (_inner_ok 看 result_json.ok) → 卡片假报
    "运行成功" (prod s_3389df456f: ref 5/6/9/10 全 404, 脚本 catch 后 exit 0, 显示成功但
    无数据). in-place 改 result; 空 input_errors → no-op (正常 run 不受影响)."""
    if not input_errors:
        return
    joined = "; ".join(input_errors)
    result["stderr"] = (result.get("stderr") or "") + "\n[runner] input errors: " + joined
    result["ok"] = False
    if not result.get("exit_code"):
        result["exit_code"] = 1
    result["error"] = result.get("error") or ("input errors: " + joined)


def _write_result(workdir: Path, result: dict) -> None:
    (workdir / "result.json").write_text(
        json.dumps(result, ensure_ascii=False),
        encoding="utf-8",
    )


if __name__ == "__main__":
    sys.exit(main())
