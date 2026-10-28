# SPDX-License-Identifier: MIT
"""Exercise independent document packages through their real CLI and public libraries."""

from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from pathlib import Path
from typing import Any

import pytest
from bowhead_skill.loader import load_skill
from bowhead_skill.runtime import SkillRuntime
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pypdf import PdfReader

from bowhead_skill.catalog import SkillCatalog, builtin_catalog_dir

CATALOG = Path(__file__).resolve().parents[1] / "catalog"
KINDS = ("docx", "pdf", "pptx", "xlsx")
OUTPUTS = {"docx": "document.docx", "pdf": "report.pdf", "pptx": "deck.pptx", "xlsx": "workbook.xlsx"}

# Public-library and schema JSON values are heterogeneous at the CLI test boundary.
JsonObject = dict[str, Any]
pytestmark = pytest.mark.integration


def example(kind: str) -> JsonObject:
    return json.loads((CATALOG / f"bowra-{kind}" / "examples/create.json").read_text(encoding="utf-8"))


def invoke(kind: str, request: JsonObject, directory: Path) -> subprocess.CompletedProcess[str]:
    directory.mkdir(parents=True, exist_ok=True)
    input_path = directory / "request.json"
    input_path.write_text(json.dumps(request, ensure_ascii=False), encoding="utf-8")
    return subprocess.run(  # noqa: S603 - fixed Python executable and independently authored CLI.
        [sys.executable, str(CATALOG / f"bowra-{kind}" / "scripts/run.py"), "--input", str(input_path), "--workdir", str(directory)],
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )


@pytest.mark.parametrize("kind", KINDS)
def test_created_documents_reopen_with_expected_content(kind: str, tmp_path: Path) -> None:
    result = invoke(kind, example(kind), tmp_path)
    assert result.returncode == 0, result.stderr
    metadata = json.loads(result.stdout)
    path = tmp_path / OUTPUTS[kind]
    raw = path.read_bytes()
    assert metadata["path"] == path.name
    assert metadata["bytes"] == len(raw)
    assert metadata["sha256"] == hashlib.sha256(raw).hexdigest()
    assert json.loads((tmp_path / "result.json").read_text()) == metadata
    if kind == "docx":
        document = Document(path)
        assert document.paragraphs[0].text == "Monthly delivery review"
        assert document.paragraphs[1].style.name == "Heading 1"
        assert document.tables[0].cell(1, 1).text == "240"
    elif kind == "pdf":
        text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
        assert "Monthly delivery review" in text
        assert "240" in text
        assert "12 minutes" in text
    elif kind == "pptx":
        presentation = Presentation(path)
        assert len(presentation.slides) == 2
        assert presentation.slides[0].shapes[0].text == "Pilot results"
        table = next(shape.table for shape in presentation.slides[1].shapes if shape.has_table)
        assert table.cell(1, 1).text == "240"
    else:
        workbook = load_workbook(path)
        try:
            assert workbook["Metrics"]["B2"].value == 240
            assert workbook["Metrics"]["B2"].data_type == "n"
            assert workbook["Metrics"].freeze_panes == "A2"
            assert workbook["Follow-up"]["B2"].value.date().isoformat() == "2026-10-08"
        finally:
            workbook.close()


@pytest.mark.parametrize("kind", KINDS)
def test_inspection_retains_source_identity_and_supported_content(kind: str, tmp_path: Path) -> None:
    create_dir = tmp_path / "create"
    assert invoke(kind, example(kind), create_dir).returncode == 0
    source_raw = (create_dir / OUTPUTS[kind]).read_bytes()
    inspect_dir = tmp_path / "inspect"
    inspect_dir.mkdir()
    source = inspect_dir / f"source.{kind}"
    source.write_bytes(source_raw)
    result = invoke(kind, {"operation": "inspect", "source": source.name}, inspect_dir)
    assert result.returncode == 0, result.stderr
    report = json.loads((inspect_dir / "inspection.json").read_text())
    assert report["source_sha256"] == hashlib.sha256(source_raw).hexdigest()
    assert source.read_bytes() == source_raw
    assert report["scope"]
    assert report["limitations"]
    assert "240" in json.dumps(report)
    if kind == "docx":
        assert [block["type"] for block in report["blocks"]][:4] == ["paragraph", "paragraph", "paragraph", "table"]
    elif kind == "pdf":
        assert report["pages"][0]["page"] == 1
    elif kind == "pptx":
        assert len(report["slides"]) == 2
    else:
        assert report["sheets"][0]["rows"][1] == ["Requests", 240]


def test_xlsx_preserves_formula_looking_text_and_leading_zero_identifiers(tmp_path: Path) -> None:
    request = {"operation": "create", "sheets": [{"name": "Inputs", "columns": ["Value"], "rows": [["=1+1"], ["00123"]]}]}
    result = invoke("xlsx", request, tmp_path)
    assert result.returncode == 0, result.stderr
    workbook = load_workbook(tmp_path / "workbook.xlsx")
    try:
        assert workbook.active["A2"].value == "=1+1"
        assert workbook.active["A2"].data_type == "s"
        assert workbook.active["A3"].value == "00123"
    finally:
        workbook.close()


def test_xlsx_inspection_labels_formula_without_claiming_calculation(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active["A1"] = "=1+1"
    workbook.save(tmp_path / "source.xlsx")
    workbook.close()
    result = invoke("xlsx", {"operation": "inspect", "source": "source.xlsx"}, tmp_path)
    assert result.returncode == 0, result.stderr
    report = json.loads((tmp_path / "inspection.json").read_text())
    assert report["sheets"][0]["rows"][0][0] == {"formula": "=1+1", "calculated": False}


@pytest.mark.parametrize("kind", KINDS)
def test_reject_mismatched_table_rows_before_delivering_output(kind: str, tmp_path: Path) -> None:
    request = example(kind)
    if kind in {"docx", "pdf"}:
        request["blocks"][2]["rows"][0] = ["Missing value"]
    elif kind == "pptx":
        request["slides"][1]["rows"][0] = ["Missing value"]
    else:
        request["sheets"][0]["rows"][0] = ["Missing value"]
    result = invoke(kind, request, tmp_path)
    assert result.returncode != 0
    assert "column count" in result.stderr
    assert not (tmp_path / OUTPUTS[kind]).exists()


@pytest.mark.parametrize("kind", KINDS)
def test_existing_outputs_are_preserved_on_retry(kind: str, tmp_path: Path) -> None:
    original = b"previous deliverable"
    target = tmp_path / OUTPUTS[kind]
    target.write_bytes(original)
    result = invoke(kind, example(kind), tmp_path)
    assert result.returncode != 0
    assert target.read_bytes() == original


@pytest.mark.parametrize("kind", KINDS)
@pytest.mark.parametrize("source_kind", ["parent", "absolute", "symlink"])
def test_inspection_rejects_sources_outside_staged_workdir(kind: str, source_kind: str, tmp_path: Path) -> None:
    outside = tmp_path / f"outside.{kind}"
    outside.write_bytes(b"private file")
    workdir = tmp_path / "workdir"
    workdir.mkdir()
    if source_kind == "symlink":
        (workdir / f"link.{kind}").symlink_to(outside)
        source = f"link.{kind}"
    elif source_kind == "absolute":
        source = str(outside)
    else:
        source = f"../{outside.name}"
    result = invoke(kind, {"operation": "inspect", "source": source}, workdir)
    assert result.returncode != 0
    assert outside.read_bytes() == b"private file"
    assert not (workdir / "inspection.json").exists()


def test_pdf_plain_text_does_not_become_markup(tmp_path: Path) -> None:
    request = {"operation": "create", "title": "Literal text", "blocks": [{"type": "paragraph", "text": "<b>not a tag</b> & 240"}]}
    result = invoke("pdf", request, tmp_path)
    assert result.returncode == 0, result.stderr
    text = PdfReader(tmp_path / "report.pdf").pages[0].extract_text()
    assert "<b>not a tag</b> & 240" in text


@pytest.mark.parametrize("kind", KINDS)
def test_creation_preserves_requested_chinese_content(kind: str, tmp_path: Path) -> None:
    request = example(kind)
    text = "季度总结"
    if kind in {"docx", "pdf"}:
        request["title"] = text
        if kind == "pdf":
            request["font_mode"] = "cjk"
    elif kind == "pptx":
        request["slides"][0]["title"] = text
    else:
        request["sheets"][0]["rows"][0][0] = text
    result = invoke(kind, request, tmp_path)
    assert result.returncode == 0, result.stderr
    if kind == "docx":
        assert Document(tmp_path / "document.docx").paragraphs[0].text == text
    elif kind == "pdf":
        assert text in PdfReader(tmp_path / "report.pdf").pages[0].extract_text()
    elif kind == "pptx":
        assert Presentation(tmp_path / "deck.pptx").slides[0].shapes[0].text == text
    else:
        workbook = load_workbook(tmp_path / "workbook.xlsx")
        try:
            assert workbook["Metrics"]["A2"].value == text
        finally:
            workbook.close()


@pytest.mark.parametrize("kind", KINDS)
@pytest.mark.anyio
async def test_existing_runtime_executes_independent_package(kind: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("BOWHEAD_SKILL_WORKDIR_ROOT", str(tmp_path))
    spec = load_skill(CATALOG / f"bowra-{kind}")
    runtime = SkillRuntime(python_bin=sys.executable)
    result = await runtime.run(spec, "run.py", example(kind))
    assert result.ok, result.stderr
    metadata = json.loads(result.stdout)
    assert metadata["path"] == OUTPUTS[kind]
    assert result.workdir is not None
    output = result.workdir / OUTPUTS[kind]
    assert output.is_file()
    assert hashlib.sha256(output.read_bytes()).hexdigest() == metadata["sha256"]


def test_independent_catalog_loads_all_four_packages() -> None:
    catalog = SkillCatalog(roots=[CATALOG / f"bowra-{kind}" for kind in KINDS])
    assert catalog.list_names() == [f"bowra-{kind}" for kind in sorted(KINDS)]
    for spec in catalog.list_specs():
        assert spec.script_names == ["run.py"]
        assert spec.kind == "tool"
        assert spec.frontmatter["license"] == "MIT"


def test_standard_builtin_discovery_contains_document_packages() -> None:
    catalog = SkillCatalog(roots=[builtin_catalog_dir()])
    for kind in KINDS:
        spec = catalog.get(f"bowra-{kind}")
        assert spec is not None
        assert spec.source_dir == (CATALOG / f"bowra-{kind}").resolve()
