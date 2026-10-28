"""Exercise the public template through the actual standalone Canvas renderer."""

from __future__ import annotations

import importlib.util
import json
import shutil
import subprocess
import sys
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from lxml import etree
from pptx import Presentation

SKILL = Path(__file__).resolve().parents[1] / "catalog" / "canvas-pptx"


def _sections() -> list[dict]:
    result = [
        {"kind": "h1", "text": "Community project review"},
        {"kind": "eyebrow", "text": "Example report"},
        {"kind": "deck", "markdown": "Synthetic data for an editable presentation."},
        {"kind": "narrative", "markdown": "The project is making steady progress."},
        {
            "kind": "kpi_row",
            "cells": [{"lbl": label, "val": str(i + 5), "sub": "Example count"} for i, label in enumerate(["Tasks", "Reviews", "Checks", "Ideas"])],
        },
        {
            "kind": "recommendations",
            "items": [
                {
                    "priority": "P1",
                    "action": "Review feedback",
                    "rationale": "Collect suggestions from contributors.",
                    "owner": "Project team",
                    "timeline": "Next week",
                },
                {
                    "priority": "P2",
                    "action": "Improve examples",
                    "rationale": "Keep the documentation easy to follow.",
                    "owner": "Contributors",
                    "timeline": "This month",
                },
                {
                    "priority": "P2",
                    "action": "Share results",
                    "rationale": "Publish a short progress summary.",
                    "owner": "Maintainers",
                    "timeline": "Next review",
                },
            ],
        },
    ]
    for kind, title in [("bar", "Completed tasks"), ("line", "Weekly progress")]:
        result.extend(
            [
                {
                    "kind": "chart",
                    "id": kind,
                    "spec": {
                        "kind": kind,
                        "title": title,
                        "columns": [{"key": "label", "type": "text"}, {"key": "count", "type": "numeric", "header": "Count"}],
                        "rows": [{"label": "Week 1", "count": 5}, {"label": "Week 2", "count": 8}],
                    },
                },
                {
                    "kind": "insight_card",
                    "title": title,
                    "thesis": "Progress is steady",
                    "narrative": "These numbers are synthetic test data.",
                    "evidence": [{"kind": "chart", "chart_section_id": kind}],
                },
            ]
        )
    return result


@pytest.fixture
def public_skill(tmp_path: Path) -> Path:
    destination = tmp_path / "canvas-pptx"
    shutil.copytree(SKILL / "scripts", destination / "scripts", ignore=shutil.ignore_patterns("__pycache__"))
    (destination / "templates").mkdir()
    shutil.copyfile(SKILL / "templates/community-neutral.pptx", destination / "templates/community-neutral.pptx")
    (destination / "templates/default.json").write_text(json.dumps({"schema_version": 1, "template": "community-neutral.pptx"}))
    return destination


def _render(
    skill: Path, workdir: Path, *, sections: list[dict] | None = None, locale: str | None = None, fallback: bool = False
) -> tuple[dict, object]:
    workdir.mkdir()
    template = skill / "templates/community-neutral.pptx"
    if fallback:
        template = workdir / "empty.pptx"
        Presentation().save(template)
    payload = {"sections": _sections() if sections is None else sections, "template_path": str(template)}
    if locale:
        payload["locale"] = locale
    (workdir / "input.json").write_text(json.dumps(payload), encoding="utf-8")
    result = subprocess.run(
        [sys.executable, str(skill / "scripts/render_from_canvas.py"), "--input", str(workdir / "input.json"), "--workdir", str(workdir)],
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    return json.loads((workdir / "result.json").read_text()), Presentation(workdir / "output.pptx")


def test_community_template_is_reproducible_and_has_only_synthetic_package_data() -> None:
    spec = importlib.util.spec_from_file_location("community_template_builder", SKILL / "scripts/build_community_template.py")
    assert spec and spec.loader
    builder = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(builder)
    actual = (SKILL / "templates/community-neutral.pptx").read_bytes()
    assert builder.build() == actual
    assert builder.build() == actual

    def inspect_package(blob: bytes) -> None:
        with ZipFile(BytesIO(blob)) as archive:
            for name in archive.namelist():
                assert "thumbnail" not in name and "/media/" not in name, name
                data = archive.read(name)
                if name.endswith(".xlsx"):
                    inspect_package(data)
                if name.endswith((".xml", ".rels")):
                    root = etree.fromstring(data)
                    assert not root.xpath('//*[@TargetMode="External"]'), name
                    if name == "docProps/core.xml":
                        for field in ("creator", "lastModifiedBy", "description", "subject", "keywords"):
                            assert not root.xpath(f'//*[local-name()="{field}"]/text()'), name
                    if name == "docProps/app.xml":
                        for field in ("Company", "Manager", "HyperlinkBase"):
                            assert not root.xpath(f'//*[local-name()="{field}"]/text()'), name

    inspect_package(actual)
    prs = Presentation(BytesIO(actual))
    assert [s.slide_layout.name for s in prs.slides] == ["COVER", "TITLE_ONLY", "KPI_GRID", "CHART_SLIDE", "CHART_SLIDE", "TITLE_ONLY", "CLOSING"]
    assert [len(s.shapes) for s in prs.slides] == [7, 16, 21, 9, 9, 34, 7]
    for index in (3, 4):
        chart = prs.slides[index].shapes[4].chart
        assert [c.label for c in chart.plots[0].categories] == ["Example A", "Example B"]
        assert list(chart.series[0].values) == [1, 2]


@pytest.mark.parametrize("fallback", [False, True], ids=["explicit", "empty-template-fallback"])
def test_community_renderer_replaces_slots_and_embedded_charts(public_skill: Path, tmp_path: Path, fallback: bool) -> None:
    result, prs = _render(public_skill, tmp_path / "rendered", fallback=fallback)
    assert result["slides"] == 7
    assert result["demo_residue_count"] == 0
    text = "\n".join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    for expected in ("Community project review", "Review feedback", "Key metrics", "Recommendations", "Start Review feedback (P1)."):
        assert expected in text
    for index in (3, 4):
        chart = prs.slides[index].shapes[4].chart
        assert [c.label for c in chart.plots[0].categories] == ["Week 1", "Week 2"]
        assert list(chart.series[0].values) == [5, 8]
    assert prs.core_properties.author == ""


def test_community_renderer_preserves_explicit_chinese_and_removes_unused_chart_slides(public_skill: Path, tmp_path: Path) -> None:
    sections = [{"kind": "h1", "text": "用户提供的标题"}, {"kind": "narrative", "markdown": "保留用户选择的语言。"}]
    result, prs = _render(public_skill, tmp_path / "localized", sections=sections, locale="zh-CN")
    assert result["slides"] == 5
    text = "\n".join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    assert "用户提供的标题" in text
    assert "总览指标" in text
    assert "行动建议" in text
    assert not any(s.has_chart for slide in prs.slides for s in slide.shapes)


@pytest.mark.parametrize(
    "manifest",
    [
        {"schema_version": 1, "template": "../outside.pptx"},
        {"schema_version": 1, "template": "canvas-default.pptx"},
        {"schema_version": 2, "template": "community-neutral.pptx"},
    ],
)
def test_invalid_or_unavailable_packaged_template_fails_closed(public_skill: Path, manifest: dict) -> None:
    (public_skill / "templates/default.json").write_text(json.dumps(manifest))
    result = subprocess.run(
        [sys.executable, str(public_skill / "scripts/render_from_canvas.py"), "--help"], capture_output=True, text=True, timeout=30, check=False
    )
    assert result.returncode != 0
    assert "template" in result.stderr.lower()
