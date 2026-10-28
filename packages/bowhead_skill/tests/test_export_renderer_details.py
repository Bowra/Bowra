"""导出渲染细节：数字格式化 + deck 标签语言（dev 1.0.77 dogfood P2-c / P2-f）。

catalog 里的 `render_from_canvas.py` / `render_deck.py` 是**独立脚本**（跑在沙箱里，
不作为包 import），所以这里按路径加载后直接测那几个纯函数。
"""

from __future__ import annotations

import importlib.util
import tempfile
from pathlib import Path
from typing import Any

import pytest

from bowhead_skill.catalog import builtin_catalog_dir

_CATALOG = builtin_catalog_dir()


def _load(rel: str, name: str) -> Any:
    path = _CATALOG / rel
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec and spec.loader, path
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(scope="module")
def docx_mod() -> Any:
    return _load("canvas-docx/scripts/render_from_canvas.py", "_r_docx")


@pytest.fixture(scope="module")
def deck_mod() -> Any:
    return _load("pptx/scripts/render_deck.py", "_r_deck")


# ── P2-c: unformatted numbers ────────────────────────────────────────────────────────


def test_ugly_float_becomes_readable(docx_mod: Any) -> None:
    """实测进用户交付物的两个值。"""
    assert docx_mod._cell_text(21778172.57949218) == "21,778,172.58"
    assert docx_mod._cell_text(18.070109577367795) == "18.07"


def test_integral_float_drops_decimals(docx_mod: Any) -> None:
    assert docx_mod._cell_text(210940000.0) == "210,940,000"


def test_int_untouched_so_years_dont_get_separators(docx_mod: Any) -> None:
    """★ 这条是防我自己：给 int 加千分位会把年份写成 `2,026`。"""
    assert docx_mod._cell_text(2026) == "2026"
    assert docx_mod._cell_text(1) == "1"


def test_tiny_float_keeps_significance(docx_mod: Any) -> None:
    """两位小数会把 0.0034 抹成 `0.00`（把真数据显示成零，比难看更糟）。"""
    assert docx_mod._cell_text(0.0034) == "0.0034"


def test_zero_and_bool_and_string_unaffected(docx_mod: Any) -> None:
    assert docx_mod._cell_text(0.0) == "0"
    assert docx_mod._cell_text(True) == "True"
    assert docx_mod._cell_text("华东区") == "华东区"


def test_nan_and_inf_do_not_crash(docx_mod: Any) -> None:
    assert docx_mod._cell_text(float("nan")) == "nan"
    assert docx_mod._cell_text(float("inf")) == "inf"


def test_rich_cell_shape_still_unwrapped(docx_mod: Any) -> None:
    """`{delta: ...}` 富结构不能因为加了数字分支而失效。"""
    assert docx_mod._cell_text({"delta": 12.5}) == "12.50"
    assert docx_mod._cell_text({"value": "持平"}) == "持平"


def test_pdf_matches_docx(docx_mod: Any) -> None:
    """两个渲染器必须同口径——不然同一份报告导 Word 和导 PDF 数字长得不一样。"""
    pdf = _load("canvas-pdf/scripts/render_from_canvas.py", "_r_pdf")
    for v in (21778172.57949218, 18.070109577367795, 2026, 0.0034, "华东区"):
        assert pdf._cell_text(v) == docx_mod._cell_text(v), v


# ── P2-f: mixed Chinese/English deck text ─────────────────────────────────────────────


def _sec(kind: str, **payload: Any) -> dict[str, Any]:
    return {"kind": kind, "payload": payload}


def test_english_deck_gets_english_labels(deck_mod: Any) -> None:
    """dogfood 实测形态：正文全英文，标签却写死中文「核心指标」「行动建议」。"""
    sections = [
        _sec("h1", text="Supermarket Retail — August 2026"),
        _sec("insight_card", title="Revenue Tracking at 26.37M/Day", thesis="August opened with strong momentum across all regions."),
    ]
    assert deck_mod._label_lang(sections) == "en"
    assert deck_mod._LABELS["en"]["kpi"] == "Key Metrics"


def test_chinese_deck_keeps_chinese_labels(deck_mod: Any) -> None:
    sections = [
        _sec("h1", text="2026 年 8 月经营汇报"),
        _sec("insight_card", title="营收日均 2637 万", thesis="八月开局稳健，各区域均实现增长。"),
    ]
    assert deck_mod._label_lang(sections) == "zh"


def test_chinese_body_with_english_metric_names_stays_chinese(deck_mod: Any) -> None:
    """★ 等权比较会判错：中文表达同样内容用字远少于英文，正文夹几个英文指标名
    就会被算成英文档。"""
    sections = [
        _sec("h1", text="八月经营汇报"),
        _sec(
            "insight_card",
            title="营收与利润",
            thesis="本月 revenue 与 gross profit margin 双双改善，daily active users 同步上行。",
        ),
    ]
    assert deck_mod._label_lang(sections) == "zh"


def test_no_judgable_text_keeps_legacy_chinese(deck_mod: Any) -> None:
    """一个字都判不出来时保持改动前的行为（标签本来就写死中文），不能崩也不该改样。"""
    assert deck_mod._label_lang([]) == "zh"
    assert deck_mod._label_lang([_sec("chart", spec={"rows": [1, 2]})]) == "zh"


# ── Render a real docx: a correct helper does not prove every call site uses it ─────────


def _docx_text(sections: list[dict[str, Any]], docx_mod: Any) -> str:
    """跑真 render()，把整份文档的文字（含表格）拼成一个串。"""
    from docx import Document

    doc = Document()
    docx_mod.render(doc, sections, Path(tempfile.mkdtemp()))
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            out += [c.text for c in row.cells]
    return "\n".join(out)


def test_data_table_renders_formatted_numbers(docx_mod: Any) -> None:
    """★ 端到端：用户真拿到的那份文件里不能再有 `21778172.57949218`。"""
    text = _docx_text(
        [
            {
                "kind": "data_table",
                "columns": [{"key": "region", "header": "区域"}, {"key": "revenue", "header": "营收"}],
                "rows": [{"region": "华东", "revenue": 21778172.57949218}],
            }
        ],
        docx_mod,
    )
    assert "21,778,172.58" in text
    assert "21778172.57949218" not in text


def test_attribution_table_renders_formatted_numbers(docx_mod: Any) -> None:
    """★ 这条抓的正是我第一版漏掉的调用点：attribution 表当时还在直接用 `_sanitize`。

    helper 单测全绿也发现不了它——只有端到端真渲一次才看得出调用点没接上。
    """
    text = _docx_text(
        [
            {
                "kind": "attribution",
                "rows": [{"rank": 1, "dim": "省份", "name": "北京", "from": 18.070109577367795, "to": 21778172.57949218, "pct": 0.0034}],
            }
        ],
        docx_mod,
    )
    assert "18.07" in text
    assert "21,778,172.58" in text
    assert "18.070109577367795" not in text


def test_unknown_kind_leaves_no_english_placeholder(docx_mod: Any, capsys: Any) -> None:
    """交付物里不许出现 `(<kind> — see the live report)`。

    116 dogfood 1.0.74 实测：用户拿到的 Word **开头 4 行**就是这种英文注记（当时漏的
    是 code_block）。补单个 kind 的分支治不了病——下一个新 kind 会原样再踩，所以改的是
    fallback 本身：文件里一个字不打，掉了什么走 stderr。
    """
    text = _docx_text([{"kind": "some_future_kind", "foo": 1}], docx_mod)
    assert "see the live report" not in text
    assert "some_future_kind" not in text
    # Do not drop silently: leave a trace on stderr.
    assert "some_future_kind" in capsys.readouterr().err


def test_unknown_kind_with_text_still_renders(docx_mod: Any) -> None:
    """有正文的未知 kind 仍要渲出来——安静降级不等于扔内容。"""
    text = _docx_text([{"kind": "some_future_kind", "text": "这段必须还在"}], docx_mod)
    assert "这段必须还在" in text


# ── 116 production 1.0.79 dogfood observation A: Markdown tables inside narratives ──────
#
# Writer rule 6 explicitly says "use Markdown tables for comparisons", and frontend remark-gfm renders them correctly on canvas;
# **only the exporter failed to recognize them**. The user's Word document had 14 paragraphs of `| 省份 | 营业额（亿元） |` and
# `|------|------|` separator lines. The system requires Markdown tables, so the renderer must recognize them.

_MD_WITH_TABLE = """下表清晰呈现了各省门店规模与产出效率的对比：

| 省份 | 门店数（家） | 店均营业额（亿元） |
|------|------------|------------------|
| 江苏 | 5 | 2.51 |
| 上海 | 11 | 2.43 |

以上为效率对比。"""


def test_markdown_table_becomes_real_word_table(docx_mod: Any) -> None:
    """★ 端到端：Word 里必须是真表格，且不留任何竖线残渣。"""
    import tempfile as _tf

    from docx import Document

    doc = Document()
    docx_mod._narrative(doc, _MD_WITH_TABLE)
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    assert len(doc.tables) == 1
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["省份", "门店数（家）", "店均营业额（亿元）"]
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["江苏", "5", "2.51"]
    # No separator lines or pipes may remain in body text.
    assert not any("|" in t for t in paras), paras
    assert not any("---" in t for t in paras), paras
    # Preserve body text before and after the table.
    assert paras[0].startswith("下表清晰呈现")
    assert paras[-1] == "以上为效率对比。"
    _tf.mkdtemp()


def test_pipe_line_without_separator_is_not_a_table(docx_mod: Any) -> None:
    """★ 反向闸：孤立一行带 `|` 的文字（如 `A | B 两种口径`）不能被吞成表格。"""
    from docx import Document

    doc = Document()
    docx_mod._narrative(doc, "口径说明：营业额 | 利润 两者不可直接相加。")
    assert len(doc.tables) == 0
    assert "营业额 | 利润" in doc.paragraphs[0].text


def test_pdf_narrative_also_renders_table() -> None:
    """PDF 那条比 Word 还糟：`_flush` 把缓冲行用空格 join 成一整段，
    整张表会摊成一行跑不完的竖线流水账。"""
    from reportlab.platypus import Table

    pdf = _load("canvas-pdf/scripts/render_from_canvas.py", "_r_pdf2")
    flow = pdf._narrative_flowables(_MD_WITH_TABLE, pdf._styles())
    tables = [f for f in flow if isinstance(f, Table)]
    assert len(tables) == 1
    assert tables[0]._cellvalues[0] == ["省份", "门店数（家）", "店均营业额（亿元）"]
    assert tables[0]._cellvalues[1] == ["江苏", "5", "2.51"]


# ── 116 production 1.1.18 dogfood: figure anchors leaked into deliverables and columns silently disappeared ──


def test_figure_anchor_does_not_leak_into_word(docx_mod: Any) -> None:
    """★ `[figure:sec_xxx]` 是画布内部的图片锚点，不是给人读的文本。

    canvas-pdf 一直有两条正则专门剥它，docx **一条都没有** ⇒ 用户的 Word 里躺着
    两行 `[figure:sec_86744592bdcb]`（实测）。与之前那批英文占位符同一类。
    """
    text = _docx_text(
        [{"kind": "narrative", "markdown": "全盘环比显著下滑。\n\n[figure:sec_86744592bdcb]\n\n**图1**：各省环比。"}],
        docx_mod,
    )
    assert "[figure:" not in text, text
    # Preserve both body text and captions.
    assert "全盘环比显著下滑" in text
    assert "图1" in text


def test_inline_figure_anchor_stripped_but_sentence_kept(docx_mod: Any) -> None:
    """锚点嵌在句中时，只剥锚点、不要把整句吞掉。"""
    text = _docx_text([{"kind": "narrative", "markdown": "如下图 [figure:sec_abc123] 所示。"}], docx_mod)
    assert "[figure:" not in text
    assert "所示" in text


def test_wide_table_notes_what_was_dropped(docx_mod: Any) -> None:
    """★ 实测：画布 11 列的环比对比表，Word 里只剩 8 列且**全程无提示**，
    而同一份导成 Excel 是完整 11 列。上限可以有，静默丢不行（导出契约③）。"""
    cols = [
        "省",
        "本月营业额",
        "上月营业额",
        "营业额环比%",
        "本月利润",
        "上月利润",
        "利润环比%",
        "本月利润率%",
        "上月利润率%",
        "本月门店数",
        "上月门店数",
    ]
    row = dict.fromkeys(cols, 1)
    row["省"] = "北京"
    text = _docx_text([{"kind": "chart", "spec": {"rows": [row]}}], docx_mod)
    assert "另有 3 列未显示" in text, text
    assert "上月门店数" in text  # Name the omitted columns explicitly.
    assert "完整数据见画布或 Excel 导出" in text


def test_narrow_table_has_no_note(docx_mod: Any) -> None:
    """★ 反向闸：没丢东西时不能平白多一行提示。"""
    text = _docx_text([{"kind": "chart", "spec": {"rows": [{"省": "北京", "营业额": 1.0}]}}], docx_mod)
    assert "未显示" not in text


def test_long_table_notes_dropped_rows(docx_mod: Any) -> None:
    rows = [{"省": f"p{i}", "v": float(i)} for i in range(60)]
    text = _docx_text([{"kind": "chart", "spec": {"rows": rows}}], docx_mod)
    assert "另有 10 行未显示" in text


# ── Continued 116 production 1.1.18 dogfood investigation: the same issues in other renderers ──


def test_pdf_chart_degrade_table_formats_numbers() -> None:
    """★ 这条降级路**就是生产路**：app 容器没有 matplotlib，`_render_chart_png` 恒返 None
    （116 实测导出的 PDF 内嵌图片数 = 0）。它当时用的是 `_sanitize` 而非 `_cell_text` ⇒
    同一份画布导 PDF 有 **50 个裸浮点**（`157216967.4831249`），导 Word 是 `157,216,967.48`。

    上一批「docx/pdf 同口径」只改到 data_table / attribution 两条路，漏了这条。
    """
    pdf = _load("canvas-pdf/scripts/render_from_canvas.py", "_r_pdf3")
    spec = {
        "title": "各省营业额",
        "columns": [{"key": "省", "header": "省"}, {"key": "v", "header": "营业额"}],
        "rows": [{"省": "北京", "v": 157216967.4831249}],
    }
    flow = pdf._chart_flowables(spec, pdf._styles(), Path(tempfile.mkdtemp()))
    from reportlab.platypus import Table

    tbl = next(f for f in flow if isinstance(f, Table))
    flat = [str(c) for row in tbl._cellvalues for c in row]
    assert "157,216,967.48" in flat, flat
    assert not any("157216967.48312" in c for c in flat), flat


def test_pptx_strips_inline_figure_anchor() -> None:
    """★ 行内锚点会漏进 deck 的 exec slide 正文（独占一行的恰好被句号切断而侥幸不漏
    ——不能靠这个侥幸）。

    ⚠ 必须**端到端渲一份 deck**：只测 `_FIG_ANCHOR` 正则本身是死测试——把调用点改回去
    它照样绿（实测踩过，与之前 docx 漏调用点是同一个坑）。
    """
    import json as _json
    import subprocess
    import sys as _sys

    from pptx import Presentation

    script = _CATALOG / "canvas-pptx/scripts/render_from_canvas.py"
    wd = Path(tempfile.mkdtemp())
    sections = [
        {"kind": "h1", "text": "标题"},
        {"kind": "narrative", "markdown": "如下图 [figure:sec_abc123] 所示，五省环比全面下滑。第二句补充。"},
    ]
    (wd / "input.json").write_text(_json.dumps({"sections": sections, "title": "T"}, ensure_ascii=False), encoding="utf-8")
    result = subprocess.run(
        [_sys.executable, str(script), "--input", str(wd / "input.json"), "--workdir", str(wd)],
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    out = wd / "output.pptx"
    assert out.is_file(), "deck 没渲出来"
    text = []
    for slide in Presentation(str(out)).slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                text.append(sh.text_frame.text)
    blob = "\n".join(text)
    assert "[figure:" not in blob, blob[:400]
    assert "五省环比全面下滑" in blob, blob[:400]


def test_pdf_chart_degrade_notes_dropped_rows() -> None:
    """12 行上限本身合理（页宽），但静默丢违反本文件自己的导出契约③。"""
    from reportlab.platypus import Paragraph

    pdf = _load("canvas-pdf/scripts/render_from_canvas.py", "_r_pdf4")
    spec = {
        "title": "长表",
        "columns": [{"key": "省", "header": "省"}, {"key": "v", "header": "v"}],
        "rows": [{"省": f"p{i}", "v": float(i)} for i in range(20)],
    }
    flow = pdf._chart_flowables(spec, pdf._styles(), Path(tempfile.mkdtemp()))
    texts = [f.text for f in flow if isinstance(f, Paragraph)]
    assert any("另有 8 行未显示" in t for t in texts), texts


def test_pdf_chart_degrade_short_table_has_no_note() -> None:
    """★ 反向闸：没丢就不许多一行提示。"""
    from reportlab.platypus import Paragraph

    pdf = _load("canvas-pdf/scripts/render_from_canvas.py", "_r_pdf5")
    spec = {
        "title": "短表",
        "columns": [{"key": "省", "header": "省"}, {"key": "v", "header": "v"}],
        "rows": [{"省": "北京", "v": 1.0}],
    }
    flow = pdf._chart_flowables(spec, pdf._styles(), Path(tempfile.mkdtemp()))
    assert not any("未显示" in f.text for f in flow if isinstance(f, Paragraph))
