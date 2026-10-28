"""Freeze synthetic planning cases, then probe an explicitly selected team model.

No template parsing, source installation or file writes occur on the remote host.
The existing deployed adapter consumes synthetic JSON over stdin; credentials
stay inside its container. This is not a chat/job/billing integration test.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import io
import json
import shlex
import shutil
from dataclasses import asdict
from pathlib import Path

import anyio
from bowhead_core.documents.package_io import OfficePackageError
from bowhead_core.documents.pptx_draft import compile_pptx_draft
from bowhead_core.documents.pptx_plan import SYSTEM_PROMPT, JsonEvidenceSource, prepare_planning_context, review_planning_response
from bowhead_core.documents.pptx_planning import PlanningReply, PlanningResult, plan_pptx_text
from bowhead_core.documents.pptx_text import inventory_pptx_text
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from verify_pptx_text_adapter import convert
from verify_pptx_text_adapter import run as run_local

# Read-only configuration and a bounded call through modules already deployed.
# Keep this a diagnostic command, not an uploaded worker or build context.
REMOTE_CALL = r"""
import asyncio, hashlib, json, sqlite3, sys, time
from bowhead_core.llm_config_crypto import decrypt_llm_config
from bowhead_core.team_llm_config import provider_configs_from_v2
from bowhead_router.registry import ProviderRegistry
from bowhead_router.types import LLMRequest, Message

async def main():
    request = json.load(sys.stdin)
    db = sqlite3.connect("file:/app/.data/sqlite/bowhead.db?mode=ro", uri=True)
    row = db.execute("SELECT llm_config_version,llm_config_encrypted FROM teams WHERE id=?", (request["team_id"],)).fetchone()
    if not row or row[0] != request["config_version"]:
        raise RuntimeError("configuration_changed")
    config = decrypt_llm_config(row[1])
    configs = [item for item in provider_configs_from_v2(config) if item.name == request["provider"]]
    if len(configs) != 1 or configs[0].type != "anthropic_compatible":
        raise RuntimeError("provider_unavailable")
    models = [item if isinstance(item,str) else item.get("id") for item in configs[0].models]
    if request["model"] not in models:
        raise RuntimeError("model_not_configured")
    registry = ProviderRegistry()
    configs[0].timeout_sec = 90
    configs[0].max_retries = 0
    registry.load(configs)
    adapter = registry.get(request["provider"])
    calls = []
    async def observe(outbound):
        body = json.loads(outbound.content)
        calls.append({key: body.get(key) for key in ("model","temperature","max_tokens","thinking")})
    adapter._http.event_hooks["request"].append(observe)
    started = time.monotonic()
    try:
        response = await adapter.invoke(LLMRequest(model=request["model"], system=request["system"], messages=[Message(role="user",content=request["context"])], temperature=0, max_output_tokens=1800, extended_thinking=False))
        same = db.execute("SELECT llm_config_version,llm_config_encrypted FROM teams WHERE id=?", (request["team_id"],)).fetchone() == row
        print(json.dumps({"content":response.content,"model":response.model,"raw_model":response.raw.get("model"),"finish_reason":response.finish_reason,"raw_usage":response.raw.get("usage"),"physical_requests":calls,"elapsed_seconds":round(time.monotonic()-started,3),"configuration_unchanged":same,"config_version":row[0]}),flush=True)
    except Exception as exc:
        print(json.dumps({"error_type":type(exc).__name__,"physical_requests":calls,"elapsed_seconds":round(time.monotonic()-started,3),"outcome_unknown":True}),flush=True)
    finally:
        await registry.close_all()
        db.close()
asyncio.run(main())
"""


def _write(path: Path, value: object) -> None:
    with path.open("x") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)


def _deck(case: dict[str, object]) -> bytes:
    """Neutral editable fixture; labels have no planner-side special meaning."""
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("F5F5EF")

    def text(value, y, size, height):
        box = slide.shapes.add_textbox(Inches(0.8), Inches(y), Inches(11.7), Inches(height))
        box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        box.text_frame.word_wrap = True
        box.text = value
        for run in box.text_frame.paragraphs[0].runs:
            run.font.name, run.font.size = "Arial", Pt(size)
            run.font.color.rgb = RGBColor.from_string("273F39")

    text(case["title"], 0.6, 32, 0.8)
    text(case["summary"], 1.7, 26, 0.7)
    rows = case["table"]
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.85), Inches(2.8), Inches(11.6), Inches(2.5)).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string("487262" if r == 0 else "E4EBE2")
            for run in cell.text_frame.paragraphs[0].runs:
                run.font.name, run.font.size = "Arial", Pt(22)
                run.font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else "273F39")
    text(case["footer"], 6.7, 12, 0.35)
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _reverse(value: object) -> object:
    if isinstance(value, dict):
        return {key: _reverse(child) for key, child in reversed(list(value.items()))}
    if isinstance(value, list):
        return [_reverse(child) for child in reversed(value)]
    return value


def prepare(corpus: Path, output: Path) -> None:
    payload = json.loads(corpus.read_text())
    cases = payload["cases"]
    permutations = ((cases[0], "c11"), (cases[9], "c12")) if payload["profile"] == "synthetic-unregistered-planning-v1" else ()
    for original, key in permutations:
        case = copy.deepcopy(original)
        case["id"] = key
        case["data"] = [{"name": item["name"], "value": _reverse(item["value"])} for item in case["data"]]
        cases.append(case)
    output.mkdir(parents=True, exist_ok=False)
    manifest = {
        "corpus_sha256": hashlib.sha256(corpus.read_bytes()).hexdigest(),
        "system_sha256": hashlib.sha256(SYSTEM_PROMPT.encode()).hexdigest(),
        "cases": [],
    }
    _write(output / "system.json", SYSTEM_PROMPT)
    for case in cases:
        directory = output / case["id"]
        directory.mkdir()
        source = _deck(case)
        with (directory / "source.pptx").open("xb") as stream:
            stream.write(source)
        _write(directory / "case.json", case)
        data = tuple(JsonEvidenceSource(item["name"], json.dumps(item["value"], ensure_ascii=False).encode()) for item in case["data"])
        context = prepare_planning_context(source, case["instruction"], data)
        _write(directory / "context.json", asdict(context))
        manifest["cases"].append(
            {
                "id": case["id"],
                "context_digest": context.digest,
                "source_sha256": hashlib.sha256(source).hexdigest(),
                "expected_sha256": hashlib.sha256(json.dumps(case["expected"], sort_keys=True).encode()).hexdigest(),
            }
        )
    _write(output / "manifest.json", manifest)
    print(json.dumps({"frozen_cases": len(cases), "output": str(output)}))  # noqa: T201 — developer CLI


def score(directory: Path) -> dict[str, object]:
    case = json.loads((directory / "case.json").read_text())
    source = (directory / "source.pptx").read_bytes()
    data = tuple(JsonEvidenceSource(item["name"], json.dumps(item["value"], ensure_ascii=False).encode()) for item in case["data"])
    context = json.loads((directory / "context.json").read_text())
    response = json.loads((directory / "response.json").read_text())
    result = {"case": case["id"], "expected_action": case["expected"]["action"], "passed": False, "candidate_created": False}
    if (
        "error_type" in response
        or response.get("finish_reason") != "stop"
        or not response.get("configuration_unchanged")
        or not isinstance(response.get("raw_usage"), dict)
    ):
        return {**result, "error": "provider_result_unusable"}
    try:
        review = review_planning_response(source, case["instruction"], data, response["content"], expected_context_digest=context["digest"])
    except OfficePackageError as exc:
        return {**result, "error": exc.code}
    result["review"] = asdict(review)
    if review.action == "propose":
        # Compile a private, unreviewed candidate before consulting offline
        # truth. Test answers score behavior; they never authorize execution.
        try:
            draft = compile_pptx_draft(source, case["instruction"], data, response["content"], expected_context_digest=context["digest"])
        except OfficePackageError as exc:
            return {**result, "error": exc.code}
        with (directory / "candidate.pptx").open("xb") as stream:
            stream.write(draft.candidate)
        _write(directory / "draft.json", asdict(draft.manifest))
        result.update(candidate_created=True, edit_receipt=asdict(draft.manifest.edit_receipt), draft_manifest_sha256=draft.manifest.digest)
    expected = case["expected"]
    if review.action != expected["action"]:
        return {**result, "error": "semantic_action_mismatch"}
    if review.action == "ask":
        slots = {slot.slot_id: slot.text for slot in inventory_pptx_text(source).slots}
        correct = review.reason == expected["reason"] and set(expected["affected"]).issubset({slots[key] for key in review.affected_slot_ids})
        return {**result, "passed": correct, "question_requires_human_review": True}
    actual = {change.before: change.after for change in review.changes}
    if actual != expected["changes"] or len(actual) != len(review.changes):
        return {**result, "error": "semantic_values_or_scope_mismatch", "actual_changes": actual}
    # Equal output numbers can hide selection of a different business metric.
    # Source expectations remain test-side, never repair feedback or model input.
    facts = {fact["id"]: fact for fact in json.loads(context["payload"])["facts"]}
    for change in review.changes:
        if change.before not in expected.get("references", {}):
            continue
        actual_refs = {(data[facts[ref]["source"]].name, facts[ref]["path"]) for ref in change.evidence_refs if ref in facts}
        if actual_refs != {tuple(item) for item in expected["references"][change.before]}:
            return {**result, "error": "semantic_references_mismatch"}
    return {**result, "passed": True}


async def _execute_case(
    args: argparse.Namespace, ssh: str, remote: str, directory: Path, context_digest: str
) -> tuple[PlanningResult, list[dict[str, object]]]:
    responses = []

    async def invoke_model(system_prompt: str, prompt: str) -> PlanningReply:
        attempt = directory / f"attempt-{len(responses) + 1}"
        attempt.mkdir()
        request = {
            "team_id": args.team_id,
            "config_version": args.config_version,
            "provider": args.provider,
            "model": args.model,
            "system": system_prompt,
            "context": prompt,
        }
        _write(attempt / "prompt.json", {"system": system_prompt, "user": prompt})
        _write(attempt / "dispatch.json", {"context_digest": context_digest, "model": args.model, "provider": args.provider})
        completed = await anyio.run_process(
            [ssh, "-o", "BatchMode=yes", "-o", "ConnectTimeout=10", "--", args.ssh_host, remote],
            input=json.dumps(request).encode(),
            check=True,
        )
        response = json.loads(completed.stdout)
        _write(attempt / "response.json", response)
        responses.append(response)
        if (
            response.get("outcome_unknown")
            or not response.get("configuration_unchanged")
            or not isinstance(response.get("raw_usage"), dict)
            or response.get("model") != args.model
            or response.get("raw_model") != args.model
        ):
            raise RuntimeError("Provider outcome or identity is unusable; no automatic retry")
        return PlanningReply(response["content"], response["finish_reason"])

    case = json.loads((directory / "case.json").read_text())
    data = tuple(JsonEvidenceSource(item["name"], json.dumps(item["value"], ensure_ascii=False).encode()) for item in case["data"])
    execution = await plan_pptx_text((directory / "source.pptx").read_bytes(), case["instruction"], data, invoke=invoke_model)
    return execution, responses


def invoke(args: argparse.Namespace) -> None:
    ssh = shutil.which("ssh")
    if ssh is None or args.ssh_host.startswith("-"):
        raise RuntimeError("SSH transport unavailable")
    output = args.run_dir.resolve()
    manifest = json.loads((output / "manifest.json").read_text())
    system = json.loads((output / "system.json").read_text())
    if hashlib.sha256(system.encode()).hexdigest() != manifest["system_sha256"] or system != SYSTEM_PROMPT:
        raise RuntimeError("Frozen prompt changed; create a new evaluation run")
    remote = shlex.join(["docker", "exec", "-i", args.container, "python", "-c", REMOTE_CALL])
    summary = []
    for item in manifest["cases"]:
        directory = output / item["id"]
        context = json.loads((directory / "context.json").read_text())
        case = json.loads((directory / "case.json").read_text())
        data = tuple(JsonEvidenceSource(value["name"], json.dumps(value["value"], ensure_ascii=False).encode()) for value in case["data"])
        current = prepare_planning_context((directory / "source.pptx").read_bytes(), case["instruction"], data)
        if (
            current.digest != item["context_digest"]
            or context != asdict(current)
            or hashlib.sha256(json.dumps(case["expected"], sort_keys=True).encode()).hexdigest() != item["expected_sha256"]
        ):
            raise RuntimeError("Frozen case changed")
        if (directory / "response.json").exists():
            raise RuntimeError("Existing response: never silently reissue an evaluation call")
        # Save dispatch identity before any network I/O. Unknown outcomes are
        # never retried automatically, including a lost SSH response.
        _write(
            directory / "dispatch.json",
            {"context_digest": current.digest, "model": args.model, "provider": args.provider, "config_version": args.config_version},
        )
        execution, responses = anyio.run(_execute_case, args, ssh, remote, directory, current.digest)
        _write(directory / "execution.json", asdict(execution))
        _write(directory / "response.json", responses[-1])
        result = score(directory)
        result["attempt_count"] = len(execution.attempts)
        _write(directory / "score.json", result)
        summary.append(result)
        print(json.dumps({key: value for key, value in result.items() if key not in ("review", "edit_receipt")}), flush=True)  # noqa: T201
    _write(output / "summary.json", summary)


def render(args: argparse.Namespace) -> None:
    """Render only hash-matched synthetic candidates in a fresh QA directory."""
    root = args.run_dir.resolve()
    selected = []
    for item in json.loads((root / "manifest.json").read_text())["cases"]:
        directory = root / item["id"]
        score_record = json.loads((directory / "score.json").read_text())
        if not score_record["passed"] or not score_record["candidate_created"]:
            continue
        for name in ("source", "candidate"):
            source = directory / (name + ".pptx")
            actual = hashlib.sha256(source.read_bytes()).hexdigest()
            if actual != score_record["edit_receipt"][name + "_sha256"]:
                raise RuntimeError("Scored source/candidate bytes changed; refusing conversion")
            selected.append((item["id"], source))
    if not selected:
        raise RuntimeError("No independently scored synthetic candidates to render")
    output = args.output_dir.resolve()
    output.mkdir(parents=True, exist_ok=False)
    records = []
    for case_id, source in selected:
        directory = output / case_id
        directory.mkdir(exist_ok=True)
        record = convert(source, soffice=args.soffice.resolve(), pdftoppm=args.pdftoppm.resolve(), output=directory)
        records.append({"case": case_id, **record})
    _write(output / "receipt.json", {"converter_version": run_local([str(args.soffice.resolve()), "--version"]), "pdfs": records})
    print(json.dumps({"pdf_count": len(records), "visual_verdict": "not_run", "output": str(output)}))  # noqa: T201


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    freeze = commands.add_parser("prepare")
    freeze.add_argument("--corpus", type=Path, required=True)
    freeze.add_argument("--output-dir", type=Path, required=True)
    call = commands.add_parser("run")
    call.add_argument("--run-dir", type=Path, required=True)
    call.add_argument("--ssh-host", required=True)
    call.add_argument("--container", required=True)
    call.add_argument("--team-id", required=True)
    call.add_argument("--config-version", type=int, required=True)
    call.add_argument("--provider", required=True)
    call.add_argument("--model", required=True)
    pdf = commands.add_parser("render")
    pdf.add_argument("--run-dir", type=Path, required=True)
    pdf.add_argument("--output-dir", type=Path, required=True)
    pdf.add_argument("--soffice", type=Path, required=True)
    pdf.add_argument("--pdftoppm", type=Path, required=True)
    args = parser.parse_args()
    if args.command == "prepare":
        prepare(args.corpus, args.output_dir.resolve())
    elif args.command == "run":
        invoke(args)
    else:
        render(args)


if __name__ == "__main__":
    main()
