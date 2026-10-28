"""Local, synthetic-only PPTX adapter / exact-source PDF smoke test.

This deliberately accepts no customer file or arbitrary model-generated plan.
Fixture expectations are test-side approvals, not a semantic mapping algorithm.
Production Office conversion still requires the isolated document worker.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import subprocess
from dataclasses import asdict
from pathlib import Path

from bowhead_core.documents.pptx_text import (
    TextEditScope,
    apply_pptx_text_update,
    inventory_pptx_text,
    parse_text_replacements,
)
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from pypdf import PdfReader


def fixture() -> bytes:
    """Three ordinary slides: multi-run title, native table, grouped narrative."""
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)

    def label(shapes, text, x, y, width, height, size, color="273F39", split=False):
        box = shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        box.text_frame.word_wrap = True
        box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        paragraph = box.text_frame.paragraphs[0]
        pieces = (text[:10], text[10:]) if split else (text,)
        for piece in pieces:
            run = paragraph.add_run()
            run.text = piece
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor.from_string(color)
        return box

    def slide(title, number):
        page = deck.slides.add_slide(deck.slide_layouts[6])
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = RGBColor.from_string("F5F5EF")
        bar = page.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(0.8), Inches(0.1), Inches(5.65))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string("487262")
        bar.line.fill.background()
        label(page.shapes, title, 1, 0.8, 11, 0.9, 32, split=True)
        label(page.shapes, "Synthetic fixture / source-preserving text updates", 1, 6.75, 10, 0.3, 11, "5C6C65")
        label(page.shapes, str(number), 11.7, 6.75, 0.7, 0.3, 11)
        return page

    cover = slide("Quarter one / Service operations", 1)
    label(cover.shapes, "Open cases", 1.1, 2.35, 5, 0.5, 20)
    label(cover.shapes, "59 cases", 1.1, 3.05, 8, 1.1, 50)
    label(cover.shapes, "First-period review", 1.1, 4.55, 10.5, 0.7, 24)
    page = slide("Quarter one / Regional workload", 2)
    table = page.shapes.add_table(3, 3, Inches(1.1), Inches(2.35), Inches(10.8), Inches(2.4)).table
    for row, values in enumerate((("Region", "Open cases", "Owner"), ("North", "42", "Support A"), ("South", "17", "Support B"))):
        for col, value in enumerate(values):
            cell = table.cell(row, col)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string("487262" if row == 0 else "E4EBE2")
            for run in cell.text_frame.paragraphs[0].runs:
                run.font.name, run.font.size = "Arial", Pt(20)
                run.font.color.rgb = RGBColor.from_string("FFFFFF" if row == 0 else "273F39")
    page = slide("Quarter one / Actions", 3)
    group = page.shapes.add_group_shape()
    label(group.shapes, "Triage the oldest unresolved cases.", 1.1, 2.4, 10.8, 0.8, 26)
    label(group.shapes, "Review ownership with regional leads.", 1.1, 3.7, 10.8, 0.8, 26)
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def update(source: bytes, values: dict[str, str]) -> tuple[bytes, dict[str, object]]:
    inventory = inventory_pptx_text(source)
    updates = [{"slot_id": slot.slot_id, "text": values[slot.text]} for slot in inventory.slots if slot.text in values]
    if len(updates) != len(values) or inventory.unsupported_objects:
        raise RuntimeError("Synthetic fixture scope is incomplete")
    scope = TextEditScope(inventory.source_sha256, tuple(item["slot_id"] for item in updates))
    payload = json.dumps({"updates": updates})
    candidate, receipt = apply_pptx_text_update(source, scope, parse_text_replacements(payload))
    return candidate, {**asdict(receipt), "proposal_bytes": len(payload.encode()), "approved_slot_ids": scope.slot_ids}


def run(command: list[str]) -> str:
    result = subprocess.run(command, capture_output=True, text=True, timeout=120, check=True)  # noqa: S603 — fixed local development tools
    return result.stdout.strip()


def convert(source: Path, *, soffice: Path, pdftoppm: Path, output: Path) -> dict[str, object]:
    content = source.read_bytes()
    inventory = inventory_pptx_text(content)
    profile = output / (source.stem + "-lo-profile")
    export_filter = 'pdf:impress_pdf_Export:{"ExportHiddenSlides":{"type":"boolean","value":"true"}}'
    run(
        [str(soffice), "-env:UserInstallation=" + profile.as_uri(), "--headless", "--convert-to", export_filter, "--outdir", str(output), str(source)]
    )
    if source.read_bytes() != content:
        raise RuntimeError("PDF conversion changed its source")
    pdf = output / (source.stem + ".pdf")
    reader = PdfReader(pdf)
    if len(reader.pages) != inventory.slide_count:
        raise RuntimeError("PDF page count does not match its source")
    deck = Presentation(io.BytesIO(content))
    dimensions = (deck.slide_width / 12700, deck.slide_height / 12700)
    for page in reader.pages:
        if any(
            abs(actual - expected) > 0.1
            for actual, expected in zip((float(page.mediabox.width), float(page.mediabox.height)), dimensions, strict=True)
        ):
            raise RuntimeError("PDF page dimensions do not match their source")
    texts = ["".join((page.extract_text() or "").split()) for page in reader.pages]
    for slot in inventory.slots:
        if "".join(slot.text.split()) not in texts[slot.slide - 1]:
            raise RuntimeError("PDF does not contain expected source text")
    run([str(pdftoppm), "-png", "-r", "110", str(pdf), str(output / source.stem)])
    return {
        "source_file": source.name,
        "source_sha256": hashlib.sha256(content).hexdigest(),
        "pdf_file": pdf.name,
        "pdf_sha256": hashlib.sha256(pdf.read_bytes()).hexdigest(),
        "page_count": len(reader.pages),
        "text_and_page_checks": "passed",
        "visual_verdict": "not_run",
        "filter": export_filter,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True, help="New directory; existing outputs are never overwritten")
    parser.add_argument("--soffice", type=Path, required=True)
    parser.add_argument("--pdftoppm", type=Path, required=True)
    args = parser.parse_args()
    output = args.output_dir.resolve()
    output.mkdir(parents=True, exist_ok=False)
    source = fixture()
    first_values = {
        "Quarter one / Service operations": "Quarter two / Service operations",
        "Quarter one / Regional workload": "Quarter two / Regional workload",
        "Quarter one / Actions": "Quarter two / Actions",
        "59 cases": "74 cases",
        "First-period review": "Second-period review",
        "42": "53",
        "17": "21",
        "Triage the oldest unresolved cases.": "Add capacity for the North support queue.",
        "Review ownership with regional leads.": "Confirm a weekly review with regional leads.",
    }
    first, first_receipt = update(source, first_values)
    second_values = {value: value.replace("Quarter two", "Quarter three").replace("Second-period", "Third-period") for value in first_values.values()}
    second_values.update(
        {
            "74 cases": "68 cases",
            "53": "49",
            "21": "19",
            "Add capacity for the North support queue.": "Retain the added North support capacity.",
            "Confirm a weekly review with regional leads.": "Track case age in the weekly regional review.",
        }
    )
    second, second_receipt = update(first, second_values)
    if first_receipt["approved_slot_ids"] != second_receipt["approved_slot_ids"]:
        raise RuntimeError("Value-only refresh changed slot identities")
    pdfs = []
    for name, content in (("reference", source), ("quarter-two", first), ("quarter-three", second)):
        path = output / (name + ".pptx")
        with path.open("xb") as stream:
            stream.write(content)
        pdfs.append(convert(path, soffice=args.soffice.resolve(), pdftoppm=args.pdftoppm.resolve(), output=output))
    report = {
        "fixture": "synthetic service operations, no semantic mapper or model call",
        "source_edit_receipts": [first_receipt, second_receipt],
        "pdfs": pdfs,
        "converter_version": run([str(args.soffice.resolve()), "--version"]),
        "llm_calls": 0,
        "production_ready": False,
        "remaining_checks": [
            "independent visual review",
            "font identity and glyph geometry",
            "native Office reopen",
            "semantic binding",
            "worker and publication",
        ],
    }
    with (output / "receipt.json").open("x") as stream:
        json.dump(report, stream, indent=2)
    print(json.dumps(report, indent=2))  # noqa: T201 — intentional developer CLI output


if __name__ == "__main__":
    main()
